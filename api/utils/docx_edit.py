#
#  Copyright 2026 The InfiniFlow Authors. All Rights Reserved.
#
#  Licensed under the Apache License, Version 2.0 (the "License");
#  you may not use this file except in compliance with the License.
#  You may obtain a copy of the License at
#
#      http://www.apache.org/licenses/LICENSE-2.0
#
#  Unless required by applicable law or agreed to in writing, software
#  distributed under the License is distributed on an "AS IS" BASIS,
#  WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
#  See the License for the specific language governing permissions and
#  limitations under the License.
#
"""docx 段落级编辑共享内核。

流程文档编辑（POST /flow/<id>/document/edit）与对话附件编辑
（POST /files/<id>/edit）共用的校验与应用逻辑。两个入口共享同一套
ops 校验规则（200 处上限/runs 一致性/控制字符清洗/编辑删除互斥）与
同一定位应用事务（先全部定位成功再统一应用），保证两条链路行为同源。

前端契约（para_index 口径）与 rag/app/naive.py Docx.to_paragraphs 一致，
即 GET /files/<id>/content 返回的段落顺序。

校验失败抛 ValueError（消息面向用户，端点可直接返回）。
"""
import hashlib
import os
import re
from io import BytesIO

from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph

from api.utils.doc_utils import doc_to_docx_via_libreoffice, is_doc_file

_CTRL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")

# 单次请求最大改动处数（edits+deletes+inserts+table_edits 合计）
MAX_OPS_PER_REQUEST = 200
# 单段/单单元格文本长度上限
MAX_TEXT_LEN = 20000


def safe_filename(name: str) -> str:
    """剥掉路径分隔符与不可打印控制字符，避免 object name 被前端文件名注入子目录；超长时保留扩展名截断。"""
    base = os.path.basename((name or "").replace("\\", "/"))
    cleaned = "".join(ch for ch in base if ch.isprintable()).strip()
    if not cleaned:
        return "unnamed"
    if len(cleaned) > 200:
        root, ext = os.path.splitext(cleaned)
        digest = hashlib.md5(cleaned.encode("utf-8")).hexdigest()[:8]
        cleaned = f"{root[:160]}_{digest}{ext}"
    return cleaned


def _build_para_map(doc):
    """复刻 rag/app/naive.py Docx.to_paragraphs 的遍历规则，建立
    para_index → ('p', DocxParagraph) / ('table', DocxTable) / ('image', None) 映射。
    w:p 空文本且无图跳过（不占 index）、有图记 image 占 index；w:tbl 整表占一个 index
    （存 DocxTable 实例，供 table_edits 定位单元格）。"""
    para_map = {}
    idx = 0
    for block in doc.element.body:
        if block.tag.endswith("p"):
            p = DocxParagraph(block, doc)
            text = _CTRL_CHARS.sub("", (p.text or "").strip())
            if text:
                para_map[idx] = ("p", p)
                idx += 1
            else:
                # 空文本段落：与 naive.get_picture 一致，检测 pic:pic + a:blip 可解析关系才算图片；
                # 空文本且无图则完全跳过（不占 index，与 naive.to_paragraphs 的 continue 一致）
                has_img = False
                for img in block.xpath(".//pic:pic"):
                    embed = img.xpath(".//a:blip/@r:embed")
                    if embed and embed[0] in doc.part.related_parts:
                        has_img = True
                        break
                if has_img:
                    para_map[idx] = ("image", None)
                    idx += 1
        elif block.tag.endswith("tbl"):
            para_map[idx] = ("table", DocxTable(block, doc))
            idx += 1
    return para_map


_COLOR_RE = re.compile(r"^#[0-9A-Fa-f]{6}$")
_RUN_BOOL_KEYS = ("bold", "italic", "underline", "strike", "superscript", "subscript")


def _parse_runs(raw):
    """解析并校验可选 runs 字段：None → None（走旧整段替换）；
    非法结构/颜色/字号抛 ValueError，由调用方转 400。"""
    if raw is None:
        return None
    if not isinstance(raw, list) or not raw:
        raise ValueError("runs 必须是非空数组或省略")
    if len(raw) > 500:
        raise ValueError("runs 片段数量超限（最多 500）")
    parsed = []
    for r in raw:
        if not isinstance(r, dict):
            raise ValueError("runs 项格式非法")
        text = _CTRL_CHARS.sub("", str(r.get("text") or ""))
        if not text:
            raise ValueError("runs 片段文本不能为空")
        item = {"text": text}
        for k in _RUN_BOOL_KEYS:
            if r.get(k):
                item[k] = True
        for k in ("color", "bg_color"):
            v = r.get(k)
            if v:
                v = str(v)
                if not _COLOR_RE.match(v):
                    raise ValueError(f"{k} 颜色值非法：{v}")
                item[k] = v
        font = r.get("font")
        if font:
            item["font"] = str(font)[:50]
        size = r.get("size")
        if size is not None:
            try:
                size = float(size)
            except (TypeError, ValueError):
                raise ValueError("size 字号必须是数字")
            if not (1 <= size <= 200):
                raise ValueError("size 字号超出范围（1-200pt）")
            item["size"] = size
        parsed.append(item)
    return parsed


def _replace_para_text(p: DocxParagraph, new_text: str):
    """整段替换文本：保留首 run 格式（沿用段落样式），其余 run 清空。
    超链接/域先移除——p.runs 不覆盖其内部 run，残留会导致新旧文本拼接。"""
    el = p._element
    for child in list(el):
        if child.tag in (qn("w:hyperlink"), qn("w:fldSimple")):
            el.remove(child)
    runs = p.runs
    if runs:
        runs[0].text = new_text
        for r in runs[1:]:
            r.text = ""
    else:
        p.add_run(new_text)


def _set_run_font(run, name: str):
    """同时设置西文（ascii/hAnsi）与中文（eastAsia）字体，否则中文不生效。"""
    run.font.name = name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:eastAsia"), name)


def _apply_runs(p: DocxParagraph, runs):
    """按 runs 重写段落文本 run（保留段落级 style/对齐）。
    runs 经 _parse_runs 校验。bg_color 用 w:shd 底纹实现。"""
    # 清空段落内联内容：除直接 run 外还要移除超链接/域（p.runs 不覆盖它们，
    # 残留会拼接进重写后的段落）；保留 pPr 段落属性
    el = p._element
    for child in list(el):
        if child.tag in (qn("w:r"), qn("w:hyperlink"), qn("w:fldSimple")):
            el.remove(child)
    for item in runs:
        run = p.add_run(item["text"])
        if item.get("bold"):
            run.bold = True
        if item.get("italic"):
            run.italic = True
        if item.get("underline"):
            run.underline = True
        if item.get("strike"):
            run.font.strike = True
        if item.get("superscript"):
            run.font.superscript = True
        if item.get("subscript"):
            run.font.subscript = True
        if item.get("color"):
            run.font.color.rgb = RGBColor.from_string(item["color"].lstrip("#"))
        if item.get("bg_color"):
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:fill"), item["bg_color"].lstrip("#"))
            run._element.get_or_add_rPr().append(shd)
        if item.get("font"):
            _set_run_font(run, item["font"])
        if item.get("size"):
            run.font.size = Pt(item["size"])


_ALIGN_VALS = ("left", "center", "right", "justify")
_ALIGN_TO_DOCX = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}


def _parse_block_attrs(e: dict) -> dict:
    """解析可选块级属性 align/indent/heading_level：只收集 payload 中提供的键
    （后端仅应用已提供的键，未提供的段落级属性原样保留）；非法值抛 ValueError。
    heading_level: null=正文 / 1-3=Heading 2-4（与前端 tag-1 约定一致）。"""
    attrs = {}
    if "align" in e:
        v = e.get("align")
        if v is not None:
            if v not in _ALIGN_VALS:
                raise ValueError(f"align 非法：{v}")
            attrs["align"] = v
    if "indent" in e:
        try:
            v = int(e.get("indent"))
        except (TypeError, ValueError):
            raise ValueError("indent 必须是整数")
        if not 0 <= v <= 8:
            raise ValueError("indent 超出范围（0-8）")
        attrs["indent"] = v
    if "heading_level" in e:
        v = e.get("heading_level")
        if v is not None:
            try:
                v = int(v)
            except (TypeError, ValueError):
                raise ValueError("heading_level 必须是整数或 null")
            if not 1 <= v <= 3:
                raise ValueError("heading_level 超出范围（1-3）")
        attrs["heading_level"] = v
    return attrs


def _apply_block_attrs(doc, p: DocxParagraph, attrs: dict):
    """应用块级属性（attrs 经 _parse_block_attrs 校验，只含提供的键）。
    标题样式按 builtin 名查找（Heading 2-4 / Normal），用户文档缺样式时
    best-effort 跳过；缩进每级 600 twips（≈ 编辑器 40px/级）。"""
    if "heading_level" in attrs:
        hl = attrs["heading_level"]
        try:
            p.style = doc.styles["Normal" if hl is None else f"Heading {hl + 1}"]
        except Exception:
            pass  # 样式缺失/文档定制样式表时跳过，文本改动不受影响
    if "align" in attrs:
        p.alignment = _ALIGN_TO_DOCX[attrs["align"]]
    if "indent" in attrs:
        p.paragraph_format.left_indent = Pt(attrs["indent"] * 30)


def _parse_table_edits(raw):
    """解析并校验 table_edits：[{para_index,row,col,new_text,runs?}]。
    与正文 edits 的差异：new_text 允许空串（清空单元格）；空串不得携带 runs。
    非法抛 ValueError（消息带格位），由调用方转 400。"""
    if not isinstance(raw, list):
        raise ValueError("table_edits 必须是数组")
    parsed = []
    for t in raw:
        if not isinstance(t, dict):
            raise ValueError("table_edits 项格式非法")
        try:
            para_index = int(t.get("para_index"))
            row = int(t.get("row"))
            col = int(t.get("col"))
        except (TypeError, ValueError):
            raise ValueError("table_edits 的 para_index/row/col 必须是整数")
        if para_index < 0 or row < 0 or col < 0:
            raise ValueError(f"表格单元格 ({row},{col}) 行列号不能为负")
        new_text = _CTRL_CHARS.sub("", str(t.get("new_text") or ""))
        if len(new_text) > MAX_TEXT_LEN:
            raise ValueError(f"表格单元格 ({row},{col}) 内容不能超过 {MAX_TEXT_LEN} 字")
        runs = _parse_runs(t.get("runs"))
        if not new_text and runs:
            raise ValueError(f"表格单元格 ({row},{col}) 清空时不能携带 runs")
        if runs and "".join(x["text"] for x in runs).strip() != new_text.strip():
            raise ValueError(f"表格单元格 ({row},{col}) runs 文本与 new_text 不一致")
        parsed.append({
            "para_index": para_index, "row": row, "col": col,
            "new_text": new_text, "runs": runs,
        })
    return parsed


def _apply_cell_text(cell, new_text: str, runs):
    """写 python-docx 单元格：runs/整段替换写入首段，其余段落清空
    （保留段落对象——docx 单元格至少需要一个段落）。\\n 由 run.text setter
    自动转 <w:br/>。"""
    paras = cell.paragraphs
    first = paras[0]
    if runs is not None:
        _apply_runs(first, runs)
    else:
        _replace_para_text(first, new_text)
    for p in paras[1:]:
        _replace_para_text(p, "")


def parse_edit_payload(body: dict) -> dict:
    """解析并校验编辑请求体（/flow/<id>/document/edit 与 /files/<id>/edit 共用）。

    body: {edits: [{para_index, new_text, runs?, ...块级属性}], deletes: [para_index],
           inserts: [{after_para_index, new_text, runs?, ...块级属性}],
           table_edits: [{para_index, row, col, new_text, runs?}]}

    返回 {edits, deletes, inserts, table_edits}（已解析校验）；
    非法抛 ValueError（消息面向用户，端点可直接 _err(str(e), 101)）。"""
    edits = body.get("edits") or []
    deletes = body.get("deletes") or []
    inserts = body.get("inserts") or []
    table_edits_raw = body.get("table_edits") or []
    if not isinstance(edits, list) or not isinstance(deletes, list) or not isinstance(inserts, list):
        raise ValueError("edits/deletes/inserts 必须是数组")
    if not isinstance(table_edits_raw, list):
        raise ValueError("table_edits 必须是数组")
    if not edits and not deletes and not inserts and not table_edits_raw:
        raise ValueError("没有需要保存的改动")
    if len(edits) + len(deletes) + len(inserts) + len(table_edits_raw) > MAX_OPS_PER_REQUEST:
        raise ValueError(f"单次最多修改 {MAX_OPS_PER_REQUEST} 处")

    parsed_edits = []
    edit_indexes = set()
    for e in edits:
        if not isinstance(e, dict):
            raise ValueError("edits 项格式非法")
        try:
            para_index = int(e.get("para_index"))
        except (TypeError, ValueError):
            raise ValueError("para_index 必须是整数")
        new_text = _CTRL_CHARS.sub("", str(e.get("new_text") or "")).strip()
        if not new_text:
            raise ValueError("段落内容不能为空")
        if len(new_text) > MAX_TEXT_LEN:
            raise ValueError(f"单段内容不能超过 {MAX_TEXT_LEN} 字")
        try:
            runs = _parse_runs(e.get("runs"))
            block_attrs = _parse_block_attrs(e)
        except ValueError as ve:
            raise ValueError(f"段落 {para_index} 格式非法：{ve}")
        # 前端 newText 为 trim 后文本而 runs 来自未 trim 的块文本，两侧 strip 后再比
        if runs and "".join(x["text"] for x in runs).strip() != new_text:
            raise ValueError(f"段落 {para_index} runs 文本与 new_text 不一致")
        parsed_edits.append((para_index, new_text, runs, block_attrs))
        edit_indexes.add(para_index)

    parsed_deletes = []
    seen_deletes = set()
    for d in deletes:
        try:
            idx = int(d)
        except (TypeError, ValueError):
            raise ValueError("deletes 项必须是整数段落号")
        if idx in seen_deletes:
            raise ValueError(f"段落 {idx} 重复删除")
        seen_deletes.add(idx)
        parsed_deletes.append(idx)
    overlap = edit_indexes & set(parsed_deletes)
    if overlap:
        raise ValueError(f"段落 {sorted(overlap)} 不能同时修改和删除")

    parsed_inserts = []
    for ins in inserts:
        if not isinstance(ins, dict):
            raise ValueError("inserts 项格式非法")
        try:
            after = int(ins.get("after_para_index"))
        except (TypeError, ValueError):
            raise ValueError("after_para_index 必须是整数（-1 表示文档开头）")
        new_text = _CTRL_CHARS.sub("", str(ins.get("new_text") or "")).strip()
        if not new_text:
            raise ValueError("新段落内容不能为空")
        if len(new_text) > MAX_TEXT_LEN:
            raise ValueError(f"新段落内容不能超过 {MAX_TEXT_LEN} 字")
        try:
            runs = _parse_runs(ins.get("runs"))
            block_attrs = _parse_block_attrs(ins)
        except ValueError as ve:
            raise ValueError(f"新段落格式非法：{ve}")
        if runs and "".join(x["text"] for x in runs).strip() != new_text:
            raise ValueError("新段落 runs 文本与 new_text 不一致")
        parsed_inserts.append((after, new_text, runs, block_attrs))

    parsed_table_edits = _parse_table_edits(table_edits_raw)

    return {
        "edits": parsed_edits,
        "deletes": parsed_deletes,
        "inserts": parsed_inserts,
        "table_edits": parsed_table_edits,
    }


def ensure_docx_blob(blob: bytes, file_name: str) -> tuple:
    """归一化编辑底稿：docx 直接返回；老格式 .doc 先经 LibreOffice 转 docx
    （与 /files/<id>/content 的解析路径一致，para_index 映射保持同源）。
    返回 (docx_blob, 去扩展名文件根, 是否由 doc 转换而来)；其他格式抛 ValueError。"""
    is_docx = (file_name or "").lower().endswith(".docx")
    if is_docx:
        return blob, os.path.splitext(file_name or "document")[0], False
    if not is_doc_file(blob, file_name or ""):
        raise ValueError("仅支持编辑 doc/docx 文档")
    converted = doc_to_docx_via_libreoffice(blob)
    if not converted:
        raise ValueError(".doc 转换失败，暂无法编辑该文档")
    return converted, os.path.splitext(file_name or "document")[0], True


def apply_ops_to_doc(doc, parsed: dict) -> None:
    """在 python-docx Document 上定位并统一应用编辑 ops（原地修改）。
    先全部定位成功，再按 删除 → 插入 → 改写 → 表格 顺序应用
    （改写持有元素引用，不受结构变化影响）；任一定位失败抛 ValueError，
    文档保持原状（半改状态不可能出现——定位全部先于应用）。"""
    para_map = _build_para_map(doc)
    keys_sorted = sorted(para_map)

    def _entry_of(idx):
        entry = para_map.get(idx)
        if not entry:
            return None
        kind, target = entry
        if kind != "p":
            label = "表格" if kind == "table" else "图片"
            return ("atomic", label, None)
        return ("p", None, target)

    located_edits = []
    for para_index, new_text, runs, block_attrs in parsed["edits"]:
        entry = _entry_of(para_index)
        if entry is None:
            raise ValueError(f"段落 {para_index} 定位失败，文档可能已变化，请刷新后重试")
        if entry[0] == "atomic":
            raise ValueError(f"段落 {para_index} 是{entry[1]}，不支持编辑")
        located_edits.append((entry[2], new_text, runs, block_attrs))

    located_deletes = []
    for para_index in parsed["deletes"]:
        entry = _entry_of(para_index)
        if entry is None:
            raise ValueError(f"段落 {para_index} 定位失败，文档可能已变化，请刷新后重试")
        if entry[0] == "atomic":
            raise ValueError(f"段落 {para_index} 是{entry[1]}，不支持删除")
        located_deletes.append(entry[2])

    located_inserts = []  # (mode, ref_paragraph|None, text, style_src|None, runs|None, block_attrs|None)
    delete_set = set(parsed["deletes"])
    for after, new_text, runs, block_attrs in parsed["inserts"]:
        if after >= 0 and para_map.get(after) is None and after != -1:
            raise ValueError(f"插入位置 {after} 无效")
        style_src = None
        if after >= 0:
            anchor = para_map.get(after)
            if anchor and anchor[0] == "p":
                style_src = anchor[1]
        # 插入参照段必须跳过本请求将删除的段落：ref 先被删会脱离文档树，
        # insert_paragraph_before 对已分离元素静默写入空气段（内容丢失）
        nxt = next(
            (
                para_map[k][1]
                for k in keys_sorted
                if k > after and k not in delete_set and para_map[k][0] == "p"
            ),
            None,
        )
        if nxt is not None:
            located_inserts.append(("before", nxt, new_text, style_src, runs, block_attrs))
        else:
            located_inserts.append(("append", None, new_text, style_src, runs, block_attrs))

    # table_edits 同样先全部定位成功，再统一应用（保持事务语义）；越界行列在此拦截，避免 IndexError
    located_table_edits = []
    for t in parsed["table_edits"]:
        entry = para_map.get(t["para_index"])
        if entry is None:
            raise ValueError(f"段落 {t['para_index']} 定位失败，文档可能已变化，请刷新后重试")
        if entry[0] != "table":
            raise ValueError(f"段落 {t['para_index']} 不是表格，table_edits 定位非法")
        table = entry[1]
        if t["row"] >= len(table.rows) or t["col"] >= len(table.columns):
            raise ValueError(
                f"表格 {t['para_index']} 单元格 ({t['row']},{t['col']}) 超出范围"
            )
        located_table_edits.append((table.cell(t["row"], t["col"]), t["new_text"], t["runs"]))

    for p in located_deletes:
        el = p._element
        el.getparent().remove(el)
    for mode, ref, text, style_src, rns, attrs in located_inserts:
        if mode == "append":
            new_p = doc.add_paragraph(text)
        else:
            new_p = ref.insert_paragraph_before(text)
        if style_src is not None:
            try:
                new_p.style = style_src.style
            except Exception:
                pass
        if attrs:
            _apply_block_attrs(doc, new_p, attrs)
        if rns is not None:
            _apply_runs(new_p, rns)
    for target, text, rns, attrs in located_edits:
        if attrs:
            _apply_block_attrs(doc, target, attrs)
        if rns is not None:
            _apply_runs(target, rns)
        else:
            _replace_para_text(target, text)
    for cell, text, rns in located_table_edits:
        _apply_cell_text(cell, text, rns)


def edit_docx_blob(blob: bytes, file_name: str, parsed: dict) -> tuple:
    """一步到位：底稿归一化（.doc→docx）+ 应用 ops + 序列化。
    返回 (new_docx_blob, 文件根, 是否由 doc 转换而来)；非法抛 ValueError。"""
    docx_blob, root, converted = ensure_docx_blob(blob, file_name)
    doc = DocxDocument(BytesIO(docx_blob))
    apply_ops_to_doc(doc, parsed)
    out = BytesIO()
    doc.save(out)
    return out.getvalue(), root, converted
