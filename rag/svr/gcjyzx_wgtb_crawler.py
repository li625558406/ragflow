#!/usr/bin/env python3
#
#  Copyright 2026 The InfiniFlow Authors. All Rights Reserved.
#
#  Licensed under the Apache License, Version 2.0 (the "License");
#  you may not use this file except in compliance with the License.
#  You may obtain a copy of the License at
#
#      http://www.apache.org/licenses/LICENSE-2.0
#
#  Unless required by applicable law or agreed to in writing, software
#  distributed under the License is distributed on an "AS IS" BASIS,
#  WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
#  See the License for the specific language governing permissions and
#  limitations under the License.
#
"""
Crawler for gcjyzx.zhangzhou.gov.cn — 违规通报 (wgtb) section.

Static server-rendered listing, currently 11 articles on a single page.

Data flow
─────────
  1. **Listing**: GET /wgtb/moreinfo.html
     → Parse <li class="list-item"> entries for title, URL, date.
  2. **Detail**: GET /wgtb/{YYYYMMDD}/{uuid}.html  (SSR HTML)
     → Parse h3.title, .toolbar metadata, .content-main body, attachment links.
  3. **Attachments** (same patterns as zcfg/jyxx):
     - /zzggzy/{infoid}/{fileuuid}/{filename} → direct binary download
     - ztbfjyz(...) onclick → same path extraction
     - /ct_outlink_displace=... → external government site redirects (link only)
  4. ZIP auto-extraction with content parsing (pdfplumber, python-docx, openpyxl).

Checkpoint/resume: state saved every 5 articles.  Time-bounded check
(default 3300s) stops gracefully before the 3600s task-timeout window.

Usage
-----
    python gcjyzx_wgtb_crawler.py \
        --tenant-id <TENANT_ID> \
        --kb-id <KB_ID> \
        --task-name <NAME>

    # Optional:
        --max-articles 100       # Limit records (0=unlimited)
        --full                   # Ignore saved state, re-crawl
        --max-runtime 3300       # Max runtime before graceful stop
"""

import argparse
import json
import logging
import os
import random
import re
import sys
import time
import urllib.parse
import urllib.request
import zipfile
from datetime import datetime

_SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
_PROJECT_ROOT = os.path.abspath(os.path.join(_SCRIPT_DIR, "..", ".."))
sys.path.insert(0, _PROJECT_ROOT)

from common import settings
from common.log_utils import init_root_logger
from common.misc_utils import get_uuid
from bs4 import BeautifulSoup


# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------

_SITE_ROOT = "https://gcjyzx.zhangzhou.gov.cn"
_LISTING_PATH = "/wgtb/moreinfo.html"
_SECTION_LABEL = "违规通报"

# Checkpoint batch size (articles)
_BATCH_SIZE = 5

# Default max runtime (55 min, 5 min margin)
_MAX_RUNTIME_DEFAULT = 3300

# Anti-crawling delays (seconds)
_REQUEST_DELAY_MIN = 0.3
_REQUEST_DELAY_MAX = 1.0

# State filename
_STATE_FILENAME = "_crawler_state.json"

_USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/120.0.0.0 Safari/537.36"
)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _safe_print(msg):
    try:
        print(msg)
    except UnicodeEncodeError:
        print(msg.encode("gbk", errors="replace").decode("gbk"))


def _sanitize_filename(text, max_len=120):
    if not text:
        return "untitled"
    name = re.sub(r'[\\/:*?"<>|]', "_", str(text).strip())
    name = re.sub(r'\s+', "_", name).strip("._ ")
    return (name or "untitled")[:max_len]


def _extract_infoid_from_url(url):
    """Extract article UUID from detail URL like /wgtb/20260403/{uuid}.html."""
    m = re.search(r"/([0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12})\.html", url)
    return m.group(1) if m else None


def _extract_onclick_path(onclick):
    """Extract file path from ztbfjyz('/zzggzy/...','0','1') onclick."""
    if not onclick:
        return None
    m = re.search(r"""ztbfjyz\s*\(\s*['"]([^'"]+)['"]""", onclick)
    return m.group(1) if m else None


# ---------------------------------------------------------------------------
# HTTP helpers
# ---------------------------------------------------------------------------

def _http_get(url, referer=None, timeout=30):
    """GET request returning bytes or None."""
    req = urllib.request.Request(url)
    req.add_header("User-Agent", _USER_AGENT)
    if referer:
        req.add_header("Referer", referer)
    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            return resp.read()
    except Exception as e:
        logging.error("GET %s failed: %s", url, e)
        return None


def _download_binary(url, referer=None, timeout=60):
    """Download binary content; returns bytes or None."""
    req = urllib.request.Request(url)
    req.add_header("User-Agent", _USER_AGENT)
    if referer:
        req.add_header("Referer", referer)
    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            return resp.read()
    except Exception as e:
        logging.error("Download failed: %s — %s", url[:120], e)
        return None


# ---------------------------------------------------------------------------
# Listing crawler
# ---------------------------------------------------------------------------

def _crawl_listing_page(page_num):
    """Fetch and parse one listing page.

    Returns (articles_list, has_more).
    articles_list: [{"title", "url", "date_str", "infoid"}, ...]
    """
    if page_num <= 1:
        url = _SITE_ROOT + _LISTING_PATH
    else:
        url = f"{_SITE_ROOT}/wgtb/{page_num}.html"

    html_bytes = _http_get(url, referer=_SITE_ROOT + "/")
    if not html_bytes:
        return [], False

    html = html_bytes.decode("utf-8", errors="replace")
    soup = BeautifulSoup(html, "html.parser")

    articles = []
    for li in soup.find_all("li", class_="list-item"):
        a = li.find("a", href=True)
        if not a:
            continue
        href = a["href"].strip()
        title = ""
        p = a.find("p", class_="name")
        if p and p.get("title"):
            title = p["title"].strip()
        if not title:
            title = (a.get("title") or a.get_text()).strip()
        if not title or len(title) < 2:
            continue

        date_str = ""
        span = li.find("span", class_="date")
        if span:
            date_str = span.get_text(strip=True)

        infoid = _extract_infoid_from_url(href)

        articles.append({
            "title": title,
            "url": _SITE_ROOT + href,
            "date_str": date_str,
            "infoid": infoid,
        })

    # Check if next page exists by fetching it
    has_more = False
    if len(articles) > 0:
        next_url = f"{_SITE_ROOT}/wgtb/{page_num + 1}.html"
        has_more = _http_get(next_url, referer=url) is not None

    return articles, has_more


# ---------------------------------------------------------------------------
# Detail page parsing
# ---------------------------------------------------------------------------

def _parse_detail_html(html_bytes, detail_url):
    """Parse a wgtb article detail page.

    Returns dict with: title, pub_date, info_source, infoid,
    content_html, content_text, attachments.
    """
    if not html_bytes:
        return None

    html = html_bytes.decode("utf-8", errors="replace")
    soup = BeautifulSoup(html, "html.parser")

    info = {
        "title": "",
        "pub_date": "",
        "info_source": "",
        "infoid": _extract_infoid_from_url(detail_url),
        "content_html": "",
        "content_text": "",
        "attachments": [],
    }

    # Title
    title_el = soup.select_one("h3.title")
    if title_el:
        info["title"] = title_el.get_text(strip=True)

    # Toolbar metadata
    toolbar = soup.select_one(".toolbar")
    if toolbar:
        tb_text = toolbar.get_text()
        src_m = re.search(r"信息来源[：:]\s*(.+?)\s*发布时间", tb_text, re.DOTALL)
        if src_m:
            info["info_source"] = src_m.group(1).strip()
        date_m = re.search(r"发布时间[：:]\s*(\d{4}-\d{2}-\d{2})", tb_text)
        if date_m:
            info["pub_date"] = date_m.group(1).strip()

    # Content area
    content_main = soup.select_one(".content-main")
    if content_main:
        for tag in content_main.find_all(["script", "style"]):
            tag.decompose()
        info["content_html"] = content_main.decode_contents()
        info["content_text"] = content_main.get_text(separator="\n", strip=True)

    # Attachments — four types:
    #   a) ztbfjyz onclick
    #   b) /zzggzy/{infoid}/{fileid}/{filename} direct download
    #   c) /ct_outlink_displace=... external government site redirect
    #   d) Absolute URLs with document extensions
    for a_tag in soup.find_all("a", href=True):
        href = (a_tag.get("href") or "").strip()
        onclick = a_tag.get("onclick", "")
        text = (a_tag.get_text() or "").strip()
        title_attr = a_tag.get("title", "")

        filename = title_attr or text or ""

        # Type a: ztbfjyz onclick
        path = _extract_onclick_path(onclick)
        if path:
            info["attachments"].append({
                "filename": filename,
                "url": _SITE_ROOT + path,
                "type": "download",
            })
            continue

        # Type b: direct /zzggzy/ link
        if href.startswith("/zzggzy/"):
            info["attachments"].append({
                "filename": filename or os.path.basename(href.split("?")[0]),
                "url": _SITE_ROOT + href,
                "type": "download",
            })
            continue

        # Type c: /ct_outlink_displace external redirect
        if "/ct_outlink_displace=" in href:
            url_full = href if href.startswith("http") else _SITE_ROOT + href
            info["attachments"].append({
                "filename": filename or "外部链接",
                "url": url_full,
                "type": "external",
            })
            continue

        # Type d: document extensions in href
        ext = os.path.splitext(href.split("?")[0])[1].lower()
        if ext in (".pdf", ".doc", ".docx", ".xls", ".xlsx",
                    ".ppt", ".pptx", ".zip", ".rar", ".7z"):
            url_full = href if href.startswith("http") else _SITE_ROOT + href
            info["attachments"].append({
                "filename": filename or os.path.basename(href.split("?")[0]),
                "url": url_full,
                "type": "download",
            })

    return info


# ---------------------------------------------------------------------------
# Attachment download + ZIP extraction
# ---------------------------------------------------------------------------

def _download_attachments(attachments, dest_dir):
    """Download all downloadable (type='download') attachments.

    Returns list of local file paths.
    """
    downloaded = []
    if not attachments:
        return downloaded

    os.makedirs(dest_dir, exist_ok=True)

    for att in attachments:
        if att.get("type") != "download":
            continue

        url = att["url"]
        filename = att["filename"]
        if not filename or filename in ("外部链接", "unknown"):
            filename = os.path.basename(url.split("?")[0].split("#")[0])
        filename = _sanitize_filename(filename, max_len=120)

        dest_path = os.path.join(dest_dir, filename)
        if os.path.exists(dest_path):
            downloaded.append(dest_path)
            continue

        # URL-encode Chinese characters in filename portion
        parts = url.rsplit("/", 1)
        if len(parts) > 1:
            url = parts[0] + "/" + urllib.parse.quote(parts[1])

        body = _download_binary(url, referer=_SITE_ROOT + "/")
        if body and len(body) >= 100:
            with open(dest_path, "wb") as f:
                f.write(body)
            downloaded.append(dest_path)
            logging.info("Downloaded %s (%d bytes)", filename, len(body))
        else:
            logging.warning("Download failed/too small: %s", filename)

    return downloaded


def _extract_zip(zip_path):
    """Extract a ZIP file; returns list of extracted paths. ZIP is removed after."""
    extracted = []
    dest_dir = os.path.dirname(zip_path)
    try:
        with zipfile.ZipFile(zip_path, "r") as zf:
            for name in zf.namelist():
                safe_name = re.sub(r'[\\/:*?"<>|]', "_", os.path.basename(name))
                if not safe_name:
                    continue
                dest_path = os.path.join(dest_dir, safe_name)
                if os.path.exists(dest_path):
                    continue
                with open(dest_path, "wb") as f:
                    f.write(zf.read(name))
                extracted.append(dest_path)
        os.remove(zip_path)
    except Exception as e:
        logging.warning("ZIP extract error for %s: %s", zip_path, e)
    return extracted


# ---------------------------------------------------------------------------
# Text extraction from binary files
# ---------------------------------------------------------------------------

def _extract_text_from_file(filepath):
    """Extract plain text from PDF / DOCX / XLSX / TXT files."""
    ext = os.path.splitext(filepath)[1].lower()
    try:
        if ext == ".txt":
            with open(filepath, "r", encoding="utf-8", errors="replace") as f:
                return f.read()
        elif ext == ".pdf":
            import pdfplumber
            parts = []
            with pdfplumber.open(filepath) as pdf:
                for pg in pdf.pages:
                    text = pg.extract_text()
                    if text:
                        parts.append(text)
            return "\n\n".join(parts)
        elif ext in (".docx", ".doc"):
            import docx
            doc = docx.Document(filepath)
            return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        elif ext in (".xlsx", ".xls"):
            import openpyxl
            wb = openpyxl.load_workbook(filepath, read_only=True)
            parts = []
            for ws in wb.worksheets:
                rows = []
                for row in ws.iter_rows(values_only=True):
                    rows.append(" | ".join(str(c) if c is not None else "" for c in row))
                if rows:
                    parts.append(f"### {ws.title}\n" + "\n".join(rows))
            wb.close()
            return "\n\n".join(parts)
    except Exception as e:
        logging.warning("Text extraction failed for %s: %s", filepath, e)
    return None


# ---------------------------------------------------------------------------
# Markdown builder
# ---------------------------------------------------------------------------

def _build_markdown(detail, download_dir, source_url):
    """Build a Markdown document from article detail + attachments."""
    info = detail or {}
    title = info.get("title", "无标题")
    pub_date = info.get("pub_date", "")
    info_source = info.get("info_source", "")

    lines = [
        f"# {title}",
        "",
        f"**数据来源:** 漳州市工程项目交易中心 — {_SECTION_LABEL}",
        f"**页面地址:** {source_url}",
        f"**抓取时间:** {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
    ]
    if pub_date:
        lines.append(f"**发布时间:** {pub_date}")
    if info_source:
        lines.append(f"**信息来源:** {info_source}")
    lines.append("")

    # Main content
    content_text = info.get("content_text", "")
    if content_text:
        lines.append("---")
        lines.append("")
        lines.append("## 正文")
        lines.append("")
        content_clean = re.sub(r"\n{3,}", "\n\n", content_text)
        lines.append(content_clean)
        lines.append("")

    # Attachments
    attachments = info.get("attachments", [])
    if attachments:
        lines.append("---")
        lines.append("")
        lines.append("## 附件")
        lines.append("")

        for att in attachments:
            fname = att.get("filename", "unknown")
            att_type = att.get("type", "")
            att_url = att.get("url", "")

            if att_type == "external":
                lines.append(f"- [{fname}]({att_url}) （外部链接）")
            else:
                lines.append(f"- {fname}")
        lines.append("")

        # Embed extracted attachment text for downloaded files
        if download_dir and os.path.isdir(download_dir):
            lines.append("### 附件内容")
            lines.append("")
            for att in attachments:
                if att.get("type") != "download":
                    continue
                fname = att.get("filename", "")
                local_path = os.path.join(download_dir, fname)
                safe_name = _sanitize_filename(fname, max_len=120)
                alt_path = os.path.join(download_dir, safe_name)
                if not os.path.exists(local_path) and os.path.exists(alt_path):
                    local_path = alt_path
                if not os.path.exists(local_path):
                    continue

                lines.append(f"#### {fname}")
                lines.append("")
                text = _extract_text_from_file(local_path)
                if text and text.strip():
                    if len(text) > 50000:
                        text = text[:50000] + "\n\n（内容过长，已截断）"
                    lines.append(text)
                else:
                    lines.append("（无法提取文本内容）")
                lines.append("")

    return "\n".join(lines)


# ---------------------------------------------------------------------------
# State management
# ---------------------------------------------------------------------------

def _load_state(output_dir):
    path = os.path.join(output_dir, _STATE_FILENAME)
    if os.path.exists(path):
        try:
            with open(path, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception as e:
            logging.warning("Failed to load crawler state: %s", e)
    return {"processed_ids": [], "completed": False}


def _save_state(output_dir, state):
    path = os.path.join(output_dir, _STATE_FILENAME)
    os.makedirs(output_dir, exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(state, f, ensure_ascii=False)


# ---------------------------------------------------------------------------
# KB upload
# ---------------------------------------------------------------------------

def _upload_to_kb(md_content, attachment_files, kb_id, tenant_id, folder_name):
    """Upload markdown + attachment files to knowledge base."""
    from api.db.services.knowledgebase_service import KnowledgebaseService
    from api.db.services.file_service import FileService
    from api.db.services.document_service import DocumentService

    ok, kb = KnowledgebaseService.get_by_id(kb_id)
    if not ok:
        raise LookupError(f"Knowledge base {kb_id} not found")

    class _FO:
        def __init__(self, fn, b):
            self.id = get_uuid()
            self.filename = fn
            self.blob = b

        def read(self):
            return self.blob

    fo = _FO(f"{folder_name}.md", md_content.encode("utf-8"))
    errs, pairs = FileService.upload_document(kb, [fo], tenant_id)
    if errs:
        logging.warning("MD upload errors: %s", errs)
    for doc, _ in pairs:
        did = doc["id"]
        try:
            DocumentService.begin2parse(did)
            DocumentService.run(tenant_id, doc, {})
        except Exception as e:
            logging.error("Queue parse for %s: %s", did, e)

    for fp in attachment_files:
        fname = os.path.basename(fp)
        with open(fp, "rb") as f:
            blob = f.read()
        fo2 = _FO(fname, blob)
        errs2, pairs2 = FileService.upload_document(kb, [fo2], tenant_id)
        if errs2:
            logging.warning("Attachment upload errors: %s", errs2)
        for doc, _ in pairs2:
            did = doc["id"]
            try:
                DocumentService.begin2parse(did)
                DocumentService.run(tenant_id, doc, {})
            except Exception as e:
                logging.error("Queue parse for %s: %s", did, e)


# ---------------------------------------------------------------------------
# Argument parsing
# ---------------------------------------------------------------------------

def parse_args():
    p = argparse.ArgumentParser(
        description="gcjyzx.zhangzhou.gov.cn wgtb crawler — 漳州市工程项目交易中心 违规通报"
    )
    p.add_argument("--tenant-id", required=True)
    p.add_argument("--target-url",
                   default="https://gcjyzx.zhangzhou.gov.cn/wgtb/moreinfo.html")
    p.add_argument("--kb-id", required=True)
    p.add_argument("--task-name", required=True)
    p.add_argument("--output-dir", default=None,
                   help="Output root directory (default: PROJECT_ROOT/rag/<task_name>)")
    p.add_argument("--full", action="store_true",
                   help="Ignore saved state and re-crawl all")
    p.add_argument("--max-runtime", type=int, default=_MAX_RUNTIME_DEFAULT,
                   help="Max runtime in seconds (default: 3300)")
    p.add_argument("--max-articles", type=int, default=0,
                   help="Max articles to fetch (0 = unlimited)")
    # Legacy/unused args
    for opt in ("--section", "--max-days", "--hours",
                "--llm-id", "--llm-model", "--access-token"):
        p.add_argument(opt, default=None)
    return p.parse_args()


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    args = parse_args()

    _safe_print("\n" + "=" * 60)
    _safe_print("[WGTB] 漳州市工程项目交易中心 — {} crawler".format(_SECTION_LABEL))
    _safe_print("[WGTB] Target: {}".format(args.target_url))
    _safe_print("[WGTB] KB: {}".format(args.kb_id))
    _safe_print("[WGTB] Task: {}".format(args.task_name))
    _safe_print("[WGTB] Max articles: {}".format(
        args.max_articles if args.max_articles else "unlimited"))
    _safe_print("[WGTB] Max runtime: {}s".format(args.max_runtime))
    _safe_print("=" * 60 + "\n")
    sys.stdout.flush()

    settings.init_settings()
    logging.info("=== WGTB crawler started ===")

    output_dir = args.output_dir or os.path.join(
        _PROJECT_ROOT, "rag", args.task_name.strip()
    )
    os.makedirs(output_dir, exist_ok=True)
    _safe_print("[WGTB] Output: {}\n".format(output_dir))
    sys.stdout.flush()

    # ── State ──────────────────────────────────────────────────────────
    state = _load_state(output_dir) if not args.full else {
        "processed_ids": [], "completed": False,
    }
    processed_ids = set(state.get("processed_ids", []))

    if state.get("completed"):
        _safe_print("[WGTB] Already completed, nothing to do.")
        sys.stdout.flush()
        return

    _safe_print("[WGTB] Already processed: {} article(s)".format(len(processed_ids)))
    sys.stdout.flush()

    crawl_start = time.time()

    # ── Step 1: Crawl all listing pages ────────────────────────────────
    _safe_print("[WGTB] Step 1/3: Crawling listing pages...")
    sys.stdout.flush()

    all_articles = []
    page_num = 1

    while True:
        _safe_print("[WGTB]   Page {}...".format(page_num))
        sys.stdout.flush()

        articles, has_more = _crawl_listing_page(page_num)
        if not articles:
            _safe_print("[WGTB]   No articles on page {}, stopping.".format(page_num))
            sys.stdout.flush()
            break

        all_articles.extend(articles)
        _safe_print("[WGTB]   -> {} articles (total: {})".format(
            len(articles), len(all_articles)))
        sys.stdout.flush()

        if not has_more:
            _safe_print("[WGTB]   Last page reached.")
            sys.stdout.flush()
            break

        if args.max_articles and len(all_articles) >= args.max_articles:
            all_articles = all_articles[:args.max_articles]
            break

        page_num += 1
        time.sleep(random.uniform(_REQUEST_DELAY_MIN, _REQUEST_DELAY_MAX))

    _safe_print("[WGTB] Total articles collected: {}".format(len(all_articles)))
    sys.stdout.flush()

    if not all_articles:
        _safe_print("[WGTB] No articles found, exiting.")
        sys.stdout.flush()
        return

    # Filter already-processed
    new_articles = [
        a for a in all_articles
        if a.get("infoid") and a["infoid"] not in processed_ids
    ]
    skipped = len(all_articles) - len(new_articles)
    if skipped:
        _safe_print("[WGTB] {} already processed, {} new".format(skipped, len(new_articles)))
        sys.stdout.flush()

    if not new_articles:
        _safe_print("[WGTB] All articles already processed.")
        sys.stdout.flush()
        state["completed"] = True
        _save_state(output_dir, state)
        return

    # ── Step 2: Process each article ───────────────────────────────────
    _safe_print("\n[WGTB] Step 2/3: Processing {} article(s)...\n".format(len(new_articles)))
    sys.stdout.flush()

    processed_count = 0
    stopped_early = False
    downloads_dir = os.path.join(output_dir, "downloads")

    for idx, article in enumerate(new_articles, 1):
        # ── Time-bounded check ─────────────────────────────────────────
        elapsed = time.time() - crawl_start
        if elapsed > args.max_runtime - 120:
            _safe_print(
                "\n[WGTB] Runtime {:.0f}s approaching limit ({}s), "
                "stopping gracefully. {} processed. "
                "Next run will resume.".format(elapsed, args.max_runtime, processed_count))
            sys.stdout.flush()
            stopped_early = True
            break

        infoid = article["infoid"]
        title = article["title"]
        detail_url = article["url"]
        date_str = article["date_str"]

        _safe_print("[WGTB] [{}/{}] {}...".format(idx, len(new_articles), title[:50]))
        sys.stdout.flush()

        # Fetch detail page
        html_bytes = _http_get(detail_url, referer=args.target_url)
        if not html_bytes:
            logging.warning("Failed to fetch detail: %s", detail_url)
            processed_ids.add(infoid)
            continue

        detail = _parse_detail_html(html_bytes, detail_url)
        if not detail:
            logging.warning("Failed to parse detail: %s", detail_url)
            processed_ids.add(infoid)
            continue

        # Download attachments
        attachments = detail.get("attachments", [])
        local_files = []
        if attachments:
            article_dl_dir = os.path.join(downloads_dir, _sanitize_filename(infoid, 40))
            local_files = _download_attachments(attachments, article_dl_dir)
            # Extract ZIPs
            for fp in list(local_files):
                is_zip = fp.lower().endswith(".zip")
                if not is_zip and os.path.exists(fp) and os.path.getsize(fp) >= 4:
                    with open(fp, "rb") as f:
                        is_zip = f.read(4) == b"PK\x03\x04"
                if is_zip:
                    extracted = _extract_zip(fp)
                    local_files.remove(fp)
                    local_files.extend(extracted)

        # Build markdown
        md_content = _build_markdown(
            detail,
            os.path.join(downloads_dir, _sanitize_filename(infoid, 40))
            if attachments else "",
            detail_url,
        )

        # Save markdown locally
        folder_name = _sanitize_filename(
            "{}_{}_{}".format(date_str, infoid[:12], title[:40]), max_len=120)
        md_path = os.path.join(output_dir, f"{folder_name}.md")
        with open(md_path, "w", encoding="utf-8") as f:
            f.write(md_content)
        _safe_print("[WGTB]   Saved ({} chars, {} attachments)".format(
            len(md_content), len(local_files)))
        sys.stdout.flush()

        # Upload to KB
        if args.kb_id:
            try:
                _upload_to_kb(md_content, local_files, args.kb_id,
                             args.tenant_id, folder_name)
                _safe_print("[WGTB]   Uploaded to KB {}".format(args.kb_id))
                sys.stdout.flush()
            except Exception as e:
                logging.error("KB upload failed: %s", e)
                _safe_print("[WGTB]   Upload error: {}".format(e))
                sys.stdout.flush()

        processed_ids.add(infoid)
        processed_count += 1

        # Checkpoint every batch
        if processed_count % _BATCH_SIZE == 0:
            _save_state(output_dir, {
                "processed_ids": list(processed_ids),
                "completed": False,
            })
            _safe_print("[WGTB]   Checkpoint ({} processed)".format(processed_count))
            sys.stdout.flush()

        # Anti-crawling delay
        time.sleep(random.uniform(_REQUEST_DELAY_MIN, _REQUEST_DELAY_MAX))

    # ── Final state ────────────────────────────────────────────────────
    if not stopped_early:
        state["completed"] = True
    state["processed_ids"] = list(processed_ids)
    _save_state(output_dir, state)

    _safe_print("\n" + "=" * 60)
    _safe_print("[WGTB] Crawl complete — {} new article(s)".format(processed_count))
    if stopped_early:
        _safe_print("[WGTB] Stopped early, will resume next run")
    _safe_print("=" * 60 + "\n")
    sys.stdout.flush()
    logging.info("=== WGTB crawler finished ===")


if __name__ == "__main__":
    CONSUMER_NAME = "gcjyzx_wgtb_crawler"
    init_root_logger(CONSUMER_NAME)
    main()
