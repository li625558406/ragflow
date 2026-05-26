#!/usr/bin/env python3
#
#  Copyright 2026 The InfiniFlow Authors. All Rights Reserved.
#
#  Licensed under the Apache License, Version 2.0 (the "License");
#  you may not use this file except in compliance with the License.
#  You may obtain a copy of the License at
#
#      http://www.apache.org/licenses/LICENSE-2.0
#
#  Unless required by applicable law or agreed to in writing, software
#  distributed under the License is distributed on an "AS IS" BASIS,
#  WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
#  See the License for the specific language governing permissions and
#  limitations under the License.
#
"""
Dedicated web crawler for ggzyjy.xzfwzx.putian.gov.cn/zcwj/ (莆田市公共资源交易中心 政策文件).

Targets the 4 modules under 政策文件:
  1. 工程交易       /zcwj/014001/subPage.html
  2. 政府采购       /zcwj/014002/subPage.html
  3. 国土资源       /zcwj/014003/subPage.html
  4. 产权交易       /zcwj/014004/subPage.html

Site characteristics
────────────────────
  - Built on Epoint WebBuilder CMS (SSR HTML, no SPA)
  - Listing pages: single-page SSR, no pagination, ~250 total articles
  - Detail pages: metadata + content + attachments (same .ewb-article structure)
  - Attachments: ztbfjyz(url, '0', '1') — NO captcha required! Direct download via
    window.location.replace(url). Downloads use plain GET requests.
  - Anti-crawling: random delays between requests, full headers

Checkpoint/resume: each module is processed independently.
If the 3600s task timeout kills the run mid-way, the next trigger resumes
from the next incomplete module.

Usage:
    python putian_zcwj_crawler.py \\
        --tenant-id <TENANT_ID> \\
        --target-url https://ggzyjy.xzfwzx.putian.gov.cn/zcwj/guidety.html \\
        --kb-id <KB_ID> \\
        --task-name <NAME>
"""

import argparse
import json
import logging
import os
import random
import re
import sys
import time
import zipfile
from datetime import datetime
from io import BytesIO
from urllib.parse import urljoin, unquote

import requests
from bs4 import BeautifulSoup

_SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
_PROJECT_ROOT = os.path.abspath(os.path.join(_SCRIPT_DIR, "..", ".."))
sys.path.insert(0, _PROJECT_ROOT)

from common import settings
from common.log_utils import init_root_logger
from common.misc_utils import get_uuid

# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------
_SITE_ROOT = "https://ggzyjy.xzfwzx.putian.gov.cn"
_SITE_NAME = "莆田市公共资源交易中心"

_USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/120.0.0.0 Safari/537.36"
)

_HEADERS = {
    "User-Agent": _USER_AGENT,
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,image/webp,*/*;q=0.8",
    "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8",
    "Accept-Encoding": "gzip, deflate",
    "Cache-Control": "no-cache",
    "Connection": "keep-alive",
}

_MODULES = [
    {"key": "gcjy", "name": "工程交易",
     "list_url": "/zcwj/014001/subPage.html"},
    {"key": "zfcg", "name": "政府采购",
     "list_url": "/zcwj/014002/subPage.html"},
    {"key": "gtzy", "name": "国土资源",
     "list_url": "/zcwj/014003/subPage.html"},
    {"key": "cqjy", "name": "产权交易",
     "list_url": "/zcwj/014004/subPage.html"},
]

# Anti-crawling: random delays between requests
_REQUEST_DELAY_MIN = 1.0
_REQUEST_DELAY_MAX = 2.5

_STATE_FILENAME = "_crawler_state.json"

# -- Article link regex -------------------------------------------------------
_ARTICLE_LINK_RE = re.compile(
    r"/zcwj/(\d+)/(\d{8})/([0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12})\.html"
)

# -- Attachment onclick regex -------------------------------------------------
_ATTACH_ONCLICK_RE = re.compile(r"ztbfjyz\('([^']+)','(\d)','(\d)'\)")


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _safe_print(msg):
    try:
        print(msg)
    except UnicodeEncodeError:
        print(msg.encode("gbk", errors="replace").decode("gbk"))


def _request_delay():
    time.sleep(random.uniform(_REQUEST_DELAY_MIN, _REQUEST_DELAY_MAX))


def _init_session():
    sess = requests.Session()
    sess.headers.update(_HEADERS)
    try:
        sess.get(_SITE_ROOT + "/", timeout=30)
        time.sleep(1)
    except Exception:
        pass
    return sess


def _fetch_page(sess, url, max_retries=3):
    for attempt in range(1, max_retries + 1):
        try:
            resp = sess.get(url, timeout=60)
            resp.encoding = "utf-8"
            if resp.status_code == 200:
                return resp.text
            if attempt < max_retries:
                time.sleep(2 * attempt)
        except Exception as e:
            logging.warning("Fetch %s failed (attempt %d): %s", url, attempt, e)
            if attempt < max_retries:
                time.sleep(2 * attempt)
    return ""


# ---------------------------------------------------------------------------
# List extraction
# ---------------------------------------------------------------------------

def _extract_list_items(html, list_url):
    soup = BeautifulSoup(html, "html.parser")
    articles = []
    seen = set()

    for a_tag in soup.find_all("a", href=True):
        href = a_tag["href"].strip()
        text = a_tag.get_text(strip=True)
        if not href or not text or len(text) < 4:
            continue
        m = _ARTICLE_LINK_RE.search(href)
        if not m:
            continue
        art_id = m.group(3)
        date_str = m.group(2)
        section = m.group(1)
        if len(date_str) == 8:
            date_str = "{}-{}-{}".format(date_str[:4], date_str[4:6], date_str[6:8])
        if art_id in seen:
            continue
        seen.add(art_id)
        articles.append({
            "id": art_id,
            "title": text,
            "url": urljoin(_SITE_ROOT, href),
            "date_str": date_str,
            "section_code": section,
        })
    return articles


# ---------------------------------------------------------------------------
# Detail extraction
# ---------------------------------------------------------------------------

def _extract_detail(html, detail_url):
    result = {"title": "", "date": "", "content_text": "", "attachments": []}
    soup = BeautifulSoup(html, "html.parser")

    # Title extraction (3-tier, same as fwzx crawler)
    GENERIC_TITLES = ("招标公告，区块链已存证",)

    # 1) <h3> inside .ewb-article
    article_el = soup.select_one(".ewb-article")
    if article_el:
        for h3 in article_el.find_all("h3"):
            text = h3.get_text(strip=True)
            if text and len(text) > 2 and len(text) < 200 \
                    and not any(kw in text for kw in GENERIC_TITLES):
                result["title"] = text
                break

    # 2) <h1>, .bt, etc.
    if not result["title"]:
        for sel in ("h1", ".detail-title", ".bt", "[class*='title']"):
            el = soup.select_one(sel)
            if el:
                text = el.get_text(strip=True)
                if text and len(text) > 2 and len(text) < 200 \
                        and not any(kw in text for kw in GENERIC_TITLES):
                    result["title"] = text
                    break

    # 3) Breadcrumb
    if not result["title"]:
        breadcrumb = soup.select_one(
            ".ewb-location, .location, [class*='location'], [class*='breadcrumb'], .chain-list")
        if breadcrumb:
            items = breadcrumb.get_text(strip=True).split(">")
            for item in reversed(items):
                item = item.strip()
                if item and len(item) > 3 and "首页" not in item \
                        and "政策文件" not in item:
                    result["title"] = item
                    break

    # Date extraction
    body_text = soup.get_text()
    date_m = re.search(r"发布时间[：:]\s*(\d{4}[-/]\d{1,2}[-/]\d{1,2})", body_text)
    if date_m:
        result["date"] = date_m.group(1)
    else:
        date_m = re.search(r"(\d{4}-\d{2}-\d{2})", body_text)
        if date_m:
            result["date"] = date_m.group(1)

    # Content extraction
    content_parts = []

    # Try .TRS_Editor for article body
    trs_el = soup.select_one(".TRS_Editor")
    if trs_el:
        text = trs_el.get_text(separator="\n", strip=True)
        if len(text) > 30:
            content_parts.append(text)

    # Also get .ewb-article text (includes title bar, date, etc.)
    if article_el:
        article_content = article_el.get_text(separator="\n", strip=True)
        if article_content and len(article_content) > 30:
            content_parts.append(article_content)

    # Fallback: general content selectors
    if not content_parts:
        for sel in (".article-content", ".detail-content", ".content", "article"):
            el = soup.select_one(sel)
            if el:
                text = el.get_text(separator="\n", strip=True)
                if len(text) > 50:
                    content_parts.append(text)
                    break

    # Last resort: body text with footer cutoff
    if not content_parts:
        for tag in soup.find_all(["nav", "header", "footer", "script", "style"]):
            tag.decompose()
        fallback = soup.get_text(separator="\n", strip=True)
        cutoff_markers = ["国家部委", "联系地址：", "主办：", "备案序号：",
                          "网站标识码", "附件下载："]
        for marker in cutoff_markers:
            idx = fallback.find(marker)
            if idx > 200:
                fallback = fallback[:idx]
                break
        if len(fallback) > 50:
            content_parts.append(fallback)

    result["content_text"] = "\n\n".join(content_parts)

    # Extract attachments from onclick handlers
    for a_tag in soup.find_all("a", onclick=True):
        onclick = a_tag["onclick"]
        name = a_tag.get("title") or a_tag.get_text(strip=True)
        m = _ATTACH_ONCLICK_RE.search(onclick)
        if m:
            result["attachments"].append({
                "name": name,
                "download_path": m.group(1),
                "ztb_attach": m.group(2),
                "use_ztb_yzm": m.group(3),
            })

    return result


# ---------------------------------------------------------------------------
# Attachment download (direct GET — NO captcha for zcwj attachments)
# ---------------------------------------------------------------------------

def _download_attachment(sess, att, dest_dir):
    """Download a single attachment via direct GET request.
    zcwj attachments use ztbfjyz(url, '0', '1') — no captcha needed.
    The JS does window.location.replace(url) which is a simple GET.
    """
    path = att.get("download_path", "")
    name = att.get("name", "unknown")
    download_url = urljoin(_SITE_ROOT, path)

    os.makedirs(dest_dir, exist_ok=True)
    safe_name = re.sub(r'[\\/:*?"<>|]', "_", name)
    dest_path = os.path.join(dest_dir, safe_name)

    if os.path.exists(dest_path) and os.path.getsize(dest_path) > 100:
        _safe_print("      (cached) {}".format(safe_name))
        return dest_path

    try:
        _request_delay()
        resp = sess.get(download_url, timeout=120, stream=True)
        if resp.status_code == 200 and len(resp.content) > 100:
            with open(dest_path, "wb") as f:
                f.write(resp.content)
            _safe_print("      Downloaded: {} ({} bytes)".format(
                safe_name, len(resp.content)))
            return dest_path
        else:
            _safe_print("      Download failed: {} (status={}, size={})".format(
                safe_name, resp.status_code, len(resp.content)))
            # Check if it's an HTML error page
            if b"<!DOCTYPE" in resp.content[:200] or b"<html" in resp.content[:200]:
                logging.warning("Download returned HTML for %s", download_url)
    except Exception as e:
        logging.warning("Download error for %s: %s", name, e)

    return None


# ---------------------------------------------------------------------------
# File extraction
# ---------------------------------------------------------------------------

def _extract_zip(zip_path):
    extracted = []
    dest_dir = os.path.dirname(zip_path)
    try:
        with zipfile.ZipFile(zip_path, "r") as zf:
            for name in zf.namelist():
                if name.startswith("__MACOSX") or name.startswith("."):
                    continue
                safe_name = re.sub(r'[\\/:*?"<>|]', "_", os.path.basename(name))
                if not safe_name:
                    continue
                dest_path = os.path.join(dest_dir, safe_name)
                if os.path.exists(dest_path):
                    continue
                with open(dest_path, "wb") as f:
                    f.write(zf.read(name))
                extracted.append(dest_path)
                _safe_print("      Extracted: {}".format(safe_name))
        os.remove(zip_path)
    except Exception as e:
        logging.warning("ZIP extract error for %s: %s",
                        os.path.basename(zip_path), e)
    return extracted


def _extract_file_text(filepath):
    ext = os.path.splitext(filepath)[1].lower()
    try:
        if ext == ".txt":
            with open(filepath, "r", encoding="utf-8", errors="replace") as f:
                return f.read()
        elif ext == ".pdf":
            try:
                import pdfplumber
                parts = []
                with pdfplumber.open(filepath) as pdf:
                    for page in pdf.pages:
                        text = page.extract_text()
                        if text:
                            parts.append(text)
                return "\n\n".join(parts)
            except ImportError:
                pass
            try:
                import fitz
                doc = fitz.open(filepath)
                parts = []
                for page in doc:
                    text = page.get_text()
                    if text:
                        parts.append(text)
                doc.close()
                return "\n\n".join(parts)
            except ImportError:
                return "(PDF file, no parser available)"
        elif ext in (".doc", ".docx"):
            # Try python-docx first (only works for .docx)
            if ext == ".docx":
                try:
                    import docx
                    doc = docx.Document(filepath)
                    return "\n".join(
                        p.text for p in doc.paragraphs if p.text.strip())
                except ImportError:
                    pass
            # Try antiword or textract for .doc
            try:
                import subprocess
                result = subprocess.run(
                    ["antiword", filepath],
                    capture_output=True, text=True, timeout=30)
                if result.returncode == 0 and result.stdout.strip():
                    return result.stdout
            except Exception:
                pass
            return "(DOC file, text extraction limited)"
        elif ext == ".xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(filepath, read_only=True)
            parts = []
            for ws in wb.worksheets:
                rows = []
                for row in ws.iter_rows(values_only=True):
                    rows.append(
                        " | ".join(str(c) if c is not None else "" for c in row))
                if rows:
                    parts.append("### {}\n".format(ws.title) + "\n".join(rows))
            wb.close()
            return "\n\n".join(parts)
    except Exception as e:
        logging.warning("Failed to extract text from %s: %s", filepath, e)
    return None


# ---------------------------------------------------------------------------
# Markdown building
# ---------------------------------------------------------------------------

def _build_markdown(art, detail, attachment_texts):
    title = detail.get("title") or art.get("title", "无标题")
    lines = [
        "# {}".format(title),
        "",
        "**栏目:** {}".format(art.get("section", "")),
    ]
    date_str = detail.get("date") or art.get("date_str", "")
    if date_str:
        lines.append("**日期:** {}".format(date_str))
    lines.append("**URL:** {}".format(art.get("url", "")))
    lines.append("")

    content = detail.get("content_text", "")
    if content:
        if len(content) > 50000:
            content = content[:50000] + "\n\n（内容过长，已截断）"
        lines.append("## 详细内容")
        lines.append("")
        lines.append(content)
        lines.append("")

    if attachment_texts:
        lines.append("## 附件内容")
        lines.append("")
        for fname, ftext in attachment_texts:
            lines.append("### {}".format(fname))
            lines.append("")
            if ftext:
                if len(ftext) > 50000:
                    ftext = ftext[:50000] + "\n\n（附件内容过长，已截断）"
                lines.append(ftext)
            else:
                lines.append("（无法提取文本内容）")
            lines.append("")

    return "\n".join(lines)


# ---------------------------------------------------------------------------
# State persistence
# ---------------------------------------------------------------------------

def _load_state(output_dir):
    path = os.path.join(output_dir, _STATE_FILENAME)
    if os.path.exists(path):
        try:
            with open(path, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception as e:
            logging.warning("Failed to load crawler state: %s", e)
    return {"processed_ids": [], "completed_sections": []}


def _save_state(output_dir, state):
    path = os.path.join(output_dir, _STATE_FILENAME)
    os.makedirs(output_dir, exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(state, f, ensure_ascii=False)
    logging.info("State saved (%d IDs, %d sections done)",
                 len(state.get("processed_ids", [])),
                 len(state.get("completed_sections", [])))


# ---------------------------------------------------------------------------
# KB upload
# ---------------------------------------------------------------------------

def _upload_to_kb(filepath, kb_id, tenant_id, parser_id="general"):
    from api.db.services.knowledgebase_service import KnowledgebaseService
    from api.db.services.file_service import FileService
    from api.db.services.document_service import DocumentService

    ok, kb = KnowledgebaseService.get_by_id(kb_id)
    if not ok:
        raise LookupError("Knowledge base {} not found".format(kb_id))

    with open(filepath, "rb") as f:
        blob = f.read()

    class _FileObj:
        def __init__(self, fn, b):
            self.id = get_uuid()
            self.filename = fn
            self.blob = b

        def read(self):
            return self.blob

    fo = _FileObj(os.path.basename(filepath), blob)
    errs, pairs = FileService.upload_document(kb, [fo], tenant_id)
    if errs:
        logging.warning("Upload errors: %s", errs)
    for doc, _ in pairs:
        did = doc["id"]
        try:
            DocumentService.update_by_id(did, {"parser_id": parser_id})
        except Exception:
            pass
        try:
            from api.db.services.task_service import queue_tasks
            from api.db.services.file2document_service import File2DocumentService
            bucket, name = File2DocumentService.get_storage_address(doc_id=did)
            queue_tasks(doc, bucket, name, 0)
        except Exception as e:
            logging.error("Failed to queue parsing for %s: %s", did, e)


# ---------------------------------------------------------------------------
# Module-level processing
# ---------------------------------------------------------------------------

def _process_module(sess, output_dir, kb_id, tenant_id,
                    module_info, processed_ids, state):
    module_key = module_info["key"]
    module_name = module_info["name"]
    list_url = urljoin(_SITE_ROOT, module_info["list_url"])

    if module_key in state.get("completed_sections", []):
        _safe_print("[{}]   Already completed, skipping.".format(module_name))
        sys.stdout.flush()
        return 0

    _safe_print("[{}]   Fetching listing: {}".format(module_name, list_url))
    sys.stdout.flush()

    list_html = _fetch_page(sess, list_url)
    if not list_html:
        _safe_print("[{}]   ERROR: Failed to fetch listing page".format(module_name))
        sys.stdout.flush()
        return 0

    articles = _extract_list_items(list_html, list_url)
    _safe_print("[{}]   Found {} total articles".format(
        module_name, len(articles)))
    sys.stdout.flush()

    new_articles = [a for a in articles if a["id"] not in processed_ids]
    _safe_print("[{}]   {} new, {} already processed".format(
        module_name, len(new_articles), len(articles) - len(new_articles)))
    sys.stdout.flush()

    if not new_articles:
        if module_key not in state.get("completed_sections", []):
            state.setdefault("completed_sections", []).append(module_key)
            _save_state(output_dir, state)
        return 0

    for a in new_articles:
        a["section"] = module_name

    BATCH_SIZE = 10
    total_processed = 0
    fail_count = 0
    downloads_dir = os.path.join(output_dir, "downloads")

    for batch_start in range(0, len(new_articles), BATCH_SIZE):
        batch = new_articles[batch_start:batch_start + BATCH_SIZE]
        batch_num = batch_start // BATCH_SIZE + 1
        md_parts = []
        batch_ids = []

        for idx, art in enumerate(batch, 1):
            global_idx = batch_start + idx
            _safe_print("[{}]   [{}/{}] {}".format(
                module_name, global_idx, len(new_articles),
                art["title"][:60]))
            sys.stdout.flush()

            detail_html = _fetch_page(sess, art["url"])
            if not detail_html:
                fail_count += 1
                detail = {
                    "title": art["title"],
                    "date": art.get("date_str", ""),
                    "content_text": "标题: {}\n日期: {}\nURL: {}".format(
                        art["title"], art.get("date_str", ""), art["url"]),
                    "attachments": [],
                }
            else:
                detail = _extract_detail(detail_html, art["url"])

            # Download attachments (direct GET, no captcha)
            attachment_texts = []
            for att in detail.get("attachments", []):
                att_name = att.get("name", "unknown")
                dest_dir = os.path.join(downloads_dir, module_key, art["id"])
                fp = _download_attachment(sess, att, dest_dir)
                if fp:
                    is_zip = (
                            fp.lower().endswith((".zip", ".rar")) or
                            (os.path.getsize(fp) >= 4 and
                             open(fp, "rb").read(4) == b"PK\x03\x04")
                    )
                    if is_zip:
                        extracted = _extract_zip(fp)
                        for ext_fp in extracted:
                            text = _extract_file_text(ext_fp)
                            attachment_texts.append(
                                (os.path.basename(ext_fp), text))
                    else:
                        text = _extract_file_text(fp)
                        attachment_texts.append((att_name, text))

            md = _build_markdown(art, detail, attachment_texts)

            articles_dir = os.path.join(output_dir, "articles", module_key)
            os.makedirs(articles_dir, exist_ok=True)
            md_path = os.path.join(articles_dir,
                                   "{}.md".format(art["id"]))
            with open(md_path, "w", encoding="utf-8") as f:
                f.write(md)

            md_parts.append(md)
            batch_ids.append(art["id"])

            _request_delay()

        # Checkpoint
        if md_parts:
            batch_path = os.path.join(
                output_dir, "{}_{:03d}.md".format(module_key, batch_num))
            with open(batch_path, "w", encoding="utf-8") as f:
                f.write("\n\n---\n\n".join(md_parts))

            processed_ids.update(batch_ids)
            state["processed_ids"] = list(processed_ids)
            _save_state(output_dir, state)

            if kb_id:
                try:
                    _upload_to_kb(batch_path, kb_id, tenant_id)
                except Exception as e:
                    _safe_print("[{}]   batch {} upload failed: {}".format(
                        module_name, batch_num, e))
                    logging.error("Upload failed for %s batch %d: %s",
                                  module_name, batch_num, e)

            total_processed += len(md_parts)
            _safe_print("[{}]   batch {} uploaded ({}/{} done)".format(
                module_name, batch_num, total_processed,
                len(new_articles)))
            sys.stdout.flush()

    if module_key not in state.get("completed_sections", []):
        state.setdefault("completed_sections", []).append(module_key)
        _save_state(output_dir, state)

    _safe_print("[{}]   Done: {} processed, {} failed\n".format(
        module_name, total_processed, fail_count))
    sys.stdout.flush()
    return total_processed


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def parse_args():
    p = argparse.ArgumentParser(
        description="putian.zcwj crawler - 莆田市公共资源交易中心 政策文件"
    )
    p.add_argument("--tenant-id", required=True)
    p.add_argument("--target-url",
                   default="https://ggzyjy.xzfwzx.putian.gov.cn/zcwj/guidety.html")
    p.add_argument("--kb-id", required=True)
    p.add_argument("--task-name", required=True)
    p.add_argument("--output-dir", default=None)
    p.add_argument("--full", action="store_true")
    p.add_argument("--section", default=None,
                   help="Comma-separated: gcjy,zfcg,gtzy,cqjy")
    p.add_argument("--max-runtime", type=int, default=3300,
                   help="Max runtime (default: 3300s=55min)")
    for opt in ("--llm-id", "--llm-model", "--access-token"):
        p.add_argument(opt, default=None)
    return p.parse_args()


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    args = parse_args()

    _safe_print("\n" + "=" * 60)
    _safe_print("[PUTIAN-ZCWJ] 莆田市公共资源交易中心 政策文件 crawler")
    _safe_print("[PUTIAN-ZCWJ] KB: {}".format(args.kb_id))
    _safe_print("[PUTIAN-ZCWJ] Max runtime: {}s".format(args.max_runtime))
    _safe_print("=" * 60 + "\n")
    sys.stdout.flush()

    settings.init_settings()
    logging.info("=== Putian ZCWJ crawler started ===")

    output_dir = args.output_dir or os.path.join(
        _PROJECT_ROOT, "rag", args.task_name.strip()
    )
    os.makedirs(output_dir, exist_ok=True)
    _safe_print("[PUTIAN-ZCWJ] Output: {}\n".format(output_dir))
    sys.stdout.flush()

    state = _load_state(output_dir) if not args.full else {
        "processed_ids": [], "completed_sections": []
    }
    processed_ids = set(state.get("processed_ids", []))
    completed_sections = set(state.get("completed_sections", []))
    _safe_print("[PUTIAN-ZCWJ] Previously processed: {}, "
                "completed modules: {}\n".format(
                    len(processed_ids), len(completed_sections)))
    sys.stdout.flush()

    if args.section:
        selected = set(args.section.split(","))
        active_modules = [m for m in _MODULES if m["key"] in selected]
    else:
        active_modules = list(_MODULES)

    _safe_print("[PUTIAN-ZCWJ] Modules: {}".format(
        ", ".join(m["name"] for m in active_modules)))
    sys.stdout.flush()

    # Init requests session
    sess = _init_session()
    _safe_print("[PUTIAN-ZCWJ] Session initialized\n")
    sys.stdout.flush()

    try:
        run_start = time.time()
        total_processed = 0
        stopped_early = False

        for mod_info in active_modules:
            elapsed = time.time() - run_start
            remaining = args.max_runtime - elapsed
            grace = min(120, args.max_runtime * 0.05)
            if remaining < grace:
                _safe_print(
                    "\n[PUTIAN-ZCWJ] Runtime {:.0f}s, remaining {:.0f}s < "
                    "grace {:.0f}s, stopping early.".format(
                        elapsed, remaining, grace))
                sys.stdout.flush()
                stopped_early = True
                break

            n = _process_module(sess, output_dir, args.kb_id,
                                args.tenant_id, mod_info, processed_ids, state)
            total_processed += n

        _safe_print("\n" + "=" * 60)
        if stopped_early:
            _safe_print("[PUTIAN-ZCWJ] Partial run: {} articles.".format(
                total_processed))
        else:
            _safe_print("[PUTIAN-ZCWJ] Done: {} articles processed.".format(
                total_processed))
        _safe_print("=" * 60 + "\n")
        sys.stdout.flush()
        logging.info("=== Putian ZCWJ crawler finished: %d articles ===",
                     total_processed)

    finally:
        sess.close()


if __name__ == "__main__":
    CONSUMER_NAME = "putian_zcwj_crawler"
    init_root_logger(CONSUMER_NAME)
    main()
