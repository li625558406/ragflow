#!/usr/bin/env python3
#
#  Copyright 2026 The InfiniFlow Authors. All Rights Reserved.
#
#  Licensed under the Apache License, Version 2.0 (the "License");
#  you may not use this file except in compliance with the License.
#  You may obtain a copy of the License at
#
#      http://www.apache.org/licenses/LICENSE-2.0
#
#  Unless required by applicable law or agreed to in writing, software
#  distributed under the License is distributed on an "AS IS" BASIS,
#  WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
#  See the License for the specific language governing permissions and
#  limitations under the License.
#
"""
Crawler for ggzyjy.quanzhou.gov.cn — 市级政策法规 (Municipal Policy / Zcfg).

Target: https://ggzyjy.quanzhou.gov.cn/articleList/getZcfgNew.do?centerId=-1

Site characteristics
────────────────────
Traditional Java/JSP application with jQuery AJAX.  Both listing and detail
data are fetched via POST requests to ``*.do`` endpoints with a pseudo-JSON
body format (keys are unquoted, e.g. ``{pageIndex:1,pageSize:10,major_id:0}``).

Listing API
   POST /articleList/getZcfgList.do → JSON
   ``{result, data: {totalPage, totalRecord, dataList: [{newsId, title, ...}]}}``

Detail API
   POST /articleList/getZcfgDetail.do → JSON
   ``{result, data: {title, article_no, agency, pubDate, typename, content (HTML)}}``

   The ``content`` field is HTML and may contain embedded attachment links
   (``<a class="ke-insertfile" href="...">``).  Attachments are downloaded,
   text-extracted, and included in the output markdown.

Usage
─────
    python quanzhou_zcfg_crawler.py \\
        --tenant-id <TENANT_ID> \\
        --kb-id <KB_ID> \\
        --task-name <NAME>

    # Optional:
        --max-runtime 3300    # Max runtime before graceful stop (default: 3300)
        --full                # Ignore saved state, re-crawl all
"""
import argparse
import html as html_mod
import json
import logging
import os
import random
import re
import ssl
import sys
import time
import urllib.parse
import urllib.request
import zipfile
from datetime import datetime

_SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
_PROJECT_ROOT = os.path.abspath(os.path.join(_SCRIPT_DIR, "..", ".."))
sys.path.insert(0, _PROJECT_ROOT)

from common import settings
from common.log_utils import init_root_logger
from common.misc_utils import get_uuid

# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------

_SITE_ROOT = "https://ggzyjy.quanzhou.gov.cn"
_TAG_PREFIX = "[QZ-ZCFG]"

_PAGE_SIZE = 10
_BATCH_SIZE = 3
_MAX_RUNTIME_DEFAULT = 3300
_REQUEST_DELAY_MIN = 0.3
_REQUEST_DELAY_MAX = 1.0
_STATE_FILENAME = "_crawler_state.json"

_USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/120.0.0.0 Safari/537.36"
)

_HEADERS_JSON = {
    "User-Agent": _USER_AGENT,
    "Accept": "application/json, text/javascript, */*; q=0.01",
    "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8",
    "Content-Type": "application/json;charset=UTF-8",
    "X-Requested-With": "XMLHttpRequest",
}

_HEADERS_HTML = {
    "User-Agent": _USER_AGENT,
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
    "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8",
}

_HEADERS_DOWNLOAD = {
    "User-Agent": _USER_AGENT,
    "Accept": "*/*",
    "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8",
}

_SSL_CTX = ssl.create_default_context()
_SSL_CTX.check_hostname = False
_SSL_CTX.verify_mode = ssl.CERT_NONE

# API endpoints
_LISTING_URL = _SITE_ROOT + "/articleList/getZcfgList.do"
_DETAIL_URL = _SITE_ROOT + "/articleList/getZcfgDetail.do"
_LISTING_PAGE_URL = _SITE_ROOT + "/articleList/getZcfgNew.do?centerId=-1"

# Listing API params: major_id=0 for municipal, type=1
_LISTING_API_PARAMS = "major_id:0,type:1,keyword:\"\""


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _safe_print(msg):
    try:
        print(msg)
    except UnicodeEncodeError:
        print(msg.encode("gbk", errors="replace").decode("gbk"))


def _request_delay():
    time.sleep(random.uniform(_REQUEST_DELAY_MIN, _REQUEST_DELAY_MAX))


def _sanitize_filename(name, max_len=120):
    if not name:
        return "unnamed"
    safe = re.sub(r'[\\/:*?"<>|\r\n\t]+', "_", name)
    safe = re.sub(r"_+", "_", safe)
    safe = safe.strip("._ ")
    if len(safe) > max_len:
        base, ext = os.path.splitext(safe)
        safe = base[:max_len - len(ext)] + (ext or "")
    return safe or "unnamed"


def _normalize_date(date_str):
    if not date_str:
        return ""
    m = re.search(r'(\d{4}-\d{1,2}-\d{1,2})', str(date_str))
    return m.group(1) if m else str(date_str)[:10]


# ---------------------------------------------------------------------------
# HTTP helpers
# ---------------------------------------------------------------------------

def _http_post(url, body_str, referer=None):
    hdrs = dict(_HEADERS_JSON)
    if referer:
        hdrs["Referer"] = referer
    data_bytes = body_str.encode("utf-8")
    req = urllib.request.Request(url, data=data_bytes, headers=hdrs)
    try:
        with urllib.request.urlopen(req, timeout=30, context=_SSL_CTX) as resp:
            return resp.read().decode("utf-8", errors="replace")
    except Exception as e:
        logging.warning("POST %s failed: %s", url, e)
        return None


def _http_download(url, referer=None):
    hdrs = dict(_HEADERS_DOWNLOAD)
    if referer:
        hdrs["Referer"] = referer
    req = urllib.request.Request(url, headers=hdrs)
    try:
        with urllib.request.urlopen(req, timeout=60, context=_SSL_CTX) as resp:
            return resp.read()
    except Exception as e:
        logging.warning("Download %s failed: %s", url, e)
        return None


# ---------------------------------------------------------------------------
# Listing API
# ---------------------------------------------------------------------------

def _crawl_listing():
    """Crawl all listing pages. Returns list[dict] of raw article items."""
    all_items = []
    page_index = 1

    while True:
        body = "{" + "pageIndex:{},pageSize:{},{}".format(
            page_index, _PAGE_SIZE, _LISTING_API_PARAMS) + "}"

        text = _http_post(_LISTING_URL, body, referer=_LISTING_PAGE_URL)
        if not text:
            break

        try:
            data = json.loads(text)
        except json.JSONDecodeError:
            logging.warning("JSON decode error for listing page %d", page_index)
            break

        if not data.get("result"):
            logging.warning("Listing API error page %d: %s",
                          page_index, data.get("error", ""))
            break

        d = data.get("data", {})
        items = d.get("dataList", [])
        if not items:
            break

        all_items.extend(items)
        total_page = d.get("totalPage", 0)
        _safe_print("[{}]   Page {}/{} — {} items (total: {})".format(
            _TAG_PREFIX, page_index, total_page, len(items), len(all_items)))
        sys.stdout.flush()

        if page_index >= total_page:
            break

        page_index += 1
        _request_delay()

    return all_items


# ---------------------------------------------------------------------------
# Detail API
# ---------------------------------------------------------------------------

def _fetch_detail(news_id):
    """Fetch detail for one article via AJAX API.

    Returns dict with keys: title, article_no, agency, pub_date, typename,
    content_html, content_text, attachments.
    """
    referer = _SITE_ROOT + "/articleList/ZcfgDetail.do?newsId=" + str(news_id)
    body = '{newsId:"%s"}' % news_id

    text = _http_post(_DETAIL_URL, body, referer=referer)
    if not text:
        return None

    try:
        data = json.loads(text)
    except json.JSONDecodeError:
        logging.warning("JSON decode error for detail newsId=%s", news_id)
        return None

    if not data.get("result"):
        logging.warning("Detail API error for newsId=%s: %s",
                      news_id, data.get("error", ""))
        return None

    d = data.get("data", {})
    if not isinstance(d, dict):
        return None

    content_html = d.get("content", "") or ""

    # Extract text from HTML content
    content_text = _html_to_text(content_html)

    # Extract attachment links from HTML content
    attachments = _extract_attachments_from_html(content_html)

    return {
        "title": d.get("title", ""),
        "article_no": d.get("article_no", ""),
        "agency": d.get("agency", ""),
        "pub_date": _normalize_date(d.get("pubDate", "")),
        "typename": d.get("typename", ""),
        "source": d.get("source", ""),
        "click_count": d.get("clickCount", ""),
        "content_html": content_html,
        "content_text": content_text,
        "attachments": attachments,
    }


# ---------------------------------------------------------------------------
# HTML / attachment parsing
# ---------------------------------------------------------------------------

def _html_to_text(html_str):
    """Strip HTML tags and decode entities, return clean text."""
    if not html_str:
        return ""

    # Check for image-only content
    has_images = bool(re.search(r'<img[^>]+>', html_str, re.I))

    # Remove script/style tags
    text = re.sub(r'<(script|style)[^>]*>.*?</\1>', '', html_str, flags=re.DOTALL | re.I)

    # Replace common block tags with newlines
    text = re.sub(r'</?(?:div|p|tr|li|br|h[1-6]|table|hr)[^>]*>', '\n', text, flags=re.I)

    # Remove all remaining tags
    text = re.sub(r'<[^>]+>', '', text)

    # Decode HTML entities
    text = html_mod.unescape(text)

    # Collapse whitespace
    text = re.sub(r'[ \t]+', ' ', text)
    text = re.sub(r'\n{3,}', '\n\n', text)
    text = text.strip()

    # If only images with no text, add a note
    if not text and has_images:
        text = "（本文为扫描件图片，正文内容请查看附件图片）"

    return text


def _extract_attachments_from_html(html_str):
    """Extract attachment links and images from HTML content.

    Handles:
      - <a> tags linking to files (PDF, DOC, XLS, ZIP, etc.)
      - <img> tags (scanned documents embedded as images)

    Returns list[dict]: [{filename, url}]
    """
    if not html_str:
        return []

    attachments = []
    seen_urls = set()

    # 1. Match <a> tags with href attributes linking to files
    for m in re.finditer(
        r'<a[^>]*href=["\']([^"\']*(?:\.pdf|\.doc|\.docx|\.xls|\.xlsx|'
        r'\.rar|\.zip|\.ppt|\.pptx|\.txt|\.jpg|\.png|\.gif|'
        r'filedown|download|UploadFile|attach|relatePath|'
        r'file/\w+)'
        r'[^"\']*)["\'][^>]*>([^<]*)</a>',
        html_str, re.I
    ):
        href = m.group(1).strip()
        link_text = m.group(2).strip()
        link_text = html_mod.unescape(link_text)

        if not href:
            continue

        url = _normalize_url(href)
        if not url or url in seen_urls:
            continue
        seen_urls.add(url)

        fname = link_text if link_text else os.path.basename(
            urllib.parse.urlparse(url).path.split("?")[0]
        )
        if not fname:
            fname = "attachment"

        attachments.append({"filename": fname, "url": url})

    # 2. Match <img> tags (scanned document images)
    for m in re.finditer(
        r'<img[^>]*src=["\']([^"\']+)["\'][^>]*>',
        html_str, re.I
    ):
        src = m.group(1).strip()
        if not src:
            continue

        url = _normalize_url(src)
        if not url or url in seen_urls:
            continue
        seen_urls.add(url)

        # Determine filename from URL path
        parsed = urllib.parse.urlparse(url)
        fname = os.path.basename(parsed.path.split("?")[0])
        if not fname:
            fname = "image.jpg"

        attachments.append({"filename": fname, "url": url})

    return attachments


def _normalize_url(url):
    """Ensure URL is absolute."""
    if not url:
        return ""
    if url.startswith("http"):
        return url
    if url.startswith("//"):
        return "https:" + url
    if url.startswith("/"):
        return _SITE_ROOT + url
    return ""


# ---------------------------------------------------------------------------
# Attachment download and text extraction
# ---------------------------------------------------------------------------

def _download_attachments(attachments, download_dir):
    os.makedirs(download_dir, exist_ok=True)
    local_files = []

    for att in attachments:
        url = att.get("url", "")
        if not url:
            continue

        fname = _sanitize_filename(att.get("filename", "attachment"), max_len=120)
        ext = os.path.splitext(urllib.parse.urlparse(url).path.split("?")[0])[1].lower()
        if ext and not fname.lower().endswith(ext):
            fname += ext
        if not ext and "." not in fname:
            fname += ".pdf"

        filepath = os.path.join(download_dir, fname)
        if os.path.exists(filepath) and os.path.getsize(filepath) > 0:
            local_files.append(filepath)
            continue

        data = _http_download(url)
        if data:
            with open(filepath, "wb") as f:
                f.write(data)
            local_files.append(filepath)
            _request_delay()

    return local_files


def _extract_zip(filepath):
    extracted = []
    extract_dir = os.path.splitext(filepath)[0] + "_extracted"
    os.makedirs(extract_dir, exist_ok=True)
    try:
        with zipfile.ZipFile(filepath, "r") as zf:
            for name in zf.namelist():
                safe_name = _sanitize_filename(name, max_len=120)
                out_path = os.path.join(extract_dir, safe_name)
                parent_dir = os.path.dirname(out_path)
                if parent_dir:
                    os.makedirs(parent_dir, exist_ok=True)
                with zf.open(name) as src, open(out_path, "wb") as dst:
                    dst.write(src.read())
                extracted.append(out_path)
    except Exception as e:
        logging.warning("ZIP extraction failed for %s: %s", filepath, e)
    return extracted


def _extract_text_from_file(filepath):
    ext = os.path.splitext(filepath)[1].lower()
    try:
        if ext == ".txt":
            with open(filepath, "r", encoding="utf-8", errors="replace") as f:
                return f.read()
        elif ext == ".pdf":
            import pdfplumber
            parts = []
            with pdfplumber.open(filepath) as pdf:
                for pg in pdf.pages:
                    text = pg.extract_text()
                    if text:
                        parts.append(text)
            return "\n\n".join(parts)
        elif ext in (".docx", ".doc"):
            import docx
            doc = docx.Document(filepath)
            return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        elif ext in (".xlsx", ".xls"):
            import openpyxl
            wb = openpyxl.load_workbook(filepath, read_only=True)
            parts = []
            for ws in wb.worksheets:
                rows = []
                for row in ws.iter_rows(values_only=True):
                    rows.append(" | ".join(str(c) if c is not None else "" for c in row))
                if rows:
                    parts.append("### {}\n".format(ws.title) + "\n".join(rows))
            return "\n\n".join(parts)
    except Exception as e:
        logging.warning("Text extraction failed for %s: %s", filepath, e)
    return ""


# ---------------------------------------------------------------------------
# Markdown builder
# ---------------------------------------------------------------------------

def _build_markdown(detail, detail_url, download_dir):
    title = detail.get("title", "") or "无标题"
    pub_date = detail.get("pub_date", "")
    article_no = detail.get("article_no", "")
    agency = detail.get("agency", "")
    typename = detail.get("typename", "")
    source = detail.get("source", "")
    click_count = detail.get("click_count", "")
    content_text = detail.get("content_text", "")
    attachments = detail.get("attachments", [])

    lines = [
        "# {}".format(title),
        "",
        "**数据来源:** 泉州市公共资源交易信息网 — 市级政策法规",
        "**页面地址:** {}".format(detail_url),
        "**抓取时间:** {}".format(datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
    ]
    if pub_date:
        lines.append("**发布时间:** {}".format(pub_date))
    if article_no:
        lines.append("**文号:** {}".format(article_no))
    if agency:
        lines.append("**发布机构:** {}".format(agency))
    if typename:
        lines.append("**分类:** {}".format(typename))
    if source:
        lines.append("**来源:** {}".format(source))
    if click_count:
        lines.append("**浏览量:** {}".format(click_count))
    lines.append("")

    if content_text:
        lines.append("---")
        lines.append("")
        lines.append("## 正文")
        lines.append("")
        content_clean = re.sub(r"\n{3,}", "\n\n", content_text)
        if len(content_clean) > 100000:
            content_clean = content_clean[:100000] + "\n\n（内容过长，已截断）"
        lines.append(content_clean)
        lines.append("")

    if attachments:
        lines.append("---")
        lines.append("")
        lines.append("## 附件")
        lines.append("")
        for att in attachments:
            fname = att.get("filename", "unknown")
            att_url = att.get("url", "")
            lines.append("- [{}]({})".format(fname, att_url))
        lines.append("")

    return "\n".join(lines), title


def _build_attachment_appendix(attachments, download_dir):
    """Build appendix markdown for attachment content."""
    if not attachments or not download_dir or not os.path.isdir(download_dir):
        return ""

    lines = ["### 附件内容", ""]
    for att in attachments:
        fname = att.get("filename", "")
        local_path = os.path.join(download_dir, fname)
        safe_name = _sanitize_filename(fname, max_len=120)
        if not os.path.exists(local_path):
            alt_path = os.path.join(download_dir, safe_name)
            if os.path.exists(alt_path):
                local_path = alt_path
        if not os.path.exists(local_path):
            for root, _, files in os.walk(download_dir):
                for fn in files:
                    if fn == safe_name or fn == fname:
                        local_path = os.path.join(root, fn)
                        break
        if not os.path.exists(local_path):
            continue

        lines.append("#### {}".format(fname))
        lines.append("")
        extracted_text = _extract_text_from_file(local_path)
        if extracted_text and extracted_text.strip():
            if len(extracted_text) > 50000:
                extracted_text = extracted_text[:50000] + "\n\n（内容过长，已截断）"
            lines.append(extracted_text)
        else:
            lines.append("（无法提取文本内容）")
        lines.append("")

    return "\n".join(lines)


# ---------------------------------------------------------------------------
# State management
# ---------------------------------------------------------------------------

def _load_state(output_dir):
    path = os.path.join(output_dir, _STATE_FILENAME)
    if os.path.exists(path):
        try:
            with open(path, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception as e:
            logging.warning("Failed to load crawler state: %s", e)
    return {"processed_ids": []}


def _save_state(output_dir, state):
    path = os.path.join(output_dir, _STATE_FILENAME)
    os.makedirs(output_dir, exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(state, f, ensure_ascii=False)


# ---------------------------------------------------------------------------
# KB upload
# ---------------------------------------------------------------------------

def _upload_to_kb(md_content, attachment_files, kb_id, tenant_id, folder_name):
    from api.db.services.knowledgebase_service import KnowledgebaseService
    from api.db.services.file_service import FileService
    from api.db.services.document_service import DocumentService

    ok, kb = KnowledgebaseService.get_by_id(kb_id)
    if not ok:
        raise LookupError("Knowledge base {} not found".format(kb_id))

    class _FO:
        def __init__(self, fn, b):
            self.id = get_uuid()
            self.filename = fn
            self.blob = b

        def read(self):
            return self.blob

    fo = _FO("{}.md".format(folder_name), md_content.encode("utf-8"))
    errs, pairs = FileService.upload_document(kb, [fo], tenant_id)
    if errs:
        logging.warning("MD upload errors: %s", errs)
    for doc, _ in pairs:
        did = doc["id"]
        try:
            DocumentService.begin2parse(did)
            DocumentService.run(tenant_id, doc, {})
        except Exception as e:
            logging.error("Queue parse for %s: %s", did, e)

    for fp in attachment_files:
        fname = os.path.basename(fp)
        with open(fp, "rb") as f:
            blob = f.read()
        fo2 = _FO(fname, blob)
        errs2, pairs2 = FileService.upload_document(kb, [fo2], tenant_id)
        if errs2:
            logging.warning("Attachment upload errors: %s", errs2)
        for doc, _ in pairs2:
            did = doc["id"]
            try:
                DocumentService.begin2parse(did)
                DocumentService.run(tenant_id, doc, {})
            except Exception as e:
                logging.error("Queue parse for %s: %s", did, e)


# ---------------------------------------------------------------------------
# Argument parsing
# ---------------------------------------------------------------------------

def parse_args():
    p = argparse.ArgumentParser(
        description="quanzhou_zcfg_crawler — 泉州市公共资源交易信息网 市级政策法规"
    )
    p.add_argument("--tenant-id", required=True)
    p.add_argument("--kb-id", required=True)
    p.add_argument("--task-name", required=True)
    p.add_argument("--output-dir", default=None,
                   help="Output root directory (default: PROJECT_ROOT/rag/<task_name>)")
    p.add_argument("--full", action="store_true",
                   help="Ignore saved state and re-crawl all")
    p.add_argument("--max-runtime", type=int, default=_MAX_RUNTIME_DEFAULT,
                   help="Max runtime in seconds (default: 3300)")
    for opt in ("--max-days", "--hours", "--max-articles",
                "--llm-id", "--llm-model", "--access-token", "--target-url"):
        p.add_argument(opt, default=None)
    return p.parse_args()


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    args = parse_args()

    _safe_print("\n" + "=" * 60)
    _safe_print("[{}] 泉州市公共资源交易信息网 — 市级政策法规 crawler".format(_TAG_PREFIX))
    _safe_print("[{}] KB: {}".format(_TAG_PREFIX, args.kb_id))
    _safe_print("[{}] Task: {}".format(_TAG_PREFIX, args.task_name))
    _safe_print("[{}] Max runtime: {}s".format(_TAG_PREFIX, args.max_runtime))
    _safe_print("=" * 60 + "\n")
    sys.stdout.flush()

    settings.init_settings()
    logging.info("=== QZ-ZCFG crawler started ===")

    output_dir = args.output_dir or os.path.join(
        _PROJECT_ROOT, "rag", args.task_name.strip()
    )
    os.makedirs(output_dir, exist_ok=True)
    _safe_print("[{}] Output: {}\n".format(_TAG_PREFIX, output_dir))
    sys.stdout.flush()

    # ── State ──────────────────────────────────────────────────────────
    state = _load_state(output_dir) if not args.full else {"processed_ids": []}
    processed_ids = set(state.get("processed_ids", []))

    _safe_print("[{}] Already processed: {} article(s)".format(_TAG_PREFIX, len(processed_ids)))
    sys.stdout.flush()

    crawl_start = time.time()
    downloads_dir = os.path.join(output_dir, "downloads")

    # ── Step 1: Crawl listing ─────────────────────────────────────────
    _safe_print("[{}] Crawling listing pages...".format(_TAG_PREFIX))
    sys.stdout.flush()

    all_items = _crawl_listing()
    _safe_print("[{}] Total articles found: {}".format(_TAG_PREFIX, len(all_items)))
    sys.stdout.flush()

    if not all_items:
        _safe_print("[{}] No articles found, exiting.".format(_TAG_PREFIX))
        return

    # Filter already-processed
    new_items = []
    for item in all_items:
        news_id = str(item.get("newsId", ""))
        if news_id not in processed_ids:
            new_items.append(item)

    skipped = len(all_items) - len(new_items)
    if skipped:
        _safe_print("[{}] {} already processed, {} new".format(
            _TAG_PREFIX, skipped, len(new_items)))
        sys.stdout.flush()

    if not new_items:
        _safe_print("[{}] No new articles to process.".format(_TAG_PREFIX))
        return

    # ── Step 2: Process each article ──────────────────────────────────
    processed_count = 0
    stopped_early = False

    for i, item in enumerate(new_items, 1):
        elapsed = time.time() - crawl_start
        remaining = args.max_runtime - elapsed
        if remaining < 120:
            _safe_print(
                "\n[{}] Runtime {:.0f}s, {:.0f}s remaining. "
                "Stopping gracefully. {} processed.".format(
                    _TAG_PREFIX, elapsed, remaining, processed_count))
            sys.stdout.flush()
            stopped_early = True
            break

        news_id = str(item.get("newsId", ""))
        list_title = item.get("title", "") or "(no title)"

        _safe_print("[{}] [{}/{}] {}".format(
            _TAG_PREFIX, i, len(new_items), list_title[:80]))
        sys.stdout.flush()

        _request_delay()

        # ── Fetch detail ─────────────────────────────────────────────
        detail = _fetch_detail(news_id)
        if not detail:
            _safe_print("[{}]   WARNING: Failed to fetch detail, using listing data".format(_TAG_PREFIX))
            detail = {
                "title": list_title,
                "article_no": "",
                "agency": "",
                "pub_date": _normalize_date(item.get("pubDate", "")),
                "typename": item.get("typename", ""),
                "source": "",
                "click_count": item.get("clickCount", ""),
                "content_html": "",
                "content_text": "",
                "attachments": [],
            }

        detail_url = _SITE_ROOT + "/articleList/ZcfgDetail.do?newsId=" + news_id
        title = detail.get("title", "") or list_title

        attachments = detail.get("attachments", [])
        content_text = detail.get("content_text", "")

        # ── Download attachments ────────────────────────────────────
        local_files = []
        article_dl_dir = ""
        if attachments:
            dl_name = "zcfg_{}_{}".format(
                news_id, _sanitize_filename(title[:30], 40)
            )
            article_dl_dir = os.path.join(downloads_dir, dl_name)
            local_files = _download_attachments(attachments, article_dl_dir)

            # Check for ZIP files and extract
            for fp in list(local_files):
                is_zip = fp.lower().endswith(".zip")
                if not is_zip and os.path.exists(fp) and os.path.getsize(fp) >= 4:
                    with open(fp, "rb") as f:
                        is_zip = f.read(4) == b"PK\x03\x04"
                if is_zip:
                    extracted = _extract_zip(fp)
                    local_files.remove(fp)
                    local_files.extend(extracted)

        # ── Build markdown ──────────────────────────────────────────
        md_content, md_title = _build_markdown(detail, detail_url, article_dl_dir)

        # Append attachment content
        if attachments and article_dl_dir:
            appendix = _build_attachment_appendix(attachments, article_dl_dir)
            if appendix:
                md_content += "\n" + appendix + "\n"

        # Save markdown locally
        pub_date = detail.get("pub_date", "")
        date_for_name = pub_date or datetime.now().strftime("%Y-%m-%d")
        folder_name = _sanitize_filename(
            "zcfg_{}_{}_{}".format(
                date_for_name, news_id, title[:40]
            ), max_len=120
        )
        md_path = os.path.join(output_dir, "{}.md".format(folder_name))
        with open(md_path, "w", encoding="utf-8") as f:
            f.write(md_content)
        _safe_print("[{}]   Saved ({} chars, {} attachments)".format(
            _TAG_PREFIX, len(md_content), len(local_files)))
        sys.stdout.flush()

        # Upload to KB
        if args.kb_id:
            try:
                _upload_to_kb(md_content, local_files, args.kb_id,
                            args.tenant_id, folder_name)
            except Exception as e:
                logging.error("KB upload failed: %s", e)
                _save_state(output_dir, {"processed_ids": list(processed_ids)})
                _safe_print("[{}]   Upload error: {}".format(_TAG_PREFIX, e))
                sys.stdout.flush()

        processed_ids.add(news_id)
        processed_count += 1

        if processed_count % _BATCH_SIZE == 0:
            _save_state(output_dir, {"processed_ids": list(processed_ids)})
            _safe_print("[{}]   Checkpoint ({} processed)".format(
                _TAG_PREFIX, processed_count))
            sys.stdout.flush()

    # ── Final state ────────────────────────────────────────────────────
    _save_state(output_dir, {"processed_ids": list(processed_ids)})

    _safe_print("\n" + "=" * 60)
    _safe_print("[{}] Crawl complete — {} new article(s)".format(_TAG_PREFIX, processed_count))
    if stopped_early:
        _safe_print("[{}] Stopped early, will resume next run".format(_TAG_PREFIX))
    _safe_print("=" * 60 + "\n")
    sys.stdout.flush()
    logging.info("=== QZ-ZCFG crawler finished ===")


if __name__ == "__main__":
    CONSUMER_NAME = "quanzhou_zcfg_crawler"
    init_root_logger(CONSUMER_NAME)
    main()
