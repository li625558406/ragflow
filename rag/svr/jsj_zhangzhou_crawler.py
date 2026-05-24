#!/usr/bin/env python3
#
#  Copyright 2026 The InfiniFlow Authors. All Rights Reserved.
#
#  Licensed under the Apache License, Version 2.0 (the "License");
#  you may not use this file except in compliance with the License.
#  You may obtain a copy of the License at
#
#      http://www.apache.org/licenses/LICENSE-2.0
#
#  Unless required by applicable law or agreed to in writing, software
#  distributed under the License is distributed on an "AS IS" BASIS,
#  WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
#  See the License for the specific language governing permissions and
#  limitations under the License.
#
"""
Crawler for jsj.zhangzhou.gov.cn — 漳州市住房和城乡建设局 政务公开 (ALL data).

Target: http://jsj.zhangzhou.gov.cn/cms/html/zzszfhcxjsj/zwgk/index.html

Site characteristics
────────────────────
Built on ylcms CMS. All listing data is server-rendered HTML — no JS
rendering needed. The main page has 6 modules with 17 tabs total. Each tab
shows 5 preview items and links to a section index page with full pagination.

Modules & tabs:
  1. 工作动态 (gzdt), 通知公告 (tzgg3), 最新文件 (zxwj)
  2. 财政资金 (zjxx), 人事信息 (rsxx), 规划计划 (ghjh)
  3. 建筑工程 (jzgc), 消防工程 (xfgc), 资质管理 (zzgl1)
  4. 城市建设 (csjs), 村镇建设, 安全管理 (aqgl)
  5. 科技设计, 房屋交易 (fwcq), 统计信息 (tjxx)
  6. 重点项目 (zdxm), 电子刊物 (dzkw)

Listing format: <UL id="resources"> <LI><A href>title</A><SPAN>date</SPAN></LI>
Pagination: /cms/sitemanage/index.shtml?siteId={siteId}&page=N

Two detail page types:
  Type A — Regular articles: /cms/html/.../YYYY-MM-DD/NUM.html
    Meta: ArticleTitle, PubDate (YYYY—MM—DD HH:MM), ContentSource
    Body: div.top-title1 (title), div#Content (article body)
    No standard attachment block

  Type B — PublicInfo articles: /cms/infopublic/publicInfo.shtml?id=...&siteId=...
    Meta: ArticleTitle, PubDate
    Body: div.title h1, ul.info (metadata), div#Content div.tsrc (body)
    Attachments: li#wjxz span#pdfattachments a (PDF), span#wordattachments a (Word)

Data flow
─────────
  1. Parse main page → discover all 17 sections (tab names + index URLs)
  2. For each section → visit index page → get siteId + page count → paginate
  3. Parse listings → extract article URLs, titles, dates
  4. Dedup by article ID across sections
  5. For each new article → fetch detail page → parse content + attachments
  6. Download attachments (PDF/DOC/XLSX/ZIP) → extract text
  7. Build markdown → save locally → upload to KB

Checkpoint/resume: state saved per batch. Time-bounded check (default 3300s)
stops gracefully before the 3600s task-timeout window.

Usage
─────
    python jsj_zhangzhou_crawler.py \
        --tenant-id <TENANT_ID> \
        --kb-id <KB_ID> \
        --task-name <NAME>

    # Optional:
        --max-runtime 3300    # Max runtime before graceful stop
        --full                # Ignore saved state, re-crawl
"""

import argparse
import hashlib
import json
import logging
import os
import random
import re
import sys
import time
import urllib.parse
import urllib.request
import zipfile
from datetime import datetime

_SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
_PROJECT_ROOT = os.path.abspath(os.path.join(_SCRIPT_DIR, "..", ".."))
sys.path.insert(0, _PROJECT_ROOT)

from common import settings
from common.log_utils import init_root_logger
from common.misc_utils import get_uuid
from bs4 import BeautifulSoup


# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------

_SITE_ROOT = "http://jsj.zhangzhou.gov.cn"
_MAIN_PAGE = "/cms/html/zzszfhcxjsj/zwgk/index.html"
_MAIN_PAGE_URL = _SITE_ROOT + _MAIN_PAGE

# Checkpoint batch size (articles)
_BATCH_SIZE = 3

# Default max runtime (55 min, 5 min margin)
_MAX_RUNTIME_DEFAULT = 3300

# Anti-crawling delays (seconds)
_REQUEST_DELAY_MIN = 0.2
_REQUEST_DELAY_MAX = 0.8

# State filename
_STATE_FILENAME = "_crawler_state.json"

# Max pagination pages per section
_MAX_PAGES = 500

_USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/120.0.0.0 Safari/537.36"
)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _safe_print(msg):
    try:
        print(msg)
    except UnicodeEncodeError:
        print(msg.encode("gbk", errors="replace").decode("gbk"))


def _sanitize_filename(text, max_len=120):
    if not text:
        return "untitled"
    name = re.sub(r'[\\/:*?"<>|]', "_", str(text).strip())
    name = re.sub(r'\s+', "_", name).strip("._ ")
    return (name or "untitled")[:max_len]


def _resolve_url(href, base_url):
    if href.startswith("http://") or href.startswith("https://"):
        return href
    if href.startswith("//"):
        return "https:" + href
    return urllib.parse.urljoin(base_url, href)


def _extract_id_from_url(url):
    """Extract a stable article ID from URL.

    Patterns:
      /YYYY-MM-DD/NUM.html          -> NUM (e.g. 1927315017)
      publicInfo.shtml?id=NUM       -> NUM (e.g. 20684996108160004)
      leafinfodirectory/NUM.html    -> NUM
    """
    # publicInfo ID
    m = re.search(r'[?&]id=(\d+)', url)
    if m:
        return m.group(1)
    # Date-based article: /YYYY-MM-DD/NUM.html
    m = re.search(r'/(\d{4}-\d{2}-\d{2})/(\d+)\.html', url)
    if m:
        return m.group(2)
    # leafinfodirectory or other numeric
    m = re.search(r'/(\d{9,})\.html', url)
    if m:
        return m.group(1)
    path = urllib.parse.urlparse(url).path
    return hashlib.md5(path.encode()).hexdigest()[:12]


def _normalize_pubdate(text):
    """Normalize PubDate from various formats to YYYY-MM-DD."""
    if not text:
        return ""
    text = text.strip()
    # "2026—05—22 16:29" -> "2026-05-22"
    text = text.replace("\u2014", "-").replace("\uff0d", "-").replace("/", "-")
    m = re.search(r'(\d{4})-(\d{2})-(\d{2})', text)
    if m:
        return "{}-{}-{}".format(m.group(1), m.group(2), m.group(3))
    return ""


# ---------------------------------------------------------------------------
# HTTP helpers
# ---------------------------------------------------------------------------

def _http_get(url, referer=None, timeout=30):
    req = urllib.request.Request(url)
    req.add_header("User-Agent", _USER_AGENT)
    req.add_header("Accept", "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8")
    req.add_header("Accept-Language", "zh-CN,zh;q=0.9,en;q=0.8")
    if referer:
        req.add_header("Referer", referer)
    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            return resp.read()
    except Exception as e:
        logging.error("GET %s failed: %s", url, e)
        return None


def _download_binary(url, referer=None, timeout=60):
    req = urllib.request.Request(url)
    req.add_header("User-Agent", _USER_AGENT)
    if referer:
        req.add_header("Referer", referer)
    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            return resp.read()
    except Exception as e:
        logging.error("Download failed: %s — %s", url[:120], e)
        return None


# ---------------------------------------------------------------------------
# Section discovery from main page
# ---------------------------------------------------------------------------

def _discover_sections(html_bytes):
    """Parse the main zwgk page to discover all sections.

    Returns list of {label, index_url, is_leafinfo}.

    For standard tabs, label comes from the "更多" link text.
    For leafinfodirectory tabs, the "更多" link text is a URL — we
    resolve the label from the tab button instead.
    """
    if not html_bytes:
        return []

    html = html_bytes.decode("utf-8", errors="replace")
    soup = BeautifulSoup(html, "html.parser")

    # Build label map from tab buttons: {suffix -> label}
    # e.g. "375615511_1" -> "村镇建设"
    button_labels = {}
    for btn in soup.find_all("li", id=re.compile(r"^tabstripButton_")):
        btn_id = btn.get("id", "")
        suffix = btn_id.replace("tabstripButton_", "")
        a_tag = btn.find("a")
        if a_tag:
            button_labels[suffix] = a_tag.get_text(strip=True)

    sections = []
    seen_urls = set()

    for more_div in soup.find_all("div", id=re.compile(r"^tabstripMore_")):
        a_tag = more_div.find("a", href=True)
        if not a_tag:
            continue

        href = a_tag["href"].strip()
        label = a_tag.get_text(strip=True)

        if not label or href in seen_urls:
            continue
        seen_urls.add(href)

        # Resolve label for leafinfodirectory (link text is a URL, not a label)
        is_leafinfo = "leafinfodirectory" in href or "directoryId" in href
        if is_leafinfo or label.startswith("index.shtml"):
            more_id = more_div.get("id", "")
            suffix = more_id.replace("tabstripMore_", "")
            resolved = button_labels.get(suffix, "")
            if resolved:
                label = resolved

        full_url = _resolve_url(href, _MAIN_PAGE_URL)

        sections.append({
            "label": label,
            "index_url": full_url,
            "is_leafinfo": is_leafinfo,
        })

    return sections


# ---------------------------------------------------------------------------
# Listing page parsing
# ---------------------------------------------------------------------------

def _parse_section_listing(html_bytes, base_url, section_label, seen_ids):
    """Parse a section listing page for article links.

    Two formats exist:
      A) <UL id="resources"> <LI><A href>title</A><SPAN>date</SPAN></LI>
      B) <UL id="infos"> <DIV class="menuli"> <LI><A>title</A><SPAN>date</SPAN></LI>
    Also handles sitemanage pagination pages.
    """
    if not html_bytes:
        return []

    html = html_bytes.decode("utf-8", errors="replace")
    soup = BeautifulSoup(html, "html.parser")
    articles = []

    ul = soup.find("ul", id="resources")
    list_mode = "li"  # standard: LI as direct children

    if not ul:
        ul = soup.find("ul", id="infos")
        list_mode = "menuli"  # leafinfodirectory: LI inside DIV.menuli

    if not ul:
        return articles

    if list_mode == "menuli":
        # Format B: find all div.menuli, extract LI from inside
        for menuli in ul.find_all("div", class_="menuli", recursive=False):
            li = menuli.find("li")
            if not li:
                continue
            a_tag = li.find("a", href=True)
            if not a_tag:
                continue
            href = a_tag["href"].strip()
            title = (a_tag.get("title") or a_tag.get_text()).strip()
            if not title or len(title) < 3:
                continue
            if href.startswith("javascript:") or href == "#":
                continue
            span = li.find("span")
            date_text = span.get_text(strip=True) if span else ""
            date_str = date_text.strip() if date_text else ""
            article_id = _extract_id_from_url(href)
            if article_id in seen_ids:
                continue
            seen_ids.add(article_id)
            articles.append({
                "title": title,
                "url": _resolve_url(href, base_url),
                "date_str": date_str,
                "article_id": article_id,
                "section": section_label,
            })
        return articles

    # Format A: direct LI children
    for li in ul.find_all("li", recursive=False):
        a_tag = li.find("a", href=True)
        if not a_tag:
            continue

        href = a_tag["href"].strip()
        title = (a_tag.get("title") or a_tag.get_text()).strip()
        if not title or len(title) < 3:
            continue
        if href.startswith("javascript:") or href == "#":
            continue

        span = li.find("span")
        date_text = span.get_text(strip=True) if span else ""
        date_str = date_text.strip() if date_text else ""

        article_id = _extract_id_from_url(href)
        if article_id in seen_ids:
            continue
        seen_ids.add(article_id)

        articles.append({
            "title": title,
            "url": _resolve_url(href, base_url),
            "date_str": date_str,
            "article_id": article_id,
            "section": section_label,
        })

    return articles


def _extract_page_info(html_bytes):
    """Extract siteId, directoryId and page count from a section index page.

    Returns (site_id, directory_id, page_count).
    directory_id is only set for leafinfodirectory sections.
    """
    if not html_bytes:
        return (None, None, 1)

    html = html_bytes.decode("utf-8", errors="replace")
    soup = BeautifulSoup(html, "html.parser")

    # siteId from meta
    site_id = None
    meta_site = soup.find("meta", attrs={"name": "siteIdMeta"})
    if meta_site and meta_site.get("content"):
        site_id = meta_site["content"].strip()

    # directoryId from pagination links (leafinfodirectory sections)
    directory_id = None
    for a in soup.find_all("a", href=re.compile(r"directoryId=")):
        m = re.search(r'directoryId=(\d+)', a.get("href", ""))
        if m:
            directory_id = m.group(1)
            break

    # Page count: find all page=N links, take max N
    page_count = 1
    for a in soup.find_all("a", href=re.compile(r"page=\d+")):
        m = re.search(r'page=(\d+)', a.get("href", ""))
        if m:
            p = int(m.group(1))
            if p > page_count:
                page_count = p

    return (site_id, directory_id, max(page_count, 1))


def _crawl_section_paginated(section, referer, seen_ids):
    """Crawl all pages of a section, returning all articles across all pages.

    Standard sections: /cms/sitemanage/index.shtml?siteId={siteId}&page=N
    Leafinfodirectory: /cms/infopublic/index.shtml?directoryId={dirId}&...&page=N
    """
    index_url = section["index_url"]
    label = section["label"]
    is_leafinfo = section.get("is_leafinfo", False)

    all_articles = []

    # Page 1: the section index URL
    time.sleep(random.uniform(_REQUEST_DELAY_MIN, _REQUEST_DELAY_MAX))
    html_bytes = _http_get(index_url, referer=referer)
    if not html_bytes:
        return all_articles

    site_id, directory_id, page_count = _extract_page_info(html_bytes)
    _safe_print("      siteId={}, dirId={}, {} page(s)".format(site_id, directory_id, page_count))
    sys.stdout.flush()

    # Parse page 1
    articles = _parse_section_listing(html_bytes, index_url, label, seen_ids)
    all_articles.extend(articles)

    if page_count <= 1:
        return all_articles

    # Pages 2..N
    if is_leafinfo and directory_id:
        # Leafinfodirectory pagination
        page_base = "{}?directoryId={}&leafInfoDirectoryPageEnable=true&siteId={}".format(
            _SITE_ROOT + "/cms/infopublic/index.shtml", directory_id, site_id or "")
    elif site_id:
        # Standard sitemanage pagination
        page_base = "{}?siteId={}".format(
            _SITE_ROOT + "/cms/sitemanage/index.shtml", site_id)
    else:
        return all_articles

    for page_num in range(2, min(page_count + 1, _MAX_PAGES + 1)):
        page_url = "{}&page={}".format(page_base, page_num)
        time.sleep(random.uniform(_REQUEST_DELAY_MIN, _REQUEST_DELAY_MAX))
        html_bytes = _http_get(page_url, referer=index_url)
        if not html_bytes:
            continue

        articles = _parse_section_listing(html_bytes, page_url, label, seen_ids)
        if not articles:
            break
        all_articles.extend(articles)

    return all_articles


# ---------------------------------------------------------------------------
# Detail page parsing — Type A: Regular HTML articles
# ---------------------------------------------------------------------------

def _parse_detail_regular(html_bytes, detail_url):
    """Parse a regular article detail page.

    Title: meta ArticleTitle or div.top-title1
    Date: meta PubDate or div.top-title2 span.rq
    Source: meta ContentSource or div.top-title2 span.ly
    Content: div#Content
    """
    if not html_bytes:
        return None

    html = html_bytes.decode("utf-8", errors="replace")
    soup = BeautifulSoup(html, "html.parser")

    info = {
        "title": "",
        "pub_date": "",
        "info_source": "",
        "article_id": _extract_id_from_url(detail_url),
        "content_html": "",
        "content_text": "",
        "attachments": [],
    }

    # Title: meta
    meta_title = soup.find("meta", attrs={"name": "ArticleTitle"})
    if meta_title and meta_title.get("content"):
        info["title"] = meta_title["content"].strip()

    # Title: div.top-title1 (fallback)
    if not info["title"]:
        h1 = soup.find("div", class_="top-title1")
        if h1:
            info["title"] = h1.get_text(strip=True)

    # PubDate: meta
    meta_pub = soup.find("meta", attrs={"name": "PubDate"})
    if meta_pub and meta_pub.get("content"):
        info["pub_date"] = _normalize_pubdate(meta_pub["content"])

    # PubDate fallback: span.rq
    if not info["pub_date"]:
        rq = soup.find("span", class_="rq")
        if rq:
            t = rq.get_text(strip=True)
            info["pub_date"] = _normalize_pubdate(t)

    # ContentSource: meta
    meta_src = soup.find("meta", attrs={"name": "ContentSource"})
    if meta_src and meta_src.get("content"):
        info["info_source"] = meta_src["content"].strip()

    # Source fallback: span.ly
    if not info["info_source"]:
        ly = soup.find("span", class_="ly")
        if ly:
            s = ly.get_text(strip=True)
            s = re.sub(r'^来源[：:]\s*', '', s).strip()
            if s:
                info["info_source"] = s

    # Content: div#Content
    content_div = soup.find("div", id="Content")
    if not content_div:
        content_div = soup.find("div", class_="content")
    if content_div:
        for tag in content_div.find_all(["script", "style"]):
            tag.decompose()
        info["content_html"] = content_div.decode_contents()
        info["content_text"] = content_div.get_text(separator="\n", strip=True)

    return info


# ---------------------------------------------------------------------------
# Detail page parsing — Type B: PublicInfo articles
# ---------------------------------------------------------------------------

def _parse_detail_publicinfo(html_bytes, detail_url):
    """Parse a publicInfo article detail page.

    Title: meta ArticleTitle or div.title h1
    Date: meta PubDate or ul.info > li (发布日期)
    Source: meta ContentSource or ul.info > li (发文机关)
    Content: div#Content div.tsrc
    Attachments: li#wjxz a links (pdf/word)
    """
    if not html_bytes:
        return None

    html = html_bytes.decode("utf-8", errors="replace")
    soup = BeautifulSoup(html, "html.parser")

    info = {
        "title": "",
        "pub_date": "",
        "info_source": "",
        "article_id": _extract_id_from_url(detail_url),
        "content_html": "",
        "content_text": "",
        "attachments": [],
    }

    # Title: meta
    meta_title = soup.find("meta", attrs={"name": "ArticleTitle"})
    if meta_title and meta_title.get("content"):
        info["title"] = meta_title["content"].strip()

    # Title: div.title h1 (fallback)
    if not info["title"]:
        title_div = soup.find("div", class_="title")
        if title_div:
            h1 = title_div.find("h1")
            if h1:
                info["title"] = h1.get_text(strip=True)

    # PubDate: meta
    meta_pub = soup.find("meta", attrs={"name": "PubDate"})
    if meta_pub and meta_pub.get("content"):
        info["pub_date"] = _normalize_pubdate(meta_pub["content"])

    # PubDate fallback: ul.info li 发布日期
    if not info["pub_date"]:
        info_ul = soup.find("ul", class_="info")
        if info_ul:
            for li in info_ul.find_all("li"):
                text = li.get_text(strip=True)
                if "发布日期" in text:
                    info["pub_date"] = _normalize_pubdate(text)
                    break

    # ContentSource: meta
    meta_src = soup.find("meta", attrs={"name": "ContentSource"})
    if meta_src and meta_src.get("content"):
        info["info_source"] = meta_src["content"].strip()

    # Source fallback: ul.info li 发文机关
    if not info["info_source"]:
        info_ul = soup.find("ul", class_="info")
        if info_ul:
            for li in info_ul.find_all("li"):
                text = li.get_text(strip=True)
                if "发文机关" in text:
                    info["info_source"] = re.sub(r'^发文机关[：:]\s*', '', text).strip()
                    break

    # Content: div#Content div.tsrc
    content_div = soup.find("div", id="Content")
    if content_div:
        # Prefer text version (tsrc[1]) over iframe version
        tsrc_divs = content_div.find_all("div", class_="tsrc")
        text_parts = []
        for tdiv in tsrc_divs:
            iframe = tdiv.find("iframe")
            if iframe:
                continue  # skip iframe / PDF viewer
            for tag in tdiv.find_all(["script", "style"]):
                tag.decompose()
            text_parts.append(tdiv.decode_contents())
        if text_parts:
            # Combine all text tsrc divs
            combined = "".join(text_parts)
            info["content_html"] = combined
            combo_soup = BeautifulSoup(combined, "html.parser")
            info["content_text"] = combo_soup.get_text(separator="\n", strip=True)

    # Attachments: li#wjxz
    wjxz = soup.find("li", id="wjxz")
    if wjxz:
        for a_tag in wjxz.find_all("a", href=True):
            href = a_tag["href"].strip()
            text = a_tag.get_text(strip=True)
            if not text or text in ("无word文件", "无pdf文件", "无文件"):
                continue
            att_url = _resolve_url(href, detail_url)
            info["attachments"].append({
                "filename": text or os.path.basename(urllib.parse.urlparse(att_url).path),
                "url": att_url,
                "type": "download",
            })

    return info


def _parse_detail(html_bytes, detail_url):
    """Dispatch to the appropriate detail parser based on URL pattern."""
    if "publicInfo.shtml" in detail_url:
        return _parse_detail_publicinfo(html_bytes, detail_url)
    return _parse_detail_regular(html_bytes, detail_url)


# ---------------------------------------------------------------------------
# Attachment download + processing
# ---------------------------------------------------------------------------

def _download_attachments(attachments, download_dir):
    os.makedirs(download_dir, exist_ok=True)
    local_files = []

    for att in attachments:
        url = att.get("url", "")
        if not url:
            continue
        fname = _sanitize_filename(att.get("filename", "attachment"), max_len=120)
        ext = os.path.splitext(urllib.parse.urlparse(url).path)[1].lower()
        if ext and not fname.lower().endswith(ext):
            fname += ext

        filepath = os.path.join(download_dir, fname)
        if os.path.exists(filepath) and os.path.getsize(filepath) > 0:
            local_files.append(filepath)
            continue

        data = _download_binary(url)
        if data:
            with open(filepath, "wb") as f:
                f.write(data)
            local_files.append(filepath)
            time.sleep(random.uniform(0.1, 0.3))

    return local_files


def _extract_zip(filepath):
    extracted = []
    extract_dir = os.path.splitext(filepath)[0] + "_extracted"
    os.makedirs(extract_dir, exist_ok=True)
    try:
        with zipfile.ZipFile(filepath, "r") as zf:
            for name in zf.namelist():
                safe_name = _sanitize_filename(name, max_len=120)
                out_path = os.path.join(extract_dir, safe_name)
                os.makedirs(os.path.dirname(out_path) or extract_dir, exist_ok=True)
                with zf.open(name) as src, open(out_path, "wb") as dst:
                    dst.write(src.read())
                extracted.append(out_path)
    except Exception as e:
        logging.warning("ZIP extraction failed for %s: %s", filepath, e)
    return extracted


# ---------------------------------------------------------------------------
# Text extraction from binary files
# ---------------------------------------------------------------------------

def _extract_text_from_file(filepath):
    ext = os.path.splitext(filepath)[1].lower()
    try:
        if ext == ".txt":
            with open(filepath, "r", encoding="utf-8", errors="replace") as f:
                return f.read()
        elif ext == ".pdf":
            import pdfplumber
            parts = []
            with pdfplumber.open(filepath) as pdf:
                for pg in pdf.pages:
                    text = pg.extract_text()
                    if text:
                        parts.append(text)
            return "\n\n".join(parts)
        elif ext in (".docx", ".doc"):
            import docx
            doc = docx.Document(filepath)
            return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        elif ext in (".xlsx", ".xls"):
            import openpyxl
            wb = openpyxl.load_workbook(filepath, read_only=True)
            parts = []
            for ws in wb.worksheets:
                rows = []
                for row in ws.iter_rows(values_only=True):
                    rows.append(" | ".join(str(c) if c is not None else "" for c in row))
                if rows:
                    parts.append("### {}\n".format(ws.title) + "\n".join(rows))
            wb.close()
            return "\n\n".join(parts)
    except Exception as e:
        logging.warning("Text extraction failed for %s: %s", filepath, e)
    return None


# ---------------------------------------------------------------------------
# Markdown builder
# ---------------------------------------------------------------------------

def _build_markdown(detail, download_dir, source_url):
    info = detail or {}
    title = info.get("title", "\u65e0\u6807\u9898")
    pub_date = info.get("pub_date", "")
    info_source = info.get("info_source", "")

    lines = [
        "# {}".format(title),
        "",
        "**\u6570\u636e\u6765\u6e90:** \u6f33\u5dde\u5e02\u4f4f\u623f\u548c\u57ce\u4e61\u5efa\u8bbe\u5c40",
        "**\u9875\u9762\u5730\u5740:** {}".format(source_url),
        "**\u6293\u53d6\u65f6\u95f4:** {}".format(datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
    ]
    if pub_date:
        lines.append("**\u53d1\u5e03\u65f6\u95f4:** {}".format(pub_date))
    if info_source:
        lines.append("**\u4fe1\u606f\u6765\u6e90:** {}".format(info_source))
    lines.append("")

    content_text = info.get("content_text", "")
    if content_text:
        lines.append("---")
        lines.append("")
        lines.append("## \u6b63\u6587")
        lines.append("")
        content_clean = re.sub(r"\n{3,}", "\n\n", content_text)
        lines.append(content_clean)
        lines.append("")

    attachments = info.get("attachments", [])
    if attachments:
        lines.append("---")
        lines.append("")
        lines.append("## \u9644\u4ef6")
        lines.append("")
        for att in attachments:
            fname = att.get("filename", "unknown")
            att_url = att.get("url", "")
            lines.append("- [{}]({})".format(fname, att_url))
        lines.append("")

        if download_dir and os.path.isdir(download_dir):
            lines.append("### \u9644\u4ef6\u5185\u5bb9")
            lines.append("")
            for att in attachments:
                if att.get("type") != "download":
                    continue
                fname = att.get("filename", "")
                local_path = os.path.join(download_dir, fname)
                safe_name = _sanitize_filename(fname, max_len=120)
                alt_path = os.path.join(download_dir, safe_name)
                if not os.path.exists(local_path) and os.path.exists(alt_path):
                    local_path = alt_path
                if not os.path.exists(local_path):
                    for root, _, files in os.walk(download_dir):
                        for fn in files:
                            if fn == safe_name or fn == fname:
                                local_path = os.path.join(root, fn)
                                break
                if not os.path.exists(local_path):
                    continue

                lines.append("#### {}".format(fname))
                lines.append("")
                text = _extract_text_from_file(local_path)
                if text and text.strip():
                    if len(text) > 50000:
                        text = text[:50000] + "\n\n\uff08\u5185\u5bb9\u8fc7\u957f\uff0c\u5df2\u622a\u65ad\uff09"
                    lines.append(text)
                else:
                    lines.append("\uff08\u65e0\u6cd5\u63d0\u53d6\u6587\u672c\u5185\u5bb9\uff09")
                lines.append("")

    return "\n".join(lines)


# ---------------------------------------------------------------------------
# State management
# ---------------------------------------------------------------------------

def _load_state(output_dir):
    path = os.path.join(output_dir, _STATE_FILENAME)
    if os.path.exists(path):
        try:
            with open(path, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception as e:
            logging.warning("Failed to load crawler state: %s", e)
    return {"processed_ids": []}


def _save_state(output_dir, state):
    path = os.path.join(output_dir, _STATE_FILENAME)
    os.makedirs(output_dir, exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(state, f, ensure_ascii=False)


# ---------------------------------------------------------------------------
# KB upload
# ---------------------------------------------------------------------------

def _upload_to_kb(md_content, attachment_files, kb_id, tenant_id, folder_name):
    from api.db.services.knowledgebase_service import KnowledgebaseService
    from api.db.services.file_service import FileService
    from api.db.services.document_service import DocumentService

    ok, kb = KnowledgebaseService.get_by_id(kb_id)
    if not ok:
        raise LookupError("Knowledge base {} not found".format(kb_id))

    class _FO:
        def __init__(self, fn, b):
            self.id = get_uuid()
            self.filename = fn
            self.blob = b

        def read(self):
            return self.blob

    fo = _FO("{}.md".format(folder_name), md_content.encode("utf-8"))
    errs, pairs = FileService.upload_document(kb, [fo], tenant_id)
    if errs:
        logging.warning("MD upload errors: %s", errs)
    for doc, _ in pairs:
        did = doc["id"]
        try:
            DocumentService.begin2parse(did)
            DocumentService.run(tenant_id, doc, {})
        except Exception as e:
            logging.error("Queue parse for %s: %s", did, e)

    for fp in attachment_files:
        fname = os.path.basename(fp)
        with open(fp, "rb") as f:
            blob = f.read()
        fo2 = _FO(fname, blob)
        errs2, pairs2 = FileService.upload_document(kb, [fo2], tenant_id)
        if errs2:
            logging.warning("Attachment upload errors: %s", errs2)
        for doc, _ in pairs2:
            did = doc["id"]
            try:
                DocumentService.begin2parse(did)
                DocumentService.run(tenant_id, doc, {})
            except Exception as e:
                logging.error("Queue parse for %s: %s", did, e)


# ---------------------------------------------------------------------------
# Argument parsing
# ---------------------------------------------------------------------------

def parse_args():
    p = argparse.ArgumentParser(
        description="jsj.zhangzhou.gov.cn crawler — 漳州市住建局 政务公开 (ALL data)"
    )
    p.add_argument("--tenant-id", required=True)
    p.add_argument("--kb-id", required=True)
    p.add_argument("--task-name", required=True)
    p.add_argument("--output-dir", default=None,
                   help="Output root directory (default: PROJECT_ROOT/rag/<task_name>)")
    p.add_argument("--full", action="store_true",
                   help="Ignore saved state and re-crawl all")
    p.add_argument("--max-runtime", type=int, default=_MAX_RUNTIME_DEFAULT,
                   help="Max runtime in seconds (default: 3300)")
    for opt in ("--max-days", "--hours", "--max-articles",
                "--llm-id", "--llm-model", "--access-token", "--target-url"):
        p.add_argument(opt, default=None)
    return p.parse_args()


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    args = parse_args()

    _safe_print("\n" + "=" * 60)
    _safe_print("[JSJ-ZZ] 漳州市住房和城乡建设局 — 政务公开 crawler")
    _safe_print("[JSJ-ZZ] KB: {}".format(args.kb_id))
    _safe_print("[JSJ-ZZ] Task: {}".format(args.task_name))
    _safe_print("[JSJ-ZZ] Max runtime: {}s".format(args.max_runtime))
    _safe_print("=" * 60 + "\n")
    sys.stdout.flush()

    settings.init_settings()
    logging.info("=== JSJ-ZhangZhou crawler started ===")

    output_dir = args.output_dir or os.path.join(
        _PROJECT_ROOT, "rag", args.task_name.strip()
    )
    os.makedirs(output_dir, exist_ok=True)
    _safe_print("[JSJ-ZZ] Output: {}\n".format(output_dir))
    sys.stdout.flush()

    # ── State ──────────────────────────────────────────────────────────
    state = _load_state(output_dir) if not args.full else {"processed_ids": []}
    processed_ids = set(state.get("processed_ids", []))

    _safe_print("[JSJ-ZZ] Already processed: {} article(s)".format(len(processed_ids)))
    sys.stdout.flush()

    crawl_start = time.time()

    # ── Step 1: Discover sections ──────────────────────────────────────
    _safe_print("\n[JSJ-ZZ] Step 1/4: Discovering sections from main page...")
    sys.stdout.flush()

    time.sleep(random.uniform(_REQUEST_DELAY_MIN, _REQUEST_DELAY_MAX))
    main_html = _http_get(_MAIN_PAGE_URL)
    sections = _discover_sections(main_html)

    _safe_print("[JSJ-ZZ]   Found {} section(s):".format(len(sections)))
    for s in sections:
        _safe_print("[JSJ-ZZ]     - {} ({})".format(s["label"], s["index_url"]))
    sys.stdout.flush()

    if not sections:
        _safe_print("[JSJ-ZZ] No sections found, aborting.")
        sys.stdout.flush()
        return

    # ── Step 2: Crawl all sections ─────────────────────────────────────
    _safe_print("\n[JSJ-ZZ] Step 2/4: Crawling section listings...")
    sys.stdout.flush()

    seen_ids = set()  # cross-section dedup during crawl
    all_articles = []

    for idx, section in enumerate(sections, 1):
        elapsed = time.time() - crawl_start
        remaining = args.max_runtime - elapsed
        if remaining < 120:
            _safe_print("[JSJ-ZZ] Runtime limit approaching, skipping remaining sections")
            sys.stdout.flush()
            break

        _safe_print("[JSJ-ZZ]   [{}/{}] {}...".format(idx, len(sections), section["label"]))
        sys.stdout.flush()

        articles = _crawl_section_paginated(section, _MAIN_PAGE_URL, seen_ids)
        _safe_print("[JSJ-ZZ]     -> {} new article(s)".format(len(articles)))
        sys.stdout.flush()
        all_articles.extend(articles)

    _safe_print("[JSJ-ZZ]   Total: {} unique article(s) across {} section(s)".format(
        len(all_articles), len(sections)))
    sys.stdout.flush()

    if not all_articles:
        _safe_print("[JSJ-ZZ] No articles found, done.")
        sys.stdout.flush()
        return

    # Section breakdown
    section_counts = {}
    for a in all_articles:
        sec = a.get("section", "unknown")
        section_counts[sec] = section_counts.get(sec, 0) + 1
    for sec, cnt in sorted(section_counts.items()):
        _safe_print("[JSJ-ZZ]     {}: {} article(s)".format(sec, cnt))
    sys.stdout.flush()

    # Filter already-processed
    new_articles = [
        a for a in all_articles
        if a.get("article_id") and a["article_id"] not in processed_ids
    ]
    skipped = len(all_articles) - len(new_articles)
    if skipped:
        _safe_print("[JSJ-ZZ] {} already processed, {} new".format(skipped, len(new_articles)))
        sys.stdout.flush()

    if not new_articles:
        _safe_print("[JSJ-ZZ] All available articles already processed.")
        sys.stdout.flush()
        return

    # ── Step 3: Process each article ───────────────────────────────────
    _safe_print("\n[JSJ-ZZ] Step 3/4: Processing {} article(s)...\n".format(len(new_articles)))
    sys.stdout.flush()

    processed_count = 0
    stopped_early = False
    downloads_dir = os.path.join(output_dir, "downloads")

    for i, article in enumerate(new_articles, 1):
        # Time-bounded check
        elapsed = time.time() - crawl_start
        remaining = args.max_runtime - elapsed
        if remaining < 120:
            _safe_print(
                "\n[JSJ-ZZ] Runtime {:.0f}s, {:.0f}s remaining (limit {}s), "
                "stopping gracefully. {} processed. "
                "Next run will resume.".format(elapsed, remaining, args.max_runtime, processed_count))
            sys.stdout.flush()
            stopped_early = True
            break

        article_id = article["article_id"]
        title = article["title"]
        detail_url = article["url"]
        date_str = article.get("date_str", "")
        section_label = article.get("section", "")

        _safe_print("[JSJ-ZZ] [{}/{}] [{}] {}...".format(
            i, len(new_articles), section_label, title[:50]))
        sys.stdout.flush()

        # Handle direct PDF links in listing
        if detail_url.lower().endswith(".pdf"):
            dl_name = "{}_{}".format(article_id[:12], _sanitize_filename(title[:30], 40))
            article_dl_dir = os.path.join(downloads_dir, dl_name)
            os.makedirs(article_dl_dir, exist_ok=True)

            fname = os.path.basename(urllib.parse.urlparse(detail_url).path)
            pdf_path = os.path.join(article_dl_dir, fname)
            pdf_data = _download_binary(detail_url, referer=_MAIN_PAGE_URL)
            if pdf_data:
                with open(pdf_path, "wb") as f:
                    f.write(pdf_data)
                pdf_text = _extract_text_from_file(pdf_path)
                md_content = "\n".join([
                    "# {}".format(title),
                    "",
                    "**数据来源:** 漳州市住房和城乡建设局",
                    "**页面地址:** {}".format(detail_url),
                    "**抓取时间:** {}".format(datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
                    "**发布时间:** {}".format(date_str),
                    "",
                    "---",
                    "",
                    "## 正文（PDF提取）",
                    "",
                    pdf_text or "（无法提取PDF文本内容）",
                ])
                folder_name = _sanitize_filename(
                    "{}_{}_{}".format(date_str or datetime.now().strftime("%Y-%m-%d"),
                                      article_id[:12], title[:40]), max_len=120)
                md_path = os.path.join(output_dir, "{}.md".format(folder_name))
                with open(md_path, "w", encoding="utf-8") as f:
                    f.write(md_content)
                _safe_print("[JSJ-ZZ]   Saved PDF ({} chars)".format(len(md_content)))
                sys.stdout.flush()

                if args.kb_id:
                    try:
                        _upload_to_kb(md_content, [pdf_path], args.kb_id,
                                     args.tenant_id, folder_name)
                    except Exception as e:
                        logging.error("KB upload failed: %s", e)

                processed_ids.add(article_id)
                processed_count += 1
                time.sleep(random.uniform(_REQUEST_DELAY_MIN, _REQUEST_DELAY_MAX))
                continue

        # Fetch detail page
        time.sleep(random.uniform(_REQUEST_DELAY_MIN, _REQUEST_DELAY_MAX))
        html_bytes = _http_get(detail_url, referer=_MAIN_PAGE_URL)
        if not html_bytes:
            logging.warning("Failed to fetch detail: %s", detail_url)
            processed_ids.add(article_id)
            continue

        detail = _parse_detail(html_bytes, detail_url)
        if not detail:
            logging.warning("Failed to parse detail: %s", detail_url)
            processed_ids.add(article_id)
            continue

        # Prefer detail page title if available
        if detail.get("title") and len(detail["title"]) > len(title):
            title = detail["title"]

        # Download attachments
        attachments = detail.get("attachments", [])
        local_files = []
        article_dl_dir = ""
        if attachments:
            dl_name = "{}_{}".format(article_id[:12], _sanitize_filename(title[:30], 40))
            article_dl_dir = os.path.join(downloads_dir, dl_name)
            local_files = _download_attachments(attachments, article_dl_dir)
            # Extract ZIPs
            for fp in list(local_files):
                is_zip = fp.lower().endswith(".zip")
                if not is_zip and os.path.exists(fp) and os.path.getsize(fp) >= 4:
                    with open(fp, "rb") as f:
                        is_zip = f.read(4) == b"PK\x03\x04"
                if is_zip:
                    extracted = _extract_zip(fp)
                    local_files.remove(fp)
                    local_files.extend(extracted)

        # Build markdown
        md_content = _build_markdown(detail, article_dl_dir, detail_url)

        # Save markdown locally
        date_for_name = date_str or datetime.now().strftime("%Y-%m-%d")
        folder_name = _sanitize_filename(
            "{}_{}_{}".format(date_for_name, article_id[:12], title[:40]), max_len=120)
        md_path = os.path.join(output_dir, "{}.md".format(folder_name))
        with open(md_path, "w", encoding="utf-8") as f:
            f.write(md_content)
        _safe_print("[JSJ-ZZ]   Saved ({} chars, {} attachments)".format(
            len(md_content), len(local_files)))
        sys.stdout.flush()

        # Upload to KB
        if args.kb_id:
            try:
                _upload_to_kb(md_content, local_files, args.kb_id,
                             args.tenant_id, folder_name)
            except Exception as e:
                logging.error("KB upload failed: %s", e)
                _safe_print("[JSJ-ZZ]   Upload error: {}".format(e))
                sys.stdout.flush()

        processed_ids.add(article_id)
        processed_count += 1

        # Checkpoint every batch
        if processed_count % _BATCH_SIZE == 0:
            _save_state(output_dir, {"processed_ids": list(processed_ids)})
            _safe_print("[JSJ-ZZ]   Checkpoint ({} processed)".format(processed_count))
            sys.stdout.flush()

    # ── Final state ────────────────────────────────────────────────────
    state["processed_ids"] = list(processed_ids)
    _save_state(output_dir, state)

    _safe_print("\n" + "=" * 60)
    _safe_print("[JSJ-ZZ] Crawl complete — {} new article(s)".format(processed_count))
    if stopped_early:
        _safe_print("[JSJ-ZZ] Stopped early, will resume next run")
    _safe_print("[JSJ-ZZ] Target: ALL data across {} sections".format(len(sections)))
    _safe_print("=" * 60 + "\n")
    sys.stdout.flush()
    logging.info("=== JSJ-ZhangZhou crawler finished ===")


if __name__ == "__main__":
    CONSUMER_NAME = "jsj_zhangzhou_crawler"
    init_root_logger(CONSUMER_NAME)
    main()
