#!/usr/bin/env python3
#
#  Copyright 2026 The InfiniFlow Authors. All Rights Reserved.
#
#  Licensed under the Apache License, Version 2.0 (the "License");
#  you may not use this file except in compliance with the License.
#  You may obtain a copy of the License at
#
#      http://www.apache.org/licenses/LICENSE-2.0
#
#  Unless required by applicable law or agreed to in writing, software
#  distributed under the License is distributed on an "AS IS" BASIS,
#  WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
#  See the License for the specific language governing permissions and
#  limitations under the License.
#
"""
Dedicated web crawler for fjnj.gov.cn 解读回应 (jdhy) page.

Crawls the 解读回应 page which contains 5 content modules, each showing
the 7 most recent items. Each module has a "更多..." link pointing to
a dedicated listing page with pagination via sitemanage/index.shtml.

Site characteristics
────────────────────
  • SSR HTML — all module data is in the initial page source.
  • 5 modules: 政策解读, 图说图解, H5解读, 音频解读, 回应关切.
  • Each module: 7 items on main page + 20/page on listing pages.
  • Pagination: /cms/sitemanage/index.shtml?siteId=XXXX&page=N (20/page).
  • Detail format: CMS html only (no publicInfo.shtml on this page).
  • Attachments: /cms/pages/{id}/ pdfattachments/, wordattachments/.
  • SSL — verify=False needed for .gov.cn certificates.

Usage (typically spawned by task_executor):
    python fjnj_jdhy_crawler.py \\
        --tenant-id <TENANT_ID> \\
        --target-url http://www.fjnj.gov.cn/cms/html/njxrmzf/jdhy/index.html \\
        --kb-id <KB_ID> \\
        --task-name <NAME>
"""

import argparse
import datetime
import json
import logging
import os
import random
import re
import sys
import time
import urllib.request
import zipfile
from urllib.parse import urljoin, urlparse

import requests as _requests
from bs4 import BeautifulSoup

import urllib3
urllib3.disable_warnings()

_SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
_PROJECT_ROOT = os.path.abspath(os.path.join(_SCRIPT_DIR, "..", ".."))
sys.path.insert(0, _PROJECT_ROOT)

from common import settings
from common.log_utils import init_root_logger
from common.misc_utils import get_uuid

# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------
_BASE = "http://www.fjnj.gov.cn"
_SITE_NAME = "南靖县人民政府"
_JDHY_URL = f"{_BASE}/cms/html/njxrmzf/jdhy/index.html"
_SITEMANAGE_BASE = f"{_BASE}/cms/sitemanage/index.shtml"

_MODULES = [
    {"name": "政策解读", "list_path": "/cms/html/njxrmzf/zcjd/index.html",
     "siteId": "60421384817130000"},
    {"name": "图说图解", "list_path": "/cms/html/njxrmzf/tstj/index.html",
     "siteId": "830546737142010006"},
    {"name": "H5解读", "list_path": "/cms/html/njxrmzf/h5jd/index.html",
     "siteId": "830621233957850000"},
    {"name": "音频解读", "list_path": "/cms/html/njxrmzf/ypjd/index.html",
     "siteId": "830621142905040000"},
    {"name": "回应关切", "list_path": "/cms/html/njxrmzf/hygq1/index.html",
     "siteId": "830546737204910015"},
]

_USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/120.0.0.0 Safari/537.36"
)

_HTML_HEADERS = {
    "User-Agent": _USER_AGENT,
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,image/webp,*/*;q=0.8",
    "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8",
}

# Anti-crawling delays (seconds)
_PAGE_DELAY = (1.0, 2.5)
_ARTICLE_DELAY = (0.3, 1.0)

_STATE_FILENAME = "_crawler_state.json"

_ATTACH_EXTENSIONS = {
    ".pdf", ".doc", ".docx", ".xls", ".xlsx",
    ".ppt", ".pptx", ".zip", ".rar", ".7z",
    ".txt", ".jpg", ".jpeg", ".png", ".tif", ".tiff",
}


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _safe_print(msg):
    try:
        print(msg)
    except UnicodeEncodeError:
        print(msg.encode("gbk", errors="replace").decode("gbk"))


def _request_delay(min_s, max_s):
    time.sleep(random.uniform(min_s, max_s))


def _sanitize_filename(text, max_len=150):
    if not text:
        return "untitled"
    name = re.sub(r'[\\/:*?"<>|]', "_", str(text).strip())
    name = re.sub(r'\s+', " ", name)
    name = name.strip("._ ")
    return name[:max_len] if name else "untitled"


def _is_attach_url(url):
    lower = url.lower().split("?")[0]
    return any(lower.endswith(ext) for ext in _ATTACH_EXTENSIONS)


# ---------------------------------------------------------------------------
# HTTP session
# ---------------------------------------------------------------------------

def _init_session():
    sess = _requests.Session()
    sess.headers.update(_HTML_HEADERS)
    sess.verify = False
    try:
        sess.get(_BASE, timeout=30, verify=False)
        logging.info("Session initialized")
    except Exception as e:
        logging.warning("Failed to init session: %s", e)
    return sess


# ---------------------------------------------------------------------------
# Listing: multi-source extraction (main page + listing pages + sitemanage)
# ---------------------------------------------------------------------------

def _extract_items_from_html(html, base_url, mod_name, seen_hrefs):
    """Extract detail links from any HTML chunk. Returns list of item dicts."""
    soup = BeautifulSoup(html, 'lxml')
    items = []

    for a in soup.find_all('a', href=True):
        href = a.get('href', '').strip()
        if not href or href.startswith('#') or href.startswith('javascript'):
            continue
        abs_href = urljoin(base_url, href)

        # Skip navigation/index links and non-local URLs
        if '/index.html' in abs_href or abs_href in seen_hrefs:
            continue
        if urlparse(abs_href).netloc != urlparse(_BASE).netloc:
            continue

        # Must be a CMS detail page
        if not re.search(r'/\d{4}-\d{2}-\d{2}/\d+\.html', abs_href):
            continue

        title = a.get_text(strip=True)
        if not title or len(title) < 4:
            continue

        seen_hrefs.add(abs_href)

        # Date from URL path
        date_str = ''
        m = re.search(r'/(\d{4})-(\d{2})-(\d{2})/', abs_href)
        if m:
            date_str = f"{m.group(1)}-{m.group(2)}-{m.group(3)}"

        # Date from parent element text
        if not date_str:
            parent = a.parent
            while parent and parent.name not in ('li', 'div', 'span'):
                parent = parent.parent
            if parent:
                parent_text = parent.get_text(strip=True)
                m = re.search(r'(\d{4}-\d{1,2}-\d{1,2})', parent_text)
                if m:
                    date_str = m.group(1)

        items.append({
            'title': title,
            'date': date_str,
            'href': abs_href,
            'module': mod_name,
        })

    return items


def _fetch_sitemanage_pages(sess, mod_cfg, seen_hrefs):
    """Paginate through sitemanage pages for a single module.

    Returns list of new items from page 2 onwards.
    """
    site_id = mod_cfg['siteId']
    items = []
    max_pages = 50  # safety limit

    for page in range(2, max_pages + 1):
        _request_delay(*_PAGE_DELAY)
        url = f"{_SITEMANAGE_BASE}?siteId={site_id}&page={page}"
        try:
            r = sess.get(url, timeout=30, verify=False)
            if r.status_code != 200:
                break
            r.encoding = 'utf-8'
        except Exception as e:
            logging.warning("sitemanage page %d error: %s", page, e)
            break

        new_items = _extract_items_from_html(r.text, url, mod_cfg['name'], seen_hrefs)
        if not new_items:
            break

        items.extend(new_items)
        _safe_print(f"    sitemanage page {page}: {len(new_items)} new items")

    return items


def _fetch_all_items(sess):
    """Extract all items from jdhy main page, module listing pages, and sitemanage.

    Three sources per module:
      1. Main jdhy page — 7 items per module
      2. Module listing page — 20 items
      3. Sitemanage pagination — 20 items/page, pages 2..N

    Returns list of dicts: {title, date, href, module}.
    """
    seen_hrefs = set()
    all_items = []

    # ---- Step 1: Main jdhy page ----
    _safe_print(f"Fetching: {_JDHY_URL}")
    r = sess.get(_JDHY_URL, timeout=60, verify=False)
    r.encoding = 'utf-8'
    soup = BeautifulSoup(r.text, 'lxml')

    # Find the mid-lm container which holds all module sections
    mid_lm = soup.find('div', class_='mid-lm')
    if mid_lm:
        # Match module headers to their list containers
        module_headers = mid_lm.find_all('div', class_='lm-level-title')
        module_lists = mid_lm.find_all('div', class_='mid-mj-list')

        # Build module name → index mapping from headers
        header_module_names = []
        for hdr in module_headers:
            text = hdr.get_text(strip=True).replace('更多...', '')
            for mod in _MODULES:
                if mod['name'] in text:
                    header_module_names.append(mod['name'])
                    break
            else:
                header_module_names.append(text)

        _safe_print(f"  Found {len(header_module_names)} module headers: {header_module_names}")
        _safe_print(f"  Found {len(module_lists)} module list containers")

        for i, list_div in enumerate(module_lists):
            mod_name = header_module_names[i] if i < len(header_module_names) else f'模块{i+1}'
            items = _extract_items_from_html(str(list_div), _JDHY_URL, mod_name, seen_hrefs)
            all_items.extend(items)
            _safe_print(f"  Main page [{mod_name}]: {len(items)} items")

    # Also scan remaining mid-mj-list divs not under mid-lm (fallback)
    for list_div in soup.find_all('div', class_='mid-mj-list'):
        if mid_lm and list_div in mid_lm.find_all('div', class_='mid-mj-list'):
            continue  # already handled
        items = _extract_items_from_html(str(list_div), _JDHY_URL, '其他', seen_hrefs)
        if items:
            all_items.extend(items)
            _safe_print(f"  Main page [其他-extra]: {len(items)} items")

    # ---- Step 2 & 3: Module listing pages + sitemanage pagination ----
    for mod_cfg in _MODULES:
        _safe_print(f"\n  Module: {mod_cfg['name']}")

        # Step 2: Listing page (page 1, 20 items)
        list_url = f"{_BASE}{mod_cfg['list_path']}"
        _request_delay(*_PAGE_DELAY)
        try:
            r = sess.get(list_url, timeout=30, verify=False)
            if r.status_code == 200:
                r.encoding = 'utf-8'
                list_items = _extract_items_from_html(r.text, list_url, mod_cfg['name'], seen_hrefs)
                all_items.extend(list_items)
                _safe_print(f"    Listing page: {len(list_items)} new items")
            else:
                _safe_print(f"    Listing page: HTTP {r.status_code}")
        except Exception as e:
            _safe_print(f"    Listing page error: {e}")

        # Step 3: Sitemanage pages 2..N
        sitemanage_items = _fetch_sitemanage_pages(sess, mod_cfg, seen_hrefs)
        all_items.extend(sitemanage_items)

    # Summary
    mod_counts = {}
    for it in all_items:
        mod_counts[it['module']] = mod_counts.get(it['module'], 0) + 1

    _safe_print(f"\n  Total items: {len(all_items)}")
    for mod, count in sorted(mod_counts.items()):
        _safe_print(f"    {mod}: {count}")

    logging.info("Fetched %d items from %d modules", len(all_items), len(mod_counts))
    return all_items


# ---------------------------------------------------------------------------
# Detail page parsing (CMS html only — no publicInfo.shtml on jdhy page)
# ---------------------------------------------------------------------------

def _parse_detail(html, detail_url):
    """Extract content from CMS html detail page."""
    soup = BeautifulSoup(html, 'lxml')
    result = {'title': '', 'pub_date': '', 'content_text': '', 'attachments': []}

    # Title: <title> tag, strip suffix like " - 政策解读"
    title_el = soup.find('title')
    if title_el:
        title_text = title_el.get_text(strip=True)
        # Remove " - <module_name>" suffix
        title_text = re.sub(r'\s*[-_|]\s*[^|_-]+$', '', title_text)
        result['title'] = title_text.strip()

    # Date from meta (PubDate format: "2026—02—09 09:23")
    for meta_name in ('PubDate', 'publishdate', 'articledate'):
        meta = soup.find('meta', attrs={'name': meta_name})
        if meta and meta.get('content'):
            raw = meta['content'].strip()
            raw = raw.replace('\u2014', '-').replace('\uff0d', '-')
            m = re.search(r'(\d{4})\D(\d{1,2})\D(\d{1,2})', raw)
            if m:
                result['pub_date'] = f"{m.group(1)}-{m.group(2).zfill(2)}-{m.group(3).zfill(2)}"
                break

    # Date from URL path
    if not result['pub_date']:
        m = re.search(r'/(\d{4})-(\d{2})-(\d{2})/', detail_url)
        if m:
            result['pub_date'] = f"{m.group(1)}-{m.group(2)}-{m.group(3)}"

    # Date from text (发布时间：YYYY-MM-DD)
    if not result['pub_date']:
        for pattern in [
            r'发布时[间日][：:]\s*(\d{4}-\d{1,2}-\d{1,2})',
            r'发布日期[：:]\s*(\d{4}-\d{1,2}-\d{1,2})',
            r'生成日期[：:]\s*(\d{4}-\d{1,2}-\d{1,2})',
        ]:
            m = re.search(pattern, html, re.I)
            if m:
                result['pub_date'] = m.group(1)[:10]
                break

    # Content
    for cls in ['main', 'article', 'content', 'detail', 'TRS_Editor', 'view', 'con']:
        div = soup.find('div', class_=re.compile(cls, re.I))
        if div:
            for tag in div.find_all(['script', 'style']):
                tag.decompose()
            text = div.get_text(separator='\n', strip=True)
            if len(text) > 100:
                text = re.sub(r'\n{3,}', '\n\n', text)
                result['content_text'] = text
                break

    # Fallback: largest text block
    if not result['content_text']:
        max_len = 0
        for div in soup.find_all('div'):
            if div.find_parent(['header', 'nav', 'footer', 'script', 'style']):
                continue
            text = div.get_text(strip=True)
            if 200 < len(text) < 50000 and len(text) > max_len:
                max_len = len(text)
                result['content_text'] = text

    # Attachments
    seen_urls = set()
    for a in soup.find_all('a', href=True):
        href = a.get('href', '')
        abs_url = urljoin(detail_url, href)
        title = a.get_text(strip=True)

        is_attach = _is_attach_url(abs_url)
        is_cms_path = any(k in href for k in [
            '/attached/', '/upload/', 'attachments/', 'wordattachments/',
            'pdfattachments/', 'fileattachments/', '/cms/pages/',
        ])

        if (is_attach or is_cms_path) and abs_url not in seen_urls:
            seen_urls.add(abs_url)
            att_name = title or os.path.basename(urlparse(abs_url).path)
            result['attachments'].append((att_name, abs_url))

    return result


# ---------------------------------------------------------------------------
# File download
# ---------------------------------------------------------------------------

def _download_file(sess, file_url, timeout=120):
    parsed = urlparse(file_url)
    main_parsed = urlparse(_BASE)

    if parsed.netloc == main_parsed.netloc or not parsed.netloc:
        try:
            resp = sess.get(file_url, timeout=timeout, stream=True, verify=False)
            if resp.status_code == 200 and len(resp.content) > 100:
                return resp.content
        except Exception as e:
            logging.error("Download error %s: %s", file_url, e)
        return None

    req = urllib.request.Request(file_url, headers=_HTML_HEADERS)
    try:
        resp = urllib.request.urlopen(req, timeout=timeout)
        data = resp.read()
        if len(data) > 100:
            return data
    except Exception as e:
        logging.error("Download error (external) %s: %s", file_url, e)
    return None


# ---------------------------------------------------------------------------
# ZIP extraction
# ---------------------------------------------------------------------------

def _extract_zip(zip_path, dest_dir):
    extracted = []
    try:
        with zipfile.ZipFile(zip_path, 'r') as zf:
            for name in zf.namelist():
                safe_name = re.sub(r'[\\/:*?"<>|]', "_", os.path.basename(name))
                if not safe_name:
                    continue
                dest_path = os.path.join(dest_dir, safe_name)
                if os.path.exists(dest_path):
                    continue
                with open(dest_path, 'wb') as f:
                    f.write(zf.read(name))
                extracted.append(dest_path)
                _safe_print(f"           Extracted: {safe_name}")
    except Exception as e:
        _safe_print(f"           ZIP extract error: {e}")
    return extracted


# ---------------------------------------------------------------------------
# Text extraction from attachments
# ---------------------------------------------------------------------------

def _extract_file_text(filepath):
    ext = os.path.splitext(filepath)[1].lower()
    text = ""

    try:
        if ext == '.pdf':
            import fitz
            with fitz.open(filepath) as doc:
                for page in doc:
                    text += page.get_text() + "\n"
        elif ext == '.docx':
            from docx import Document
            doc = Document(filepath)
            for para in doc.paragraphs:
                text += para.text + "\n"
        elif ext in ('.xls', '.xlsx'):
            import openpyxl
            wb = openpyxl.load_workbook(filepath, data_only=True)
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    line = '\t'.join(str(c) if c is not None else '' for c in row)
                    text += line + "\n"
                text += "\n"
        elif ext == '.txt':
            with open(filepath, 'r', encoding='utf-8', errors='replace') as f:
                text = f.read()
    except Exception as e:
        logging.warning("Failed to extract text from %s: %s", filepath, e)

    return text.strip()


# ---------------------------------------------------------------------------
# State persistence
# ---------------------------------------------------------------------------

def _load_state(output_dir):
    path = os.path.join(output_dir, _STATE_FILENAME)
    if os.path.exists(path):
        try:
            with open(path, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception as e:
            logging.warning("Failed to load state: %s", e)
    return {"processed_ids": []}


def _save_state(output_dir, state):
    os.makedirs(output_dir, exist_ok=True)
    with open(os.path.join(output_dir, _STATE_FILENAME), "w", encoding="utf-8") as f:
        json.dump(state, f, ensure_ascii=False)
    logging.info("State saved (%d IDs)", len(state.get("processed_ids", [])))


def _save_markdown(content, output_dir, article_id):
    articles_dir = os.path.join(output_dir, "articles")
    os.makedirs(articles_dir, exist_ok=True)
    safe_id = re.sub(r'[\\/:*?"<>|]', "_", article_id)
    path = os.path.join(articles_dir, f"{safe_id}.md")
    with open(path, "w", encoding="utf-8") as f:
        f.write(content)
    return path


# ---------------------------------------------------------------------------
# KB upload
# ---------------------------------------------------------------------------

def _upload_to_kb(filepath, kb_id, tenant_id, parser_id="laws"):
    from api.db.services.knowledgebase_service import KnowledgebaseService
    from api.db.services.file_service import FileService
    from api.db.services.document_service import DocumentService

    ok, kb = KnowledgebaseService.get_by_id(kb_id)
    if not ok:
        raise LookupError(f"Knowledge base {kb_id} not found")

    with open(filepath, "rb") as f:
        blob = f.read()

    class _FileObj:
        def __init__(self, filename, blob):
            self.id = get_uuid()
            self.filename = filename
            self.blob = blob

        def read(self):
            return self.blob

    file_obj = _FileObj(os.path.basename(filepath), blob)
    errs, doc_pairs = FileService.upload_document(kb, [file_obj], tenant_id)

    if errs:
        logging.warning("Upload errors: %s", errs)

    for doc, _ in doc_pairs:
        doc_id = doc["id"]
        logging.info("Document %s uploaded to KB %s", doc_id, kb_id)
        try:
            DocumentService.update_by_id(doc_id, {"parser_id": parser_id})
        except Exception as e:
            logging.error("Failed to update parser_id for %s: %s", doc_id, e)
        try:
            DocumentService.begin2parse(doc_id)
            DocumentService.run(tenant_id, doc, {})
        except Exception as e:
            logging.error("Failed to queue parsing for %s: %s", doc_id, e)
    return doc_pairs


# ---------------------------------------------------------------------------
# Markdown builder
# ---------------------------------------------------------------------------

def _build_markdown(item, detail, attachments_dir):
    module = item.get('module', '')
    lines = [
        f"# {detail.get('title', item['title'])}",
        "",
        f"**来源**: {_SITE_NAME}",
        f"**栏目**: {module}",
    ]
    if detail.get('pub_date'):
        lines.append(f"**发布日期**: {detail['pub_date']}")
    elif item.get('date'):
        lines.append(f"**发布日期**: {item['date']}")
    lines.append(f"**原文链接**: {item['href']}")

    lines.append("")
    lines.append("---")
    lines.append("")

    if detail.get('content_text'):
        lines.append(detail['content_text'])
    else:
        lines.append("(无法提取正文内容)")

    if attachments_dir and os.path.isdir(attachments_dir):
        lines.append("")
        lines.append("---")
        lines.append("")
        lines.append("## 附件")
        lines.append("")
        for fname in sorted(os.listdir(attachments_dir)):
            fpath = os.path.join(attachments_dir, fname)
            fsize = os.path.getsize(fpath)
            lines.append(f"- **{fname}** ({fsize:,} bytes)")

    return "\n".join(lines)


# ---------------------------------------------------------------------------
# Single item processing
# ---------------------------------------------------------------------------

def _process_item(sess, item, output_dir, kb_id, tenant_id):
    """Fetch detail, save markdown, upload to KB."""
    _safe_print(f"\n  [{item['module']}] {item['title'][:80]}")
    _safe_print(f"  {item['href'][:150]}")

    attachments_dir = None

    if _is_attach_url(item['href']):
        _safe_print("  Type: direct attachment")
        data = _download_file(sess, item['href'])
        if not data:
            _safe_print("  Download FAILED")
            return False

        fname = os.path.basename(urlparse(item['href']).path) or "attachment"
        safe_key = _sanitize_filename(item['href'], 100)
        attachments_dir = os.path.join(output_dir, "attachments", safe_key)
        os.makedirs(attachments_dir, exist_ok=True)
        fpath = os.path.join(attachments_dir, fname)
        with open(fpath, 'wb') as f:
            f.write(data)
        _safe_print(f"  Downloaded: {fname} ({len(data):,} bytes)")

        ext_text = _extract_file_text(fpath)
        _safe_print(f"  Extracted: {len(ext_text)} chars")

        if fname.lower().endswith('.zip'):
            _extract_zip(fpath, attachments_dir)
            for extracted_file in sorted(os.listdir(attachments_dir)):
                epath = os.path.join(attachments_dir, extracted_file)
                if os.path.isfile(epath) and epath != fpath:
                    extracted_text = _extract_file_text(epath)
                    if extracted_text:
                        ext_text += f"\n\n--- {extracted_file} ---\n\n{extracted_text}"

        detail = {
            'title': item['title'],
            'pub_date': item.get('date', ''),
            'content_text': ext_text or f"附件: {fname}",
            'attachments': [(fname, item['href'])],
        }
    else:
        try:
            r = sess.get(item['href'], timeout=60, verify=False)
            if r.status_code != 200:
                _safe_print(f"  HTTP {r.status_code} — skipped")
                return False
            r.encoding = 'utf-8'
            html = r.text
        except Exception as e:
            _safe_print(f"  Fetch error: {e}")
            return False

        detail = _parse_detail(html, item['href'])
        _safe_print(f"  Date: {detail.get('pub_date', item.get('date', ''))} | "
                    f"Content: {len(detail['content_text'])} chars")

        if detail['attachments']:
            safe_key = _sanitize_filename(item['href'], 100)
            attachments_dir = os.path.join(output_dir, "attachments", safe_key)
            os.makedirs(attachments_dir, exist_ok=True)
            for att_name, att_url in detail['attachments']:
                _safe_print(f"  Downloading: {att_name[:60]}")
                data = _download_file(sess, att_url)
                if data:
                    fname = _sanitize_filename(att_name, 100)
                    fpath = os.path.join(attachments_dir, fname)
                    with open(fpath, 'wb') as f:
                        f.write(data)
                    _safe_print(f"    OK ({len(data):,} bytes)")
                    if fname.lower().endswith('.zip'):
                        _extract_zip(fpath, attachments_dir)

    md_content = _build_markdown(item, detail, attachments_dir)
    md_path = _save_markdown(md_content, output_dir, item['href'])
    _safe_print(f"  Markdown: {md_path}")

    if kb_id:
        try:
            _upload_to_kb(md_path, kb_id, tenant_id)
        except Exception as e:
            logging.error("Upload failed: %s", e)

    return True


# ---------------------------------------------------------------------------
# Main crawl logic
# ---------------------------------------------------------------------------

def crawl(output_dir, kb_id, tenant_id, max_runtime=3300):
    start_time = time.time()

    _safe_print("=" * 60)
    _safe_print(f"FJNJ JDHY Crawler — {_SITE_NAME} 解读回应")
    _safe_print(f"Target: {_JDHY_URL}")
    _safe_print("=" * 60)

    os.makedirs(output_dir, exist_ok=True)
    state = _load_state(output_dir)
    processed_ids = set(state.get("processed_ids", []))

    sess = _init_session()

    # Fetch all items from main page + listing pages + sitemanage
    _safe_print("\n--- Parsing jdhy page + module listings + sitemanage ---")
    items = _fetch_all_items(sess)

    # Count new vs already processed
    new_items = [it for it in items if it['href'] not in processed_ids]
    _safe_print(f"\n{'='*60}")
    _safe_print(f"Total items: {len(items)} | New: {len(new_items)} | "
                f"Already processed: {len(items) - len(new_items)}")
    _safe_print(f"{'='*60}")

    total_processed = 0
    for item in new_items:
        remaining = max_runtime - (time.time() - start_time)
        grace = min(120, max_runtime * 0.05)
        if remaining < grace:
            _safe_print(f"\nTimeout approaching ({remaining:.0f}s < {grace:.0f}s) — stopping")
            break

        success = _process_item(sess, item, output_dir, kb_id, tenant_id)
        if success:
            processed_ids.add(item['href'])
            state["processed_ids"] = list(processed_ids)
            total_processed += 1
            if total_processed % 10 == 0:
                _save_state(output_dir, state)

        _request_delay(*_ARTICLE_DELAY)

    _save_state(output_dir, state)

    elapsed = time.time() - start_time
    _safe_print(f"\n{'='*60}")
    _safe_print(f"Crawl complete. {total_processed} new items in {elapsed:.0f}s")
    _safe_print(f"{'='*60}")


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(
        description="FJNJ JDHY Crawler — fjnj.gov.cn 解读回应"
    )
    parser.add_argument("--tenant-id", required=True, help="Tenant ID")
    parser.add_argument("--target-url", default=_JDHY_URL,
                        help="Target URL (default: %(default)s)")
    parser.add_argument("--kb-id", default=None, help="Knowledge base ID")
    parser.add_argument("--task-name", default="fjnj_jdhy_crawler",
                        help="Task name")
    parser.add_argument("--output-dir", default=None, help="Output directory")
    parser.add_argument("--max-runtime", type=int, default=3300,
                        help="Maximum runtime in seconds (default: 3300)")
    parser.add_argument("--project-root", default=None, help="Project root")

    args = parser.parse_args()

    if args.project_root:
        sys.path.insert(0, args.project_root)
        os.chdir(args.project_root)

    output_dir = args.output_dir or os.path.join(_SCRIPT_DIR, args.task_name)
    os.makedirs(output_dir, exist_ok=True)

    init_root_logger("fjnj_jdhy_crawler")
    logging.info("FJNJ JDHY Crawler | task=%s | output=%s", args.task_name, output_dir)

    try:
        crawl(
            output_dir=output_dir,
            kb_id=args.kb_id,
            tenant_id=args.tenant_id,
            max_runtime=args.max_runtime,
        )
    except KeyboardInterrupt:
        _safe_print("\nInterrupted by user")
        logging.info("Interrupted by user")
    except Exception as e:
        logging.exception("Fatal error: %s", e)
        _safe_print(f"\nFATAL: {e}")
        raise


if __name__ == "__main__":
    CONSUMER_NAME = "fjnj_jdhy_crawler"
    init_root_logger(CONSUMER_NAME)
    main()
