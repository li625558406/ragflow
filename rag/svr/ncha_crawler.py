#!/usr/bin/env python3
#
#  Copyright 2026 The InfiniFlow Authors. All Rights Reserved.
#
#  Licensed under the Apache License, Version 2.0 (the "License");
#  you may not use this file except in compliance with the License.
#  You may obtain a copy of the License at
#
#      http://www.apache.org/licenses/LICENSE-2.0
#
#  Unless required by applicable law or agreed to in writing, software
#  distributed under the License is distributed on an "AS IS" BASIS,
#  WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
#  See the License for the specific language governing permissions and
#  limitations under the License.
#
"""
Dedicated web crawler for www.ncha.gov.cn (国家文物局 — 政府信息公开).

Crawls three sections:
  1. /col/col2398/index.html — 政府信息公开指南 (single static page)
  2. /col/col2378/index.html — 政府信息公开制度 (article list, 11 items)
  3. /col/col2445/index.html — 高级检索 (search-based, first page 25 items)

Site characteristics
────────────────────
  • Dahan JCMS (大汉版通) — Java-based content management system.
  • Static HTML listing pages (col2378, col2398) — no AJAX needed.
  • col2445 uses a search form (POST /module/search/index.jsp) that
    returns paginated results (25 items/page, 968 pages, ~24k total).
  • Detail pages: /art/YYYY/M/D/art_{colid}_{artid}.html
    - Metadata table (发文机构, 信息分类, 发文字号, 日期)
    - Content div (.scroll_cont) with inline HTML
    - Attachments as <a> links to pdf/doc files within content
  • No WAF/captcha — standard HTTP requests with proper User-Agent.
  • Session not required — stateless pages.

Usage (typically spawned by task_executor):
    python ncha_crawler.py \
        --tenant-id <TENANT_ID> \
        --target-url http://www.ncha.gov.cn/col/col2378/index.html \
        --kb-id <KB_ID> \
        --task-name <NAME>
"""

import argparse
import datetime
import json
import logging
import os
import random
import re
import sys
import time
import zipfile
from urllib.parse import urljoin, urlparse, parse_qs

import requests as _requests
from bs4 import BeautifulSoup

import urllib3
urllib3.disable_warnings()

_SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
_PROJECT_ROOT = os.path.abspath(os.path.join(_SCRIPT_DIR, "..", ".."))
sys.path.insert(0, _PROJECT_ROOT)

from common import settings
from common.log_utils import init_root_logger
from common.misc_utils import get_uuid

# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------
_BASE = "http://www.ncha.gov.cn"
_SITE_NAME = "国家文物局"

_USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/120.0.0.0 Safari/537.36"
)

_HTML_HEADERS = {
    "User-Agent": _USER_AGENT,
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,image/webp,*/*;q=0.8",
    "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8",
}

_SEARCH_URL = f"{_BASE}/module/search/index.jsp"
_SEARCH_PAGE_SIZE = 25

# Anti-crawling delays (seconds)
_PAGE_DELAY = (1.0, 2.5)
_ITEM_DELAY = (0.3, 1.0)

_STATE_FILENAME = "_crawler_state.json"

_ATTACH_EXTENSIONS = {
    ".pdf", ".doc", ".docx", ".xls", ".xlsx",
    ".ppt", ".pptx", ".zip", ".rar", ".7z",
    ".txt", ".jpg", ".jpeg", ".png", ".tif", ".tiff",
}


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _safe_print(msg):
    try:
        print(msg)
    except UnicodeEncodeError:
        print(msg.encode("gbk", errors="replace").decode("gbk"))


def _request_delay(min_s, max_s):
    time.sleep(random.uniform(min_s, max_s))


def _sanitize_filename(text, max_len=150):
    if not text:
        return "untitled"
    name = re.sub(r'[\\/:*?"<>|]', "_", str(text).strip())
    name = re.sub(r'\s+', " ", name)
    name = name.strip("._ ")
    return name[:max_len] if name else "untitled"


def _is_attach_url(url):
    lower = url.lower().split("?")[0]
    return any(lower.endswith(ext) for ext in _ATTACH_EXTENSIONS)


def _now_str():
    return datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")


# ---------------------------------------------------------------------------
# HTTP session
# ---------------------------------------------------------------------------

def _init_session():
    sess = _requests.Session()
    sess.headers.update(_HTML_HEADERS)
    sess.verify = False
    try:
        r = sess.get(_BASE, timeout=30, verify=False)
        r.encoding = 'utf-8'
        logging.info("Session initialized (%d bytes)", len(r.text))
    except Exception as e:
        logging.warning("Failed to init session: %s", e)
    return sess


# ---------------------------------------------------------------------------
# Section 1: col2398 — 政府信息公开指南 (single static page)
# ---------------------------------------------------------------------------

def _fetch_col2398(sess):
    """Extract the single informational page content from col2398."""
    url = f"{_BASE}/col/col2398/index.html"
    _safe_print(f"  Fetching: {url}")

    try:
        r = sess.get(url, timeout=30, verify=False)
        r.encoding = 'utf-8'
    except Exception as e:
        logging.error("Failed to fetch col2398: %s", e)
        return []

    soup = BeautifulSoup(r.text, 'lxml')

    # Extract the title
    title_tag = soup.find('h1')
    title = title_tag.get_text(strip=True) if title_tag else "国家文物局政府信息公开指南"

    # Extract all paragraph content
    content_parts = []
    for p in soup.find_all('p'):
        text = p.get_text(strip=True)
        if text and len(text) > 5:
            content_parts.append(text)

    content_text = "\n\n".join(content_parts)

    return [{
        'article_id': 'col2398_guide',
        'title': title,
        'content_html': content_text,
        'source_url': url,
        'section': '政府信息公开指南',
        'is_static_page': True,
        'metadata': {},
    }]


# ---------------------------------------------------------------------------
# Section 2: col2378 — 政府信息公开制度 (article list, ~11 items)
# ---------------------------------------------------------------------------

def _fetch_col2378_items(sess):
    """Parse the article list from col2378."""
    url = f"{_BASE}/col/col2378/index.html"
    _safe_print(f"  Fetching: {url}")

    try:
        r = sess.get(url, timeout=30, verify=False)
        r.encoding = 'utf-8'
    except Exception as e:
        logging.error("Failed to fetch col2378: %s", e)
        return []

    soup = BeautifulSoup(r.text, 'lxml')

    items = []
    for a in soup.find_all('a', href=True):
        href = a.get('href', '').strip()
        title = a.get('title', '').strip() or a.get_text(strip=True)
        if '/art/' in href and 'art_2378_' in href and title:
            abs_url = urljoin(_BASE, href)
            # Extract article ID from URL: art_2378_XXXXX
            m = re.search(r'art_2378_(\d+)', href)
            article_id = m.group(1) if m else href

            items.append({
                'article_id': f'col2378_{article_id}',
                'title': title,
                'detail_url': abs_url,
                'section': '政府信息公开制度',
                'is_static_page': False,
            })

    return items


# ---------------------------------------------------------------------------
# Section 3: col2445 — 高级检索 (search form, paginated results)
# ---------------------------------------------------------------------------

def _fetch_col2445_items(sess, max_pages=1):
    """Submit the search form and parse results from col2445.

    The search form auto-submits when col2445 loads. We POST directly to the
    search endpoint with empty filters to get ALL results sorted by date desc.

    Args:
        max_pages: Maximum number of result pages to fetch (default 1 = 25 items).
    """
    _safe_print(f"  Searching via {_SEARCH_URL} (max {max_pages} page(s))...")

    all_items = []

    for page in range(1, max_pages + 1):
        # Build search params matching the form on col2445
        data = {
            'vc_name': '',
            'field_908': '',       # 信息分类
            'field_855': '',       # 索引号
            'field_856': '',       # 发文机构
            'field_857': '',       # 文章内容
            'field_858': '',       # 发文字号
            'c_createtime_start': '',
            'c_createtime_end': '',
            'i_columnid': 'style_10',
            'field': 'vc_name:1,c_createtime:3,field_855:1,field_856:12,field_857:7,field_858:1,field_908:1',
            'fullpath': '0',
            'splitflag': '',
            'currentplace': '',
            'currpage': str(page),
        }
        headers = {
            **_HTML_HEADERS,
            "Referer": f"{_BASE}/col/col2445/index.html",
            "Content-Type": "application/x-www-form-urlencoded",
        }

        try:
            r = sess.post(_SEARCH_URL, data=data, headers=headers, timeout=60, verify=False)
            r.encoding = 'utf-8'
        except Exception as e:
            logging.error("Search API error (page %d): %s", page, e)
            continue

        soup = BeautifulSoup(r.text, 'lxml')

        # Parse results: <li> elements with <a> and <b> (date)
        for li in soup.find_all('li'):
            a_tag = li.find('a', href=True)
            if not a_tag:
                continue
            href = a_tag.get('href', '').strip()
            title = a_tag.get('title', '').strip() or a_tag.get_text(strip=True)
            if not title:
                continue

            # Resolve relative URL
            abs_url = urljoin(_BASE, href) if href.startswith('../') or href.startswith('/') else href
            abs_url = urljoin(f"{_BASE}/module/search/", href)

            # Extract article ID from URL: art_{colid}_{artid}
            m = re.search(r'art_(\d+)_(\d+)', href)
            if m:
                article_id = f"{m.group(1)}_{m.group(2)}"
            else:
                article_id = re.sub(r'[^a-zA-Z0-9]', '_', href)

            if not any(it['detail_url'] == abs_url for it in all_items):
                all_items.append({
                    'article_id': f'col2445_{article_id}',
                    'title': title,
                    'detail_url': abs_url,
                    'section': '高级检索',
                    'is_static_page': False,
                })

        _safe_print(f"    Page {page}: {len(all_items)} items collected so far")

        if page < max_pages:
            _request_delay(*_PAGE_DELAY)

    return all_items


# ---------------------------------------------------------------------------
# Detail page parsing
# ---------------------------------------------------------------------------

def _parse_detail(html, detail_url):
    """Parse a detail/article page and extract metadata + content + attachments.

    Detail pages have this structure:
      - table.tska: metadata (发文机构, 信息分类, 标题, 发文字号, 日期)
      - h1: article title
      - span.dna: publish date
      - div.scroll_cont: main content with <p> tags
      - Content may contain attachment links (pdf/doc/docx)
    """
    soup = BeautifulSoup(html, 'lxml')

    # Extract metadata from table.tska
    metadata = {}
    table = soup.find('table', class_='tska')
    if table:
        rows = table.find_all('tr')
        for row in rows:
            cells = row.find_all('td')
            for i in range(0, len(cells) - 1, 2):
                key = cells[i].get_text(strip=True).rstrip('：')
                val = cells[i + 1].get_text(strip=True) if i + 1 < len(cells) else ''
                if key and val:
                    metadata[key] = val

    # Extract title from h1
    title_tag = soup.find('h1')
    title = title_tag.get_text(strip=True) if title_tag else metadata.get('标题', '')

    # Extract publish date
    date_span = soup.find('span', class_='dna')
    pub_date = date_span.get_text(strip=True) if date_span else metadata.get('发布日期', '')

    # Extract main content from div.scroll_cont
    content_div = soup.find('div', class_='scroll_cont')
    content_html = str(content_div) if content_div else ''
    content_text = _html_to_text(content_html) if content_html else ''

    # Find attachment links within the content
    attach_links = []
    if content_div:
        for a in content_div.find_all('a', href=True):
            href = a.get('href', '').strip()
            if not href or href.startswith('#') or href.startswith('javascript'):
                continue
            abs_url = urljoin(detail_url, href)
            if _is_attach_url(abs_url):
                att_name = a.get_text(strip=True) or os.path.basename(urlparse(abs_url).path)
                attach_links.append((att_name, abs_url))

    # Also scan entire page for attachment links (some may be outside content div)
    for a in soup.find_all('a', href=True):
        href = a.get('href', '').strip()
        if not href or href.startswith('#') or href.startswith('javascript'):
            continue
        abs_url = urljoin(detail_url, href)
        if _is_attach_url(abs_url) and abs_url not in [u for _, u in attach_links]:
            att_name = a.get_text(strip=True) or os.path.basename(urlparse(abs_url).path)
            attach_links.append((att_name, abs_url))

    return {
        'title': title,
        'pub_date': pub_date,
        'metadata': metadata,
        'content_text': content_text,
        'content_html': content_html,
        'attach_links': attach_links,
    }


def _html_to_text(html_content):
    """Convert HTML content to plain text."""
    if not html_content:
        return ""
    soup = BeautifulSoup(html_content, 'lxml')
    for tag in soup.find_all(['script', 'style']):
        tag.decompose()
    for tag_name in ['p', 'div', 'br', 'h1', 'h2', 'h3', 'h4', 'h5', 'h6', 'li', 'tr']:
        for tag in soup.find_all(tag_name):
            tag.insert_after(soup.new_string('\n'))
    text = soup.get_text(separator=' ', strip=True)
    text = re.sub(r' {2,}', ' ', text)
    text = re.sub(r'\n{3,}', '\n\n', text)
    text = re.sub(r'\n +', '\n', text)
    return text.strip()


# ---------------------------------------------------------------------------
# File download
# ---------------------------------------------------------------------------

def _download_file(sess, file_url, timeout=120):
    from urllib.parse import quote as _url_quote
    encoded_url = _url_quote(file_url, safe=":/?&=#%")
    try:
        resp = sess.get(encoded_url, timeout=timeout, stream=True, verify=False)
        if resp.status_code == 200 and len(resp.content) > 100:
            return resp.content
    except Exception as e:
        logging.error("Download error %s: %s", encoded_url, e)
    return None


# ---------------------------------------------------------------------------
# ZIP extraction
# ---------------------------------------------------------------------------

def _extract_zip(zip_path, dest_dir):
    extracted = []
    try:
        with zipfile.ZipFile(zip_path, 'r') as zf:
            for name in zf.namelist():
                safe_name = re.sub(r'[\\/:*?"<>|]', "_", os.path.basename(name))
                if not safe_name:
                    continue
                dest_path = os.path.join(dest_dir, safe_name)
                if os.path.exists(dest_path):
                    continue
                with open(dest_path, 'wb') as f:
                    f.write(zf.read(name))
                extracted.append(dest_path)
                _safe_print(f"           Extracted: {safe_name}")
    except Exception as e:
        _safe_print(f"           ZIP extract error: {e}")
    return extracted


# ---------------------------------------------------------------------------
# Text extraction from attachments
# ---------------------------------------------------------------------------

def _extract_file_text(filepath):
    ext = os.path.splitext(filepath)[1].lower()
    text = ""
    try:
        if ext == '.pdf':
            import fitz
            with fitz.open(filepath) as doc:
                for page in doc:
                    text += page.get_text() + "\n"
        elif ext == '.docx':
            from docx import Document
            doc = Document(filepath)
            for para in doc.paragraphs:
                text += para.text + "\n"
        elif ext in ('.xls', '.xlsx'):
            import openpyxl
            wb = openpyxl.load_workbook(filepath, data_only=True)
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    line = '\t'.join(str(c) if c is not None else '' for c in row)
                    text += line + "\n"
                text += "\n"
        elif ext == '.txt':
            with open(filepath, 'r', encoding='utf-8', errors='replace') as f:
                text = f.read()
    except Exception as e:
        logging.warning("Failed to extract text from %s: %s", filepath, e)
    return text.strip()


# ---------------------------------------------------------------------------
# State persistence
# ---------------------------------------------------------------------------

def _load_state(output_dir):
    path = os.path.join(output_dir, _STATE_FILENAME)
    if os.path.exists(path):
        try:
            with open(path, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception as e:
            logging.warning("Failed to load state: %s", e)
    return {"processed_ids": []}


def _save_state(output_dir, state):
    os.makedirs(output_dir, exist_ok=True)
    with open(os.path.join(output_dir, _STATE_FILENAME), "w", encoding="utf-8") as f:
        json.dump(state, f, ensure_ascii=False, indent=2)
    logging.info("State saved (%d IDs)", len(state.get("processed_ids", [])))


# ---------------------------------------------------------------------------
# Markdown builder
# ---------------------------------------------------------------------------

def _build_markdown(item, detail):
    lines = [
        f"# {item.get('title') or detail.get('title', '无标题')}",
        "",
        f"**来源**: {_SITE_NAME} — {item.get('section', '')}",
    ]
    if detail.get('pub_date'):
        lines.append(f"**发布日期**: {detail['pub_date']}")

    metadata = detail.get('metadata', {})
    if metadata.get('发文机构'):
        lines.append(f"**发文机构**: {metadata['发文机构']}")
    if metadata.get('发文字号'):
        lines.append(f"**发文字号**: {metadata['发文字号']}")
    if metadata.get('信息分类'):
        lines.append(f"**信息分类**: {metadata['信息分类']}")
    if metadata.get('成文日期'):
        lines.append(f"**成文日期**: {metadata['成文日期']}")

    source_url = item.get('detail_url') or item.get('source_url', '')
    if source_url:
        lines.append(f"**原文链接**: {source_url}")

    lines.append("")
    lines.append("---")
    lines.append("")

    if detail.get('content_text'):
        lines.append(detail['content_text'])
    else:
        lines.append("(无法提取正文内容)")

    if detail.get('attach_texts'):
        lines.append("")
        lines.append("---")
        lines.append("")
        lines.append("## 附件内容")
        lines.append("")
        for att_name, att_text in detail['attach_texts']:
            if att_text:
                lines.append(f"### {att_name}")
                lines.append("")
                lines.append(att_text)
                lines.append("")

    return "\n".join(lines)


def _save_markdown(content, output_dir, article_id):
    articles_dir = os.path.join(output_dir, "articles")
    os.makedirs(articles_dir, exist_ok=True)
    safe_id = re.sub(r'[\\/:*?"<>|]', "_", article_id)
    path = os.path.join(articles_dir, f"{safe_id}.md")
    with open(path, "w", encoding="utf-8") as f:
        f.write(content)
    return path


# ---------------------------------------------------------------------------
# KB upload
# ---------------------------------------------------------------------------

def _upload_to_kb(filepath, kb_id, tenant_id, parser_id="laws"):
    from api.db.services.knowledgebase_service import KnowledgebaseService
    from api.db.services.file_service import FileService
    from api.db.services.document_service import DocumentService

    ok, kb = KnowledgebaseService.get_by_id(kb_id)
    if not ok or kb is None:
        logging.warning("KB %s not found, skip upload %s", kb_id, os.path.basename(filepath))
        return []

    # KB-level dedup
    existing = set()
    try:
        for d in DocumentService.select(DocumentService.model.name).where(
            DocumentService.model.kb_id == kb_id
        ):
            existing.add(d.name)
    except Exception:
        pass
    fname = os.path.basename(filepath)
    if fname in existing:
        logging.info("Skip duplicate: %s", fname)
        return []

    with open(filepath, "rb") as f:
        blob = f.read()

    class _FileObj:
        def __init__(self, filename, blob):
            self.id = get_uuid()
            self.filename = filename
            self.blob = blob

        def read(self):
            return self.blob

    file_obj = _FileObj(fname, blob)
    errs, doc_pairs = FileService.upload_document(kb, [file_obj], tenant_id)

    if errs:
        logging.warning("Upload errors: %s", errs)

    for doc, _ in doc_pairs:
        doc_id = doc["id"]
        logging.info("Document %s uploaded to KB %s", doc_id, kb_id)
        try:
            DocumentService.update_by_id(doc_id, {"parser_id": parser_id})
        except Exception as e:
            logging.error("Failed to update parser_id for %s: %s", doc_id, e)
        try:
            from api.db.services.task_service import queue_tasks
            from api.db.services.file2document_service import File2DocumentService
            bucket, name = File2DocumentService.get_storage_address(doc_id=doc_id)
            queue_tasks(doc, bucket, name, 0)
        except Exception as e:
            logging.error("Failed to queue parsing for %s: %s", doc_id, e)
    return doc_pairs


# ---------------------------------------------------------------------------
# Single item processing
# ---------------------------------------------------------------------------

def _process_item(sess, item, output_dir, kb_id, tenant_id):
    """Process one item: fetch detail page, download attachments, save, upload."""
    article_id = item['article_id']
    is_static = item.get('is_static_page', False)

    if is_static:
        # Static page — content already in the item
        _safe_print(f"\n  [{article_id}] {item['title'][:80]}")
        detail = {
            'title': item['title'],
            'pub_date': '',
            'metadata': {},
            'content_text': item.get('content_html', ''),
            'content_html': '',
            'attach_links': [],
        }
    else:
        detail_url = item['detail_url']
        _safe_print(f"\n  [{article_id}] {item['title'][:80]}")
        _safe_print(f"  URL: {detail_url}")

        try:
            r = sess.get(detail_url, timeout=30, verify=False)
            r.encoding = 'utf-8'
        except Exception as e:
            logging.error("Failed to fetch detail %s: %s", detail_url, e)
            return False

        detail = _parse_detail(r.text, detail_url)
        _safe_print(f"  Content: {len(detail.get('content_text', ''))} chars")

    # Download attachments
    attach_texts = []
    if detail.get('attach_links'):
        safe_key = _sanitize_filename(article_id, 80)
        attachments_dir = os.path.join(output_dir, "attachments", safe_key)
        os.makedirs(attachments_dir, exist_ok=True)

        for att_name, att_url in detail['attach_links']:
            _safe_print(f"  Downloading: {att_name[:60]}")
            data = _download_file(sess, att_url)
            if data:
                fname = _sanitize_filename(att_name, 100)
                fpath = os.path.join(attachments_dir, fname)
                with open(fpath, 'wb') as f:
                    f.write(data)
                _safe_print(f"    OK ({len(data):,} bytes)")

                ext_text = _extract_file_text(fpath)
                if ext_text:
                    attach_texts.append((att_name, ext_text))
                    _safe_print(f"    Extracted: {len(ext_text)} chars")

                if fname.lower().endswith('.zip'):
                    extracted = _extract_zip(fpath, attachments_dir)
                    for ext_file in extracted:
                        if os.path.isfile(ext_file):
                            ext_text2 = _extract_file_text(ext_file)
                            if ext_text2:
                                basename = os.path.basename(ext_file)
                                attach_texts.append((basename, ext_text2))
            else:
                _safe_print(f"    Download FAILED")

    detail['attach_texts'] = attach_texts
    md_content = _build_markdown(item, detail)
    md_path = _save_markdown(md_content, output_dir, article_id)
    _safe_print(f"  Markdown: {md_path}")

    if kb_id:
        try:
            _upload_to_kb(md_path, kb_id, tenant_id)
        except Exception as e:
            logging.error("Upload failed: %s", e)

    return True


# ---------------------------------------------------------------------------
# Main crawl logic
# ---------------------------------------------------------------------------

def crawl(output_dir, kb_id, tenant_id, target_url, max_runtime=3300):
    start_time = time.time()

    # Determine which section to crawl based on URL
    if 'col2398' in target_url:
        section = 'col2398'
        section_label = '政府信息公开指南'
    elif 'col2378' in target_url:
        section = 'col2378'
        section_label = '政府信息公开制度'
    elif 'col2445' in target_url:
        section = 'col2445'
        section_label = '高级检索'
    else:
        _safe_print(f"ERROR: Unsupported target URL: {target_url}")
        _safe_print("Must be one of: col2398, col2378, col2445")
        return

    _safe_print("=" * 60)
    _safe_print(f"NCHA Crawler — {_SITE_NAME}")
    _safe_print(f"Section: {section_label} ({section})")
    _safe_print(f"Target: {target_url}")
    _safe_print(f"Start: {_now_str()}")
    _safe_print("=" * 60)

    os.makedirs(output_dir, exist_ok=True)
    state = _load_state(output_dir)
    processed_ids = set(state.get("processed_ids", []))

    sess = _init_session()

    # ── Fetch items based on section ────────────────────────────────────
    _safe_print("\n--- Fetching items ---")
    if section == 'col2398':
        all_items = _fetch_col2398(sess)
    elif section == 'col2378':
        all_items = _fetch_col2378_items(sess)
    elif section == 'col2445':
        all_items = _fetch_col2445_items(sess, max_pages=1)

    _safe_print(f"\nTotal items: {len(all_items)}")

    # ── Deduplicate ─────────────────────────────────────────────────────
    new_items = [it for it in all_items if it['article_id'] not in processed_ids]
    _safe_print(f"New: {len(new_items)} | Already processed: {len(all_items) - len(new_items)}")

    if not new_items:
        _safe_print("No new items — done.")
        return

    # ── Process each item with timeout awareness ────────────────────────
    _safe_print(f"\n{'='*60}")
    _safe_print(f"Processing {len(new_items)} items...")
    _safe_print(f"{'='*60}")

    total_processed = 0
    for idx, item in enumerate(new_items, 1):
        # ── Time-bounded check ─────────────────────────────────────────
        elapsed = time.time() - start_time
        remaining = max_runtime - elapsed
        grace = min(120, max_runtime * 0.05)
        if remaining < grace:
            _safe_print(
                f"\nTimeout approaching ({elapsed:.0f}s elapsed, "
                f"{remaining:.0f}s remaining < {grace:.0f}s grace) — "
                f"stopping. Processed {total_processed} items."
            )
            break

        _safe_print(f"\n[{idx}/{len(new_items)}] ({elapsed:.0f}s)")
        success = _process_item(sess, item, output_dir, kb_id, tenant_id)

        if success:
            processed_ids.add(item['article_id'])
            state["processed_ids"] = list(processed_ids)
            total_processed += 1
            if total_processed % 10 == 0:
                _save_state(output_dir, state)

        _request_delay(*_ITEM_DELAY)

    _save_state(output_dir, state)

    elapsed = time.time() - start_time
    _safe_print(f"\n{'='*60}")
    _safe_print(f"Crawl complete. {total_processed} items in {elapsed:.0f}s")
    _safe_print(f"Total IDs tracked: {len(processed_ids)}")
    _safe_print(f"{'='*60}")


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(
        description="NCHA Crawler — ncha.gov.cn 政府信息公开"
    )
    parser.add_argument("--tenant-id", required=True, help="Tenant ID")
    parser.add_argument("--target-url", required=True,
                        help="Target URL (col2398, col2378, or col2445)")
    parser.add_argument("--kb-id", default=None, help="Knowledge base ID")
    parser.add_argument("--task-name", default="ncha_crawler",
                        help="Task name")
    parser.add_argument("--output-dir", default=None, help="Output directory")
    parser.add_argument("--max-runtime", type=int, default=3300,
                        help="Maximum runtime in seconds (default: 3300)")
    parser.add_argument("--project-root", default=None, help="Project root")
    parser.add_argument("--llm-id", default=None, help="LLM ID (unused)")
    parser.add_argument("--llm-model", default=None, help="LLM model (unused)")
    parser.add_argument("--access-token", default=None, help="Access token (unused)")

    args = parser.parse_args()

    if args.project_root:
        sys.path.insert(0, args.project_root)
        os.chdir(args.project_root)

    output_dir = args.output_dir or os.path.join(_PROJECT_ROOT, "rag", args.task_name.strip())
    os.makedirs(output_dir, exist_ok=True)

    init_root_logger("ncha_crawler")
    logging.info("NCHA Crawler | task=%s | output=%s", args.task_name, output_dir)

    settings.init_settings()

    try:
        crawl(
            output_dir=output_dir,
            kb_id=args.kb_id,
            tenant_id=args.tenant_id,
            target_url=args.target_url,
            max_runtime=args.max_runtime,
        )
    except KeyboardInterrupt:
        _safe_print("\nInterrupted by user")
        logging.info("Interrupted by user")
    except Exception as e:
        logging.exception("Fatal error: %s", e)
        _safe_print(f"\nFATAL: {e}")
        raise


if __name__ == "__main__":
    CONSUMER_NAME = "ncha_crawler"
    init_root_logger(CONSUMER_NAME)
    main()
