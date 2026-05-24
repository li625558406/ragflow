#!/usr/bin/env python3
#
#  Copyright 2026 The InfiniFlow Authors. All Rights Reserved.
#
#  Licensed under the Apache License, Version 2.0 (the "License");
#  you may not use this file except in compliance with the License.
#  You may obtain a copy of the License at
#
#      http://www.apache.org/licenses/LICENSE-2.0
#
#  Unless required by applicable law or agreed to in writing, software
#  distributed under the License is distributed on an "AS IS" BASIS,
#  WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
#  See the License for the specific language governing permissions and
#  limitations under the License.
#
"""
Crawler for ggzyfw.fujian.gov.cn — 政策法规 (policies) section.

The site is a Vue 2.6 SPA with AES-256-CBC encrypted API responses and
portal-sign MD5 request signing.  All API calls go to /FwPortalApi/*.

Data flow
─────────
  1. **Listing**: POST /FwPortalApi/Article/PageList
     → {pageSize, pageNo, type: "92,93,94,96,97,98,922"}
     → Returns 1,047 records across 53 pages.
  2. **Detail**: POST /FwPortalApi/Article/Detail
     → {id, type}
     → Returns {ID, TITLE, SOURCES, TM, CONTENTS (HTML), ...}
  3. **Attachments**: <a> tags in CONTENTS HTML linking to .pdf/.doc/.docx/.xlsx
     files on external government domains.  Standard HTTP download.
  4. ZIP auto-extraction with content parsing (pdfplumber, python-docx, openpyxl).

Authentication
──────────────
  - portal-sign header: MD5(SECRET + sorted_key_value_concat)
  - ts (epoch ms) included in sign calculation and request body
  - AES-256-CBC decryption (PKCS7 padding) on all API responses

Checkpoint/resume: state saved every 5 pages.  Time-bounded check
(default 3300s) stops gracefully before the 3600s task-timeout window.

Usage
-----
    python ggzyfw_policies_crawler.py \
        --tenant-id <TENANT_ID> \
        --kb-id <KB_ID> \
        --task-name <NAME>

    # Optional:
        --max-articles 100       # Limit records (0=unlimited)
        --full                   # Ignore saved state, re-crawl
        --max-runtime 3300       # Max runtime before graceful stop
"""

import argparse
import base64
import hashlib
import json
import logging
import os
import random
import re
import sys
import time
import urllib.parse
import urllib.request
import zipfile
from datetime import datetime

_SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
_PROJECT_ROOT = os.path.abspath(os.path.join(_SCRIPT_DIR, "..", ".."))
sys.path.insert(0, _PROJECT_ROOT)

from common import settings
from common.log_utils import init_root_logger
from common.misc_utils import get_uuid
from bs4 import BeautifulSoup

# Crypto (for AES decryption)
try:
    from Crypto.Cipher import AES
    from Crypto.Util.Padding import unpad
    CRYPTO_AVAILABLE = True
except ImportError:
    CRYPTO_AVAILABLE = False


# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------

_SITE_ROOT = "https://ggzyfw.fujian.gov.cn"
_API_BASE = _SITE_ROOT + "/FwPortalApi"
_LISTING_PATH = "/Article/PageList"
_DETAIL_PATH = "/Article/Detail"
_SECTION_LABEL = "政策法规"

# Policy types to crawl (all categories)
_POLICY_TYPES = "92,93,94,96,97,98,922"

# AES-256-CBC keys
_AES_KEY = b"EB444973714E4A40876CE66BE45D5930"
_AES_IV = b"B5A8904209931867"
_SIGN_SECRET = "B3978D054A72A7002063637CCDF6B2E5"

_PAGE_SIZE = 20

# Checkpoint batch size (pages)
_BATCH_PAGES = 5

# Default max runtime (55 min, 5 min margin)
_MAX_RUNTIME_DEFAULT = 3300

# Anti-crawling delays
_REQUEST_DELAY_MIN = 0.5
_REQUEST_DELAY_MAX = 1.5

# State filename
_STATE_FILENAME = "_crawler_state.json"

_USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/120.0.0.0 Safari/537.36"
)

# Document extensions that we can parse
_TEXT_EXTS = {".pdf", ".docx", ".doc", ".xlsx", ".xls", ".txt"}


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _safe_print(msg):
    try:
        print(msg)
    except UnicodeEncodeError:
        print(msg.encode("gbk", errors="replace").decode("gbk"))


def _sanitize_filename(text, max_len=120):
    if not text:
        return "untitled"
    name = re.sub(r'[\\/:*?"<>|]', "_", str(text).strip())
    name = re.sub(r'\s+', "_", name).strip("._ ")
    return (name or "untitled")[:max_len]


# ---------------------------------------------------------------------------
# AES decryption
# ---------------------------------------------------------------------------

def _decrypt(encrypted_b64):
    """Decrypt AES-256-CBC encrypted base64 string from API response."""
    if not CRYPTO_AVAILABLE:
        raise RuntimeError("pycryptodome not installed (pip install pycryptodome)")
    cipher = AES.new(_AES_KEY, AES.MODE_CBC, _AES_IV)
    encrypted_bytes = base64.b64decode(encrypted_b64)
    decrypted = cipher.decrypt(encrypted_bytes)
    return unpad(decrypted, AES.block_size).decode("utf-8")


# ---------------------------------------------------------------------------
# portal-sign calculation
# ---------------------------------------------------------------------------

def _get_sign(params):
    """Calculate portal-sign header value.

    Mirrors the JS getSign():
      1. Remove empty params
      2. Sort keys case-insensitively
      3. Concat: SECRET + key1+val1 + key2+val2 + ...
      4. For object/array values, use JSON.stringify (no spaces)
      5. MD5 -> lowercase hex
    """
    clean = {}
    for k, v in params.items():
        if v == "" or v is None:
            continue
        clean[k] = v

    sorted_keys = sorted(clean.keys(), key=lambda x: x.upper())
    parts = []
    for k in sorted_keys:
        v = clean[k]
        if isinstance(v, (dict, list)):
            parts.append(k + json.dumps(v, separators=(",", ":"), ensure_ascii=False))
        else:
            parts.append(k + str(v))

    concat = _SIGN_SECRET + "".join(parts)
    return hashlib.md5(concat.encode("utf-8")).hexdigest().lower()


# ---------------------------------------------------------------------------
# API helpers
# ---------------------------------------------------------------------------

def _call_api(endpoint, params, timeout=30):
    """POST to /FwPortalApi/endpoint with signing and decryption."""
    if not CRYPTO_AVAILABLE:
        raise RuntimeError("pycryptodome not installed")

    ts = int(time.time() * 1000)
    params["ts"] = ts
    sign = _get_sign(params)

    url = _API_BASE + endpoint
    body = json.dumps(params, ensure_ascii=False).encode("utf-8")

    req = urllib.request.Request(url, data=body, method="POST")
    req.add_header("User-Agent", _USER_AGENT)
    req.add_header("Content-Type", "application/json;charset=UTF-8")
    req.add_header("Accept", "application/json, text/plain, */*")
    req.add_header("portal-sign", sign)
    req.add_header("Referer", _SITE_ROOT + "/policies/list")

    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            raw = resp.read()
    except Exception as e:
        logging.error("API call %s failed: %s", endpoint, e)
        return None

    try:
        data = json.loads(raw.decode("utf-8"))
    except Exception as e:
        logging.error("JSON parse error for %s: %s", endpoint, e)
        return None

    if data.get("State") == "200" and data.get("Data"):
        try:
            return json.loads(_decrypt(data["Data"]))
        except Exception as e:
            logging.error("Decrypt error for %s: %s", endpoint, e)
            return None

    logging.error("API error for %s: State=%s, Msg=%s",
                   endpoint, data.get("State"), data.get("Msg", ""))
    return None


# ---------------------------------------------------------------------------
# Listing fetcher
# ---------------------------------------------------------------------------

def _fetch_listing_page(page_no):
    """Fetch one page of the policies listing.

    Returns dict: {PageTotal, PageNo, Total, PageSize, Table}.
    Table rows: {ID, TITLE, TYPE, TM, CONTENTS: null}.
    """
    return _call_api(_LISTING_PATH, {
        "pageSize": _PAGE_SIZE,
        "pageNo": page_no,
        "type": _POLICY_TYPES,
        "name": "",
    })


# ---------------------------------------------------------------------------
# Detail fetcher
# ---------------------------------------------------------------------------

def _fetch_detail(article_id, article_type):
    """Fetch article detail including HTML content.

    Returns dict: {ID, TITLE, SOURCES, TM, CONTENTS, ...}.
    """
    return _call_api(_DETAIL_PATH, {
        "id": article_id,
        "type": article_type,
    })


# ---------------------------------------------------------------------------
# Attachment extraction from HTML
# ---------------------------------------------------------------------------

def _extract_attachments(html_content):
    """Parse attachment links from detail HTML content.

    Looks for <a> tags with href ending in document extensions
    (.pdf, .doc, .docx, .xls, .xlsx, .zip, .rar, .7z, .ppt, .pptx).

    Returns list[dict]: {filename, url, type}.
    """
    attachments = []
    if not html_content:
        return attachments

    soup = BeautifulSoup(html_content, "html.parser")
    for a_tag in soup.find_all("a", href=True):
        href = (a_tag.get("href") or "").strip()
        if not href or href.startswith("#") or href.startswith("javascript"):
            continue

        text = (a_tag.get_text() or "").strip()
        title_attr = a_tag.get("title", "")
        filename = title_attr or text or ""

        # Check for document extensions
        path_part = href.lower().rsplit("?", 1)[0]
        ext = os.path.splitext(path_part)[1]
        if ext not in (".pdf", ".doc", ".docx", ".xls", ".xlsx",
                        ".ppt", ".pptx", ".zip", ".rar", ".7z", ".txt"):
            # Not a file download — skip external reference links
            continue

        url_full = href if href.startswith("http") else _SITE_ROOT + href
        if not filename or len(filename) < 2:
            filename = os.path.basename(href.split("?")[0])

        attachments.append({
            "filename": filename,
            "url": url_full,
            "type": "download",
        })

    return attachments


# ---------------------------------------------------------------------------
# Attachment download + ZIP extraction
# ---------------------------------------------------------------------------

def _download_attachments(attachments, dest_dir):
    """Download all file attachments. Returns list of local paths."""
    downloaded = []
    if not attachments:
        return downloaded

    os.makedirs(dest_dir, exist_ok=True)

    for att in attachments:
        url = att["url"]
        filename = _sanitize_filename(att.get("filename", ""), max_len=120)
        if not filename or filename == "unknown":
            filename = os.path.basename(url.split("?")[0].split("#")[0])
            filename = _sanitize_filename(filename, max_len=120)

        dest_path = os.path.join(dest_dir, filename)
        if os.path.exists(dest_path) and os.path.getsize(dest_path) >= 100:
            downloaded.append(dest_path)
            continue

        # URL-encode Chinese characters in filename portion
        parts = url.rsplit("/", 1)
        if len(parts) > 1:
            url_encoded = parts[0] + "/" + urllib.parse.quote(parts[1])
        else:
            url_encoded = url

        body = _download_binary(url_encoded)
        if body and len(body) >= 100:
            with open(dest_path, "wb") as f:
                f.write(body)
            downloaded.append(dest_path)
            logging.info("Downloaded %s (%d bytes)", filename, len(body))
        else:
            logging.warning("Download failed/too small: %s (%s)", filename, url[:120])

    return downloaded


def _download_binary(url, timeout=60):
    """Download binary content; returns bytes or None."""
    req = urllib.request.Request(url)
    req.add_header("User-Agent", _USER_AGENT)
    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            return resp.read()
    except Exception as e:
        logging.error("Download failed: %s — %s", url[:120], e)
        return None


def _extract_zip(zip_path):
    """Extract ZIP file; returns list of extracted paths. ZIP is removed."""
    extracted = []
    dest_dir = os.path.dirname(zip_path)
    try:
        with zipfile.ZipFile(zip_path, "r") as zf:
            for name in zf.namelist():
                safe_name = re.sub(r'[\\/:*?"<>|]', "_", os.path.basename(name))
                if not safe_name:
                    continue
                dest_path = os.path.join(dest_dir, safe_name)
                if os.path.exists(dest_path):
                    continue
                with open(dest_path, "wb") as f:
                    f.write(zf.read(name))
                extracted.append(dest_path)
        os.remove(zip_path)
    except Exception as e:
        logging.warning("ZIP extract error for %s: %s", zip_path, e)
    return extracted


# ---------------------------------------------------------------------------
# Text extraction from binary files
# ---------------------------------------------------------------------------

def _extract_text_from_file(filepath):
    """Extract plain text from PDF / DOCX / XLSX / TXT files."""
    ext = os.path.splitext(filepath)[1].lower()
    try:
        if ext == ".txt":
            with open(filepath, "r", encoding="utf-8", errors="replace") as f:
                return f.read()
        elif ext == ".pdf":
            import pdfplumber
            parts = []
            with pdfplumber.open(filepath) as pdf:
                for pg in pdf.pages:
                    text = pg.extract_text()
                    if text:
                        parts.append(text)
            return "\n\n".join(parts)
        elif ext in (".docx", ".doc"):
            import docx
            doc = docx.Document(filepath)
            return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        elif ext in (".xlsx", ".xls"):
            import openpyxl
            wb = openpyxl.load_workbook(filepath, read_only=True)
            parts = []
            for ws in wb.worksheets:
                rows = []
                for row in ws.iter_rows(values_only=True):
                    rows.append(" | ".join(
                        str(c) if c is not None else "" for c in row))
                if rows:
                    parts.append(f"### {ws.title}\n" + "\n".join(rows))
            wb.close()
            return "\n\n".join(parts)
    except Exception as e:
        logging.warning("Text extraction failed for %s: %s", filepath, e)
    return None


# ---------------------------------------------------------------------------
# Markdown builder
# ---------------------------------------------------------------------------

def _build_markdown(detail, download_dir, source_url):
    """Build Markdown from article detail + downloaded attachments."""
    info = detail or {}
    title = info.get("TITLE", "无标题")
    pub_date = (info.get("TM") or "").strip()
    info_source = (info.get("SOURCES") or "").strip()

    lines = [
        f"# {title}",
        "",
        f"**数据来源:** 福建省公共资源交易电子公共服务平台 — {_SECTION_LABEL}",
        f"**页面地址:** {source_url}",
        f"**抓取时间:** {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
    ]
    if pub_date:
        lines.append(f"**发布时间:** {pub_date}")
    if info_source:
        lines.append(f"**信息来源:** {info_source}")
    lines.append("")

    # Main content (HTML -> plain text)
    contents = info.get("CONTENTS", "")
    if contents:
        lines.append("---")
        lines.append("")
        lines.append("## 正文")
        lines.append("")

        soup = BeautifulSoup(contents, "html.parser")
        # Remove attachment links from body text (they're listed separately)
        for a in soup.find_all("a", href=True):
            href = (a.get("href") or "").lower()
            path_part = href.rsplit("?", 1)[0]
            if any(path_part.endswith(e) for e in _TEXT_EXTS):
                a.decompose()

        text = soup.get_text(separator="\n", strip=True)
        text = re.sub(r"\n{3,}", "\n\n", text)
        lines.append(text)
        lines.append("")

    # Attachments
    attachments = info.get("_attachments", [])
    if attachments:
        lines.append("---")
        lines.append("")
        lines.append("## 附件")
        lines.append("")

        for att in attachments:
            fname = att.get("filename", "unknown")
            att_url = att.get("url", "")
            lines.append(f"- [{fname}]({att_url})")
        lines.append("")

        # Embed extracted attachment text
        if download_dir and os.path.isdir(download_dir):
            lines.append("### 附件内容")
            lines.append("")
            for att in attachments:
                fname = att.get("filename", "")
                local_path = os.path.join(download_dir,
                                           _sanitize_filename(fname, 120))
                if not os.path.exists(local_path):
                    # Try without sanitization
                    local_path = os.path.join(download_dir, fname)
                if not os.path.exists(local_path):
                    continue

                lines.append(f"#### {fname}")
                lines.append("")
                text = _extract_text_from_file(local_path)
                if text and text.strip():
                    if len(text) > 50000:
                        text = text[:50000] + "\n\n（内容过长，已截断）"
                    lines.append(text)
                else:
                    lines.append("（无法提取文本内容）")
                lines.append("")

    return "\n".join(lines)


# ---------------------------------------------------------------------------
# State management
# ---------------------------------------------------------------------------

def _load_state(output_dir):
    path = os.path.join(output_dir, _STATE_FILENAME)
    if os.path.exists(path):
        try:
            with open(path, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception as e:
            logging.warning("Failed to load crawler state: %s", e)
    return {"processed_ids": [], "completed": False, "last_page": 1}


def _save_state(output_dir, state):
    path = os.path.join(output_dir, _STATE_FILENAME)
    os.makedirs(output_dir, exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(state, f, ensure_ascii=False)


# ---------------------------------------------------------------------------
# KB upload
# ---------------------------------------------------------------------------

def _upload_to_kb(md_content, attachment_files, kb_id, tenant_id, folder_name):
    """Upload markdown + attachment files to knowledge base."""
    from api.db.services.knowledgebase_service import KnowledgebaseService
    from api.db.services.file_service import FileService
    from api.db.services.document_service import DocumentService

    ok, kb = KnowledgebaseService.get_by_id(kb_id)
    if not ok:
        raise LookupError("Knowledge base {} not found".format(kb_id))

    class _FO:
        def __init__(self, fn, b):
            self.id = get_uuid()
            self.filename = fn
            self.blob = b

        def read(self):
            return self.blob

    fo = _FO(f"{folder_name}.md", md_content.encode("utf-8"))
    errs, pairs = FileService.upload_document(kb, [fo], tenant_id)
    if errs:
        logging.warning("MD upload errors: %s", errs)
    for doc, _ in pairs:
        did = doc["id"]
        try:
            DocumentService.begin2parse(did)
            DocumentService.run(tenant_id, doc, {})
        except Exception as e:
            logging.error("Queue parse for %s: %s", did, e)

    for fp in attachment_files:
        fname = os.path.basename(fp)
        with open(fp, "rb") as f:
            blob = f.read()
        fo2 = _FO(fname, blob)
        errs2, pairs2 = FileService.upload_document(kb, [fo2], tenant_id)
        if errs2:
            logging.warning("Attachment upload errors: %s", errs2)
        for doc, _ in pairs2:
            did = doc["id"]
            try:
                DocumentService.begin2parse(did)
                DocumentService.run(tenant_id, doc, {})
            except Exception as e:
                logging.error("Queue parse for %s: %s", did, e)


# ---------------------------------------------------------------------------
# Argument parsing
# ---------------------------------------------------------------------------

def parse_args():
    p = argparse.ArgumentParser(
        description="ggzyfw.fujian.gov.cn policies crawler — 政策法规"
    )
    p.add_argument("--tenant-id", required=True)
    p.add_argument("--target-url",
                   default="https://ggzyfw.fujian.gov.cn/policies/list")
    p.add_argument("--kb-id", required=True)
    p.add_argument("--task-name", required=True)
    p.add_argument("--output-dir", default=None,
                   help="Output root directory (default: PROJECT_ROOT/rag/<task_name>)")
    p.add_argument("--full", action="store_true",
                   help="Ignore saved state and re-crawl all")
    p.add_argument("--max-runtime", type=int, default=_MAX_RUNTIME_DEFAULT,
                   help="Max runtime in seconds (default: 3300)")
    p.add_argument("--max-articles", type=int, default=0,
                   help="Max records to fetch (0 = unlimited)")
    # Legacy/unused args
    for opt in ("--section", "--max-days", "--hours",
                "--llm-id", "--llm-model", "--access-token"):
        p.add_argument(opt, default=None)
    return p.parse_args()


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    args = parse_args()

    _safe_print("\n" + "=" * 60)
    _safe_print("[GGFW-POL] 福建省公共资源交易平台 — {} crawler".format(_SECTION_LABEL))
    _safe_print("[GGFW-POL] Target: {}".format(args.target_url))
    _safe_print("[GGFW-POL] KB: {}".format(args.kb_id))
    _safe_print("[GGFW-POL] Task: {}".format(args.task_name))
    _safe_print("[GGFW-POL] Max articles: {}".format(
        args.max_articles if args.max_articles else "unlimited"))
    _safe_print("[GGFW-POL] Max runtime: {}s".format(args.max_runtime))
    _safe_print("[GGFW-POL] Policy types: {}".format(_POLICY_TYPES))
    _safe_print("=" * 60 + "\n")
    sys.stdout.flush()

    if not CRYPTO_AVAILABLE:
        _safe_print("[GGFW-POL] ERROR: pycryptodome not installed.")
        _safe_print("[GGFW-POL]   pip install pycryptodome")
        sys.stdout.flush()
        return

    settings.init_settings()
    logging.info("=== GGFW-POL crawler started ===")

    output_dir = args.output_dir or os.path.join(
        _PROJECT_ROOT, "rag", args.task_name.strip()
    )
    os.makedirs(output_dir, exist_ok=True)
    _safe_print("[GGFW-POL] Output: {}\n".format(output_dir))
    sys.stdout.flush()

    # ── State ──────────────────────────────────────────────────────────
    state = _load_state(output_dir) if not args.full else {
        "processed_ids": [], "completed": False, "last_page": 1,
    }
    processed_ids = set(state.get("processed_ids", []))

    if state.get("completed"):
        _safe_print("[GGFW-POL] Already completed, nothing to do.")
        sys.stdout.flush()
        return

    _safe_print("[GGFW-POL] Already processed: {} record(s)".format(
        len(processed_ids)))
    sys.stdout.flush()

    crawl_start = time.time()

    # ── Step 1: Get total count ────────────────────────────────────────
    _safe_print("[GGFW-POL] Step 1/3: Fetching first page to get total...")
    sys.stdout.flush()

    first_page = _fetch_listing_page(1)
    if not first_page:
        _safe_print("[GGFW-POL] Failed to fetch listing, exiting.")
        sys.stdout.flush()
        return

    total_records = first_page.get("Total", 0)
    total_pages = first_page.get("PageTotal", 0)
    _safe_print("[GGFW-POL] Total: {} records across {} pages".format(
        total_records, total_pages))
    sys.stdout.flush()

    if total_records == 0:
        _safe_print("[GGFW-POL] No data, exiting.")
        sys.stdout.flush()
        return

    max_records = args.max_articles if args.max_articles else total_records

    # ── Step 2: Paginate through listing ───────────────────────────────
    _safe_print("\n[GGFW-POL] Step 2/3: Paginating through listing...")
    sys.stdout.flush()

    start_page = state.get("last_page", 1)
    if start_page > 1:
        _safe_print("[GGFW-POL] Resuming from page {}".format(start_page))
        sys.stdout.flush()

    all_rows = []
    stopped_early = False

    for page_no in range(start_page, total_pages + 1):
        # ── Time-bounded check ─────────────────────────────────────
        elapsed = time.time() - crawl_start
        if elapsed > args.max_runtime - 300:
            _safe_print(
                "\n[GGFW-POL] Runtime {:.0f}s approaching limit ({}s), "
                "stopping pagination. {} rows collected. "
                "Next run resumes from page {}.".format(
                    elapsed, args.max_runtime, len(all_rows), page_no))
            sys.stdout.flush()
            stopped_early = True
            break

        if page_no == 1:
            rows = first_page.get("Table", [])
        else:
            page_data = _fetch_listing_page(page_no)
            if not page_data:
                logging.warning("Failed page %d, skipping", page_no)
                continue
            rows = page_data.get("Table", [])

        all_rows.extend(rows)

        if page_no % 10 == 0 or page_no == total_pages:
            _safe_print("[GGFW-POL]   Page {} of {} ({} rows)".format(
                page_no, total_pages, len(all_rows)))
            sys.stdout.flush()

        # Checkpoint every batch of pages
        if page_no % _BATCH_PAGES == 0:
            state["last_page"] = page_no + 1
            state["processed_ids"] = list(processed_ids)
            _save_state(output_dir, state)

        if args.max_articles and len(all_rows) >= max_records:
            all_rows = all_rows[:max_records]
            break

        time.sleep(random.uniform(_REQUEST_DELAY_MIN, _REQUEST_DELAY_MAX))

    _safe_print("[GGFW-POL] Collected {} rows total.\n".format(len(all_rows)))
    sys.stdout.flush()

    if not all_rows:
        _safe_print("[GGFW-POL] No rows, exiting.")
        sys.stdout.flush()
        return

    # ── Filter already-processed ───────────────────────────────────────
    new_rows = [r for r in all_rows if str(r.get("ID", "")) not in processed_ids]
    skipped = len(all_rows) - len(new_rows)
    if skipped:
        _safe_print("[GGFW-POL] {} already processed, {} new".format(
            skipped, len(new_rows)))
        sys.stdout.flush()

    if not new_rows:
        _safe_print("[GGFW-POL] All rows already processed.")
        sys.stdout.flush()
        if not stopped_early:
            state["completed"] = True
            _save_state(output_dir, state)
        return

    # ── Step 3: Process each record ────────────────────────────────────
    _safe_print("\n[GGFW-POL] Step 3/3: Processing {} new record(s)...\n".format(
        len(new_rows)))
    sys.stdout.flush()

    processed_count = 0
    downloads_dir = os.path.join(output_dir, "downloads")

    for idx, row in enumerate(new_rows, 1):
        # ── Time-bounded check ─────────────────────────────────────────
        elapsed = time.time() - crawl_start
        if elapsed > args.max_runtime - 120:
            _safe_print(
                "\n[GGFW-POL] Runtime {:.0f}s approaching limit ({}s), "
                "stopping gracefully. {} processed. "
                "Next run will resume.".format(
                    elapsed, args.max_runtime, processed_count))
            sys.stdout.flush()
            stopped_early = True
            break

        article_id = str(row.get("ID", ""))
        article_type = str(row.get("TYPE", ""))
        title = (row.get("TITLE") or "").strip()
        pub_date = (row.get("TM") or "").strip()

        _safe_print("[GGFW-POL] [{}/{}] {}...".format(
            idx, len(new_rows), title[:50]))
        sys.stdout.flush()

        # Fetch detail
        detail = _fetch_detail(article_id, article_type)
        if not detail or not isinstance(detail, dict):
            logging.warning("Failed detail for ID=%s", article_id)
            processed_ids.add(article_id)
            continue

        # Extract attachments from content HTML
        attachments = _extract_attachments(detail.get("CONTENTS", ""))
        detail["_attachments"] = attachments

        # Download attachments
        local_files = []
        if attachments:
            art_dl_dir = os.path.join(downloads_dir,
                                       _sanitize_filename(article_id, 40))
            local_files = _download_attachments(attachments, art_dl_dir)

            # Extract ZIPs
            for fp in list(local_files):
                is_zip = fp.lower().endswith(".zip")
                if not is_zip and os.path.exists(fp) and os.path.getsize(fp) >= 4:
                    with open(fp, "rb") as f:
                        is_zip = f.read(4) == b"PK\x03\x04"
                if is_zip:
                    extracted = _extract_zip(fp)
                    local_files.remove(fp)
                    local_files.extend(extracted)

        # Build markdown
        source_url = "{}/index/newDetail?id={}".format(_SITE_ROOT, article_id)
        md_content = _build_markdown(
            detail,
            os.path.join(downloads_dir, _sanitize_filename(article_id, 40))
            if attachments else "",
            source_url,
        )

        # Save markdown locally
        folder_name = _sanitize_filename(
            "{}_{}_{}".format(pub_date[:10] if pub_date else "nodate",
                              article_id, title[:40]),
            max_len=120)
        md_path = os.path.join(output_dir, f"{folder_name}.md")
        with open(md_path, "w", encoding="utf-8") as f:
            f.write(md_content)
        _safe_print("[GGFW-POL]   Saved ({} chars, {} attachments)".format(
            len(md_content), len(local_files)))
        sys.stdout.flush()

        # Upload to KB
        if args.kb_id:
            try:
                _upload_to_kb(md_content, local_files, args.kb_id,
                             args.tenant_id, folder_name)
                _safe_print("[GGFW-POL]   Uploaded to KB {}".format(args.kb_id))
                sys.stdout.flush()
            except Exception as e:
                logging.error("KB upload failed: %s", e)
                _safe_print("[GGFW-POL]   Upload error: {}".format(e))
                sys.stdout.flush()

        processed_ids.add(article_id)
        processed_count += 1

        # Checkpoint every batch
        if processed_count % (_BATCH_PAGES * _PAGE_SIZE) == 0:
            state["last_page"] = state.get("last_page", 1)
            state["processed_ids"] = list(processed_ids)
            _save_state(output_dir, state)
            _safe_print("[GGFW-POL]   Checkpoint ({} processed)".format(
                processed_count))
            sys.stdout.flush()

        # Anti-crawling delay
        time.sleep(random.uniform(_REQUEST_DELAY_MIN, _REQUEST_DELAY_MAX))

    # ── Final state ────────────────────────────────────────────────────
    if not stopped_early:
        state["completed"] = True
        state["last_page"] = 1
    state["processed_ids"] = list(processed_ids)
    _save_state(output_dir, state)

    _safe_print("\n" + "=" * 60)
    _safe_print("[GGFW-POL] Crawl complete — {} new record(s)".format(
        processed_count))
    if stopped_early:
        _safe_print("[GGFW-POL] Stopped early, will resume next run")
    _safe_print("=" * 60 + "\n")
    sys.stdout.flush()
    logging.info("=== GGFW-POL crawler finished ===")


if __name__ == "__main__":
    CONSUMER_NAME = "ggzyfw_policies_crawler"
    init_root_logger(CONSUMER_NAME)
    main()
