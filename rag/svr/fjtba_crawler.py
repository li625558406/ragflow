#!/usr/bin/env python3
#
#  Copyright 2026 The InfiniFlow Authors. All Rights Reserved.
#
#  Licensed under the Apache License, Version 2.0 (the "License");
#  you may not use this file except in compliance with the License.
#  You may obtain a copy of the License at
#
#      http://www.apache.org/licenses/LICENSE-2.0
#
#  Unless required by applicable law or agreed to in writing, software
#  distributed under the License is distributed on an "AS IS" BASIS,
#  WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
#  See the License for the specific language governing permissions and
#  limitations under the License.
#
"""
Crawler for www.fjtba.com (福建省招标投标协会) — 政策法规.

Target: http://www.fjtba.com/PortalPage/ISDInfo.aspx?type=2
Section: 政策法规 (type=2), 632 articles across 64 pages.

Site characteristics
────────────────────
ASP.NET WebForms site with AJAX-based data loading.  All content is loaded
via POST requests to ``/PortalPage/AjaxHandler/DataHandler.ashx``.

API endpoints (form-encoded POST)
─────────────────────────────────
Listing:
    OPtype=GetArticleList  type=2  pageNo=N  pageSize=10
    Returns ``{"count": 632, "ds": {"ds": [...]}}``

    Each article::
        RN, ID, TITLE, TYPE, TM (ISO date), IMG_AID, CONTENTS (full HTML)

The listing API already returns the full CONTENTS HTML (up to 200KB+),
so we do NOT need to call the detail API for each article.

Data flow
─────────
  1. POST listing API page by page → extract article IDs, titles, dates, content
  2. Dedup by article ID via ``_crawler_state.json``
  3. For each new article → parse HTML content → extract text + attachment links
  4. Download attachments if any (PDF/DOC/XLSX/ZIP) → extract text
  5. Build markdown → save locally → upload to KB

Usage
─────
    python fjtba_crawler.py \\
        --tenant-id <TENANT_ID> \\
        --kb-id <KB_ID> \\
        --task-name <NAME>

    # Optional:
        --max-runtime 3300    # Max runtime before graceful stop
        --full                # Ignore saved state, re-crawl
"""

import argparse
import hashlib
import json
import logging
import os
import random
import re
import sys
import time
import urllib.parse
import urllib.request
import zipfile
from datetime import datetime

_SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
_PROJECT_ROOT = os.path.abspath(os.path.join(_SCRIPT_DIR, "..", ".."))
sys.path.insert(0, _PROJECT_ROOT)

from common import settings
from common.log_utils import init_root_logger
from common.misc_utils import get_uuid
from bs4 import BeautifulSoup


# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------

_SITE_ROOT = "http://www.fjtba.com"
_AJAX_URL = _SITE_ROOT + "/PortalPage/AjaxHandler/DataHandler.ashx"
_LISTING_URL = _SITE_ROOT + "/PortalPage/ISDInfo.aspx?type=2"
_SECTION_TYPE = "2"
_SECTION_LABEL = "政策法规"

_PAGE_SIZE = 10
_BATCH_SIZE = 5
_MAX_RUNTIME_DEFAULT = 3300
_REQUEST_DELAY_MIN = 0.2
_REQUEST_DELAY_MAX = 0.8
_STATE_FILENAME = "_crawler_state.json"
_MAX_PAGES = 100

_USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/120.0.0.0 Safari/537.36"
)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _safe_print(msg):
    try:
        print(msg)
    except UnicodeEncodeError:
        print(msg.encode("gbk", errors="replace").decode("gbk"))


def _sanitize_filename(text, max_len=120):
    if not text:
        return "untitled"
    name = re.sub(r'[\\/:*?"<>|]', "_", str(text).strip())
    name = re.sub(r'\s+', "_", name).strip("._ ")
    return (name or "untitled")[:max_len]


def _resolve_url(href, base_url):
    if href.startswith("http://") or href.startswith("https://"):
        return href
    if href.startswith("//"):
        return "https:" + href
    return urllib.parse.urljoin(base_url, href)


def _normalize_pubdate(text):
    """Normalize date from ISO format or various formats to YYYY-MM-DD."""
    if not text:
        return ""
    text = text.strip()
    text = text.replace("\u2014", "-").replace("\uff0d", "-").replace("/", "-")
    m = re.search(r'(\d{4})-(\d{2})-(\d{2})', text)
    if m:
        return "{}-{}-{}".format(m.group(1), m.group(2), m.group(3))
    return ""


# ---------------------------------------------------------------------------
# HTTP helpers
# ---------------------------------------------------------------------------

def _http_post(url, data_dict, referer=None, timeout=30):
    """POST form-encoded data to a URL, return response bytes."""
    data_bytes = urllib.parse.urlencode(data_dict).encode("utf-8")
    req = urllib.request.Request(url, data=data_bytes)
    req.add_header("User-Agent", _USER_AGENT)
    req.add_header("Content-Type", "application/x-www-form-urlencoded")
    req.add_header("Accept", "application/json, text/javascript, */*; q=0.01")
    req.add_header("Accept-Language", "zh-CN,zh;q=0.9,en;q=0.8")
    req.add_header("X-Requested-With", "XMLHttpRequest")
    if referer:
        req.add_header("Referer", referer)
    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            return resp.read()
    except Exception as e:
        logging.error("POST %s failed: %s", url, e)
        return None


def _download_binary(url, referer=None, timeout=60):
    req = urllib.request.Request(url)
    req.add_header("User-Agent", _USER_AGENT)
    if referer:
        req.add_header("Referer", referer)
    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            return resp.read()
    except Exception as e:
        logging.error("Download failed: %s — %s", url[:120], e)
        return None


# ---------------------------------------------------------------------------
# Listing API
# ---------------------------------------------------------------------------

def _fetch_listing_page(page_no, referer=None):
    """POST the listing API for one page.  Returns (articles_list, total_count)."""
    params = {
        "OPtype": "GetArticleList",
        "type": _SECTION_TYPE,
        "pageNo": str(page_no),
        "pageSize": str(_PAGE_SIZE),
        "r": str(random.random()),
    }
    data_bytes = _http_post(_AJAX_URL, params, referer=referer)
    if not data_bytes:
        return [], 0

    text = data_bytes.decode("utf-8", errors="replace")
    try:
        parsed = json.loads(text)
    except json.JSONDecodeError:
        logging.error("Listing page %d: invalid JSON", page_no)
        return [], 0

    total = int(parsed.get("count", 0))
    articles = parsed.get("ds", {})
    if isinstance(articles, dict):
        articles = articles.get("ds", [])
    if not isinstance(articles, list):
        articles = []

    return articles, total


def _crawl_all_listings(seen_ids):
    """Iterate all listing pages and return deduped articles."""
    all_articles = []

    time.sleep(random.uniform(_REQUEST_DELAY_MIN, _REQUEST_DELAY_MAX))
    first_page, total = _fetch_listing_page(1, referer=_LISTING_URL)

    if not first_page and total == 0:
        _safe_print("[FJTBA] Empty first page, aborting.")
        return all_articles

    total_pages = (total + _PAGE_SIZE - 1) // _PAGE_SIZE if total else 1
    total_pages = min(total_pages, _MAX_PAGES)
    _safe_print("[FJTBA] Total: {} article(s), {} page(s)".format(total, total_pages))
    sys.stdout.flush()

    for item in first_page:
        aid = str(item.get("ID", ""))
        if aid and aid not in seen_ids:
            seen_ids.add(aid)
            all_articles.append({
                "article_id": aid,
                "title": (item.get("TITLE") or "").strip(),
                "tm": (item.get("TM") or "").strip(),
                "content_html": (item.get("CONTENTS") or "").strip(),
            })

    for page_no in range(2, total_pages + 1):
        time.sleep(random.uniform(_REQUEST_DELAY_MIN, _REQUEST_DELAY_MAX))
        articles, _ = _fetch_listing_page(page_no, referer=_LISTING_URL)
        if not articles:
            _safe_print("[FJTBA] Page {} returned no articles, stopping.".format(page_no))
            break
        for item in articles:
            aid = str(item.get("ID", ""))
            if aid and aid not in seen_ids:
                seen_ids.add(aid)
                all_articles.append({
                    "article_id": aid,
                    "title": (item.get("TITLE") or "").strip(),
                    "tm": (item.get("TM") or "").strip(),
                    "content_html": (item.get("CONTENTS") or "").strip(),
                })

    return all_articles


# ---------------------------------------------------------------------------
# Content parsing (from CONTENTS HTML)
# ---------------------------------------------------------------------------

def _parse_content_html(html, page_url):
    """Extract text and attachment links from the CONTENTS HTML.

    Returns (text, attachment_list).
    """
    if not html:
        return "", []

    soup = BeautifulSoup(html, "html.parser")

    # Remove script/style tags
    for tag in soup.find_all(["script", "style"]):
        tag.decompose()

    text = soup.get_text(separator="\n", strip=True)
    text = re.sub(r"\n{3,}", "\n\n", text)

    # Find attachment links in the HTML
    attachments = []
    seen_urls = set()
    for a_tag in soup.find_all("a", href=True):
        href = a_tag["href"].strip()
        if not href or href.startswith("javascript:") or href.startswith("#"):
            continue
        link_text = a_tag.get_text(strip=True)
        href_lower = href.lower()

        # Check for attachment indicators
        is_attachment = False
        if re.search(r'\.(pdf|doc|docx|xls|xlsx|zip|rar|7z)(\?|$)', href_lower):
            is_attachment = True
        elif link_text and any(kw in link_text for kw in ("附件", "下载", "attachment", "download")):
            is_attachment = True
        elif any(kw in href_lower for kw in ("/upload/", "/Upload/", "/attachments/", "/files/")):
            is_attachment = True

        if is_attachment and href not in seen_urls:
            seen_urls.add(href)
            fname = link_text if link_text else os.path.basename(urllib.parse.urlparse(href).path)
            if not fname:
                fname = "attachment"
            attachments.append({
                "filename": fname,
                "url": _resolve_url(href, page_url),
            })

    return text, attachments


# ---------------------------------------------------------------------------
# Attachment download + processing
# ---------------------------------------------------------------------------

def _download_attachments(attachments, download_dir, referer=None):
    os.makedirs(download_dir, exist_ok=True)
    local_files = []

    for att in attachments:
        url = att.get("url", "")
        if not url:
            continue
        fname = _sanitize_filename(att.get("filename", "attachment"), max_len=120)
        ext = os.path.splitext(urllib.parse.urlparse(url).path)[1].lower()
        if ext and re.match(r'\.(pdf|doc|docx|xls|xlsx|zip|rar|7z|txt)$', ext):
            if not fname.lower().endswith(ext):
                fname += ext

        filepath = os.path.join(download_dir, fname)
        if os.path.exists(filepath) and os.path.getsize(filepath) > 0:
            local_files.append(filepath)
            continue

        data = _download_binary(url, referer=referer)
        if data:
            if data[:100].strip().startswith(b"<!DOCTYPE") or data[:100].strip().startswith(b"<html"):
                logging.warning("Attachment %s returned HTML, skipping.", url[:120])
                continue
            with open(filepath, "wb") as f:
                f.write(data)
            local_files.append(filepath)
            time.sleep(random.uniform(0.1, 0.3))

    return local_files


def _extract_zip(filepath):
    extracted = []
    extract_dir = os.path.splitext(filepath)[0] + "_extracted"
    os.makedirs(extract_dir, exist_ok=True)
    try:
        with zipfile.ZipFile(filepath, "r") as zf:
            for name in zf.namelist():
                safe_name = _sanitize_filename(name, max_len=120)
                out_path = os.path.join(extract_dir, safe_name)
                os.makedirs(os.path.dirname(out_path) or extract_dir, exist_ok=True)
                with zf.open(name) as src, open(out_path, "wb") as dst:
                    dst.write(src.read())
                extracted.append(out_path)
    except Exception as e:
        logging.warning("ZIP extraction failed for %s: %s", filepath, e)
    return extracted


# ---------------------------------------------------------------------------
# Text extraction from binary files
# ---------------------------------------------------------------------------

def _extract_text_from_file(filepath):
    ext = os.path.splitext(filepath)[1].lower()
    try:
        if ext == ".txt":
            with open(filepath, "r", encoding="utf-8", errors="replace") as f:
                return f.read()
        elif ext == ".pdf":
            import pdfplumber
            parts = []
            with pdfplumber.open(filepath) as pdf:
                for pg in pdf.pages:
                    t = pg.extract_text()
                    if t:
                        parts.append(t)
            return "\n\n".join(parts)
        elif ext in (".docx", ".doc"):
            import docx
            doc = docx.Document(filepath)
            return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        elif ext in (".xlsx", ".xls"):
            import openpyxl
            wb = openpyxl.load_workbook(filepath, read_only=True)
            parts = []
            for ws in wb.worksheets:
                rows = []
                for row in ws.iter_rows(values_only=True):
                    rows.append(" | ".join(str(c) if c is not None else "" for c in row))
                if rows:
                    parts.append("### {}\n".format(ws.title) + "\n".join(rows))
            wb.close()
            return "\n\n".join(parts)
    except Exception as e:
        logging.warning("Text extraction failed for %s: %s", filepath, e)
    return None


# ---------------------------------------------------------------------------
# Markdown builder
# ---------------------------------------------------------------------------

def _build_markdown(title, pub_date, content_text, attachments, download_dir, detail_url):
    lines = [
        "# {}".format(title or "无标题"),
        "",
        "**数据来源:** 福建省招标投标协会 — {}".format(_SECTION_LABEL),
        "**页面地址:** {}".format(detail_url),
        "**抓取时间:** {}".format(datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
    ]
    if pub_date:
        lines.append("**发布时间:** {}".format(pub_date))
    lines.append("")

    if content_text:
        lines.append("---")
        lines.append("")
        lines.append("## 正文")
        lines.append("")
        content_clean = re.sub(r"\n{3,}", "\n\n", content_text)
        lines.append(content_clean)
        lines.append("")

    if attachments:
        lines.append("---")
        lines.append("")
        lines.append("## 附件")
        lines.append("")
        for att in attachments:
            fname = att.get("filename", "unknown")
            att_url = att.get("url", "")
            lines.append("- [{}]({})".format(fname, att_url))
        lines.append("")

        if download_dir and os.path.isdir(download_dir):
            lines.append("### 附件内容")
            lines.append("")
            for att in attachments:
                fname = att.get("filename", "")
                local_path = os.path.join(download_dir, fname)
                safe_name = _sanitize_filename(fname, max_len=120)
                alt_path = os.path.join(download_dir, safe_name)
                if not os.path.exists(local_path) and os.path.exists(alt_path):
                    local_path = alt_path
                if not os.path.exists(local_path):
                    for root, _, files in os.walk(download_dir):
                        for fn in files:
                            if fn == safe_name or fn == fname:
                                local_path = os.path.join(root, fn)
                                break
                if not os.path.exists(local_path):
                    continue

                lines.append("#### {}".format(fname))
                lines.append("")
                extracted_text = _extract_text_from_file(local_path)
                if extracted_text and extracted_text.strip():
                    if len(extracted_text) > 50000:
                        extracted_text = extracted_text[:50000] + "\n\n（内容过长，已截断）"
                    lines.append(extracted_text)
                else:
                    lines.append("（无法提取文本内容）")
                lines.append("")

    return "\n".join(lines)


# ---------------------------------------------------------------------------
# State management
# ---------------------------------------------------------------------------

def _load_state(output_dir):
    path = os.path.join(output_dir, _STATE_FILENAME)
    if os.path.exists(path):
        try:
            with open(path, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception as e:
            logging.warning("Failed to load crawler state: %s", e)
    return {"processed_ids": []}


def _save_state(output_dir, state):
    path = os.path.join(output_dir, _STATE_FILENAME)
    os.makedirs(output_dir, exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(state, f, ensure_ascii=False)


# ---------------------------------------------------------------------------
# KB upload
# ---------------------------------------------------------------------------

def _upload_to_kb(md_content, attachment_files, kb_id, tenant_id, folder_name):
    from api.db.services.knowledgebase_service import KnowledgebaseService
    from api.db.services.file_service import FileService
    from api.db.services.document_service import DocumentService

    ok, kb = KnowledgebaseService.get_by_id(kb_id)
    if not ok:
        raise LookupError("Knowledge base {} not found".format(kb_id))

    class _FO:
        def __init__(self, fn, b):
            self.id = get_uuid()
            self.filename = fn
            self.blob = b

        def read(self):
            return self.blob

    fo = _FO("{}.md".format(folder_name), md_content.encode("utf-8"))
    errs, pairs = FileService.upload_document(kb, [fo], tenant_id)
    if errs:
        logging.warning("MD upload errors: %s", errs)
    for doc, _ in pairs:
        did = doc["id"]
        try:
            DocumentService.begin2parse(did)
            DocumentService.run(tenant_id, doc, {})
        except Exception as e:
            logging.error("Queue parse for %s: %s", did, e)

    for fp in attachment_files:
        fname = os.path.basename(fp)
        with open(fp, "rb") as f:
            blob = f.read()
        fo2 = _FO(fname, blob)
        errs2, pairs2 = FileService.upload_document(kb, [fo2], tenant_id)
        if errs2:
            logging.warning("Attachment upload errors: %s", errs2)
        for doc, _ in pairs2:
            did = doc["id"]
            try:
                DocumentService.begin2parse(did)
                DocumentService.run(tenant_id, doc, {})
            except Exception as e:
                logging.error("Queue parse for %s: %s", did, e)


# ---------------------------------------------------------------------------
# Argument parsing
# ---------------------------------------------------------------------------

def parse_args():
    p = argparse.ArgumentParser(
        description="fjtba_crawler — 福建省招标投标协会 政策法规"
    )
    p.add_argument("--tenant-id", required=True)
    p.add_argument("--kb-id", required=True)
    p.add_argument("--task-name", required=True)
    p.add_argument("--output-dir", default=None,
                   help="Output root directory (default: PROJECT_ROOT/rag/<task_name>)")
    p.add_argument("--full", action="store_true",
                   help="Ignore saved state and re-crawl all")
    p.add_argument("--max-runtime", type=int, default=_MAX_RUNTIME_DEFAULT,
                   help="Max runtime in seconds (default: 3300)")
    for opt in ("--max-days", "--hours", "--max-articles",
                "--llm-id", "--llm-model", "--access-token", "--target-url"):
        p.add_argument(opt, default=None)
    return p.parse_args()


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    args = parse_args()

    _safe_print("\n" + "=" * 60)
    _safe_print("[FJTBA] 福建省招标投标协会 — {} crawler".format(_SECTION_LABEL))
    _safe_print("[FJTBA] KB: {}".format(args.kb_id))
    _safe_print("[FJTBA] Task: {}".format(args.task_name))
    _safe_print("[FJTBA] Max runtime: {}s".format(args.max_runtime))
    _safe_print("[FJTBA] Target: {}".format(_LISTING_URL))
    _safe_print("=" * 60 + "\n")
    sys.stdout.flush()

    settings.init_settings()
    logging.info("=== FJTBA crawler started ===")

    output_dir = args.output_dir or os.path.join(
        _PROJECT_ROOT, "rag", args.task_name.strip()
    )
    os.makedirs(output_dir, exist_ok=True)
    _safe_print("[FJTBA] Output: {}\n".format(output_dir))
    sys.stdout.flush()

    # ── State ──────────────────────────────────────────────────────────
    state = _load_state(output_dir) if not args.full else {"processed_ids": []}
    processed_ids = set(state.get("processed_ids", []))

    _safe_print("[FJTBA] Already processed: {} article(s)".format(len(processed_ids)))
    sys.stdout.flush()

    crawl_start = time.time()

    # ── Step 1: Crawl listing via AJAX API ─────────────────────────────
    _safe_print("\n[FJTBA] Step 1/3: Crawling listing via AJAX API...")
    sys.stdout.flush()

    seen_ids = set()
    all_articles = _crawl_all_listings(seen_ids)

    _safe_print("[FJTBA]   Total: {} article(s) from listing".format(len(all_articles)))
    sys.stdout.flush()

    if not all_articles:
        _safe_print("[FJTBA] No articles found, done.")
        sys.stdout.flush()
        return

    # Filter already-processed
    new_articles = [
        a for a in all_articles
        if a.get("article_id") and a["article_id"] not in processed_ids
    ]
    skipped = len(all_articles) - len(new_articles)
    if skipped:
        _safe_print("[FJTBA] {} already processed, {} new".format(skipped, len(new_articles)))
        sys.stdout.flush()

    if not new_articles:
        _safe_print("[FJTBA] All available articles already processed.")
        sys.stdout.flush()
        return

    # ── Step 2: Process each article ───────────────────────────────────
    _safe_print("\n[FJTBA] Step 2/3: Processing {} article(s)...\n".format(len(new_articles)))
    sys.stdout.flush()

    processed_count = 0
    stopped_early = False
    downloads_dir = os.path.join(output_dir, "downloads")

    for i, article in enumerate(new_articles, 1):
        elapsed = time.time() - crawl_start
        remaining = args.max_runtime - elapsed
        if remaining < 120:
            _safe_print(
                "\n[FJTBA] Runtime {:.0f}s, {:.0f}s remaining (limit {}s), "
                "stopping gracefully. {} processed. "
                "Next run will resume.".format(elapsed, remaining, args.max_runtime, processed_count))
            sys.stdout.flush()
            stopped_early = True
            break

        article_id = article["article_id"]
        title = article["title"]
        date_str = _normalize_pubdate(article.get("tm", ""))
        content_html = article.get("content_html", "")

        full_url = "{}/PortalPage/ISDInfoDetail.aspx?type=2&id={}".format(
            _SITE_ROOT, article_id)

        _safe_print("[FJTBA] [{}/{}] {}...".format(
            i, len(new_articles), title[:60]))
        sys.stdout.flush()

        # Parse HTML content
        content_text, attachments = _parse_content_html(content_html, full_url)

        # Download attachments
        local_files = []
        article_dl_dir = ""
        if attachments:
            dl_name = "{}_{}".format(article_id[:12], _sanitize_filename(title[:30], 40))
            article_dl_dir = os.path.join(downloads_dir, dl_name)
            local_files = _download_attachments(attachments, article_dl_dir, referer=full_url)
            for fp in list(local_files):
                is_zip = fp.lower().endswith(".zip")
                if not is_zip and os.path.exists(fp) and os.path.getsize(fp) >= 4:
                    with open(fp, "rb") as f:
                        is_zip = f.read(4) == b"PK\x03\x04"
                if is_zip:
                    extracted = _extract_zip(fp)
                    local_files.remove(fp)
                    local_files.extend(extracted)

        # Build markdown
        md_content = _build_markdown(title, date_str, content_text,
                                     attachments, article_dl_dir, full_url)

        # Save markdown locally
        date_for_name = date_str or datetime.now().strftime("%Y-%m-%d")
        folder_name = _sanitize_filename(
            "{}_{}_{}".format(date_for_name, article_id[:12], title[:40]), max_len=120)
        md_path = os.path.join(output_dir, "{}.md".format(folder_name))
        with open(md_path, "w", encoding="utf-8") as f:
            f.write(md_content)
        _safe_print("[FJTBA]   Saved ({} chars, {} attachments)".format(
            len(md_content), len(local_files)))
        sys.stdout.flush()

        # Upload to KB
        if args.kb_id:
            try:
                _upload_to_kb(md_content, local_files, args.kb_id,
                             args.tenant_id, folder_name)
            except Exception as e:
                logging.error("KB upload failed: %s", e)
                _save_state(output_dir, {"processed_ids": list(processed_ids)})
                _safe_print("[FJTBA]   Upload error: {}".format(e))
                sys.stdout.flush()

        processed_ids.add(article_id)
        processed_count += 1

        if processed_count % _BATCH_SIZE == 0:
            _save_state(output_dir, {"processed_ids": list(processed_ids)})
            _safe_print("[FJTBA]   Checkpoint ({} processed)".format(processed_count))
            sys.stdout.flush()

    # ── Final state ────────────────────────────────────────────────────
    state["processed_ids"] = list(processed_ids)
    _save_state(output_dir, state)

    _safe_print("\n" + "=" * 60)
    _safe_print("[FJTBA] Crawl complete — {} new article(s)".format(processed_count))
    if stopped_early:
        _safe_print("[FJTBA] Stopped early, will resume next run")
    _safe_print("=" * 60 + "\n")
    sys.stdout.flush()
    logging.info("=== FJTBA crawler finished ===")


if __name__ == "__main__":
    CONSUMER_NAME = "fjtba_crawler"
    init_root_logger(CONSUMER_NAME)
    main()
