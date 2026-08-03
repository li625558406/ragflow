"""
龙岩市公共资源交易中心 — 交易信息采集 (custom_runner)
====================================================

Target: https://ggzy.longyan.gov.cn/lyztb/

6 大栏目 51 个页签 (TABS 矩阵经首页导航 + getInfoMationList AllCount 双重验证):
  工程建设(30) / 政府采购(6) / 企业采购(4) / 产权交易(4) / 土地矿业(4, 12位码) / 林权水权(3)

Site characteristics
────────────────────
Epoint WebBuilder 平台。列表 API getInfoMationList 为 form-urlencoded POST,
第 1 页无需 token/验证码; 响应 {AllCount, custom:[{title,title2,infourl,
infodate,zhuanzai,strcomment,index}]}, 按 infodate 倒序, 无 id 字段。

详情页 = SITE_ROOT + infourl (服务端渲染 HTML):
  标题 <h3 class="bigtitle"> / 信息时间 / <p class="sub-cp"> 信息来源 /
  正文 div#mainContent / 附件 ztbfjyz('...downloadztbattach?attachGuid=<guid>...')

附件直连: ztbAttachDownloadAction.action?cmd=getContent&attachGuid=..&
appUrlFlag=ztb001&siteGuid=.. — 可能被 pageVerify 验证码拦截, 代码优雅降级
(记录 captcha_blocked_count, 附件 URL 仍写入 markdown/extracted_json)。

反爬: WAF 短时 ~30 快速请求→403(~20s 自愈); 深分页(8+)出验证码。
策略: 请求间隔 0.8-2.0s; full 模式每页签只取第 1 页(pageSize=30);
日常(date_filter)模式每页签最多翻 5 页且遇日期早于目标日即停。

Call entry: unified_crawler.py dispatches via custom_runner → run()
"""

import html as html_mod
import json
import logging
import os
import random
import re
import shutil
import ssl
import sys
import tempfile
import time
import urllib.parse
import urllib.request
import zipfile
from datetime import datetime
from typing import Any, Dict, List, Optional, Tuple

_THIS_DIR = os.path.dirname(os.path.abspath(__file__))
_PROJECT_ROOT = os.path.abspath(os.path.join(_THIS_DIR, "..", ".."))
if _PROJECT_ROOT not in sys.path:
    sys.path.insert(0, _PROJECT_ROOT)

from rag.svr.crawler_engine.models import item_from_dict

try:
    from rag.utils.redis_conn import REDIS_CONN, RedisDistributedLock
except ImportError:
    REDIS_CONN = None  # type: ignore
    RedisDistributedLock = None  # type: ignore

logger = logging.getLogger("longyan_ggzy_collection_crawler")

# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------

SITE_ID = "longyan_ggzy"
SITE_DISPLAY = "龙岩市公共资源交易中心 ggzy.longyan.gov.cn"
CATEGORY = "bid"
INFO_TYPE = "龙岩市-交易信息"          # 需求#10: type 字段按列表主题统一填写
KB_ID_DEFAULT = "3b4f619c85c211f198269135a1db216c"

_SITE_ROOT = "https://ggzy.longyan.gov.cn"
_TAG_PREFIX = "[LY-GGZY]"

_PAGE_SIZE = 30
_MAX_DAILY_PAGES = 5                  # date_filter 模式每页签最大翻页数(深分页有验证码)
_REQUEST_DELAY_MIN = 0.8
_REQUEST_DELAY_MAX = 2.0
_LISTING_RETRIES = 2          # 列表请求重试次数 (WAF 瞬时 403 自愈约 20s)
_RETRY_BACKOFF_MIN = 15       # 重试退避下限 (秒)
_RETRY_BACKOFF_MAX = 25       # 重试退避上限 (秒)
_LOCK_TIMEOUT = 10800                 # 3h: 首次全量 51 页签耗时可能超 30min 默认 TTL

_USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/125.0.0.0 Safari/537.36"
)

_SITE_GUID = "7eb5f7f1-9041-43ad-8e13-8fcb82ea831a"
_PROJECT_NAME = "/EpointWebBuilder"
_APP_URL_FLAG = "ztb001"

_API_LISTING = _SITE_ROOT + _PROJECT_NAME + "/rest/GgSearchAction/getInfoMationList"
_API_ATTACH = _SITE_ROOT + _PROJECT_NAME + "/webbuildermis/attach/ztbAttachDownloadAction.action"

# ---------------------------------------------------------------------------
# TABS 矩阵 — (categoryNum, section_name, tab_path)
# section_name = 6 大栏目(需求#4 确认口径); tab_path 进 extracted_json
# ---------------------------------------------------------------------------

_SEC_GCJS = "龙岩市-工程建设"
_SEC_ZFCG = "龙岩市-政府采购"
_SEC_QYCG = "龙岩市-企业采购"
_SEC_CQJY = "龙岩市-产权交易"
_SEC_TDKY = "龙岩市-土地矿业"
_SEC_LKSQ = "龙岩市-林权水权交易"

TABS: List[Tuple[str, str, str]] = []

# 工程建设: 房建/市政/交通/水利 各 5 页签
for _cat, _cname in (("007004", "房屋建筑"), ("007005", "市政基础设施"),
                     ("007006", "交通"), ("007007", "水利")):
    for _tab, _tname in (("001", "招标公告信息"), ("002", "答疑结果"),
                         ("003", "中标候选人公示"), ("004", "中标结果公示"),
                         ("005", "合同签署")):
        TABS.append((_cat + _tab, _SEC_GCJS, "{}/{}".format(_cname, _tname)))
# 信息化: 仅 招标+答疑
for _tab, _tname in (("001", "招标公告信息"), ("002", "答疑结果")):
    TABS.append(("007008" + _tab, _SEC_GCJS, "信息化/{}".format(_tname)))
# 工业项目: 招标+答疑+中标结果公示(站点实际存在 004, 为需求超集, 按站点采集)
for _tab, _tname in (("001", "招标公告信息"), ("002", "答疑结果"), ("004", "中标结果公示")):
    TABS.append(("007010" + _tab, _SEC_GCJS, "工业项目/{}".format(_tname)))
# 其他工程: 全 5 页签
for _tab, _tname in (("001", "招标公告信息"), ("002", "答疑结果"),
                     ("003", "中标候选人公示"), ("004", "中标结果公示"),
                     ("005", "合同签署")):
    TABS.append(("007009" + _tab, _SEC_GCJS, "其他工程/{}".format(_tname)))

# 政府采购: 008001 政府采购 / 008003 其他采购 各 3 页签
for _cat, _cname in (("008001", "政府采购"), ("008003", "其他采购")):
    for _tab, _tname in (("001", "预公告"), ("002", "招标公告"), ("003", "中标、成交公告")):
        TABS.append((_cat + _tab, _SEC_ZFCG, "{}/{}".format(_cname, _tname)))

# 企业采购 (008005)
for _tab, _tname in (("001", "招标公告"), ("002", "变更公告"),
                     ("003", "中标、成交公告"), ("004", "预公告")):
    TABS.append(("008005" + _tab, _SEC_QYCG, _tname))

# 产权交易 (011006)
for _tab, _tname in (("001", "企业和行政事业单位产权资产交易"), ("002", "无形资产交易"),
                     ("003", "其他项目"), ("004", "成交结果公告")):
    TABS.append(("011006" + _tab, _SEC_CQJY, _tname))

# 土地矿业 (011007) — 12 位叶码
for _tab, _tname in (("011007001001", "土地/出让公告"), ("011007001002", "土地/出让结果"),
                     ("011007002001", "矿业/出让公告"), ("011007002002", "矿业/出让结果")):
    TABS.append((_tab, _SEC_TDKY, _tname))

# 林权水权 — 注意林权编号反转: 011010002=交易公告, 011010001=成交公告
TABS.append(("011010002", _SEC_LKSQ, "林权/交易公告"))
TABS.append(("011010001", _SEC_LKSQ, "林权/成交公告"))
TABS.append(("011011001", _SEC_LKSQ, "水权/交易公告"))

_HEADERS_JSON = {
    "User-Agent": _USER_AGENT,
    "Accept": "application/json, text/javascript, */*; q=0.01",
    "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8",
    "Content-Type": "application/x-www-form-urlencoded; charset=UTF-8",
    "X-Requested-With": "XMLHttpRequest",
}

_HEADERS_HTML = {
    "User-Agent": _USER_AGENT,
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
    "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8",
}

_HEADERS_DOWNLOAD = {
    "User-Agent": _USER_AGENT,
    "Accept": "*/*",
    "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8",
}

_SSL_CTX = ssl.create_default_context()
_SSL_CTX.check_hostname = False
_SSL_CTX.verify_mode = ssl.CERT_NONE

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _safe_print(msg: str) -> None:
    try:
        print(msg)
    except UnicodeEncodeError:
        print(msg.encode("gbk", errors="replace").decode("gbk"))


def _request_delay() -> None:
    time.sleep(random.uniform(_REQUEST_DELAY_MIN, _REQUEST_DELAY_MAX))


def _sanitize_filename(name: str, max_len: int = 120) -> str:
    if not name:
        return "unnamed"
    safe = re.sub(r'[\\/:*?"<>|\r\n\t]+', "_", name)
    safe = re.sub(r"_+", "_", safe)
    safe = safe.strip("._ ")
    if len(safe) > max_len:
        base, ext = os.path.splitext(safe)
        safe = base[: max_len - len(ext)] + (ext or "")
    return safe or "unnamed"


def _normalize_date(date_str: str) -> str:
    if not date_str:
        return ""
    m = re.search(r"(\d{4}-\d{1,2}-\d{1,2})", str(date_str))
    if not m:
        return str(date_str)[:10]
    y, mo, d = m.group(1).split("-")
    return "{}-{:02d}-{:02d}".format(y, int(mo), int(d))


# ---------------------------------------------------------------------------
# HTTP helpers
# ---------------------------------------------------------------------------

def _http_post_form(url: str, form: dict, extra_headers: Optional[dict] = None) -> Optional[str]:
    """form-urlencoded POST. Returns response text or None."""
    hdrs = dict(_HEADERS_JSON)
    if extra_headers:
        hdrs.update(extra_headers)
    body = urllib.parse.urlencode(form).encode("utf-8")
    req = urllib.request.Request(url, data=body, headers=hdrs)
    try:
        with urllib.request.urlopen(req, timeout=30, context=_SSL_CTX) as resp:
            return resp.read().decode("utf-8", errors="replace")
    except Exception as e:
        logging.warning("POST %s failed: %s", url, e)
        return None


def _http_get_html(url: str, referer: Optional[str] = None) -> Optional[str]:
    """GET HTML page. Returns response text or None."""
    hdrs = dict(_HEADERS_HTML)
    if referer:
        hdrs["Referer"] = referer
    req = urllib.request.Request(url, headers=hdrs)
    try:
        with urllib.request.urlopen(req, timeout=30, context=_SSL_CTX) as resp:
            return resp.read().decode("utf-8", errors="replace")
    except Exception as e:
        logging.warning("GET %s failed: %s", url, e)
        return None


def _http_download(url: str, referer: Optional[str] = None) -> Optional[bytes]:
    """GET binary download. Returns bytes or None."""
    hdrs = dict(_HEADERS_DOWNLOAD)
    if referer:
        hdrs["Referer"] = referer
    req = urllib.request.Request(url, headers=hdrs)
    try:
        with urllib.request.urlopen(req, timeout=60, context=_SSL_CTX) as resp:
            return resp.read()
    except Exception as e:
        logging.warning("Download %s failed: %s", url, e)
        return None


# ---------------------------------------------------------------------------
# Listing API
# ---------------------------------------------------------------------------

def _fetch_tab_rows(category_num: str, page_index: int) -> List[dict]:
    """Fetch one page (0-based) of a tab listing. Returns row dicts.

    验证码/反爬响应不是 JSON → 返回 [] (上层按空页处理, 不崩)。
    网络/WAF 瞬时失败(text 为空)重试一次: WAF 约 20s 自愈, 退避 15-25s。
    """
    form = {
        "siteGuid": _SITE_GUID,
        "categoryNum": category_num,
        "kw": "",
        "pageIndex": str(page_index),
        "pageSize": str(_PAGE_SIZE),
        "YZM": "",
        "ImgGuid": "",
    }
    text = None
    for attempt in range(_LISTING_RETRIES):
        text = _http_post_form(
            _API_LISTING, form, extra_headers={"Referer": _SITE_ROOT + "/lyztb/"}
        )
        if text:
            break
        if attempt == 0:
            logging.warning(
                "longyan listing empty response, retrying: cat=%s page=%d",
                category_num, page_index,
            )
            time.sleep(random.uniform(_RETRY_BACKOFF_MIN, _RETRY_BACKOFF_MAX))
    if not text:
        logging.warning(
            "longyan listing still empty after %d attempts, treating as empty page: cat=%s page=%d",
            _LISTING_RETRIES, category_num, page_index,
        )
        return []
    try:
        data = json.loads(text)
    except json.JSONDecodeError:
        logging.warning(
            "longyan listing not JSON (captcha/WAF?) cat=%s page=%d head=%r",
            category_num, page_index, text[:80],
        )
        return []
    if not isinstance(data, dict):
        logging.warning(
            "longyan listing JSON root not object: cat=%s page=%d type=%s",
            category_num, page_index, type(data).__name__,
        )
        return []
    rows = data.get("custom")
    if not isinstance(rows, list):
        logging.warning(
            "longyan listing missing 'custom' list: cat=%s page=%d keys=%s",
            category_num, page_index, sorted(data.keys()),
        )
        return []
    return [r for r in rows if isinstance(r, dict)]


# ---------------------------------------------------------------------------
# Detail page
# ---------------------------------------------------------------------------

def _html_to_text(html_str: str) -> str:
    """Strip HTML tags and decode entities, return clean text."""
    if not html_str:
        return ""
    text = re.sub(r"<(script|style)[^>]*>.*?</\1>", "", html_str, flags=re.DOTALL | re.I)
    text = re.sub(
        r"</?(?:div|p|tr|li|h[1-6]|table|hr|section|article|header|footer)[^>]*>",
        "\n", text, flags=re.I,
    )
    text = re.sub(r"<br\s*/?>", "\n", text, flags=re.I)
    text = re.sub(r"</td>", " ", text, flags=re.I)
    text = re.sub(r"</tr>", "\n", text, flags=re.I)
    text = re.sub(r"<[^>]+>", "", text)
    text = html_mod.unescape(text)
    text = text.replace("\xa0", " ").replace("　", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _extract_attachments_from_html(html_str: str, page_url: str) -> List[dict]:
    """Extract attachment links from detail page HTML.

    两种来源:
      1. ztbfjyz('...downloadztbattach?attachGuid=<guid>...') onclick —
         构造 ztbAttachDownloadAction 直连 URL (appUrlFlag=ztb001)
      2. 直接文件链接 (.pdf/.doc/.xls/.zip/.rar/.ppt)
    Returns [{name, guid, url}] — url 均为绝对地址(需求#13: 原文链接完整可访问)。
    """
    attachments: List[dict] = []
    seen: set = set()

    for m in re.finditer(
        r"ztbfjyz\('([^']+/downloadztbattach\?attachGuid="
        r"([0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12})[^']*)'",
        html_str, re.I,
    ):
        guid = m.group(2)
        if guid in seen:
            continue
        seen.add(guid)
        name = "{}.unknown".format(guid)
        post_context = html_str[m.end():m.end() + 500]
        name_m = re.search(r'title="([^"]+)"', post_context)
        if not name_m:
            name_m = re.search(r'>([^<]{3,120})</a>', post_context)
        if name_m:
            name = html_mod.unescape(name_m.group(1).strip())
        real_url = (
            "{}?cmd=getContent&attachGuid={}&appUrlFlag={}&siteGuid={}".format(
                _API_ATTACH, guid, _APP_URL_FLAG, _SITE_GUID
            )
        )
        attachments.append({
            "name": _sanitize_filename(name),
            "guid": guid,
            "url": real_url,
        })

    for m in re.finditer(
        r'<a[^>]*href=["\']([^"\']*(?:\.pdf|\.doc[x]?|\.xls[x]?|\.rar|\.zip|\.ppt[x]?)'
        r'[^"\']*)["\'][^>]*>([^<]*)</a>',
        html_str, re.I,
    ):
        href = m.group(1).strip()
        link_text = html_mod.unescape(m.group(2).strip())
        if not href:
            continue
        if not href.startswith("http"):
            if href.startswith("//"):
                href = "https:" + href
            elif href.startswith("/"):
                href = _SITE_ROOT + href
            else:
                href = page_url.rsplit("/", 1)[0] + "/" + href
        if href in seen:
            continue
        seen.add(href)
        fname = link_text or os.path.basename(
            urllib.parse.urlparse(href).path.split("?")[0]
        ) or "attachment"
        attachments.append({"name": _sanitize_filename(fname), "guid": "", "url": href})

    return attachments


def _fetch_detail(full_url: str) -> Optional[dict]:
    """Fetch + parse one detail page. Returns None on network failure.

    Returns {title, info_time, source, content_text, attachments}.
    """
    html = _http_get_html(full_url, referer=_SITE_ROOT + "/lyztb/")
    if not html:
        return None

    result: Dict[str, Any] = {
        "title": "", "info_time": "", "source": "",
        "content_text": "", "attachments": [],
    }

    m = re.search(r'<h3\s+class="bigtitle">(.+?)</h3>', html, re.DOTALL)
    if m:
        result["title"] = html_mod.unescape(re.sub(r"<[^>]+>", "", m.group(1))).strip()

    m = re.search(r"信息时间：(\d{4}-\d{2}-\d{2})", html)
    if m:
        result["info_time"] = m.group(1)

    m = re.search(r'<p\s+class="sub-cp">(.*?)</p>', html, re.DOTALL)
    if m:
        src_m = re.search(r'信息来源：<span[^>]*>([^<]*)</span>', m.group(1))
        if src_m:
            result["source"] = src_m.group(1).strip()

    idx = html.find('id="mainContent"')
    if idx >= 0:
        content_start = html.find(">", idx) + 1
        content_end = len(html)
        for marker in ("ewb-page-lookup", "footrt", "subfooter"):
            tag_m = re.search(
                r"<[^>]*" + re.escape(marker) + r"[^>]*>", html[content_start:]
            )
            if tag_m:
                content_end = content_start + tag_m.start()
                break
        content_html = html[content_start:content_end].strip()
        content_html = re.sub(r"<[^>]*\Z", "", content_html)
        no_chain = re.sub(
            r'<div\s+class="chain"[^>]*>.*?</div>\s*</div>\s*</div>\s*</div>\s*</div>',
            "", content_html, flags=re.DOTALL,
        )
        if no_chain.strip():
            content_html = no_chain
        result["content_text"] = _html_to_text(content_html)

    result["attachments"] = _extract_attachments_from_html(html, full_url)
    if not result["title"] or not result["content_text"]:
        logging.warning(
            "longyan detail parse incomplete (title=%r content_len=%d): %s",
            result["title"][:40], len(result["content_text"]), full_url,
        )
    return result


# ---------------------------------------------------------------------------
# Attachment download
# ---------------------------------------------------------------------------

def _download_attachments(attachments: List[dict], download_dir: str) -> Tuple[List[str], int]:
    """Download attachments to local dir.

    Returns (local_file_paths, captcha_blocked_count).
    网关/验证码 HTML 响应 → 计数跳过(附件 URL 仍保留在 markdown/extracted_json)。
    """
    os.makedirs(download_dir, exist_ok=True)
    local_files: List[str] = []
    captcha_blocked = 0

    for att in attachments:
        url = att.get("url", "")
        if not url:
            continue

        fname = _sanitize_filename(att.get("name", "attachment"), max_len=120)
        ext = os.path.splitext(urllib.parse.urlparse(url).path.split("?")[0])[1].lower()
        if ext and not fname.lower().endswith(ext):
            fname += ext

        filepath = os.path.join(download_dir, fname)
        if os.path.exists(filepath) and os.path.getsize(filepath) > 0:
            local_files.append(filepath)
            continue

        data = _http_download(url, referer=_SITE_ROOT + "/lyztb/")
        if not data:
            continue

        # HTML 响应 = 网关页或验证码页, 不是真实文件
        if data[:15].lower().startswith(b"<!doctype") or data[:6].lower().startswith(b"<html"):
            gateway = data.decode("utf-8", errors="replace")
            low = gateway.lower()
            if "pageverify" in low or "验证码" in gateway:
                logging.warning("Attachment captcha-blocked: %s", url[:100])
                captcha_blocked += 1
                continue
            # 宁德同款网关: 从 form action 提取真实下载地址再试一次
            action_m = re.search(
                r'action\s*=\s*["\']([^"\']*ztbAttachDownloadAction[^"\']*)["\']',
                gateway, re.I,
            )
            if action_m:
                action = action_m.group(1)
                if not action.startswith("http"):
                    action = (_SITE_ROOT + _PROJECT_NAME +
                              "/webbuildermis/attach/" + action.lstrip("/"))
                data = _http_download(action, referer=url)
                if not data:
                    continue
                if data[:15].lower().startswith(b"<!doctype") or \
                   data[:6].lower().startswith(b"<html"):
                    captcha_blocked += 1
                    continue
            else:
                captcha_blocked += 1
                continue

        # magic bytes 修正扩展名
        if data[:4] == b"%PDF" and not fname.lower().endswith(".pdf"):
            fname += ".pdf"
        elif data[:4] == b"PK\x03\x04" and \
                not any(fname.lower().endswith(e) for e in (".zip", ".docx", ".xlsx", ".pptx")):
            fname += ".zip"

        filepath = os.path.join(download_dir, fname)
        with open(filepath, "wb") as f:
            f.write(data)
        local_files.append(filepath)
        _request_delay()

    return local_files, captcha_blocked


def _extract_zip(filepath: str) -> List[str]:
    """Extract ZIP members (需求#9: ZIP 解压后入库/入KB)。返回解压文件路径列表。"""
    extracted: List[str] = []
    extract_dir = os.path.splitext(filepath)[0] + "_extracted"
    os.makedirs(extract_dir, exist_ok=True)
    try:
        with zipfile.ZipFile(filepath, "r") as zf:
            for name in zf.namelist():
                if name.endswith("/"):
                    continue
                safe_name = _sanitize_filename(os.path.basename(name), max_len=120)
                out_path = os.path.join(extract_dir, safe_name)
                # 同名冲突加序号, 避免覆盖
                n = 1
                while os.path.exists(out_path):
                    base, ext = os.path.splitext(safe_name)
                    out_path = os.path.join(extract_dir, "{}_{}{}".format(base, n, ext))
                    n += 1
                with zf.open(name) as src, open(out_path, "wb") as dst:
                    dst.write(src.read())
                extracted.append(out_path)
    except Exception as e:
        logging.warning("ZIP extraction failed for %s: %s", filepath, e)
    return extracted


def _extract_text_from_file(filepath: str) -> str:
    """Extract text from TXT/PDF/DOCX/XLSX for the markdown appendix.

    扫描件 PDF 由 KB naive parser 的 OCR 兜底(需求 ocr 口径)。
    """
    ext = os.path.splitext(filepath)[1].lower()
    try:
        if ext == ".txt":
            with open(filepath, "r", encoding="utf-8", errors="replace") as f:
                return f.read()
        elif ext == ".pdf":
            import pdfplumber

            parts: List[str] = []
            with pdfplumber.open(filepath) as pdf:
                for pg in pdf.pages:
                    text = pg.extract_text()
                    if text:
                        parts.append(text)
            return "\n\n".join(parts)
        elif ext in (".docx", ".doc"):
            import docx

            doc = docx.Document(filepath)
            return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        elif ext in (".xlsx", ".xls"):
            import openpyxl

            wb = openpyxl.load_workbook(filepath, read_only=True)
            parts = []
            for ws in wb.worksheets:
                rows: List[str] = []
                for row in ws.iter_rows(values_only=True):
                    rows.append(" | ".join(str(c) if c is not None else "" for c in row))
                if rows:
                    parts.append("### {}\n".format(ws.title) + "\n".join(rows))
            return "\n\n".join(parts)
    except Exception as e:
        logging.warning("Text extraction failed for %s: %s", filepath, e)
    return ""
