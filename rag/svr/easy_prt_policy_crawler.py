#!/usr/bin/env python3
#
#  Copyright 2026 The InfiniFlow Authors. All Rights Reserved.
#
#  Licensed under the Apache License, Version 2.0 (the "License");
#  you may not use this file except in compliance with the License.
#  You may obtain a copy of the License at
#
#      http://www.apache.org/licenses/LICENSE-2.0
#
#  Unless required by applicable law or agreed to in writing, software
#  distributed under the License is distributed on an "AS IS" BASIS,
#  WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
#  See the License for the specific language governing permissions and
#  limitations under the License.
#
"""
Crawler for easy-prt.com/policy (政策法规) — policy & regulation articles.

Site characteristics
────────────────────
Vue.js SPA with SM4-ECB encrypted API communication via CMS endpoints.
Plain ``urllib`` is sufficient — no browser rendering required.

CMS APIs
────────
  • Listing →  POST /jeecg-system/cms/content/portal/list
               Body: JSON-wrapped hex-encrypted payload
  • Detail  →  GET  /jeecg-system/cms/content/portal/detail
               Query: ?encryptParams=<hex>

Every request must include header ``X-Platform-Id: gct``.  The SM4 key
is identical to the rest of the site.

The policy section (type=02) contains 21 articles.  A few have PDF
attachments served from the /business/ CDN path.

Checkpoint/resume: articles processed in batches of 10, state saved after
each batch.  The 3600 s task timeout is handled via --max-runtime.

Usage:
    python easy_prt_policy_crawler.py \
        --tenant-id <TENANT_ID> \
        --target-url https://easy-prt.com/policy \
        --kb-id <KB_ID> \
        --task-name <NAME>
"""

import argparse
import json
import logging
import os
import random
import re
import sys
import time
import urllib.error
import urllib.request
from datetime import datetime

from bs4 import BeautifulSoup

_SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
_PROJECT_ROOT = os.path.abspath(os.path.join(_SCRIPT_DIR, "..", ".."))
sys.path.insert(0, _PROJECT_ROOT)

from common import settings
from common.log_utils import init_root_logger
from common.misc_utils import get_uuid

try:
    from gmssl.sm4 import CryptSM4, SM4_ENCRYPT, SM4_DECRYPT
except ImportError:
    logging.error("gmssl package is required. Run: pip install gmssl")
    sys.exit(1)

# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------
_API_BASE = "https://easy-prt.com/gykjbid"
_SM4_KEY = "90bdd291004611ef87fc52540023e781"
_SITE_ROOT = "https://easy-prt.com"

_LISTING_PATH = "/jeecg-system/cms/content/portal/list"
_DETAIL_PATH = "/jeecg-system/cms/content/portal/detail"

# CMS content type — "02" = 政策法规 (policy & regulation)
_POLICY_TYPE = "02"

_USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/120.0.0.0 Safari/537.36"
)

_HEADERS = {
    "Content-Type": "application/json",
    "X-Platform-Id": "gct",
    "User-Agent": _USER_AGENT,
}

_GET_HEADERS = {
    "X-Platform-Id": "gct",
    "User-Agent": _USER_AGENT,
}

# Anti-crawling: random delays between requests (seconds)
_REQUEST_DELAY_MIN = 1.5
_REQUEST_DELAY_MAX = 3.0

# Batch checkpoint
BATCH_SIZE = 10

# Max runtime before graceful stop (seconds)
_MAX_RUNTIME_DEFAULT = 3300

# File extensions to download
_ATTACHMENT_EXTS = {".pdf", ".doc", ".docx", ".xls", ".xlsx", ".ppt", ".pptx",
                    ".zip", ".rar", ".7z", ".txt", ".jpg", ".jpeg", ".png"}

_POLICY_SECTION = "政策法规"


# ---------------------------------------------------------------------------
# SM4 helpers
# ---------------------------------------------------------------------------

def _sm4_encrypt(data):
    """SM4-ECB encrypt dict -> hex string."""
    cipher = CryptSM4()
    cipher.set_key(bytes.fromhex(_SM4_KEY), SM4_ENCRYPT)
    payload = json.dumps(data, separators=(",", ":")).encode("utf-8")
    return cipher.crypt_ecb(payload).hex()


def _sm4_decrypt(hex_data):
    """SM4-ECB decrypt hex string -> dict."""
    cipher = CryptSM4()
    cipher.set_key(bytes.fromhex(_SM4_KEY), SM4_DECRYPT)
    decrypted = cipher.crypt_ecb(bytes.fromhex(hex_data)).decode("utf-8")
    return json.loads(decrypted)


# ---------------------------------------------------------------------------
# HTTP helpers
# ---------------------------------------------------------------------------

def _request_delay():
    """Random delay between requests to avoid rate limiting."""
    time.sleep(random.uniform(_REQUEST_DELAY_MIN, _REQUEST_DELAY_MAX))


def _api_post(path, params, retries=3):
    """POST with SM4-encrypted body.  Returns decoded dict or None."""
    enc = _sm4_encrypt(params)
    body = json.dumps(enc).encode("utf-8")
    url = _API_BASE + path

    for attempt in range(1, retries + 1):
        try:
            req = urllib.request.Request(url, data=body, headers=_HEADERS,
                                        method="POST")
            with urllib.request.urlopen(req, timeout=30) as resp:
                raw = resp.read().decode("utf-8")
            stripped = raw.strip().strip('"')
            return _sm4_decrypt(stripped)
        except Exception as e:
            logging.warning("POST %s attempt %d/%d: %s",
                          path, attempt, retries, e)
            if attempt < retries:
                time.sleep(2 ** attempt)
    return None


def _api_get(path, params, retries=3):
    """GET with SM4-encrypted query param.  Returns decoded dict or None."""
    enc = _sm4_encrypt(params)
    url = "{}?encryptParams={}".format(_API_BASE + path, enc)

    for attempt in range(1, retries + 1):
        try:
            req = urllib.request.Request(url, headers=_GET_HEADERS, method="GET")
            with urllib.request.urlopen(req, timeout=30) as resp:
                raw = resp.read().decode("utf-8")
            stripped = raw.strip().strip('"')
            return _sm4_decrypt(stripped)
        except Exception as e:
            logging.warning("GET %s attempt %d/%d: %s",
                          path, attempt, retries, e)
            if attempt < retries:
                time.sleep(2 ** attempt)
    return None


def _download_file(url, retries=2):
    """Download a binary file.  Returns (content_bytes, filename) or (None, None)."""
    for attempt in range(1, retries + 1):
        try:
            req = urllib.request.Request(
                url, headers={"User-Agent": _USER_AGENT},
            )
            with urllib.request.urlopen(req, timeout=120) as resp:
                content = resp.read()
            filename = url.rstrip("/").split("/")[-1]
            if "?" in filename:
                filename = filename.split("?")[0]
            return content, filename
        except Exception as e:
            logging.warning("Download %s attempt %d/%d: %s",
                          url, attempt, retries, e)
            if attempt < retries:
                time.sleep(2 ** attempt)
    return None, None


# ---------------------------------------------------------------------------
# Listing / detail fetchers
# ---------------------------------------------------------------------------

def _fetch_policy_listing():
    """Fetch all policy article entries (type=02).

    Returns list[dict] with keys: id, title, releaseTime.
    """
    all_items = []
    page = 1

    while True:
        result = _api_post(_LISTING_PATH, {
            "pageNo": page,
            "pageSize": 20,
            "type": _POLICY_TYPE,
        })
        if not result or result.get("code") != 200:
            logging.error("Policy listing page %d failed", page)
            break

        res = result.get("result", {})
        records = res.get("records", [])
        if not records:
            break

        all_items.extend(records)
        total = res.get("total", 0)
        if len(all_items) >= total:
            break
        page += 1
        _request_delay()

    return all_items


def _fetch_policy_detail(article_id):
    """Fetch full detail for a policy article.

    Returns dict with keys: id, title, content (HTML), releaseTime,
    releaseStatus, annexes (list), images (list), categoryId.
    """
    result = _api_get(_DETAIL_PATH, {
        "id": article_id,
        "type": _POLICY_TYPE,
    })
    if not result or result.get("code") != 200:
        logging.warning("Detail API failed for %s", article_id)
        return None
    return result.get("result") or {}


# ---------------------------------------------------------------------------
# HTML → plain text
# ---------------------------------------------------------------------------

def _html_to_text(html):
    """Convert HTML content to plain text via BeautifulSoup."""
    if not html:
        return ""
    soup = BeautifulSoup(html, "lxml")
    for tag in soup.find_all(["script", "style", "noscript"]):
        tag.decompose()
    return soup.get_text(separator="\n", strip=True)


# ---------------------------------------------------------------------------
# Attachment handling
# ---------------------------------------------------------------------------

def _collect_attachment_urls(detail_data, content_html):
    """Collect attachment info from detail data and content HTML.

    Returns list of (url, filename) tuples.
    """
    attachments = []

    # From annexes field
    annexes = detail_data.get("annexes") or []
    for annex in annexes:
        file_url = annex.get("fileUrl", "") or annex.get("url", "")
        if not file_url:
            continue
        if file_url.startswith("/"):
            file_url = _SITE_ROOT + file_url
        filename = annex.get("fileName", "") or os.path.basename(
            file_url.split("?")[0],
        )
        attachments.append((file_url, filename))

    # From content HTML embedded links
    if content_html:
        soup = BeautifulSoup(content_html, "html.parser")
        for a_tag in soup.find_all("a", href=True):
            href = a_tag["href"].strip()
            ext = os.path.splitext(href.split("?")[0].split("#")[0])[1].lower()
            if ext in _ATTACHMENT_EXTS:
                if href.startswith("/"):
                    href = _SITE_ROOT + href
                fname = a_tag.get_text(strip=True) or os.path.basename(
                    href.split("?")[0],
                )
                attachments.append((href, fname))

    return attachments


def _save_attachment(content, output_dir, filename):
    """Save downloaded attachment to disk.  Returns file path."""
    att_dir = os.path.join(output_dir, "attachments")
    os.makedirs(att_dir, exist_ok=True)
    safe_name = re.sub(r'[\\/:*?"<>|]', "_", filename)
    path = os.path.join(att_dir, safe_name)
    if os.path.exists(path):
        return path
    with open(path, "wb") as f:
        f.write(content)
    return path


def _extract_zip(content, output_dir, article_id):
    """Extract ZIP bytes, return list of extracted file paths."""
    import zipfile
    import io

    extracted = []
    dest_dir = os.path.join(output_dir, "attachments", article_id)
    os.makedirs(dest_dir, exist_ok=True)

    try:
        with zipfile.ZipFile(io.BytesIO(content)) as zf:
            for name in zf.namelist():
                safe_name = re.sub(r'[\\/:*?"<>|]', "_", os.path.basename(name))
                if not safe_name:
                    continue
                dest_path = os.path.join(dest_dir, safe_name)
                if os.path.exists(dest_path):
                    continue
                with open(dest_path, "wb") as f:
                    f.write(zf.read(name))
                extracted.append(dest_path)
    except Exception as e:
        logging.warning("ZIP extract error: %s", e)
    return extracted


def _extract_text_from_file(filepath):
    """Extract plain text from PDF, DOCX, XLSX, TXT files.

    Returns text string or None.
    """
    ext = os.path.splitext(filepath)[1].lower()
    try:
        if ext == ".txt":
            with open(filepath, "r", encoding="utf-8", errors="replace") as f:
                return f.read()
        elif ext == ".pdf":
            try:
                import pdfplumber
                parts = []
                with pdfplumber.open(filepath) as pdf:
                    for p in pdf.pages:
                        text = p.extract_text()
                        if text:
                            parts.append(text)
                return "\n\n".join(parts)
            except ImportError:
                return None
        elif ext in (".docx", ".doc"):
            try:
                import docx
                doc = docx.Document(filepath)
                return "\n".join(
                    p.text for p in doc.paragraphs if p.text.strip()
                )
            except ImportError:
                return None
        elif ext in (".xlsx", ".xls"):
            try:
                import openpyxl
                wb = openpyxl.load_workbook(filepath, read_only=True)
                parts = []
                for ws in wb.worksheets:
                    rows = []
                    for row in ws.iter_rows(values_only=True):
                        rows.append(" | ".join(
                            str(c) if c is not None else "" for c in row
                        ))
                    if rows:
                        parts.append(f"### {ws.title}\n" + "\n".join(rows))
                wb.close()
                return "\n\n".join(parts)
            except ImportError:
                return None
    except Exception as e:
        logging.warning("Text extraction failed for %s: %s", filepath, e)
    return None


# ---------------------------------------------------------------------------
# Markdown building
# ---------------------------------------------------------------------------

def _build_markdown(article, detail):
    """Build a Markdown document from article listing + detail data."""
    title = detail.get("title", "") or article.get("title", "无标题")
    release_time = detail.get("releaseTime", "") or article.get("releaseTime", "")
    content_html = detail.get("content", "") or ""

    content_text = _html_to_text(content_html)

    lines = [
        f"# {title}",
        "",
        f"**来源：** {_SITE_ROOT}/policy",
        f"**发布日期：** {release_time}",
        "",
        "## 正文",
        "",
        content_text if content_text else "（暂无内容）",
    ]

    return "\n".join(lines)


# ---------------------------------------------------------------------------
# Persistence
# ---------------------------------------------------------------------------

_STATE_FILENAME = "_crawler_state.json"


def _load_state(output_dir):
    path = os.path.join(output_dir, _STATE_FILENAME)
    if os.path.exists(path):
        try:
            with open(path, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception as e:
            logging.warning("Failed to load crawler state: %s", e)
    return {"processed_ids": [], "completed": False}


def _save_state(output_dir, state):
    path = os.path.join(output_dir, _STATE_FILENAME)
    os.makedirs(output_dir, exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(state, f, ensure_ascii=False)
    logging.info("Crawler state saved (%d processed IDs)",
                 len(state.get("processed_ids", [])))


def _save_markdown(content, output_dir, article_id):
    articles_dir = os.path.join(output_dir, "articles")
    os.makedirs(articles_dir, exist_ok=True)
    path = os.path.join(articles_dir, "{}.md".format(article_id))
    with open(path, "w", encoding="utf-8") as f:
        f.write(content)
    return path


# ---------------------------------------------------------------------------
# KB upload
# ---------------------------------------------------------------------------

def _upload_to_kb(filepath, kb_id, tenant_id, parser_id="laws"):
    """Upload a file to KB and queue parsing."""
    from api.db.services.knowledgebase_service import KnowledgebaseService
    from api.db.services.file_service import FileService
    from api.db.services.document_service import DocumentService

    ok, kb = KnowledgebaseService.get_by_id(kb_id)
    if not ok:
        raise LookupError("Knowledge base {} not found".format(kb_id))

    with open(filepath, "rb") as f:
        blob = f.read()

    class _FileObj:
        def __init__(self, filename, blob):
            self.id = get_uuid()
            self.filename = filename
            self.blob = blob
        def read(self):
            return self.blob

    file_obj = _FileObj(os.path.basename(filepath), blob)
    errs, doc_pairs = FileService.upload_document(kb, [file_obj], tenant_id)
    if errs:
        logging.warning("Upload errors: %s", errs)
    for doc, _ in doc_pairs:
        doc_id = doc["id"]
        try:
            DocumentService.update_by_id(doc_id, {"parser_id": parser_id})
        except Exception:
            pass
        try:
            DocumentService.begin2parse(doc_id)
            DocumentService.run(tenant_id, doc, {})
        except Exception as e:
            logging.error("Queue parse for %s: %s", doc_id, e)
    return doc_pairs


# ---------------------------------------------------------------------------
# Main processing
# ---------------------------------------------------------------------------

def _process_section(output_dir, kb_id, tenant_id, state,
                     max_runtime, start_time):
    """Fetch listing, process articles in batches, upload to KB."""
    processed_ids = set(state.get("processed_ids", []))
    if state.get("completed"):
        _safe_print("[POLICY] Already completed, skip.\n")
        return

    # ── Fetch listing ──
    _safe_print("[POLICY] Fetching listing...")
    sys.stdout.flush()

    all_items = _fetch_policy_listing()
    _safe_print("[POLICY] {} items total from listing".format(len(all_items)))
    sys.stdout.flush()

    new_items = [a for a in all_items if a.get("id") not in processed_ids]
    _safe_print("[POLICY] {} new ({} skipped)".format(
        len(new_items), len(all_items) - len(new_items)))
    sys.stdout.flush()

    if not new_items:
        state["completed"] = True
        _save_state(output_dir, state)
        _safe_print("[POLICY] Nothing new to process.\n")
        return

    total = len(new_items)
    success_count = 0
    batch_num = 0
    stopped_early = False

    for batch_start in range(0, total, BATCH_SIZE):
        # ── Timeout check ──
        elapsed = time.time() - start_time
        if elapsed > max_runtime - 120:
            _safe_print(
                "\n[POLICY] Runtime {:.0f}s, stopping gracefully "
                "({} / {} done).".format(elapsed, success_count, total))
            sys.stdout.flush()
            stopped_early = True
            break

        batch = new_items[batch_start:batch_start + BATCH_SIZE]
        batch_num += 1
        batch_files = []
        batch_ids = []

        for idx, item in enumerate(batch, 1):
            global_idx = batch_start + idx
            item_id = item.get("id", "")
            title = item.get("title", "无标题")
            _safe_print("[POLICY]   [{}/{}] {}".format(global_idx, total, title))
            sys.stdout.flush()

            # Fetch detail
            detail = _fetch_policy_detail(item_id)
            if not detail:
                _safe_print("      -> detail fetch failed")
                sys.stdout.flush()
                continue

            content_html = detail.get("content", "") or ""

            # Collect & download attachments
            att_urls = _collect_attachment_urls(detail, content_html)
            attachment_paths = []
            seen = set()
            for file_url, filename in att_urls:
                if file_url in seen:
                    continue
                seen.add(file_url)
                _safe_print("        downloading: {}".format(filename))
                sys.stdout.flush()

                content, dl_name = _download_file(file_url)
                if not content:
                    continue
                fn = filename or dl_name or "untitled"

                ext = os.path.splitext(fn)[1].lower()
                if ext in (".zip", ".rar", ".7z"):
                    extracted = _extract_zip(content, output_dir, item_id)
                    attachment_paths.extend(extracted)
                    _safe_print("        extracted {} files".format(len(extracted)))
                else:
                    saved = _save_attachment(content, output_dir, fn)
                    attachment_paths.append(saved)

                _request_delay()

            # Build markdown
            md = _build_markdown(item, detail)

            # Append attachment text
            if attachment_paths:
                md += "\n## 附件内容\n\n"
                for ap in attachment_paths:
                    fname = os.path.basename(ap)
                    text = _extract_text_from_file(ap)
                    if text and text.strip():
                        if len(text) > 50000:
                            text = text[:50000] + "\n\n（内容过长，已截断）"
                        md += "### {}\n\n{}\n\n".format(fname, text)
                    else:
                        md += "- {}\n".format(fname)

            _save_markdown(md, output_dir, item_id)
            batch_files.append(os.path.join(
                output_dir, "articles", "{}.md".format(item_id)))
            batch_ids.append(item_id)
            success_count += 1
            _request_delay()

        # ── Checkpoint ──
        if batch_ids:
            processed_ids.update(batch_ids)
            state["processed_ids"] = list(processed_ids)
            _save_state(output_dir, state)

            if kb_id:
                for fp in batch_files:
                    if os.path.exists(fp):
                        try:
                            _upload_to_kb(fp, kb_id, tenant_id)
                        except Exception as e:
                            logging.error("Upload %s: %s", fp, e)
                # Also upload attachments
                att_root = os.path.join(output_dir, "attachments")
                if os.path.isdir(att_root):
                    for root, _, files in os.walk(att_root):
                        for f in files:
                            fp = os.path.join(root, f)
                            try:
                                _upload_to_kb(fp, kb_id, tenant_id)
                            except Exception as e:
                                logging.error("Upload att %s: %s", fp, e)

            _safe_print("[POLICY]   batch {} done ({}/{} total)\n".format(
                batch_num, success_count, total))
            sys.stdout.flush()

    if not stopped_early:
        state["completed"] = True
        _save_state(output_dir, state)

    _safe_print("[POLICY] Section complete: {} articles processed".format(
        success_count))
    sys.stdout.flush()


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def parse_args():
    p = argparse.ArgumentParser(
        description="easy-prt.com/policy (政策法规) crawler",
    )
    p.add_argument("--tenant-id", required=True, help="Tenant ID for KB upload")
    p.add_argument("--target-url", required=True,
                   help="Policy page URL (e.g. https://easy-prt.com/policy)")
    p.add_argument("--kb-id", required=True, help="Target knowledge-base ID")
    p.add_argument("--task-name", required=True,
                   help="Task name used as output sub-directory")
    p.add_argument("--output-dir", default=None,
                   help="Output root directory (default: project root)")
    p.add_argument("--full", action="store_true",
                   help="Ignore saved state and re-crawl all articles")
    p.add_argument("--max-runtime", type=int, default=3300,
                   help="Max runtime in seconds before graceful stop (default: 3300)")
    for opt in ("--section", "--max-articles", "--max-days",
                "--llm-id", "--llm-model", "--access-token"):
        p.add_argument(opt, default=None)
    return p.parse_args()


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def _safe_print(msg):
    try:
        print(msg)
    except UnicodeEncodeError:
        print(msg.encode("gbk", errors="replace").decode("gbk"))


def main():
    args = parse_args()

    _safe_print("\n" + "=" * 60)
    _safe_print("[POLICY] easy-prt.com/policy crawler")
    _safe_print("[POLICY] KB: {}".format(args.kb_id))
    _safe_print("[POLICY] Task: {}".format(args.task_name))
    _safe_print("=" * 60 + "\n")
    sys.stdout.flush()

    settings.init_settings()
    logging.info("=== EASYPRT-POLICY crawler started ===")

    output_dir = args.output_dir or os.path.join(
        _PROJECT_ROOT, "rag", args.task_name.strip(),
    )
    os.makedirs(output_dir, exist_ok=True)
    _safe_print("[POLICY] Output: {}\n".format(output_dir))
    sys.stdout.flush()

    # State
    state = _load_state(output_dir) if not args.full else {
        "processed_ids": [], "completed": False,
    }
    processed_ids = set(state.get("processed_ids", []))
    _safe_print("[POLICY] Already processed: {} article(s)\n".format(
        len(processed_ids)))
    sys.stdout.flush()

    start_time = time.time()
    _process_section(output_dir, args.kb_id, args.tenant_id, state,
                     args.max_runtime, start_time)

    _safe_print("\n" + "=" * 60)
    _safe_print("[POLICY] Crawl complete: {} articles processed".format(
        len(state.get("processed_ids", []))))
    _safe_print("=" * 60 + "\n")
    sys.stdout.flush()
    logging.info("=== EASYPRT-POLICY crawler finished ===")


if __name__ == "__main__":
    CONSUMER_NAME = "easy_prt_policy_crawler"
    init_root_logger(CONSUMER_NAME)
    main()
