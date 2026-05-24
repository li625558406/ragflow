#!/usr/bin/env python3
#
#  Copyright 2026 The InfiniFlow Authors. All Rights Reserved.
#
#  Licensed under the Apache License, Version 2.0 (the "License");
#  you may not use this file except in compliance with the License.
#  You may obtain a copy of the License at
#
#      http://www.apache.org/licenses/LICENSE-2.0
#
#  Unless required by applicable law or agreed to in writing, software
#  distributed under the License is distributed on an "AS IS" BASIS,
#  WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
#  See the License for the specific language governing permissions and
#  limitations under the License.
#
"""
Crawler for gcjyzx.zhangzhou.gov.cn 工程信息 (jyxx) section.

Site: EpointWebBuilder v5.0 — SSR listing + REST API pagination.

Data flow
─────────
  1. **Listing**: POST /EpointWebBuilder/rest/secaction/getSecInfoListYzm
     → { custom: { count, infodata: [...] } }
     Filter by today's date to get published articles.
  2. **Detail**: GET /gcxx/{path}/{date}/{uuid}.html  (SSR HTML)
     → Parse title, metadata, content table, project lifecycle tabs, attachments.
  3. **Related**: GET /json/{relationguid}.json
     → All project stages (招标公告..合同签署) for the same project.
  4. **Attachments**: GET /zzggzy/{infoid}/{fileuuid}/{filename}
     → Binary download (PDF, DOC, etc.), ZIP auto-extraction.

Project lifecycle "tabs" on detail pages are links to separate detail pages
(not in-page tabs).  The relationguid groups them into one project.

Usage
-----
    python gcjyzx_jyxx_crawler.py \
        --tenant-id <TENANT_ID> \
        --kb-id <KB_ID> \
        --task-name <NAME>

    # Optional:
        --max-articles 50       # Limit records (0=unlimited)
        --full                  # Ignore saved state, re-crawl
        --max-runtime 3300      # Max runtime before graceful stop (default: 55 min)
"""

import argparse
import json
import logging
import os
import random
import re
import sys
import time
import urllib.parse
import urllib.request
import zipfile
from datetime import datetime, date, timedelta

_SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
_PROJECT_ROOT = os.path.abspath(os.path.join(_SCRIPT_DIR, "..", ".."))
sys.path.insert(0, _PROJECT_ROOT)

from common import settings
from common.log_utils import init_root_logger
from common.misc_utils import get_uuid
from bs4 import BeautifulSoup


# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------

_SITE_ROOT = "https://gcjyzx.zhangzhou.gov.cn"
_LISTING_API = _SITE_ROOT + "/EpointWebBuilder/rest/secaction/getSecInfoListYzm"
_RELATED_JSON_URL = _SITE_ROOT + "/json"
_SITE_GUID = "7eb5f7f1-9041-43ad-8e13-8fcb82ea831a"
_PAGE_SIZE = 20

# Checkpoint batch size (articles)
_BATCH_SIZE = 10

# Default max runtime (55 min, 5 min margin)
_MAX_RUNTIME_DEFAULT = 3300

# Anti-crawling delays (seconds)
_REQUEST_DELAY_MIN = 0.3
_REQUEST_DELAY_MAX = 1.0

# State filename
_STATE_FILENAME = "_crawler_state.json"

_USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/120.0.0.0 Safari/537.36"
)

# Category display labels
_CATEGORY_LABELS = {
    "001001001": "招标公告",
    "001001002": "招标答疑",
    "001001003": "最高限价",
    "001001004": "开标一览表",
    "001001005": "中标候选人公示",
    "001001006": "中标公告",
    "001001007": "合同签署",
}

_ATTACH_EXTENSIONS = {
    ".pdf", ".doc", ".docx", ".xls", ".xlsx", ".ppt", ".pptx",
    ".zip", ".rar", ".7z", ".txt",
}


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _safe_print(msg):
    try:
        print(msg)
    except UnicodeEncodeError:
        print(msg.encode("gbk", errors="replace").decode("gbk"))


def _sanitize_filename(text, max_len=120):
    if not text:
        return "untitled"
    name = re.sub(r'[\\/:*?"<>|]', "_", str(text).strip())
    name = re.sub(r'\s+', "_", name).strip("._ ")
    return (name or "untitled")[:max_len]


def _extract_onclick_path(onclick):
    """Extract the file path from ztbfjyz('/zzggzy/...','0','1') onclick handler."""
    if not onclick:
        return None
    m = re.search(r"""ztbfjyz\s*\(\s*['"]([^'"]+)['"]""", onclick)
    return m.group(1) if m else None


# ---------------------------------------------------------------------------
# HTTP helpers
# ---------------------------------------------------------------------------

def _http_get(url, referer=None, timeout=30):
    """GET request returning bytes or None."""
    req = urllib.request.Request(url)
    req.add_header("User-Agent", _USER_AGENT)
    if referer:
        req.add_header("Referer", referer)
    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            return resp.read()
    except Exception as e:
        logging.error("GET %s failed: %s", url, e)
        return None


def _http_get_json(url, referer=None, timeout=30):
    """GET returning parsed JSON or None."""
    body = _http_get(url, referer=referer, timeout=timeout)
    if body is None:
        return None
    try:
        return json.loads(body.decode("utf-8"))
    except Exception as e:
        logging.error("JSON parse error for %s: %s", url, e)
        return None


def _http_post_form(url, data_dict, referer=None, timeout=30):
    """POST form-urlencoded, return parsed JSON or None."""
    params = urllib.parse.urlencode(data_dict).encode()
    req = urllib.request.Request(url, data=params, method="POST")
    req.add_header("Content-Type", "application/x-www-form-urlencoded")
    req.add_header("User-Agent", _USER_AGENT)
    if referer:
        req.add_header("Referer", referer)
    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            body = resp.read().decode("utf-8")
            return json.loads(body)
    except Exception as e:
        logging.error("POST %s failed: %s", url, e)
        return None


def _download_binary(url, referer=None, timeout=60):
    """Download binary content, return bytes or None."""
    req = urllib.request.Request(url)
    req.add_header("User-Agent", _USER_AGENT)
    if referer:
        req.add_header("Referer", referer)
    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            return resp.read()
    except Exception as e:
        logging.error("Download failed: %s — %s", url[:120], e)
        return None


# ---------------------------------------------------------------------------
# Listing API
# ---------------------------------------------------------------------------

def _fetch_listing_page(startdate="", enddate="", category_num="",
                        page_index=0, page_size=_PAGE_SIZE):
    """Fetch one page of engineering-info listing.

    Returns (total_count, items_list).
    Each item: {categorynum, infoid, title, infourl, infodate, projectaddressname, ...}
    """
    data = {
        "siteGuid": _SITE_GUID,
        "categoryNum": category_num,
        "pageIndex": str(page_index),
        "pageSize": str(page_size),
        "startdate": startdate,
        "enddate": enddate,
        "projectaddress": "",
        "bdsecondtype": "",
    }
    result = _http_post_form(
        _LISTING_API, data,
        referer=_SITE_ROOT + "/gcxx/jyxx.html",
    )
    if not result:
        return 0, []
    custom = result.get("custom", {})
    return custom.get("count", 0), custom.get("infodata", []) or []


# ---------------------------------------------------------------------------
# Detail page parsing
# ---------------------------------------------------------------------------

def _parse_detail_html(html_bytes, detail_url):
    """Parse server-rendered detail page.

    Returns dict with:
        title, pub_date, info_source, relationguid, infoid,
        content_html, content_text, tabs, attachments
    """
    if not html_bytes:
        return None

    html = html_bytes.decode("utf-8", errors="replace")
    soup = BeautifulSoup(html, "html.parser")

    info = {
        "title": "",
        "pub_date": "",
        "info_source": "",
        "relationguid": "",
        "infoid": "",
        "content_html": "",
        "content_text": "",
        "tabs": [],
        "attachments": [],
    }

    # Title
    title_el = soup.select_one("h3.title")
    if title_el:
        info["title"] = title_el.get_text(strip=True)

    # Hidden fields
    rel_el = soup.select_one("#relationguid")
    if rel_el:
        info["relationguid"] = rel_el.get_text(strip=True)
    infoid_el = soup.select_one("#infoid")
    if infoid_el:
        info["infoid"] = infoid_el.get_text(strip=True)

    # Toolbar metadata
    for col in soup.select(".toolbar .col"):
        text = col.get_text(strip=True)
        if "发布时间" in text:
            info["pub_date"] = text.replace("发布时间：", "").strip()
        elif "信息来源" in text:
            info["info_source"] = text.replace("信息来源：", "").strip()

    # Content area (strip tab navigation + scripts)
    content_main = soup.select_one("#infoContentM")
    if content_main:
        for tag in content_main.select(".link"):
            tag.decompose()
        for tag in content_main.select("script"):
            tag.decompose()
        info["content_html"] = content_main.decode_contents()
        info["content_text"] = content_main.get_text(separator="\n", strip=True)

    # Project lifecycle tabs (环节)
    for a_tag in soup.select("#link-info a.link-box"):
        stage_name = a_tag.get_text(strip=True)
        stage_url = a_tag.get("data-value", "")
        stage_status = a_tag.get("data-status", "0")  # 0=not started, 1=done, 2=current
        classes = a_tag.get("class", []) or []
        is_current = "current" in classes
        info["tabs"].append({
            "name": stage_name,
            "url": stage_url,
            "status": stage_status,
            "current": is_current,
        })

    # Attachments
    for a_tag in soup.select(".article-fj .attachment-list a[onclick]"):
        onclick = a_tag.get("onclick", "")
        path = _extract_onclick_path(onclick)
        filename = a_tag.get("title", "") or a_tag.get_text(strip=True)
        if path and filename:
            info["attachments"].append({
                "path": path,
                "filename": filename,
                "download_url": f"{_SITE_ROOT}{path}",
            })

    return info


# ---------------------------------------------------------------------------
# Related articles
# ---------------------------------------------------------------------------

def _fetch_related_articles(relationguid):
    """Fetch project-stage articles linked to the same project.

    Returns list[dict] with: categorynum, infoid, infourl, title, infodate
    """
    if not relationguid:
        return []
    url = f"{_RELATED_JSON_URL}/{relationguid}.json"
    data = _http_get_json(url, referer=_SITE_ROOT + "/")
    if not data:
        return []
    return data.get("custom", {}).get("list", []) or []


# ---------------------------------------------------------------------------
# Attachment download + ZIP extraction
# ---------------------------------------------------------------------------

def _download_attachments(attachments, dest_dir, infoid):
    """Download all attachments for an article. Returns list of local file paths."""
    downloaded = []
    if not attachments:
        return downloaded

    os.makedirs(dest_dir, exist_ok=True)

    for att in attachments:
        path = att["path"]
        filename = att["filename"]

        # URL-encode the filename portion
        parts = path.rsplit("/", 1)
        encoded_path = (
            parts[0] + "/" + urllib.parse.quote(parts[1])
            if len(parts) > 1 else path
        )
        url = f"{_SITE_ROOT}{encoded_path}"
        dest_path = os.path.join(dest_dir, filename)

        if os.path.exists(dest_path):
            downloaded.append(dest_path)
            continue

        referer = _SITE_ROOT + "/"
        body = _download_binary(url, referer=referer)
        if body and len(body) >= 100:
            with open(dest_path, "wb") as f:
                f.write(body)
            downloaded.append(dest_path)
            logging.info("Downloaded %s (%d bytes)", filename, len(body))
        else:
            logging.warning("Download failed/too small: %s", filename)

    return downloaded


def _extract_zip(zip_path):
    """Extract a ZIP file; returns list of extracted paths. ZIP is removed after."""
    extracted = []
    dest_dir = os.path.dirname(zip_path)
    try:
        with zipfile.ZipFile(zip_path, "r") as zf:
            for name in zf.namelist():
                safe_name = re.sub(r'[\\/:*?"<>|]', "_", os.path.basename(name))
                if not safe_name:
                    continue
                dest_path = os.path.join(dest_dir, safe_name)
                if os.path.exists(dest_path):
                    continue
                with open(dest_path, "wb") as f:
                    f.write(zf.read(name))
                extracted.append(dest_path)
        os.remove(zip_path)
    except Exception as e:
        logging.warning("ZIP extract error for %s: %s", zip_path, e)
    return extracted


# ---------------------------------------------------------------------------
# Text extraction from binary files
# ---------------------------------------------------------------------------

def _extract_text_from_file(filepath):
    """Extract plain text from PDF / DOCX / XLSX / TXT files."""
    ext = os.path.splitext(filepath)[1].lower()
    try:
        if ext == ".txt":
            with open(filepath, "r", encoding="utf-8", errors="replace") as f:
                return f.read()
        elif ext == ".pdf":
            import pdfplumber
            parts = []
            with pdfplumber.open(filepath) as pdf:
                for pg in pdf.pages:
                    text = pg.extract_text()
                    if text:
                        parts.append(text)
            return "\n\n".join(parts)
        elif ext in (".docx", ".doc"):
            import docx
            doc = docx.Document(filepath)
            return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        elif ext in (".xlsx", ".xls"):
            import openpyxl
            wb = openpyxl.load_workbook(filepath, read_only=True)
            parts = []
            for ws in wb.worksheets:
                rows = []
                for row in ws.iter_rows(values_only=True):
                    rows.append(" | ".join(str(c) if c is not None else "" for c in row))
                if rows:
                    parts.append(f"### {ws.title}\n" + "\n".join(rows))
            wb.close()
            return "\n\n".join(parts)
    except Exception as e:
        logging.warning("Text extraction failed for %s: %s", filepath, e)
    return None


# ---------------------------------------------------------------------------
# Markdown builder
# ---------------------------------------------------------------------------

def _build_markdown(detail, related_articles, download_dir, source_url):
    """Build a Markdown document from article detail + related + attachments."""
    info = detail or {}
    title = info.get("title", "无标题")
    pub_date = info.get("pub_date", "")
    info_source = info.get("info_source", "")

    lines = [
        f"# {title}",
        "",
        f"**数据来源:** 漳州市工程项目交易中心",
        f"**页面地址:** {source_url}",
        f"**抓取时间:** {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
    ]
    if pub_date:
        lines.append(f"**发布时间:** {pub_date}")
    if info_source:
        lines.append(f"**信息来源:** {info_source}")
    lines.append("")

    # ── Project lifecycle (tabs) ──
    tabs = info.get("tabs", [])
    if tabs:
        lines.append("---")
        lines.append("")
        lines.append("## 项目环节")
        lines.append("")
        for tab in tabs:
            stage_mark = " ← 当前" if tab["current"] else ""
            status_text = {
                "0": "未完成", "1": "已完成", "2": "进行中",
            }.get(tab["status"], "")
            if tab["url"]:
                lines.append(
                    f"- [{tab['name']}]({_SITE_ROOT}{tab['url']})"
                    f"（{status_text}）{stage_mark}"
                    if status_text
                    else f"- [{tab['name']}]({_SITE_ROOT}{tab['url']}){stage_mark}"
                )
            else:
                lines.append(f"- {tab['name']}（{status_text}）{stage_mark}")
        lines.append("")

    # ── Related articles ──
    if related_articles:
        lines.append("---")
        lines.append("")
        lines.append("## 关联文章")
        lines.append("")
        for ra in related_articles:
            cat_label = _CATEGORY_LABELS.get(
                ra.get("categorynum", ""), ra.get("categorynum", "")
            )
            ra_title = ra.get("title", "")
            ra_url = ra.get("infourl", "")
            ra_date = (ra.get("infodate") or "")[:10]
            if ra_url:
                lines.append(
                    f"- [{cat_label}] [{ra_title}]({_SITE_ROOT}{ra_url}) ({ra_date})"
                )
            else:
                lines.append(f"- [{cat_label}] {ra_title} ({ra_date})")
        lines.append("")

    # ── Main content ──
    content_text = info.get("content_text", "")
    if content_text:
        lines.append("---")
        lines.append("")
        lines.append("## 详细内容")
        lines.append("")
        content_clean = re.sub(r"\n{3,}", "\n\n", content_text)
        lines.append(content_clean)
        lines.append("")

    # ── Attachments ──
    attachments = info.get("attachments", [])
    if attachments:
        lines.append("---")
        lines.append("")
        lines.append("## 附件")
        lines.append("")
        for att in attachments:
            fname = att.get("filename", "unknown")
            lines.append(f"- {fname}")
        lines.append("")

        # Embed extracted attachment text
        if download_dir and os.path.isdir(download_dir):
            lines.append("### 附件内容")
            lines.append("")
            for att in attachments:
                fname = att.get("filename", "")
                local_path = os.path.join(download_dir, fname)
                if not os.path.exists(local_path):
                    continue
                lines.append(f"#### {fname}")
                lines.append("")
                text = _extract_text_from_file(local_path)
                if text and text.strip():
                    if len(text) > 50000:
                        text = text[:50000] + "\n\n（内容过长，已截断）"
                    lines.append(text)
                else:
                    lines.append("（无法提取文本内容）")
                lines.append("")

    return "\n".join(lines)


# ---------------------------------------------------------------------------
# State management
# ---------------------------------------------------------------------------

def _load_state(output_dir):
    path = os.path.join(output_dir, _STATE_FILENAME)
    if os.path.exists(path):
        try:
            with open(path, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception as e:
            logging.warning("Failed to load crawler state: %s", e)
    return {"processed_ids": [], "completed": False}


def _save_state(output_dir, state):
    path = os.path.join(output_dir, _STATE_FILENAME)
    os.makedirs(output_dir, exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(state, f, ensure_ascii=False)


# ---------------------------------------------------------------------------
# KB upload
# ---------------------------------------------------------------------------

def _upload_to_kb(md_content, attachment_files, kb_id, tenant_id, folder_name):
    """Upload markdown + attachment files to knowledge base."""
    from api.db.services.knowledgebase_service import KnowledgebaseService
    from api.db.services.file_service import FileService
    from api.db.services.document_service import DocumentService

    ok, kb = KnowledgebaseService.get_by_id(kb_id)
    if not ok:
        raise LookupError(f"Knowledge base {kb_id} not found")

    # Upload MD
    class _FO:
        def __init__(self, fn, b):
            self.id = get_uuid()
            self.filename = fn
            self.blob = b

        def read(self):
            return self.blob

    fo = _FO(f"{folder_name}.md", md_content.encode("utf-8"))
    errs, pairs = FileService.upload_document(kb, [fo], tenant_id)
    if errs:
        logging.warning("MD upload errors: %s", errs)
    for doc, _ in pairs:
        did = doc["id"]
        try:
            DocumentService.begin2parse(did)
            DocumentService.run(tenant_id, doc, {})
        except Exception as e:
            logging.error("Queue parse for %s: %s", did, e)

    # Upload attachment files
    for fp in attachment_files:
        fname = os.path.basename(fp)
        with open(fp, "rb") as f:
            blob = f.read()
        fo2 = _FO(fname, blob)
        errs2, pairs2 = FileService.upload_document(kb, [fo2], tenant_id)
        if errs2:
            logging.warning("Attachment upload errors: %s", errs2)
        for doc, _ in pairs2:
            did = doc["id"]
            try:
                DocumentService.begin2parse(did)
                DocumentService.run(tenant_id, doc, {})
            except Exception as e:
                logging.error("Queue parse for %s: %s", did, e)


# ---------------------------------------------------------------------------
# Argument parsing
# ---------------------------------------------------------------------------

def parse_args():
    p = argparse.ArgumentParser(
        description="gcjyzx.zhangzhou.gov.cn jyxx crawler — 漳州市工程项目交易信息"
    )
    p.add_argument("--tenant-id", required=True)
    p.add_argument("--target-url",
                   default="https://gcjyzx.zhangzhou.gov.cn/gcxx/jyxx.html")
    p.add_argument("--kb-id", required=True)
    p.add_argument("--task-name", required=True)
    p.add_argument("--output-dir", default=None,
                   help="Output root directory (default: PROJECT_ROOT/rag/<task_name>)")
    p.add_argument("--full", action="store_true",
                   help="Ignore saved state and re-crawl all")
    p.add_argument("--max-runtime", type=int, default=_MAX_RUNTIME_DEFAULT,
                   help="Max runtime in seconds (default: 3300)")
    p.add_argument("--max-articles", type=int, default=0,
                   help="Max articles to fetch (0 = unlimited)")
    # Legacy/unused args
    for opt in ("--section", "--max-days", "--hours",
                "--llm-id", "--llm-model", "--access-token"):
        p.add_argument(opt, default=None)
    return p.parse_args()


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    args = parse_args()

    _safe_print("\n" + "=" * 60)
    _safe_print("[GCJYZX-JYXX] 漳州市工程项目交易中心 - 工程信息 crawler")
    _safe_print("[GCJYZX-JYXX] Target: {}".format(args.target_url))
    _safe_print("[GCJYZX-JYXX] KB: {}".format(args.kb_id))
    _safe_print("[GCJYZX-JYXX] Task: {}".format(args.task_name))
    _safe_print("[GCJYZX-JYXX] Max articles: {}".format(
        args.max_articles if args.max_articles else "unlimited"))
    _safe_print("[GCJYZX-JYXX] Max runtime: {}s".format(args.max_runtime))
    _safe_print("=" * 60 + "\n")
    sys.stdout.flush()

    settings.init_settings()
    logging.info("=== GCJYZX-JYXX crawler started ===")

    output_dir = args.output_dir or os.path.join(
        _PROJECT_ROOT, "rag", args.task_name.strip()
    )
    os.makedirs(output_dir, exist_ok=True)
    _safe_print("[GCJYZX-JYXX] Output: {}\n".format(output_dir))
    sys.stdout.flush()

    # ── State ──────────────────────────────────────────────────────────
    state = _load_state(output_dir) if not args.full else {
        "processed_ids": [], "completed": False,
    }
    processed_ids = set(state.get("processed_ids", []))

    if state.get("completed"):
        _safe_print("[GCJYZX-JYXX] Already completed, nothing to do.")
        sys.stdout.flush()
        return

    _safe_print("[GCJYZX-JYXX] Already processed: {} article(s)".format(
        len(processed_ids)))
    sys.stdout.flush()

    crawl_start = time.time()

    # If started before 8 AM, crawl yesterday's data (today's may not be published yet)
    now = datetime.now()
    if now.hour < 8:
        target_date = (now - timedelta(days=1)).strftime("%Y-%m-%d")
        _safe_print("[GCJYZX-JYXX] Before 8 AM, targeting yesterday: {}".format(target_date))
    else:
        target_date = now.strftime("%Y-%m-%d")

    startdate = f"{target_date} 00:00:00"
    enddate = f"{target_date} 23:59:59"

    # ── Step 1: Fetch all listing pages ────────────────────────────────
    _safe_print("[GCJYZX-JYXX] Fetching listing for {}...".format(target_date))
    sys.stdout.flush()

    total, first_page = _fetch_listing_page(
        startdate=startdate, enddate=enddate, page_index=0,
    )

    if total == 0:
        _safe_print("[GCJYZX-JYXX] No articles published for {} . Exiting.".format(target_date))
        sys.stdout.flush()
        return

    _safe_print("[GCJYZX-JYXX] Total: {} article(s) for {}".format(total, target_date))
    sys.stdout.flush()

    # Collect all listing items
    all_items = list(first_page)
    page_index = 1
    while len(all_items) < total:
        if args.max_articles and len(all_items) >= args.max_articles:
            break
        _, page_items = _fetch_listing_page(
            startdate=startdate, enddate=enddate, page_index=page_index,
        )
        if not page_items:
            break
        all_items.extend(page_items)
        page_index += 1
        time.sleep(random.uniform(_REQUEST_DELAY_MIN, _REQUEST_DELAY_MAX))

    if args.max_articles and len(all_items) > args.max_articles:
        all_items = all_items[:args.max_articles]

    _safe_print("[GCJYZX-JYXX] Fetched {} items across {} page(s)".format(
        len(all_items), page_index))
    sys.stdout.flush()

    # Filter already-processed
    new_items = []
    for item in all_items:
        infoid = item.get("infoid", "")
        if infoid and infoid not in processed_ids:
            new_items.append(item)

    skipped = len(all_items) - len(new_items)
    if skipped:
        _safe_print("[GCJYZX-JYXX] {} already processed, {} new".format(
            skipped, len(new_items)))
        sys.stdout.flush()

    if not new_items:
        _safe_print("[GCJYZX-JYXX] No new articles to process.")
        sys.stdout.flush()
        if not state.get("completed"):
            state["completed"] = True
            _save_state(output_dir, state)
        return

    # ── Step 2: Process each article ───────────────────────────────────
    _safe_print("[GCJYZX-JYXX] Processing {} article(s)...\n".format(len(new_items)))
    sys.stdout.flush()

    processed_count = 0
    stopped_early = False
    downloads_dir = os.path.join(output_dir, "downloads")

    for idx, item in enumerate(new_items, 1):
        # ── Time-bounded check ─────────────────────────────────────────
        elapsed = time.time() - crawl_start
        if elapsed > args.max_runtime - 120:
            _safe_print(
                "\n[GCJYZX-JYXX] Runtime {:.0f}s approaching limit ({}s), "
                "stopping gracefully. {} processed. "
                "Next run will resume.".format(elapsed, args.max_runtime, processed_count))
            sys.stdout.flush()
            stopped_early = True
            break

        infoid = item.get("infoid", "")
        title = item.get("title", "无标题")
        infourl = item.get("infourl", "")
        detail_url = _SITE_ROOT + infourl if infourl else ""

        _safe_print("[GCJYZX-JYXX] [{}/{}] {}...".format(
            idx, len(new_items), title[:50]))
        sys.stdout.flush()

        if not detail_url:
            logging.warning("No detail URL for %s, skipping", infoid)
            processed_ids.add(infoid)
            continue

        # Fetch detail page
        html_bytes = _http_get(detail_url, referer=args.target_url)
        if not html_bytes:
            logging.warning("Failed to fetch detail: %s", detail_url)
            processed_ids.add(infoid)
            continue

        detail = _parse_detail_html(html_bytes, detail_url)
        if not detail:
            logging.warning("Failed to parse detail: %s", detail_url)
            processed_ids.add(infoid)
            continue

        # Fetch related articles (same project, different stages)
        relguid = detail.get("relationguid", "")
        related = _fetch_related_articles(relguid) if relguid else []

        # Download attachments
        attachments = detail.get("attachments", [])
        local_files = []
        if attachments:
            article_dl_dir = os.path.join(downloads_dir, _sanitize_filename(infoid, 40))
            local_files = _download_attachments(attachments, article_dl_dir, infoid)
            # Extract ZIPs
            for fp in list(local_files):
                is_zip = fp.lower().endswith(".zip")
                if not is_zip and os.path.exists(fp) and os.path.getsize(fp) >= 4:
                    with open(fp, "rb") as f:
                        is_zip = f.read(4) == b"PK\x03\x04"
                if is_zip:
                    extracted = _extract_zip(fp)
                    local_files.remove(fp)
                    local_files.extend(extracted)

        # Build markdown
        md_content = _build_markdown(
            detail, related,
            os.path.join(downloads_dir, _sanitize_filename(infoid, 40))
            if attachments else "",
            detail_url,
        )

        # Save markdown locally
        folder_name = _sanitize_filename(
            "{}_{}_{}".format(target_date, infoid[:12], title[:40]), max_len=120)
        md_path = os.path.join(output_dir, f"{folder_name}.md")
        with open(md_path, "w", encoding="utf-8") as f:
            f.write(md_content)
        _safe_print("[GCJYZX-JYXX]   Saved MD ({} chars, {} attachments)".format(
            len(md_content), len(local_files)))
        sys.stdout.flush()

        # Upload to KB
        if args.kb_id and md_content:
            try:
                _upload_to_kb(md_content, local_files, args.kb_id,
                             args.tenant_id, folder_name)
                _safe_print("[GCJYZX-JYXX]   Uploaded to KB {}".format(args.kb_id))
                sys.stdout.flush()
            except Exception as e:
                logging.error("KB upload failed: %s", e)
                _safe_print("[GCJYZX-JYXX]   Upload error: {}".format(e))
                sys.stdout.flush()

        processed_ids.add(infoid)
        processed_count += 1

        # Checkpoint every batch
        if processed_count % _BATCH_SIZE == 0:
            _save_state(output_dir, {
                "processed_ids": list(processed_ids),
                "completed": False,
            })
            _safe_print("[GCJYZX-JYXX]   Checkpoint ({} processed)".format(
                processed_count))
            sys.stdout.flush()

        # Anti-crawling delay
        time.sleep(random.uniform(_REQUEST_DELAY_MIN, _REQUEST_DELAY_MAX))

    # ── Final state ────────────────────────────────────────────────────
    if not stopped_early:
        state["completed"] = True
    state["processed_ids"] = list(processed_ids)
    _save_state(output_dir, state)

    _safe_print("\n" + "=" * 60)
    _safe_print("[GCJYZX-JYXX] Crawl complete — {} new article(s)".format(
        processed_count))
    if stopped_early:
        _safe_print("[GCJYZX-JYXX] Stopped early, will resume next run")
    _safe_print("=" * 60 + "\n")
    sys.stdout.flush()
    logging.info("=== GCJYZX-JYXX crawler finished ===")


if __name__ == "__main__":
    CONSUMER_NAME = "gcjyzx_jyxx_crawler"
    init_root_logger(CONSUMER_NAME)
    main()
