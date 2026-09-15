"""docx 模板工具：段落遍历（含表格内段落）、填写点候选提取、锚文本→占位符替换。

addr 定位约定：正文段落 para:<idx>；表格内段落 cell:<tbl_no>:<row>:<col>:<para_idx>
（tbl_no 为正文第几个**顶层**表格，0 起——存量模板 addr 兼容依赖此规则不变）。
缺 tbl_no 时多表格文档的 cell(r,c,p) 会跨表撞号（不同表格同位置段落共享 addr，
validate/render 按 addr 查到的段落错位），故必须带表序号。嵌套表格在父单元格 addr
后追加 ":t<j>" 段并递归（cell:0:1:2:t0:0:1:0 = 顶层表 0 的 (1,2) 单元格内第 0 个
嵌套表的 (0,1) 单元格第 0 段）。合并单元格（gridSpan）同一 tc 只按首现坐标编址一次
（此前重复编址会让同一文本多处进候选，anchor 反查报「匹配到多处」）。
文本框段落 <父addr>:tx<k>:<pi>、内容控件段落 body 直系 sdt:<k>:<pi> / 嵌套
<父addr>:sdt<k>:<pi>——新区域（文本框/sdt）只消耗扁平 index，不挤占存量
para: 序号（否则含文本框/sdt 的存量模板升级后 addr 整体错位、同形 anchor
静默错填）。注意表格 cell 段落**照旧消耗** para: 计数——存量基线用单一扁平
序号编正文段，cell 段落本就消耗它，「表格在前、段落在后」的范本升级后段落
编号才不会前移（walker 用 legacy 标志区分这两类计数消耗）。
index 为全文档扁平序号，与 addr 一一对应。

跨 run 替换取舍：普通段落跨 run 时只重写 anchor 覆盖的 run 区间，区间外格式保留
（`_replace_cross_run_in_place`）；但含超链接（w:hyperlink）/简单域（w:fldSimple）的段落**不可**整段重写——
python-docx 1.1+ 的 Paragraph.text 包含超链接内文本，而 Paragraph.runs 不包含，
整段重写会把超链接文本复制进首 run 且原节点仍在（内容重复）。这类段落改走
run 拼接替换路径（见 _replace_via_run_concat）。
"""
import io
import logging
import re

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

logger = logging.getLogger(__name__)

# 常见填写点特征：下划线空位 / 中文括号空位 / 【】空位 / ×× 占位 / 括号内"填写"提示 / "：" 结尾（冒号后留白）。
# 注意："填写"仅在括号内（如"（请填写）"）才算特征，避免正文说明文字（"应如实填写"）误报污染候选集。
# 补充模式（标准范本实测漏报修复）：
# - ":[ \t\u3000]{3,}\S"：冒号+留白+后续文字（"编制日期：　　年　月　日"、"招标人：　　（盖单位电子公章）"等冒号不在行尾的填写点）
# - "[ \u3000]{2,}年[ \u3000]*月[ \u3000]*日"：「　年　月　日」空白日期占位（要求年前有留白，
#   真实日期"2026年9月11日"不含留白不误报）
# - "\S[ \u3000]{6,}\S"：行内长空白占位（"本招标项目　　（项目名称）　已由　　（审批机关）"跨栏留白）
FILL_HINT_RE = re.compile(
    r"(_{2,}|＿{2,}|（\s*）|\(\s*\)|【\s*】|×{2,}|XX{1,}|xx{1,}|□"
    r"|[（(][^（）()]*填写[^（）()]*[)）]|：\s*$|:\s*$"
    r"|[:：][ \t\u3000]{3,}\S"
    r"|[ \u3000]{2,}年[ \u3000]*月[ \u3000]*日"
    r"|\S[ \u3000]{6,}\S)"
)

# 手动占位符：用户在模板正文里手写的 {{snake_key}}（与 renderer/docxtpl、前端实时预览同口径）
PH_RE = re.compile(r"\{\{([a-z][a-z0-9_]*)\}\}")

TXBX_DEPTH_LIMIT = 8  # 文本框嵌套深度上限：超限子树跳过（畸形 XML 防爆栈）

# python-docx nsmap 不含 mc 前缀（qn("mc:*") KeyError），markup-compatibility 命名空间自备常量
NS_MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"


def _iter_txbx_content(el) -> list:
    """深度优先收集 el 下全部 w:txbxContent（文本框内容根元素）。

    mc:AlternateContent 只下钻 mc:Choice、跳过 mc:Fallback——Word 对浮动文本框
    常存双份（Choice=wps、Fallback=VML），双份都收会让同一段落重复进候选
    （anchor 反查「匹配到多处」歧义 + LLM token 浪费）。"""
    out = []
    fallback_tag = f"{{{NS_MC}}}Fallback"

    def _walk(node):
        for child in node:
            tag = child.tag
            if tag == fallback_tag:
                continue
            if tag == qn("w:txbxContent"):
                out.append(child)
                continue  # 框内段落由 _walk_txbx 逐段处理（更深嵌套文本框随之发现）
            _walk(child)

    _walk(el)
    return out


def _walk_hf_blocks(hf_el, prefix, doc, walk_paragraph, walk_table):
    """页眉/页脚 part 编址：直系 w:p → <prefix>:<pi>（part 内局部序号）；
    直系 w:tbl → <prefix>:cell:<t>:...。hf_el 为该 part 根元素
    （python-docx _BaseHeaderFooter._element，is_linked_to_previous 为 False
    时必有定义，不会触发自动补建副作用）。walk_paragraph/walk_table 由调用方
    闭包注入，均保持默认 legacy=False——页眉页脚是新区域，零消耗存量
    para:/cell: 计数器（红线：存量模板加页眉页脚后正文 addr 序列逐字节不变）。
    模块级以便单测 monkeypatch 打桩（畸形 part 注入异常验证隔离）。"""
    pi = 0
    tbl_no = 0
    for block in hf_el.iterchildren():
        if block.tag == qn("w:p"):
            walk_paragraph(Paragraph(block, doc), f"{prefix}:{pi}")
            pi += 1
        elif block.tag == qn("w:tbl"):
            walk_table(Table(block, doc), f"{prefix}:cell:{tbl_no}")
            tbl_no += 1


def _build_addr_map(doc):
    """遍历文档全部可填写段落（正文/内容控件/表格/文本框/页眉页脚），
    返回 ({addr: Paragraph}, items)。

    编址（存量规则不变，新增前缀均为增量，旧 addr 永不复用）：
    - 正文段落 para:<idx>；表格内段落 cell:<tbl_no>:<row>:<col>:<para_idx>
      （tbl_no 只数顶层表格；嵌套表格在父单元格 addr 后追加 ":t<j>" 段递归）；
      合并单元格（gridSpan 横向合并）同一 tc 只按首现坐标编址一次。
      body 直系 w:p 与存量表格 cell 段落（legacy=True）共同消耗 para: 计数器
      ——存量基线用单一扁平序号编正文段，cell 段落本就消耗它；嵌套 ":t<j>"
      表内段落属本次新增编址范围（存量基线不编址嵌套表），不消耗。
    - 文本框段落 <父addr>:tx<k>:<pi>（k 为该段落内第 k 个文本框）；
      框内表格 <父addr>:tx<k>:cell:<t>:...；嵌套文本框继续追加 :tx<j> 段。
      mc:AlternateContent 只取 mc:Choice（见 _iter_txbx_content）。
      文本框/内容控件段落（新区域，legacy=False）只消耗扁平 index，不碰存量计数器
      （I-1 红线：含文本框/sdt 的存量模板升级后 para:/cell: 序列逐字节不变）。
    - w:sdt 内容控件独立前缀：body 直系 sdt:<k>:<pi>（内含表格 sdt:<k>:cell:...）；
      cell/文本框/嵌套内 sdt 追加 :sdt<k>: 段。
    - 页眉/页脚段落 hdr:<sec>:<pi> / ftr:<sec>:<pi>（sec 为节序号，pi 为 part 内
      局部段号）；part 内表格 …:cell:<t>:…；首页/偶数页变体追加 :first / :even 段
      （同节多类页眉同时 unlinked 若共用前缀会撞号，addr 反查段落错位）。
      linked 节与共享 part（partname 去重）只编址一次；页眉页脚段落属新区域，
      不消耗存量 para: 计数器。
    items: [{"index": 扁平序号, "text": 段落文本, "addr": 定位串}]，按文档顺序。
    """
    addr_map = {}
    items = []
    idx = 0        # 全文档扁平序号（items 排序用；新区域段落一并递增）
    para_seq = -1  # 存量段号：body 直系 w:p 与存量表格 cell 段落共同消耗（legacy 标志区分）
    tbl_no = -1    # 存量顶层表格号：只随 body 直系 w:tbl 递增
    sdt_no = -1    # body 直系 sdt 序号（独立 sdt: 前缀）

    def _emit(p, addr):
        nonlocal idx
        addr_map[addr] = p
        items.append({"index": idx, "text": p.text, "addr": addr})
        idx += 1

    def _walk_sdt(sdt_el, prefix, depth):
        """内容控件展开：body 直系 → sdt:<k>:...；cell/文本框/嵌套内 →
        <父addr>:sdt<k>:...。内部段落/表格不消耗任何存量计数器（legacy=False，
        sdt 是新区域）；sdt 无 sdtContent（畸形）→ debug 日志跳过。"""
        if depth > TXBX_DEPTH_LIMIT:
            logger.warning("sdt nesting deeper than %d at %s, skipped",
                           TXBX_DEPTH_LIMIT, prefix)
            return
        content = sdt_el.find(qn("w:sdtContent"))
        if content is None:
            logger.debug("w:sdt without sdtContent at %s, skipped", prefix)
            return
        pi = 0
        tbl_k = 0
        sdt_k = 0
        for block in content.iterchildren():
            if block.tag == qn("w:p"):
                _walk_paragraph(Paragraph(block, doc), f"{prefix}:{pi}", depth)
                pi += 1
            elif block.tag == qn("w:tbl"):
                _walk_table(Table(block, doc), f"{prefix}:cell:{tbl_k}", depth,
                            legacy=False)
                tbl_k += 1
            elif block.tag == qn("w:sdt"):
                _walk_sdt(block, f"{prefix}:sdt{sdt_k}", depth + 1)
                sdt_k += 1

    def _walk_table(tbl, prefix, depth=0, legacy=False):
        """表格编址。legacy=True（仅 body 直系存量表格，调用点须显式传参）：cell
        直系段落消耗 para_seq——存量基线用单一扁平序号编正文段，cell 段落本就
        消耗它，表格前置范本升级后段落编号不变。默认 legacy=False 是安全侧取值：
        Task 2 页眉页脚 walker 复用本函数时若漏传 legacy，只会让新区域段落不消耗
        存量计数（安全），而非静默前移存量编号。嵌套 ":t<j>" 表一律 legacy=False：
        其段落属本次新增编址范围（存量基线不编址嵌套表、不计数），透传 legacy
        会让含嵌套表的存量范本段落编号前移。cell 直系 sdt 一律 legacy=False
        （新区域）且深度 +1；嵌套表同样 depth +1（sdt→表格→cell→sdt / 表套表
        互嵌路径必须递增，否则 TXBX_DEPTH_LIMIT 防爆栈失效）。"""
        seen_tc = set()  # lxml 元素按底层 XML 节点判等：横向合并重复返回的 cell 去重
        for r, row in enumerate(tbl.rows):
            for c, cell in enumerate(row.cells):
                if cell._tc in seen_tc:
                    continue
                seen_tc.add(cell._tc)
                cell_addr = f"{prefix}:{r}:{c}"
                for pi, p in enumerate(cell.paragraphs):
                    _walk_paragraph(p, f"{cell_addr}:{pi}", depth, legacy)
                for j, sub in enumerate(cell.tables):
                    _walk_table(sub, f"{cell_addr}:t{j}", depth + 1, legacy=False)
                # cell 直系内容控件：局部 sdt 序号从 0 起，不消耗任何存量计数器
                for k, sdt in enumerate(cell._tc.findall(qn("w:sdt"))):
                    _walk_sdt(sdt, f"{cell_addr}:sdt{k}", depth + 1)

    def _walk_paragraph(p, base_addr, txbx_depth=0, legacy=False):
        """编址段落自身及其内嵌文本框（正文/表格 cell/文本框内通用）。
        legacy=True（body 直系 w:p 与存量表格 cell 段落）先消耗 para_seq 再 emit
        ——与升级前编号逐字节一致；新区域段落只消耗扁平 idx。"""
        nonlocal para_seq
        if legacy:
            para_seq += 1
        _emit(p, base_addr)
        for k, txbx in enumerate(_iter_txbx_content(p._p)):
            _walk_txbx(txbx, f"{base_addr}:tx{k}", txbx_depth + 1)

    def _walk_txbx(txbx_el, prefix, depth):
        if depth > TXBX_DEPTH_LIMIT:
            logger.warning("textbox nesting deeper than %d at %s, skipped",
                           TXBX_DEPTH_LIMIT, prefix)
            return
        pi = 0
        tbl_k = 0
        sdt_k = 0
        for block in txbx_el.iterchildren():
            if block.tag == qn("w:p"):
                _walk_paragraph(Paragraph(block, doc), f"{prefix}:{pi}", depth)
                pi += 1
            elif block.tag == qn("w:tbl"):
                _walk_table(Table(block, doc), f"{prefix}:cell:{tbl_k}", depth,
                            legacy=False)
                tbl_k += 1
            elif block.tag == qn("w:sdt"):
                _walk_sdt(block, f"{prefix}:sdt{sdt_k}", depth + 1)
                sdt_k += 1

    def _walk_body_blocks(parent_el):
        """body 直系块编址。计数器职责：tbl_no/sdt_no 在此递增（唯一入口）；
        para_seq 不在此递增——由 _walk_paragraph 的 legacy 标志驱动（body 直系
        w:p 与 legacy 表格 cell 直系段落共同消耗）。w:p → para:<para_seq+1>
        （legacy）；w:tbl → cell:<tbl_no>:...（legacy=True 显式传导至全部 cell
        直系段落）；w:sdt → sdt:<sdt_no>:...（新区域，零消耗存量计数器）。"""
        nonlocal tbl_no, sdt_no
        for block in parent_el.iterchildren():
            if block.tag == qn("w:p"):
                _walk_paragraph(Paragraph(block, doc), f"para:{para_seq + 1}",
                                legacy=True)
            elif block.tag == qn("w:tbl"):
                tbl_no += 1
                _walk_table(Table(block, doc), f"cell:{tbl_no}", legacy=True)
            elif block.tag == qn("w:sdt"):
                sdt_no += 1
                _walk_sdt(block, f"sdt:{sdt_no}", 0)

    _walk_body_blocks(doc.element.body)

    # 页眉/页脚：先判 linked 再访问 part/element——对无定义的 header 访问 .part
    # 会触发 _get_or_add_definition「自动补建定义」副作用（给无页眉文档凭空造出
    # 空页眉 part）。partname 字符串去重：多节/多类共享同一 part 时只编址一次。
    # 首页/偶数页变体带 :first/:even 后缀段——同节多类同时 unlinked 时若共用
    # hdr:<sec>:<pi> 前缀会撞号（两个不同段落同 addr，render 反查错位）。
    seen_hf = set()
    hf_slots = (
        ("hdr", "", lambda s: s.header),
        ("hdr", ":first", lambda s: s.first_page_header),
        ("hdr", ":even", lambda s: s.even_page_header),
        ("ftr", "", lambda s: s.footer),
        ("ftr", ":first", lambda s: s.first_page_footer),
        ("ftr", ":even", lambda s: s.even_page_footer),
    )
    for sec_no, section in enumerate(doc.sections):
        for kind, suffix, get_hf in hf_slots:
            hf = get_hf(section)
            if hf is None or hf.is_linked_to_previous:
                continue
            try:
                partname = str(hf.part.partname)
            except Exception:
                # 残缺引用（r:id 悬空等）等价于该 part 不可用：跳过，不拖垮整文档
                logger.exception("resolve header/footer part failed at %s:%s, skipped",
                                 kind, sec_no)
                continue
            if partname in seen_hf:
                continue
            seen_hf.add(partname)
            try:
                _walk_hf_blocks(hf._element, f"{kind}:{sec_no}{suffix}", doc,
                                _walk_paragraph, _walk_table)
            except Exception:
                # 单 part 编址失败（畸形节点等）跳过该 part，不拖垮整文档
                logger.exception("walk header/footer part %s failed, skipped", partname)
    return addr_map, items


def iter_docx_paragraphs(file_bytes: bytes) -> list:
    """返回 [{"index": 扁平序号, "text": 段落文本, "addr": 定位串}]，按文档顺序。"""
    _, items = _build_addr_map(Document(io.BytesIO(file_bytes)))
    return items


# 段落哈希定位（B端保真预览直定位用）────────────────────────────────────
# 与前端 web/src/pages/c-chat/docx-highlight.ts 保持同口径：两侧语言自带的
# \s 语义有差异（JS 含 \uFEFF、Python 含 \x1c-\x1f 等），空白归一化必须用
# 显式字符类，否则段落哈希逐字节对不上、直定位全部回退顺序匹配。

NORM_WS_RE = re.compile(
    r"[ \t\n\r\f\v\x1c-\x1f\x85\u00a0\u1680\u2000-\u200f"
    r"\u2028\u2029\u205f\u3000\ufeff]+"
)

_FNV_BASIS_1 = 0x811C9DC5
_FNV_PRIME_1 = 0x01000193
_FNV_BASIS_2 = 0x1F2BE47C  # 0x811C9DC5 ^ 0x9E3779B9
_FNV_PRIME_2 = 0x84EBCBF8  # 0x01000193 ^ 0x85EBCA6B


def norm_ws(text: str) -> str:
    """去空白归一化（与前端 NORM_WS_RE 同一显式字符类）。"""
    return NORM_WS_RE.sub("", text or "")


def para_hash32x2(norm: str) -> str:
    """段落规范化文本的指纹：UTF-8 字节流跑两趟不同常量的 32 位 FNV-1a，
    拼 16 位 hex。只用于相等性对齐（前端 fnvHash32x2 完全同构），非密码学。
    双通道让 3600 段量级的文档碰撞概率可忽略（单 32 位约 1.5e-3）。"""
    data = (norm or "").encode("utf-8")

    def fnv1a(basis: int, prime: int) -> int:
        h = basis
        for b in data:
            h = ((h ^ b) * prime) & 0xFFFFFFFF
        return h

    return f"{fnv1a(_FNV_BASIS_1, _FNV_PRIME_1):08x}{fnv1a(_FNV_BASIS_2, _FNV_PRIME_2):08x}"


def _count_raw_occurrences(raw: str, sub: str, limit: int) -> int:
    """sub 在 raw 中第 limit 次出现的序号校验：返回实际可分配出现次数。
    完整空白 run 校验与前端 raw 通道一致：候选窗口前后须为非空白字符或边界
    （留白段由可见文本分隔，更长空白串内部不存在更短 anchor 的合法出现）。"""
    found = 0
    i = raw.find(sub)
    while i >= 0:
        prev_ok = i == 0 or not NORM_WS_RE.match(raw[i - 1])
        end = i + len(sub)
        next_ok = end >= len(raw) or not NORM_WS_RE.match(raw[end])
        if prev_ok and next_ok:
            found += 1
            if found >= limit:
                return found
        i = raw.find(sub, i + 1)
    return found


def compute_anchor_positions(file_bytes: bytes, placeholders: list) -> None:
    """为带 addr 的占位符就地补充段落定位元数据（B端保真预览直定位用）：
    - p_idx：锚点所在段落扁平序号（items 文档序，重复文本兜底排序用）
    - p_hash：段落规范化文本指纹（前端渲染 DOM 按段落文本相等对齐）
    - a_occ：锚文本在该段落内的出现序号（1-based；norm 通道按去空白文本
      indexOf 计数；纯空白 anchor 按原文精确出现+完整空白 run 校验计数，
      与前端 highlightDocxRanges 两通道同口径）
    - p_total：段落总数（items 长度；前端重复文本兜底按比例就近用）
    解析失败（addr 不存在/段内出现序号超界/空段落）跳过该占位符的补充，
    前端对缺字段项回退全文顺序匹配；整文档级失败静默返回不拖垮 detail 响应。"""
    try:
        items = _build_addr_map(Document(io.BytesIO(file_bytes)))[1]
    except Exception:
        logger.exception("compute_anchor_positions: build addr map failed, skipped")
        return
    addr_items = {it["addr"]: it for it in items}
    total = len(items)
    # 同 (addr, 通道键) 组内按列表顺序分配出现序号（检测序≈文档序，与
    # detector._preassign_occ 同思路）；组内超出段内实际出现数的项不写字段
    occ_groups = {}
    for ph in placeholders:
        if not isinstance(ph, dict):
            continue
        it = addr_items.get(ph.get("addr") or "")
        if not it:
            continue
        anchor = ph.get("anchor") or ""
        a_norm = norm_ws(anchor)
        p_raw = it.get("text") or ""
        p_norm = norm_ws(p_raw)
        if not p_norm:
            continue  # 空段落无指纹可对齐
        if a_norm:
            ch_key = "N:" + a_norm
            # indexOf step+1 同口径：允许重叠计数（同字符下划线串 "＿＿＿＿"
            # 含 "＿＿＿" 在 step+1 下是 2 次，str.count 非重叠只有 1 次）
            occurrences = 0
            j = p_norm.find(a_norm)
            while j >= 0:
                occurrences += 1
                j = p_norm.find(a_norm, j + 1)
                if occurrences >= 64:
                    break  # 防御：病态重复段落截断计数（超出后 seq 校验自然跳过）
        elif len(anchor) >= 2:
            ch_key = "R:" + anchor
            occurrences = _count_raw_occurrences(p_raw, anchor, 10**9)
        else:
            continue  # 过短 anchor 两通道都无法稳定匹配
        seq = occ_groups.get((ph.get("addr"), ch_key), 0) + 1
        occ_groups[(ph.get("addr"), ch_key)] = seq
        if seq > occurrences:
            continue
        ph["p_idx"] = it["index"]
        ph["p_hash"] = para_hash32x2(p_norm)
        ph["a_occ"] = seq
        ph["p_total"] = total


def extract_docx_candidates(file_bytes: bytes) -> list:
    """提取疑似含填写点的段落（供 LLM 识别，降低 token）。
    含手动占位符 {{key}} 的段落无条件纳入（用户显式标注，不经特征猜测）。"""
    return [
        it for it in iter_docx_paragraphs(file_bytes)
        if it["text"].strip() and (FILL_HINT_RE.search(it["text"]) or PH_RE.search(it["text"]))
    ]


def _nth_index(text: str, sub: str, n: int) -> int:
    """sub 在 text 中第 n（1-based）次**非重叠**出现的起始下标；不足 n 次返回 -1。

    推进步长必须是 len(sub)（从上一次命中结束位之后继续 find）：若只 +1 会命中
    重叠出现（如 "aaaa" 里按 +1 找到的第 2 个 "aa" 在下标 1，而 str.count 的
    非重叠计数下是下标 2），与 run.text.count / joined.count 的出现序号对不上，
    occ 定位会错位。"""
    search_from = 0
    idx = -1
    for _ in range(n):
        idx = text.find(sub, search_from)
        if idx < 0:
            return -1
        search_from = idx + len(sub)
    return idx


def _has_link_or_field(p: Paragraph) -> bool:
    """段落是否含超链接（w:hyperlink）或简单域（w:fldSimple）直系子节点。

    这两类节点的文本会进入 Paragraph.text（python-docx 1.1+ 对 hyperlink 生效），
    但不在 Paragraph.runs 中，整段重写会导致文本被复制。
    """
    return bool(p._p.findall(qn("w:hyperlink")) or p._p.findall(qn("w:fldSimple")))


def _replace_via_run_concat(p: Paragraph, anchor: str, repl: str,
                            occ: int | None = None) -> bool:
    """含超链接/域段落的替换：把所有 run 的文本按序拼接，在拼接串上替换，
    再按原 run 长度切分写回各 run。

    超链接/域内文本不参与（不在 runs 中），故不会产生复制。替换导致拼接串
    长度变化时，差值并入最后一个非空 run 的文本——简单可行即可，该路径只为
    避免超链接文本复制，不追求精确保持 run 边界。锚文本在拼接串中不存在
    （跨界没拼上）时返回 False（no-op）。occ=None 为存量 replace-all 语义；
    occ=N 只替换拼接串中第 N 次非重叠出现（不足 N 次返回 False）。"""
    runs = p.runs
    joined = "".join(r.text for r in runs)
    if anchor not in joined:
        return False
    if occ is None:
        replaced = joined.replace(anchor, repl)
    else:
        start = _nth_index(joined, anchor, occ)
        if start < 0:
            return False
        replaced = joined[:start] + repl + joined[start + len(anchor):]
    # 按原 run 文本长度在替换后的拼接串上切分写回；总长度差（repl 与 anchor 不等长）
    # 并入最后一个非空 run——简单可行即可，不追求精确保持 run 边界。
    pos = 0
    last_nonempty = -1
    for i, r in enumerate(runs):
        n = len(r.text)
        r.text = replaced[pos:pos + n]
        pos += n
        if n > 0:
            last_nonempty = i
    tail = replaced[pos:]
    if tail:
        target = runs[max(last_nonempty, 0)]
        target.text = (target.text or "") + tail
    return True


def _locate_run_span(runs: list, start: int, end: int) -> tuple:
    """在 run 文本拼接坐标系里返回覆盖 [start, end) 的
    (首run下标 i, 尾run下标 j, anchor在首run内偏移, anchor在尾run内偏移)。"""
    pos = 0
    i = j = -1
    off_i = off_j = 0
    for idx, r in enumerate(runs):
        n = len(r.text)
        if i < 0 and pos + n > start:
            i, off_i = idx, start - pos
        if pos + n >= end:
            j, off_j = idx, end - pos
            break
        pos += n
    if i < 0 or j < 0:  # 理论不可达（anchor 必在拼接串覆盖范围内），兜底安全值防越界
        i, j, off_i, off_j = max(i, 0), max(j, 0), 0, 0
    return i, j, off_i, off_j


def _replace_cross_run_in_place(p: Paragraph, anchor: str, repl: str,
                                occ: int | None = None) -> bool:
    """跨 run 区间替换（格式保真）：只重写 anchor 覆盖的 run 区间——
    首 run 保留 anchor 前文本并接替换值，尾 run 保留 anchor 后文本，中间 run 清空；
    区间外 run 原样不动（段内其他位置格式完整保留）。

    扫描策略：按替换前出现次数循环，每轮从上一轮替换终点之后继续 find
    （scan_from = start + len(repl)），严格复刻 str.replace「不重扫替换产物」
    的语义——anchor 是 repl 子串时（如 anchor="name"、repl="{{name}}"），
    刚写入的 repl 不会被再次命中，避免文本腐坏与后续真实出现漏替。
    中段 run 仅在有文本时置空：空文本 run 的置空唯一效果是经 Run.text setter
    销毁 rPr 外的结构子节点（w:fldChar/w:drawing 等），跳过即保住结构。
    occ=None 为存量 replace-all 语义（remaining=全部出现次数，逐次替换）；
    occ=N 先把扫描起点推进过前 N-1 次出现（步长 len(anchor)，非重叠）、
    remaining=1，只替换第 N 次出现，不足 N 次返回 False。"""
    runs = p.runs
    joined_all = "".join(r.text for r in runs)
    total = joined_all.count(anchor)
    scan_from = 0
    if occ is not None:
        if occ > total:
            return False
        # 跳过前 occ-1 次出现：扫描起点推进到第 occ 次出现的起始下标
        for _ in range(occ - 1):
            scan_from = joined_all.find(anchor, scan_from) + len(anchor)
        remaining = 1
    else:
        remaining = total
    replaced = False
    while remaining > 0:
        remaining -= 1
        joined = "".join(r.text for r in runs)
        start = joined.find(anchor, scan_from)
        if start < 0:
            break
        end = start + len(anchor)
        i, j, off_i, off_j = _locate_run_span(runs, start, end)
        if i == j:
            runs[i].text = runs[i].text[:off_i] + repl + runs[i].text[off_j:]
        else:
            runs[i].text = runs[i].text[:off_i] + repl
            for mid in range(i + 1, j):
                if runs[mid].text:
                    runs[mid].text = ""
            runs[j].text = runs[j].text[off_j:]
        scan_from = start + len(repl)
        replaced = True
    return replaced


def _replace_in_paragraph(p: Paragraph, anchor: str, repl: str,
                          occ: int | None = None) -> bool:
    """段内替换锚文本为 repl。occ=None 为存量 replace-all 语义：优先单 run 内
    完成（首个含 anchor 的 run 全替换）；跨 run 时：普通段落只重写 anchor 覆盖的
    run 区间，区间外格式保留（见 _replace_cross_run_in_place）；含超链接/域的
    段落走 run 拼接替换（见 _replace_via_run_concat），避免超链接文本被复制进
    正文 run。occ=N 只替换段落拼接文本中第 N 次非重叠出现，不足 N 次返回 False
    （no-op）。

    occ 路径**不走**单 run 快路径：单 run 的 count 是 run 局部计数，当更早的
    出现跨界（anchor 被切成两半分属相邻 run）时，局部序号与段落级出现序号错位，
    会命中错误出现位。跨 run 路径对完全落在单个 run 内的命中同样只重写该 run
    （i==j 分支），格式保真效果与快路径等价，故 occ 统一走拼接文本坐标系。
    """
    if not anchor:
        return False
    if anchor not in p.text:
        return False
    if occ is None:
        for run in p.runs:
            if anchor in run.text:
                # str.replace 语义：同段多次出现全部替换（存量行为逐字节不变）
                run.text = run.text.replace(anchor, repl)
                return True
    if _has_link_or_field(p):
        return _replace_via_run_concat(p, anchor, repl, occ)
    if not p.runs:
        return False
    return _replace_cross_run_in_place(p, anchor, repl, occ)


def apply_docx_placeholders(file_bytes: bytes, replacements: list) -> bytes:
    """replacements: [{"addr", "anchor", "key", "occ"?}]，把 anchor 替换为 {{key}}。

    occ 语义：None（缺省/存量数据）= replace-all（存量行为完全不变，红线）；
    正整数 N = 只替换段落文本中第 N 次非重叠出现（同段同形留白第 N 处定位）。
    occ 存在但非法（非整数/布尔/小于 1）→ 跳过该条——宁可不填也不能退化成
    replace-all 误伤其他出现位。
    同段多条 occ 条目按 occ **降序**应用：若按原序逐条应用，前面的替换会改变
    后面条目所见的出现序号（原文 occ=2 在 occ=1 应用后变成「当前第 1 次」）而
    错位；降序时先替换的第 N 次不影响更小序号出现的位置。occ=None 条目排在
    occ 条目之后且相互保持原相对顺序（replace-all 先行会先吞掉 occ 的目标出现）。
    存量数据（全部无 occ）经此排序后应用顺序与原来逐字节一致。

    addr 不存在时静默跳过；anchor 不在该段落时为 no-op（幂等）；
    addr/anchor/key 任一缺失或为空时跳过该条（LLM 脏输入健壮性，与 addr
    不存在静默跳过的语义一致）。
    """
    doc = Document(io.BytesIO(file_bytes))
    addr_map, _ = _build_addr_map(doc)
    cleaned = []
    for rep in replacements:
        occ = rep.get("occ")
        if occ is not None and (isinstance(occ, bool) or not isinstance(occ, int) or occ < 1):
            logger.warning(
                "apply_docx_placeholders: invalid occ %r, skip key=%s", occ, rep.get("key"))
            continue
        cleaned.append((rep, occ))
    # 稳定排序：occ 条目降序在前，None 条目保持原序在后（见 docstring 推导）
    cleaned.sort(key=lambda t: (t[1] is None, -(t[1] or 0)))
    for rep, occ in cleaned:
        addr = rep.get("addr")
        anchor = rep.get("anchor")
        key = rep.get("key")
        if not addr or not anchor or not key:
            continue
        p = addr_map.get(addr)
        if p is not None:
            # 替换层 no-op 可观测性（设计 §8：no-op + 告警，旧数据回放）：返回 False
            # 的情形有二——anchor 不在段落文本（模板版本漂移/LLM 脏锚）、occ 超界
            # （段内出现次数不足 N）。二者此前均静默，导致「填写点没生效」无从定位；
            # addr 悬空告警在 else 分支（互斥路径），此处不会重复。
            if not _replace_in_paragraph(p, anchor, f"{{{{{key}}}}}", occ):
                logger.warning(
                    "apply_docx_placeholders: replace no-op addr=%s key=%s anchor=%r occ=%s",
                    addr, key, anchor, occ)
        else:
            # 静默跳过不变（脏输入健壮性），但留告警痕迹：存量模板的落库 addr
            # 因编址规则演进（如合并单元格去重）悬空时，可凭此定位「填写点没生效」
            logger.warning("apply_docx_placeholders: addr %r not found, skip key=%s", addr, key)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
