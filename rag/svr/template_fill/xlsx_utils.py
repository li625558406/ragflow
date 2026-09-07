"""xlsx 模板工具：非空单元格遍历与占位符替换。

addr 约定 "<sheet>!<coordinate>"（如 封面!B1）。iter 产物结构与 docx_utils 一致
（index/text/addr + sheet/coord），供 detector 复用；apply 只改 cell.value，
openpyxl 对 value 赋值不触碰样式对象（字体/填充/边框等随 cell 保留）。

load_workbook 统一使用默认 data_only=False：公式按公式文本保留，重存后不丢公式。
"""
import io

from openpyxl import load_workbook


def iter_xlsx_cells(file_bytes: bytes) -> list:
    """遍历所有工作表的非空单元格（None 或纯空白文本跳过）。

    返回 [{"index": 全簿扁平序号, "text": str(值), "addr": "<sheet>!<coord>",
    "sheet": 表名, "coord": 坐标}]，按表顺序 + 行列顺序。
    """
    wb = load_workbook(io.BytesIO(file_bytes), data_only=False)
    items, idx = [], 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None or not str(cell.value).strip():
                    continue
                items.append({
                    "index": idx,
                    "text": str(cell.value),
                    "addr": f"{ws.title}!{cell.coordinate}",
                    "sheet": ws.title,
                    "coord": cell.coordinate,
                })
                idx += 1
    return items


def extract_xlsx_candidates(file_bytes: bytes) -> list:
    """提取疑似含填写点的单元格（供 LLM 识别，降低 token）。"""
    from rag.svr.template_fill.docx_utils import FILL_HINT_RE
    return [it for it in iter_xlsx_cells(file_bytes) if FILL_HINT_RE.search(it["text"])]


def _find_cell(wb, rep: dict):
    """按 sheet/coord 定位单元格；sheet 不存在（KeyError）或 coord 非法
    （openpyxl 转坐标抛 ValueError）时返回 None，交由调用方跳过该条。
    coord 合法但越界时 openpyxl 会创建空单元格，由调用方"值为 None 不替换"兜住。"""
    try:
        return wb[rep["sheet"]][rep["coord"]]
    except (KeyError, ValueError):
        return None


def apply_xlsx_placeholders(file_bytes: bytes, replacements: list) -> bytes:
    """replacements 元素含 sheet/coord/anchor/key；把 anchor 替换为 {{key}}。

    只改 cell.value，样式保留。脏输入健壮性（与 docx_utils 语义一致）：
    sheet/coord/anchor/key 任一缺失或为空 → 预检查跳过；sheet 不存在或 coord
    非法 → _find_cell 返回 None 跳过该条而不中断整批；anchor 不在单元格值中
    → no-op（幂等）。
    """
    wb = load_workbook(io.BytesIO(file_bytes))
    for rep in replacements:
        sheet = rep.get("sheet")
        coord = rep.get("coord")
        anchor = rep.get("anchor")
        key = rep.get("key")
        if not sheet or not coord or not anchor or not key:
            continue
        cell = _find_cell(wb, rep)
        if cell is None:
            continue
        val = cell.value
        if val is not None and anchor in str(val):
            cell.value = str(val).replace(anchor, f"{{{{{key}}}}}")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
