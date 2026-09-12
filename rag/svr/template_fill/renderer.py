#
#  Copyright 2026 The InfiniFlow Authors. All Rights Reserved.
#
#  Licensed under the Apache License, Version 2.0 (the "License");
#  you may not use this file except in compliance with the License.
#  You may obtain a copy of the License at
#
#      http://www.apache.org/licenses/LICENSE-2.0
#
#  Unless required by applicable law or agreed to in writing, software
#  distributed under the License is distributed on an "AS IS" BASIS,
#  WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
#  See the License for the specific language governing permissions and
#  limitations under the License.
#
"""模板填写：渲染层。LLM 只产值，格式回填全由本层程序完成——
Word 用 docxtpl（Jinja2 语法与 {{key}} 占位符天然兼容，接受 file-like，
无需落盘临时文件），Excel 用 openpyxl 按注册表 addr（"<sheet>!<coord>"，
约定同 xlsx_utils）坐标直写：合并区写左上角天然合法，非左上角/定位失败的
格跳过不中断整批；只改 cell.value，样式随 cell 保留。"""
import copy
import io
import logging

logger = logging.getLogger(__name__)

# AI 填入值标蓝色（成稿可辨识哪些内容是 AI 写入的）：docxtpl 渲染后值文本
# 落在 {{key}} 占位 run 的位置并继承其 rPr，故只需渲染前把占位 run 标蓝。
FILL_BLUE = "0000FF"


def _set_run_color(r, val: str):
    """给 w:r 节点的 rPr 设置 w:color（已有时覆盖）。rPr 必须是 w:r 首子元素。
    用 CT_RPr.get_or_add_color 按 OOXML schema sequence 插入（w:color 须位于
    w:sz/w:u/highlight 等之前，裸 append 产出乱序 XML 会被严格校验器拒收）。"""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    rPr = r.find(qn("w:rPr"))
    if rPr is None:
        rPr = OxmlElement("w:rPr")
        r.insert(0, rPr)
    rPr.get_or_add_color().set(qn("w:val"), val)


def _split_ph_segments(text: str) -> list:
    """把 run 文本按占位符切成 [(片段, 是否占位符)] 序列（占位符外的原文保持顺序）。"""
    from rag.svr.template_fill.docx_utils import PH_RE

    parts, last = [], 0
    for m in PH_RE.finditer(text):
        if m.start() > last:
            parts.append((text[last:m.start()], False))
        parts.append((m.group(0), True))
        last = m.end()
    if last < len(text):
        parts.append((text[last:], False))
    return parts


def _colorize_placeholder_runs(blob: bytes) -> bytes:
    """渲染前置：把工作副本中 {{key}} 占位符隔离成独立 run 并标蓝。

    工作副本生成路径（docx_utils._replace_in_paragraph）常把锚文本连同前后文
    重写进同一个 run，占位符不隔离会把标签原文一起染蓝。无占位符的 blob 原样
    返回（免序列化，省一次 save；解析仍需执行才能判定有无占位符）。防御：含 w:br/w:tab/w:drawing 等非
    (rPr|t) 子节点的 run 拆分会丢结构，整体跳过——该处占位符保持原样渲染
    （值仍正常回填，只是不标蓝）。覆盖面与 docx_utils._build_addr_map 一致
    （正文段落 + 顶层表格 cell），占位符只会出现在这些位置。"""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    from rag.svr.template_fill.docx_utils import _build_addr_map, PH_RE

    doc = Document(io.BytesIO(blob))
    addr_map, _ = _build_addr_map(doc)
    changed = False
    for p in addr_map.values():
        for run in list(p.runs):
            text = run.text or ""
            if not PH_RE.search(text):
                continue
            r = run._r
            if any(ch.tag not in (qn("w:rPr"), qn("w:t")) for ch in r):
                continue
            parts = _split_ph_segments(text)
            if len(parts) == 1 and parts[0][1]:
                # 整 run 即占位符：原地标蓝即可，无需拆分
                _set_run_color(r, FILL_BLUE)
                changed = True
                continue
            parent = r.getparent()
            idx = list(parent).index(r)
            new_nodes = []
            for seg, is_ph in parts:
                if not seg:
                    continue
                nr = copy.deepcopy(r)
                for t in nr.findall(qn("w:t")):
                    nr.remove(t)
                t = OxmlElement("w:t")
                t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                t.text = seg
                nr.append(t)
                if is_ph:
                    _set_run_color(nr, FILL_BLUE)
                new_nodes.append(nr)
            parent.remove(r)
            for off, nr in enumerate(new_nodes):
                parent.insert(idx + off, nr)
            changed = True
    if not changed:
        return blob
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def render_docx(blob: bytes, values: dict) -> bytes:
    """用 docxtpl 渲染 Word 模板：values 的 key 对应文档内 {{key}} 占位符，
    未出现的 key 忽略；values 缺失的占位符按 Jinja2 默认 Undefined 渲染为
    空串（不是保持 {{key}} 原样）——executors 侧保证产值覆盖所有注册 key。
    渲染前先做占位符标蓝（_colorize_placeholder_runs），成稿中 AI 填入值为蓝色。"""
    from docxtpl import DocxTemplate
    blob = _colorize_placeholder_runs(blob)
    doc = DocxTemplate(io.BytesIO(blob))
    doc.render(values or {})
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def render_xlsx(blob: bytes, values: dict, addr_by_key: dict) -> bytes:
    """按注册表 addr（"<sheet>!<coord>"，按最后一个 "!" 拆分，兼容表名含 "!"）
    直写 Excel 单元格。脏输入健壮性：key 无产值 / addr 缺失或无 "!" → 预检查
    跳过；sheet 不存在（KeyError）/ coord 非法、空串或行越界（IndexError/
    ValueError）/ range 语法（AttributeError）→ 跳过该格不中断整批；值为
    None 不替换（openpyxl 会把 None 落成空单元格，且超边界 coord 会被惰性
    创建，均非预期写入）。注意：列越界（如 XFE1）openpyxl 不抛异常、会静默
    创建脏格——由注册表侧（xlsx_utils 从真实 workbook 生成）保证不出现。"""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(blob))
    for key, addr in (addr_by_key or {}).items():
        val = (values or {}).get(key)
        if val is None or not addr or "!" not in addr:
            continue
        sheet, _, coord = addr.rpartition("!")
        try:
            wb[sheet][coord] = val
        except (KeyError, ValueError, AttributeError, IndexError):
            # KeyError: sheet 不存在；IndexError/ValueError: coord 空/非法/行越界；
            # AttributeError: coord 是 range 语法 → 赋值落在 tuple 上
            logger.warning("xlsx render skip cell key=%s addr=%s", key, addr)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def render(file_type: str, blob: bytes, values: dict, addr_by_key: dict | None = None) -> bytes:
    """按文件类型分发：docx → docxtpl 模板渲染；其余（xlsx）→ 坐标直写。"""
    if file_type == "docx":
        return render_docx(blob, values)
    return render_xlsx(blob, values, addr_by_key or {})
