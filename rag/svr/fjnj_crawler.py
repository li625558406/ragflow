#!/usr/bin/env python3
#
#  Copyright 2026 The InfiniFlow Authors. All Rights Reserved.
#
#  Licensed under the Apache License, Version 2.0 (the "License");
#  you may not use this file except in compliance with the License.
#  You may obtain a copy of the License at
#
#      http://www.apache.org/licenses/LICENSE-2.0
#
#  Unless required by applicable law or agreed to in writing, software
#  distributed under the License is distributed on an "AS IS" BASIS,
#  WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
#  See the License for the specific language governing permissions and
#  limitations under the License.
#
"""
Dedicated web crawler for fjnj.gov.cn (南靖县人民政府).

Crawls the 政务公开 (zwgk) page which contains multiple content modules,
each with one or more tabs. All tab content is server-rendered in the
initial HTML (no AJAX — tabstrip.js only toggles visibility via CSS).

Site characteristics
────────────────────
  • SSR HTML — all module/tab data is in the initial page source.
  • Modules: 政策文件 (single list), 通知公告 (4 tabs across 2 columns).
  • Tabs: 三农服务/财政资金 (left), 人事信息/规划计划 (right).
  • Detail formats: CMS html and publicInfo.shtml.
  • Attachments: /cms/pages/{id}/pdfattachments/, wordattachments/, attachments/.
  • SSL — verify=False needed for .gov.cn certificates.

Usage (typically spawned by task_executor):
    python fjnj_crawler.py \\
        --tenant-id <TENANT_ID> \\
        --target-url http://www.fjnj.gov.cn/cms/html/njxrmzf/zwgk/index.html \\
        --kb-id <KB_ID> \\
        --task-name <NAME>
"""

import argparse
import datetime
import json
import logging
import os
import random
import re
import sys
import time
import urllib.request
import zipfile
from urllib.parse import urljoin, urlparse

import requests as _requests
from bs4 import BeautifulSoup

import urllib3
urllib3.disable_warnings()

_SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
_PROJECT_ROOT = os.path.abspath(os.path.join(_SCRIPT_DIR, "..", ".."))
sys.path.insert(0, _PROJECT_ROOT)

from common import settings
from common.log_utils import init_root_logger
from common.misc_utils import get_uuid

# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------
_BASE = "http://www.fjnj.gov.cn"
_SITE_NAME = "南靖县人民政府"
_ZWGK_URL = f"{_BASE}/cms/html/njxrmzf/zwgk/index.html"

_USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/120.0.0.0 Safari/537.36"
)

_HTML_HEADERS = {
    "User-Agent": _USER_AGENT,
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,image/webp,*/*;q=0.8",
    "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8",
}

# Anti-crawling delays (seconds)
_PAGE_DELAY = (1.0, 2.5)
_ARTICLE_DELAY = (0.3, 1.0)

_STATE_FILENAME = "_crawler_state.json"

_ATTACH_EXTENSIONS = {
    ".pdf", ".doc", ".docx", ".xls", ".xlsx",
    ".ppt", ".pptx", ".zip", ".rar", ".7z",
    ".txt", ".jpg", ".jpeg", ".png", ".tif", ".tiff",
}


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _safe_print(msg):
    try:
        print(msg)
    except UnicodeEncodeError:
        print(msg.encode("gbk", errors="replace").decode("gbk"))


def _request_delay(min_s, max_s):
    time.sleep(random.uniform(min_s, max_s))


def _sanitize_filename(text, max_len=150):
    if not text:
        return "untitled"
    name = re.sub(r'[\\/:*?"<>|]', "_", str(text).strip())
    name = re.sub(r'\s+', " ", name)
    name = name.strip("._ ")
    return name[:max_len] if name else "untitled"


def _is_attach_url(url):
    lower = url.lower().split("?")[0]
    return any(lower.endswith(ext) for ext in _ATTACH_EXTENSIONS)


# ---------------------------------------------------------------------------
# HTTP session
# ---------------------------------------------------------------------------

def _init_session():
    sess = _requests.Session()
    sess.headers.update(_HTML_HEADERS)
    sess.verify = False
    try:
        sess.get(_BASE, timeout=30, verify=False)
        logging.info("Session initialized")
    except Exception as e:
        logging.warning("Failed to init session: %s", e)
    return sess


# ---------------------------------------------------------------------------
# Listing: single-page SSR extraction
# ---------------------------------------------------------------------------

def _fetch_all_items(sess):
    """Parse the zwgk page and extract all items from all modules and tabs.

    All tab content is SSR-rendered in the initial HTML — hidden panels
    use display:none and tabstrip.js only toggles visibility.

    Returns list of dicts: {title, date, href, module}.
    """
    _safe_print(f"Fetching: {_ZWGK_URL}")
    r = sess.get(_ZWGK_URL, timeout=60, verify=False)
    r.encoding = 'utf-8'
    soup = BeautifulSoup(r.text, 'lxml')

    items = []
    seen_hrefs = set()

    def _add_items_from_container(container, mod_name):
        """Extract li > a items from a container div."""
        if not container:
            return
        for a in container.find_all('a', href=True):
            href = a.get('href', '').strip()
            if not href or href.startswith('#') or href.startswith('javascript'):
                continue
            abs_href = urljoin(_ZWGK_URL, href)

            # Skip navigation/index links
            if '/index.html' in abs_href or abs_href in seen_hrefs:
                continue
            if urlparse(abs_href).netloc != urlparse(_BASE).netloc:
                continue

            title = a.get_text(strip=True)
            if not title or len(title) < 4:
                continue

            seen_hrefs.add(abs_href)

            # Date from URL path
            date_str = ''
            m = re.search(r'/(\d{4})-(\d{2})-(\d{2})/', abs_href)
            if m:
                date_str = f"{m.group(1)}-{m.group(2)}-{m.group(3)}"

            # Date from parent li text
            if not date_str:
                parent = a.parent
                while parent and parent.name not in ('li', 'div'):
                    parent = parent.parent
                if parent:
                    parent_text = parent.get_text(strip=True)
                    m = re.search(r'(\d{4}-\d{1,2}-\d{1,2})', parent_text)
                    if m:
                        date_str = m.group(1)

            items.append({
                'title': title,
                'date': date_str,
                'href': abs_href,
                'module': mod_name,
            })

    # ---- Module 1: 政策文件 (single list) ----
    zfwj = soup.find('div', class_='zfwj')
    if zfwj:
        mid_list = zfwj.find('div', class_='mid-list')
        _add_items_from_container(mid_list, '政策文件')
        _safe_print(f"  政策文件: found items in zfwj section")

    # ---- Module 2: 通知公告 — left column (2 tabs) ----
    tzgg_left = soup.find('div', class_='tzgg-left')
    left_tab_names = {}
    if tzgg_left:
        tabbar = tzgg_left.find('div', class_='tzgg-tabbar')
        if tabbar:
            tab_texts = [a.get_text(strip=True) for a in tabbar.find_all('a') if a.get_text(strip=True)]
            for i, name in enumerate(tab_texts):
                left_tab_names[i] = name
            _safe_print(f"  tzgg-left tabs: {list(left_tab_names.values())}")

        # Find all content panels: direct div children + display:none panels
        found_panels = []
        for child_div in tzgg_left.find_all('div', recursive=False):
            lis = child_div.find_all('li')
            if len(lis) >= 3:
                found_panels.append(child_div)
        # Also check for hidden panels deeper in the DOM
        for hidden_div in tzgg_left.find_all('div', style=re.compile(r'display\s*:\s*none', re.I)):
            lis = hidden_div.find_all('li')
            if len(lis) >= 3 and hidden_div not in found_panels:
                found_panels.append(hidden_div)

        for i, panel in enumerate(found_panels):
            tab_name = left_tab_names.get(i, f'通知公告-左{i}')
            _add_items_from_container(panel, tab_name)
            _safe_print(f"  tzgg-left panel '{tab_name}': {len(panel.find_all('li'))} items")

    # ---- Module 3: 通知公告 — right column (2 tabs) ----
    tzgg_right = soup.find('div', class_='tzgg-right')
    right_tab_names = {}
    if tzgg_right:
        tabbar = tzgg_right.find('div', class_='tzgg-tabbar')
        if tabbar:
            tab_texts = [a.get_text(strip=True) for a in tabbar.find_all('a') if a.get_text(strip=True)]
            for i, name in enumerate(tab_texts):
                right_tab_names[i] = name
            _safe_print(f"  tzgg-right tabs: {list(right_tab_names.values())}")

        found_panels = []
        for child_div in tzgg_right.find_all('div', recursive=False):
            lis = child_div.find_all('li')
            if len(lis) >= 3:
                found_panels.append(child_div)
        for hidden_div in tzgg_right.find_all('div', style=re.compile(r'display\s*:\s*none', re.I)):
            lis = hidden_div.find_all('li')
            if len(lis) >= 3 and hidden_div not in found_panels:
                found_panels.append(hidden_div)

        for i, panel in enumerate(found_panels):
            tab_name = right_tab_names.get(i, f'通知公告-右{i}')
            _add_items_from_container(panel, tab_name)
            _safe_print(f"  tzgg-right panel '{tab_name}': {len(panel.find_all('li'))} items")

    # ---- Module 4: Flatten all remaining mid-list divs not already covered ----
    for mid_list in soup.find_all('div', class_='mid-list'):
        parent = mid_list.parent
        parent_cls = ' '.join(parent.get('class', [])) if parent and isinstance(parent.get('class'), list) else ''
        # Only process if not already covered by zfwj/tzgg-left/tzgg-right
        if 'zfwj' not in parent_cls and 'tzgg-left' not in parent_cls and 'tzgg-right' not in parent_cls:
            _add_items_from_container(mid_list, '其他')

    # Count by module
    mod_counts = {}
    for it in items:
        mod_counts[it['module']] = mod_counts.get(it['module'], 0) + 1

    _safe_print(f"\n  Total items: {len(items)}")
    for mod, count in sorted(mod_counts.items()):
        _safe_print(f"    {mod}: {count}")

    logging.info("Fetched %d items from %d modules", len(items), len(mod_counts))
    return items


# ---------------------------------------------------------------------------
# Detail page parsing
# ---------------------------------------------------------------------------

def _parse_cms_detail(html, detail_url):
    """Extract content from CMS html detail page."""
    soup = BeautifulSoup(html, 'lxml')
    result = {'title': '', 'pub_date': '', 'content_text': '', 'attachments': []}

    # Title
    title_el = soup.find('title')
    if title_el:
        result['title'] = title_el.get_text(strip=True).split('_')[0].strip()

    # Date from meta
    for meta_name in ('PubDate', 'publishdate', 'articledate'):
        meta = soup.find('meta', attrs={'name': meta_name})
        if meta and meta.get('content'):
            result['pub_date'] = meta['content'].strip()[:10]
            break

    # Date from URL path
    if not result['pub_date']:
        m = re.search(r'/(\d{4})-(\d{2})-(\d{2})/', detail_url)
        if m:
            result['pub_date'] = f"{m.group(1)}-{m.group(2)}-{m.group(3)}"

    # Date from text
    if not result['pub_date']:
        for pattern in [
            r'发布日[期间][：:]\s*(\d{4}-\d{1,2}-\d{1,2})',
            r'生成日期[：:]\s*(\d{4}-\d{1,2}-\d{1,2})',
        ]:
            m = re.search(pattern, html, re.I)
            if m:
                result['pub_date'] = m.group(1)[:10]
                break

    # Content
    for cls in ['article', 'content', 'detail', 'TRS_Editor', 'view', 'con', 'main']:
        div = soup.find('div', class_=re.compile(cls, re.I))
        if div:
            for tag in div.find_all(['script', 'style']):
                tag.decompose()
            text = div.get_text(separator='\n', strip=True)
            if len(text) > 100:
                text = re.sub(r'\n{3,}', '\n\n', text)
                result['content_text'] = text
                break

    # Fallback
    if not result['content_text']:
        max_len = 0
        for div in soup.find_all('div'):
            if div.find_parent(['header', 'nav', 'footer', 'script', 'style']):
                continue
            text = div.get_text(strip=True)
            if 200 < len(text) < 50000 and len(text) > max_len:
                max_len = len(text)
                result['content_text'] = text

    # Attachments
    seen_urls = set()
    for a in soup.find_all('a', href=True):
        href = a.get('href', '')
        abs_url = urljoin(detail_url, href)
        title = a.get_text(strip=True)

        if _is_attach_url(abs_url) and abs_url not in seen_urls:
            seen_urls.add(abs_url)
            att_name = title or os.path.basename(urlparse(abs_url).path)
            result['attachments'].append((att_name, abs_url))

        if ('/attached/' in href or '/upload/' in href or
            'attachments/' in href or 'wordattachments/' in href or
            'pdfattachments/' in href or 'fileattachments/' in href or
            '/cms/pages/' in href):
            if abs_url not in seen_urls:
                seen_urls.add(abs_url)
                att_name = title or os.path.basename(urlparse(abs_url).path)
                result['attachments'].append((att_name, abs_url))

    return result


def _parse_publicinfo_detail(html, detail_url):
    """Extract content from publicInfo.shtml detail page."""
    soup = BeautifulSoup(html, 'lxml')
    result = {'title': '', 'pub_date': '', 'content_text': '', 'attachments': []}

    # Title
    title_el = soup.find('title')
    if title_el:
        result['title'] = title_el.get_text(strip=True).split('_')[0].strip()

    # Date from metadata table
    for table in soup.find_all('table'):
        trs = table.find_all('tr')
        if 3 <= len(trs) <= 20:
            row_data = {}
            for tr in trs:
                cells = [td.get_text(strip=True) for td in tr.find_all(['td', 'th'])]
                if len(cells) >= 2:
                    row_data[cells[0]] = cells[1]
            for key, val in row_data.items():
                if '发布日期' in key:
                    result['pub_date'] = val[:10]
                    break
                elif '生成日期' in key:
                    result['pub_date'] = val[:10]
                    break
            if result['pub_date']:
                break

    # Date fallback: URL path
    if not result['pub_date']:
        m = re.search(r'/(\d{4})-(\d{2})-(\d{2})/', detail_url)
        if m:
            result['pub_date'] = f"{m.group(1)}-{m.group(2)}-{m.group(3)}"

    # Date fallback: text regex
    if not result['pub_date']:
        m = re.search(r'(\d{4}-\d{1,2}-\d{1,2})', html)
        if m:
            result['pub_date'] = m.group(1)

    # Content
    for cls in ['article', 'content', 'detail', 'TRS_Editor', 'view', 'con', 'main', 'body']:
        div = soup.find('div', class_=re.compile(cls, re.I))
        if div:
            for tag in div.find_all(['script', 'style']):
                tag.decompose()
            text = div.get_text(separator='\n', strip=True)
            if len(text) > 100:
                text = re.sub(r'\n{3,}', '\n\n', text)
                result['content_text'] = text
                break

    # Fallback
    if not result['content_text']:
        body = soup.find('body')
        if body:
            text = body.get_text(separator='\n', strip=True)
            if len(text) > 200:
                result['content_text'] = text

    # Attachments
    seen_urls = set()
    for a in soup.find_all('a', href=True):
        href = a.get('href', '')
        abs_url = urljoin(detail_url, href)
        title = a.get_text(strip=True)

        if _is_attach_url(abs_url) and abs_url not in seen_urls:
            seen_urls.add(abs_url)
            att_name = title or os.path.basename(urlparse(abs_url).path)
            result['attachments'].append((att_name, abs_url))

        if ('/attached/' in href or '/upload/' in href or
            'attachments/' in href or 'wordattachments/' in href or
            'pdfattachments/' in href or 'fileattachments/' in href or
            '/cms/pages/' in href):
            if abs_url not in seen_urls:
                seen_urls.add(abs_url)
                att_name = title or os.path.basename(urlparse(abs_url).path)
                result['attachments'].append((att_name, abs_url))

    return result


def _parse_detail(html, detail_url):
    """Route to correct detail parser based on URL."""
    if 'publicInfo.shtml' in detail_url:
        return _parse_publicinfo_detail(html, detail_url)
    return _parse_cms_detail(html, detail_url)


# ---------------------------------------------------------------------------
# File download
# ---------------------------------------------------------------------------

def _download_file(sess, file_url, timeout=120):
    parsed = urlparse(file_url)
    main_parsed = urlparse(_BASE)

    if parsed.netloc == main_parsed.netloc or not parsed.netloc:
        try:
            resp = sess.get(file_url, timeout=timeout, stream=True, verify=False)
            if resp.status_code == 200 and len(resp.content) > 100:
                return resp.content
        except Exception as e:
            logging.error("Download error %s: %s", file_url, e)
        return None

    req = urllib.request.Request(file_url, headers=_HTML_HEADERS)
    try:
        resp = urllib.request.urlopen(req, timeout=timeout)
        data = resp.read()
        if len(data) > 100:
            return data
    except Exception as e:
        logging.error("Download error (external) %s: %s", file_url, e)
    return None


# ---------------------------------------------------------------------------
# ZIP extraction
# ---------------------------------------------------------------------------

def _extract_zip(zip_path, dest_dir):
    extracted = []
    try:
        with zipfile.ZipFile(zip_path, 'r') as zf:
            for name in zf.namelist():
                safe_name = re.sub(r'[\\/:*?"<>|]', "_", os.path.basename(name))
                if not safe_name:
                    continue
                dest_path = os.path.join(dest_dir, safe_name)
                if os.path.exists(dest_path):
                    continue
                with open(dest_path, 'wb') as f:
                    f.write(zf.read(name))
                extracted.append(dest_path)
                _safe_print(f"           Extracted: {safe_name}")
    except Exception as e:
        _safe_print(f"           ZIP extract error: {e}")
    return extracted


# ---------------------------------------------------------------------------
# Text extraction from attachments
# ---------------------------------------------------------------------------

def _extract_file_text(filepath):
    ext = os.path.splitext(filepath)[1].lower()
    text = ""

    try:
        if ext == '.pdf':
            import fitz
            with fitz.open(filepath) as doc:
                for page in doc:
                    text += page.get_text() + "\n"
        elif ext == '.docx':
            from docx import Document
            doc = Document(filepath)
            for para in doc.paragraphs:
                text += para.text + "\n"
        elif ext in ('.xls', '.xlsx'):
            import openpyxl
            wb = openpyxl.load_workbook(filepath, data_only=True)
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    line = '\t'.join(str(c) if c is not None else '' for c in row)
                    text += line + "\n"
                text += "\n"
        elif ext == '.txt':
            with open(filepath, 'r', encoding='utf-8', errors='replace') as f:
                text = f.read()
    except Exception as e:
        logging.warning("Failed to extract text from %s: %s", filepath, e)

    return text.strip()


# ---------------------------------------------------------------------------
# State persistence
# ---------------------------------------------------------------------------

def _load_state(output_dir):
    path = os.path.join(output_dir, _STATE_FILENAME)
    if os.path.exists(path):
        try:
            with open(path, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception as e:
            logging.warning("Failed to load state: %s", e)
    return {"processed_ids": []}


def _save_state(output_dir, state):
    os.makedirs(output_dir, exist_ok=True)
    with open(os.path.join(output_dir, _STATE_FILENAME), "w", encoding="utf-8") as f:
        json.dump(state, f, ensure_ascii=False)
    logging.info("State saved (%d IDs)", len(state.get("processed_ids", [])))


def _save_markdown(content, output_dir, article_id):
    articles_dir = os.path.join(output_dir, "articles")
    os.makedirs(articles_dir, exist_ok=True)
    safe_id = re.sub(r'[\\/:*?"<>|]', "_", article_id)
    path = os.path.join(articles_dir, f"{safe_id}.md")
    with open(path, "w", encoding="utf-8") as f:
        f.write(content)
    return path


# ---------------------------------------------------------------------------
# KB upload
# ---------------------------------------------------------------------------

def _upload_to_kb(filepath, kb_id, tenant_id, parser_id="laws"):
    from api.db.services.knowledgebase_service import KnowledgebaseService
    from api.db.services.file_service import FileService
    from api.db.services.document_service import DocumentService

    ok, kb = KnowledgebaseService.get_by_id(kb_id)
    if not ok or kb is None:
        logging.warning("KB %s not found, skip upload %s", kb_id, os.path.basename(filepath))
        return []

    # KB-level dedup
    existing = set()
    try:
        for d in DocumentService.select(DocumentService.model.name).where(
            DocumentService.model.kb_id == kb_id
        ):
            existing.add(d.name)
    except Exception:
        pass
    fname = os.path.basename(filepath)
    if fname in existing:
        logging.info("Skip duplicate: %s", fname)
        return []

    with open(filepath, "rb") as f:
        blob = f.read()

    class _FileObj:
        def __init__(self, filename, blob):
            self.id = get_uuid()
            self.filename = filename
            self.blob = blob

        def read(self):
            return self.blob

    file_obj = _FileObj(fname, blob)
    errs, doc_pairs = FileService.upload_document(kb, [file_obj], tenant_id)

    if errs:
        logging.warning("Upload errors: %s", errs)

    for doc, _ in doc_pairs:
        doc_id = doc["id"]
        logging.info("Document %s uploaded to KB %s", doc_id, kb_id)
        try:
            DocumentService.update_by_id(doc_id, {"parser_id": parser_id})
        except Exception as e:
            logging.error("Failed to update parser_id for %s: %s", doc_id, e)
        try:
            from api.db.services.task_service import queue_tasks
            from api.db.services.file2document_service import File2DocumentService
            bucket, name = File2DocumentService.get_storage_address(doc_id=doc_id)
            queue_tasks(doc, bucket, name, 0)
        except Exception as e:
            logging.error("Failed to queue parsing for %s: %s", doc_id, e)
    return doc_pairs


# ---------------------------------------------------------------------------
# Markdown builder
# ---------------------------------------------------------------------------

def _build_markdown(item, detail, attachments_dir):
    module = item.get('module', '')
    lines = [
        f"# {detail.get('title', item['title'])}",
        "",
        f"**来源**: {_SITE_NAME}",
        f"**栏目**: {module}",
    ]
    if detail.get('pub_date'):
        lines.append(f"**发布日期**: {detail['pub_date']}")
    elif item.get('date'):
        lines.append(f"**发布日期**: {item['date']}")
    lines.append(f"**原文链接**: {item['href']}")

    lines.append("")
    lines.append("---")
    lines.append("")

    if detail.get('content_text'):
        lines.append(detail['content_text'])
    else:
        lines.append("(无法提取正文内容)")

    if attachments_dir and os.path.isdir(attachments_dir):
        lines.append("")
        lines.append("---")
        lines.append("")
        lines.append("## 附件")
        lines.append("")
        for fname in sorted(os.listdir(attachments_dir)):
            fpath = os.path.join(attachments_dir, fname)
            fsize = os.path.getsize(fpath)
            lines.append(f"- **{fname}** ({fsize:,} bytes)")

    return "\n".join(lines)


# ---------------------------------------------------------------------------
# Single item processing
# ---------------------------------------------------------------------------

def _process_item(sess, item, output_dir, kb_id, tenant_id):
    """Fetch detail or download attachment, save markdown, upload to KB."""
    _safe_print(f"\n  [{item['module']}] {item['title'][:80]}")
    _safe_print(f"  {item['href'][:150]}")

    attachments_dir = None

    if _is_attach_url(item['href']):
        _safe_print("  Type: direct attachment")
        data = _download_file(sess, item['href'])
        if not data:
            _safe_print("  Download FAILED")
            return False

        fname = os.path.basename(urlparse(item['href']).path) or "attachment"
        safe_key = _sanitize_filename(item['href'], 100)
        attachments_dir = os.path.join(output_dir, "attachments", safe_key)
        os.makedirs(attachments_dir, exist_ok=True)
        fpath = os.path.join(attachments_dir, fname)
        with open(fpath, 'wb') as f:
            f.write(data)
        _safe_print(f"  Downloaded: {fname} ({len(data):,} bytes)")

        ext_text = _extract_file_text(fpath)
        _safe_print(f"  Extracted: {len(ext_text)} chars")

        if fname.lower().endswith('.zip'):
            _extract_zip(fpath, attachments_dir)
            for extracted_file in sorted(os.listdir(attachments_dir)):
                epath = os.path.join(attachments_dir, extracted_file)
                if os.path.isfile(epath) and epath != fpath:
                    extracted_text = _extract_file_text(epath)
                    if extracted_text:
                        ext_text += f"\n\n--- {extracted_file} ---\n\n{extracted_text}"

        detail = {
            'title': item['title'],
            'pub_date': item.get('date', ''),
            'content_text': ext_text or f"附件: {fname}",
            'attachments': [(fname, item['href'])],
        }
    else:
        try:
            r = sess.get(item['href'], timeout=60, verify=False)
            if r.status_code != 200:
                _safe_print(f"  HTTP {r.status_code} — skipped")
                return False
            r.encoding = 'utf-8'
            html = r.text
        except Exception as e:
            _safe_print(f"  Fetch error: {e}")
            return False

        detail = _parse_detail(html, item['href'])
        _safe_print(f"  Date: {detail.get('pub_date', item.get('date', ''))} | "
                    f"Content: {len(detail['content_text'])} chars")

        if detail['attachments']:
            safe_key = _sanitize_filename(item['href'], 100)
            attachments_dir = os.path.join(output_dir, "attachments", safe_key)
            os.makedirs(attachments_dir, exist_ok=True)
            for att_name, att_url in detail['attachments']:
                _safe_print(f"  Downloading: {att_name[:60]}")
                data = _download_file(sess, att_url)
                if data:
                    fname = _sanitize_filename(att_name, 100)
                    fpath = os.path.join(attachments_dir, fname)
                    with open(fpath, 'wb') as f:
                        f.write(data)
                    _safe_print(f"    OK ({len(data):,} bytes)")
                    if fname.lower().endswith('.zip'):
                        _extract_zip(fpath, attachments_dir)

    md_content = _build_markdown(item, detail, attachments_dir)
    md_path = _save_markdown(md_content, output_dir, item['href'])
    _safe_print(f"  Markdown: {md_path}")

    if kb_id:
        try:
            _upload_to_kb(md_path, kb_id, tenant_id)
        except Exception as e:
            logging.error("Upload failed: %s", e)

    return True


# ---------------------------------------------------------------------------
# Main crawl logic
# ---------------------------------------------------------------------------

def crawl(output_dir, kb_id, tenant_id, max_runtime=3300):
    start_time = time.time()

    _safe_print("=" * 60)
    _safe_print(f"FJNJ Crawler — {_SITE_NAME}")
    _safe_print(f"Target: {_ZWGK_URL}")
    _safe_print("=" * 60)

    os.makedirs(output_dir, exist_ok=True)
    state = _load_state(output_dir)
    processed_ids = set(state.get("processed_ids", []))

    sess = _init_session()

    # Fetch all items from the SSR page
    _safe_print("\n--- Parsing zwgk page (multi-module + multi-tab) ---")
    items = _fetch_all_items(sess)

    # Count new vs already processed
    new_items = [it for it in items if it['href'] not in processed_ids]
    _safe_print(f"\n{'='*60}")
    _safe_print(f"Total items: {len(items)} | New: {len(new_items)} | "
                f"Already processed: {len(items) - len(new_items)}")
    _safe_print(f"{'='*60}")

    total_processed = 0
    for item in new_items:
        remaining = max_runtime - (time.time() - start_time)
        grace = min(120, max_runtime * 0.05)
        if remaining < grace:
            _safe_print(f"\nTimeout approaching ({remaining:.0f}s < {grace:.0f}s) — stopping")
            break

        success = _process_item(sess, item, output_dir, kb_id, tenant_id)
        if success:
            processed_ids.add(item['href'])
            state["processed_ids"] = list(processed_ids)
            total_processed += 1
            if total_processed % 10 == 0:
                _save_state(output_dir, state)

        _request_delay(*_ARTICLE_DELAY)

    _save_state(output_dir, state)

    elapsed = time.time() - start_time
    _safe_print(f"\n{'='*60}")
    _safe_print(f"Crawl complete. {total_processed} new items in {elapsed:.0f}s")
    _safe_print(f"{'='*60}")


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(
        description="FJNJ Crawler — fjnj.gov.cn (南靖县人民政府)"
    )
    parser.add_argument("--tenant-id", required=True, help="Tenant ID")
    parser.add_argument("--target-url", default=_ZWGK_URL,
                        help="Target URL (default: %(default)s)")
    parser.add_argument("--kb-id", default=None, help="Knowledge base ID")
    parser.add_argument("--task-name", default="fjnj_crawler",
                        help="Task name")
    parser.add_argument("--output-dir", default=None, help="Output directory")
    parser.add_argument("--max-runtime", type=int, default=3300,
                        help="Maximum runtime in seconds (default: 3300)")
    parser.add_argument("--project-root", default=None, help="Project root")
    parser.add_argument("--llm-id", default=None, help="LLM ID (unused)")
    parser.add_argument("--llm-model", default=None, help="LLM model (unused)")
    parser.add_argument("--access-token", default=None, help="Access token (unused)")

    args = parser.parse_args()

    if args.project_root:
        sys.path.insert(0, args.project_root)
        os.chdir(args.project_root)

    output_dir = args.output_dir or os.path.join(_PROJECT_ROOT, "rag", args.task_name.strip())
    os.makedirs(output_dir, exist_ok=True)

    init_root_logger("fjnj_crawler")
    logging.info("FJNJ Crawler | task=%s | output=%s", args.task_name, output_dir)

    settings.init_settings()

    try:
        crawl(
            output_dir=output_dir,
            kb_id=args.kb_id,
            tenant_id=args.tenant_id,
            max_runtime=args.max_runtime,
        )
    except KeyboardInterrupt:
        _safe_print("\nInterrupted by user")
        logging.info("Interrupted by user")
    except Exception as e:
        logging.exception("Fatal error: %s", e)
        _safe_print(f"\nFATAL: {e}")
        raise


if __name__ == "__main__":
    CONSUMER_NAME = "fjnj_crawler"
    init_root_logger(CONSUMER_NAME)
    main()
