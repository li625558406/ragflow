#!/usr/bin/env python3
#
#  Copyright 2026 The InfiniFlow Authors. All Rights Reserved.
#
#  Licensed under the Apache License, Version 2.0 (the "License");
#  you may not use this file except in compliance with the License.
#  You may obtain a copy of the License at
#
#      http://www.apache.org/licenses/LICENSE-2.0
#
#  Unless required by applicable law or agreed to in writing, software
#  distributed under the License is distributed on an "AS IS" BASIS,
#  WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
#  See the License for the specific language governing permissions and
#  limitations under the License.
#
"""
Crawler for www.ndrc.gov.cn — 国家发展改革委 政策发布 (ALL data).

Targets (5 sections):
  1. /xxgk/zcfb/fzggwl/ — 发展改革委令
  2. /xxgk/zcfb/ghxwj/  — 规范性文件
  3. /xxgk/zcfb/ghwb/   — 规划文本
  4. /xxgk/zcfb/gg/     — 公告
  5. /xxgk/zcfb/tz/     — 通知

Each section has index_N.html pagination (page 1 = index.html, page N = index_{N-1}.html).
Gets ALL list data. Cross-section dedup by article_id.
Server-rendered HTML — no JS rendering needed.

Data flow
---------
  1. Listing: GET each section URL -> parse .u-list > li for title, URL, date (YYYY/MM/DD)
  2. Pagination: extract page count from createPageHTML(), fetch index_N.html pages
  3. Detail: GET each .html detail page -> parse h2.article_title, meta tags,
     div.article_con content, div.attachment links
  4. Attachments: direct download (relative or absolute URLs)
  5. ZIP auto-extraction with content parsing (pdfplumber, python-docx, openpyxl)

Checkpoint/resume: state saved per batch. Time-bounded check
(default 3300s) stops gracefully before the 3600s task-timeout window.

Usage
-----
    python ndrc_fzggwl_crawler.py \
        --tenant-id <TENANT_ID> \
        --kb-id <KB_ID> \
        --task-name <NAME>

    # Optional:
        --max-runtime 3300    # Max runtime before graceful stop
        --full                # Ignore saved state, re-crawl
"""

import argparse
import hashlib
import json
import logging
import os
import random
import re
import sys
import time
import urllib.parse
import urllib.request
import zipfile
from datetime import datetime, timedelta

_SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
_PROJECT_ROOT = os.path.abspath(os.path.join(_SCRIPT_DIR, "..", ".."))
sys.path.insert(0, _PROJECT_ROOT)

from common import settings
from common.log_utils import init_root_logger
from common.misc_utils import get_uuid
from bs4 import BeautifulSoup


# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------

_SITE_ROOT = "https://www.ndrc.gov.cn"

# Five section targets with their labels
_LISTING_SECTIONS = [
    ("/xxgk/zcfb/fzggwl/", "发展改革委令"),
    ("/xxgk/zcfb/ghxwj/",  "规范性文件"),
    ("/xxgk/zcfb/ghwb/",   "规划文本"),
    ("/xxgk/zcfb/gg/",     "公告"),
    ("/xxgk/zcfb/tz/",     "通知"),
]

# Checkpoint batch size (articles)
_BATCH_SIZE = 3

# Default max runtime (55 min, 5 min margin)
_MAX_RUNTIME_DEFAULT = 3300

# Anti-crawling delays (seconds)
_REQUEST_DELAY_MIN = 0.3
_REQUEST_DELAY_MAX = 1.0

# State filename
_STATE_FILENAME = "_crawler_state.json"

# Max pagination pages per section
_MAX_PAGES = 50

_USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/120.0.0.0 Safari/537.36"
)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _safe_print(msg):
    try:
        print(msg)
    except UnicodeEncodeError:
        print(msg.encode("gbk", errors="replace").decode("gbk"))


def _sanitize_filename(text, max_len=120):
    if not text:
        return "untitled"
    name = re.sub(r'[\\/:*?"<>|]', "_", str(text).strip())
    name = re.sub(r'\s+', "_", name).strip("._ ")
    return (name or "untitled")[:max_len]


def _extract_article_id(url):
    """Extract article ID from URL.

    Patterns:
      .../t20260409_1404577.html  -> 1404577
      .../P020260409348498606128.pdf -> P020260409348498606128
    """
    m = re.search(r'/t(\d{8})_(\d+)\.html', url)
    if m:
        return m.group(2)
    m = re.search(r'/([A-Z]\d{15,})\.(?:pdf|doc|docx|xlsx|xls|zip|rar|ofd)', url)
    if m:
        return m.group(1)
    path = urllib.parse.urlparse(url).path
    m = re.search(r'/(\d{7,})\.(?:htm|html)', path)
    if m:
        return m.group(1)
    return hashlib.md5(path.encode()).hexdigest()[:12]


def _resolve_url(href, base_url):
    """Resolve a relative href against a base URL."""
    if href.startswith("http://") or href.startswith("https://"):
        return href
    if href.startswith("//"):
        return "https:" + href
    return urllib.parse.urljoin(base_url, href)


# ---------------------------------------------------------------------------
# HTTP helpers
# ---------------------------------------------------------------------------

def _http_get(url, referer=None, timeout=30):
    """GET request returning bytes or None."""
    req = urllib.request.Request(url)
    req.add_header("User-Agent", _USER_AGENT)
    if referer:
        req.add_header("Referer", referer)
    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            return resp.read()
    except Exception as e:
        logging.error("GET %s failed: %s", url, e)
        return None


def _download_binary(url, referer=None, timeout=60):
    """Download binary content; returns bytes or None."""
    req = urllib.request.Request(url)
    req.add_header("User-Agent", _USER_AGENT)
    if referer:
        req.add_header("Referer", referer)
    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            return resp.read()
    except Exception as e:
        logging.error("Download failed: %s — %s", url[:120], e)
        return None


# ---------------------------------------------------------------------------
# Listing page parsing
# ---------------------------------------------------------------------------

def _parse_listing(html_bytes, base_url, section_label, seen_ids):
    """Parse a listing page for articles.

    Format: .list > ul.u-list > li with:
      - <a href="..." title="..."> for title/URL
      - <span>YYYY/MM/DD</span> for date
      - .popbox for related interpretations (skipped)
    """
    if not html_bytes:
        return []

    html = html_bytes.decode("utf-8", errors="replace")
    soup = BeautifulSoup(html, "html.parser")
    articles = []

    ul = soup.find("ul", class_="u-list")
    if not ul:
        return articles

    for li in ul.find_all("li", recursive=False):
        a_tag = li.find("a", href=True, title=True)
        if not a_tag:
            # The first a has title, try any a with href
            a_tag = li.find("a", href=True)
        if not a_tag:
            continue

        href = a_tag["href"].strip()
        title = (a_tag.get("title") or a_tag.get_text()).strip()
        if not title or len(title) < 2:
            continue

        if href.startswith("javascript:") or href == "#":
            continue

        # Date in YYYY/MM/DD format
        span = li.find("span")
        date_text = span.get_text(strip=True) if span else ""
        if not date_text:
            continue  # skip items without date (e.g., nested .popbox links)
        # Normalize to YYYY-MM-DD
        date_str = date_text.replace("/", "-")

        article_id = _extract_article_id(href)
        if article_id in seen_ids:
            continue
        seen_ids.add(article_id)

        articles.append({
            "title": title,
            "url": _resolve_url(href, base_url),
            "date_str": date_str,
            "article_id": article_id,
            "section": section_label,
        })

    return articles


def _extract_page_count(html_bytes):
    """Extract total page count from createPageHTML(total, current, ...)."""
    if not html_bytes:
        return 1
    html = html_bytes.decode("utf-8", errors="replace")
    m = re.search(r'createPageHTML\((\d+),\s*\d+,\s*"index",\s*"html"\)', html)
    if m:
        return int(m.group(1))
    return 1


def _crawl_section_paginated(section_path, section_label, referer, seen_ids, max_pages=_MAX_PAGES):
    """Crawl a paginated section listing, returning all articles across all pages."""
    base_url = _SITE_ROOT + section_path
    all_articles = []

    # Page 1: index.html (the base URL)
    time.sleep(random.uniform(_REQUEST_DELAY_MIN, _REQUEST_DELAY_MAX))
    html_bytes = _http_get(base_url, referer=referer)
    if not html_bytes:
        return all_articles

    page_count = _extract_page_count(html_bytes)
    _safe_print("      {} page(s) total".format(page_count))
    sys.stdout.flush()

    articles = _parse_listing(html_bytes, base_url, section_label, seen_ids)
    all_articles.extend(articles)

    # Pages 2..N: index_{N-1}.html
    for page_num in range(1, min(page_count, max_pages + 1)):
        page_url = base_url.rstrip("/") + "/index_{}.html".format(page_num)
        time.sleep(random.uniform(_REQUEST_DELAY_MIN, _REQUEST_DELAY_MAX))
        html_bytes = _http_get(page_url, referer=base_url)
        if not html_bytes:
            continue

        articles = _parse_listing(html_bytes, page_url, section_label, seen_ids)
        if not articles:
            # Empty page — probably hit a real boundary
            break
        all_articles.extend(articles)

    return all_articles


# ---------------------------------------------------------------------------
# Detail page parsing
# ---------------------------------------------------------------------------

def _parse_detail_html(html_bytes, detail_url):
    """Parse an ndrc.gov.cn article detail page.

    Title: h2.article_title or meta ArticleTitle
    Metadata: meta PubDate (YYYY-MM-DD HH:MM:SS), ContentSource
    Content: div.article_con (contains TRS_Editor)
    Attachments: div.attachment div.attachment_r a links
    """
    if not html_bytes:
        return None

    html = html_bytes.decode("utf-8", errors="replace")
    soup = BeautifulSoup(html, "html.parser")

    info = {
        "title": "",
        "pub_date": "",
        "info_source": "",
        "article_id": _extract_article_id(detail_url),
        "content_html": "",
        "content_text": "",
        "attachments": [],
    }

    # Title: h2.article_title
    h2 = soup.find("h2", class_="article_title")
    if h2:
        info["title"] = h2.get_text(strip=True)

    # Fallback: meta ArticleTitle
    if not info["title"]:
        meta_title = soup.find("meta", attrs={"name": "ArticleTitle"})
        if meta_title and meta_title.get("content"):
            info["title"] = meta_title["content"].strip()

    # Metadata: meta PubDate
    meta_pubdate = soup.find("meta", attrs={"name": "PubDate"})
    if meta_pubdate and meta_pubdate.get("content"):
        info["pub_date"] = meta_pubdate["content"].strip()

    # Metadata: meta ContentSource
    meta_source = soup.find("meta", attrs={"name": "ContentSource"})
    if meta_source and meta_source.get("content"):
        info["info_source"] = meta_source["content"].strip()

    # Also try div.time and div.ly
    shezhi = soup.find("div", class_="shezhi")
    if shezhi:
        time_div = shezhi.find("div", class_="time")
        if time_div:
            t = time_div.get_text(strip=True)
            m = re.search(r'(\d{4}/\d{2}/\d{2})', t)
            if m and not info["pub_date"]:
                info["pub_date"] = m.group(1).replace("/", "-")

        ly_div = shezhi.find("div", class_="ly")
        if ly_div:
            s = ly_div.get_text(strip=True)
            s = re.sub(r'^来源[：:]\s*', '', s).strip()
            if s and not info["info_source"]:
                info["info_source"] = s

    # Content: div.article_con
    content_div = soup.find("div", class_="article_con")
    if content_div:
        for tag in content_div.find_all(["script", "style"]):
            tag.decompose()
        info["content_html"] = content_div.decode_contents()
        info["content_text"] = content_div.get_text(separator="\n", strip=True)

    # Attachments: div.attachment div.attachment_r a
    att_div = soup.find("div", class_="attachment")
    if att_div:
        att_r = att_div.find("div", class_="attachment_r")
        if att_r:
            for a_tag in att_r.find_all("a", href=True):
                href = a_tag["href"].strip()
                text = (a_tag.get_text() or "").strip()
                att_url = _resolve_url(href, detail_url)
                info["attachments"].append({
                    "filename": text or os.path.basename(urllib.parse.urlparse(att_url).path),
                    "url": att_url,
                    "type": "download",
                })

    return info


# ---------------------------------------------------------------------------
# Attachment download + processing
# ---------------------------------------------------------------------------

def _download_attachments(attachments, download_dir):
    """Download attachments to local directory. Returns list of local file paths."""
    os.makedirs(download_dir, exist_ok=True)
    local_files = []

    for att in attachments:
        url = att.get("url", "")
        if not url:
            continue
        fname = _sanitize_filename(att.get("filename", "attachment"), max_len=120)
        ext = os.path.splitext(urllib.parse.urlparse(url).path)[1].lower()
        if ext and not fname.lower().endswith(ext):
            fname += ext

        filepath = os.path.join(download_dir, fname)
        if os.path.exists(filepath) and os.path.getsize(filepath) > 0:
            local_files.append(filepath)
            continue

        data = _download_binary(url)
        if data:
            with open(filepath, "wb") as f:
                f.write(data)
            local_files.append(filepath)
            time.sleep(random.uniform(0.1, 0.3))

    return local_files


def _extract_zip(filepath):
    """Extract ZIP file to same directory. Returns list of extracted file paths."""
    extracted = []
    extract_dir = os.path.splitext(filepath)[0] + "_extracted"
    os.makedirs(extract_dir, exist_ok=True)
    try:
        with zipfile.ZipFile(filepath, "r") as zf:
            for name in zf.namelist():
                safe_name = _sanitize_filename(name, max_len=120)
                out_path = os.path.join(extract_dir, safe_name)
                os.makedirs(os.path.dirname(out_path) or extract_dir, exist_ok=True)
                with zf.open(name) as src, open(out_path, "wb") as dst:
                    dst.write(src.read())
                extracted.append(out_path)
    except Exception as e:
        logging.warning("ZIP extraction failed for %s: %s", filepath, e)
    return extracted


# ---------------------------------------------------------------------------
# Text extraction from binary files
# ---------------------------------------------------------------------------

def _extract_text_from_file(filepath):
    """Extract plain text from PDF / DOCX / XLSX / TXT files."""
    ext = os.path.splitext(filepath)[1].lower()
    try:
        if ext == ".txt":
            with open(filepath, "r", encoding="utf-8", errors="replace") as f:
                return f.read()
        elif ext == ".pdf":
            import pdfplumber
            parts = []
            with pdfplumber.open(filepath) as pdf:
                for pg in pdf.pages:
                    text = pg.extract_text()
                    if text:
                        parts.append(text)
            return "\n\n".join(parts)
        elif ext in (".docx", ".doc"):
            import docx
            doc = docx.Document(filepath)
            return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        elif ext in (".xlsx", ".xls"):
            import openpyxl
            wb = openpyxl.load_workbook(filepath, read_only=True)
            parts = []
            for ws in wb.worksheets:
                rows = []
                for row in ws.iter_rows(values_only=True):
                    rows.append(" | ".join(str(c) if c is not None else "" for c in row))
                if rows:
                    parts.append(f"### {ws.title}\n" + "\n".join(rows))
            wb.close()
            return "\n\n".join(parts)
    except Exception as e:
        logging.warning("Text extraction failed for %s: %s", filepath, e)
    return None


# ---------------------------------------------------------------------------
# Markdown builder
# ---------------------------------------------------------------------------

def _build_markdown(detail, download_dir, source_url):
    """Build a Markdown document from article detail + attachments."""
    info = detail or {}
    title = info.get("title", "无标题")
    pub_date = info.get("pub_date", "")
    info_source = info.get("info_source", "")

    lines = [
        f"# {title}",
        "",
        f"**数据来源:** 国家发展和改革委员会",
        f"**页面地址:** {source_url}",
        f"**抓取时间:** {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
    ]
    if pub_date:
        lines.append(f"**发布时间:** {pub_date}")
    if info_source:
        lines.append(f"**信息来源:** {info_source}")
    lines.append("")

    # Main content
    content_text = info.get("content_text", "")
    if content_text:
        lines.append("---")
        lines.append("")
        lines.append("## 正文")
        lines.append("")
        content_clean = re.sub(r"\n{3,}", "\n\n", content_text)
        lines.append(content_clean)
        lines.append("")

    # Attachments
    attachments = info.get("attachments", [])
    if attachments:
        lines.append("---")
        lines.append("")
        lines.append("## 附件")
        lines.append("")

        for att in attachments:
            fname = att.get("filename", "unknown")
            att_url = att.get("url", "")
            lines.append(f"- [{fname}]({att_url})")
        lines.append("")

        # Embed extracted attachment text
        if download_dir and os.path.isdir(download_dir):
            lines.append("### 附件内容")
            lines.append("")
            for att in attachments:
                if att.get("type") != "download":
                    continue
                fname = att.get("filename", "")
                local_path = os.path.join(download_dir, fname)
                safe_name = _sanitize_filename(fname, max_len=120)
                alt_path = os.path.join(download_dir, safe_name)
                if not os.path.exists(local_path) and os.path.exists(alt_path):
                    local_path = alt_path
                if not os.path.exists(local_path):
                    for root, _, files in os.walk(download_dir):
                        for fn in files:
                            if fn == safe_name or fn == fname:
                                local_path = os.path.join(root, fn)
                                break
                if not os.path.exists(local_path):
                    continue

                lines.append(f"#### {fname}")
                lines.append("")
                text = _extract_text_from_file(local_path)
                if text and text.strip():
                    if len(text) > 50000:
                        text = text[:50000] + "\n\n（内容过长，已截断）"
                    lines.append(text)
                else:
                    lines.append("（无法提取文本内容）")
                lines.append("")

    return "\n".join(lines)


# ---------------------------------------------------------------------------
# State management
# ---------------------------------------------------------------------------

def _load_state(output_dir):
    path = os.path.join(output_dir, _STATE_FILENAME)
    if os.path.exists(path):
        try:
            with open(path, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception as e:
            logging.warning("Failed to load crawler state: %s", e)
    return {"processed_ids": []}


def _save_state(output_dir, state):
    path = os.path.join(output_dir, _STATE_FILENAME)
    os.makedirs(output_dir, exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(state, f, ensure_ascii=False)


# ---------------------------------------------------------------------------
# KB upload
# ---------------------------------------------------------------------------

def _upload_to_kb(md_content, attachment_files, kb_id, tenant_id, folder_name):
    """Upload markdown + attachment files to knowledge base."""
    from api.db.services.knowledgebase_service import KnowledgebaseService
    from api.db.services.file_service import FileService
    from api.db.services.document_service import DocumentService

    ok, kb = KnowledgebaseService.get_by_id(kb_id)
    if not ok:
        raise LookupError(f"Knowledge base {kb_id} not found")

    class _FO:
        def __init__(self, fn, b):
            self.id = get_uuid()
            self.filename = fn
            self.blob = b

        def read(self):
            return self.blob

    fo = _FO(f"{folder_name}.md", md_content.encode("utf-8"))
    errs, pairs = FileService.upload_document(kb, [fo], tenant_id)
    if errs:
        logging.warning("MD upload errors: %s", errs)
    for doc, _ in pairs:
        did = doc["id"]
        try:
            DocumentService.begin2parse(did)
            DocumentService.run(tenant_id, doc, {})
        except Exception as e:
            logging.error("Queue parse for %s: %s", did, e)

    for fp in attachment_files:
        fname = os.path.basename(fp)
        with open(fp, "rb") as f:
            blob = f.read()
        fo2 = _FO(fname, blob)
        errs2, pairs2 = FileService.upload_document(kb, [fo2], tenant_id)
        if errs2:
            logging.warning("Attachment upload errors: %s", errs2)
        for doc, _ in pairs2:
            did = doc["id"]
            try:
                DocumentService.begin2parse(did)
                DocumentService.run(tenant_id, doc, {})
            except Exception as e:
                logging.error("Queue parse for %s: %s", did, e)


# ---------------------------------------------------------------------------
# Argument parsing
# ---------------------------------------------------------------------------

def parse_args():
    p = argparse.ArgumentParser(
        description="ndrc.gov.cn fzggwl crawler — 国家发展改革委 政策发布 (ALL data)"
    )
    p.add_argument("--tenant-id", required=True)
    p.add_argument("--kb-id", required=True)
    p.add_argument("--task-name", required=True)
    p.add_argument("--output-dir", default=None,
                   help="Output root directory (default: PROJECT_ROOT/rag/<task_name>)")
    p.add_argument("--full", action="store_true",
                   help="Ignore saved state and re-crawl all")
    p.add_argument("--max-runtime", type=int, default=_MAX_RUNTIME_DEFAULT,
                   help="Max runtime in seconds (default: 3300)")
    for opt in ("--max-days", "--hours", "--max-articles",
                "--llm-id", "--llm-model", "--access-token", "--target-url"):
        p.add_argument(opt, default=None)
    return p.parse_args()


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    args = parse_args()

    _safe_print("\n" + "=" * 60)
    _safe_print("[NDRC-FZGGWL] 国家发展改革委 — 政策发布 crawler")
    _safe_print("[NDRC-FZGGWL] KB: {}".format(args.kb_id))
    _safe_print("[NDRC-FZGGWL] Task: {}".format(args.task_name))
    _safe_print("[NDRC-FZGGWL] Max runtime: {}s".format(args.max_runtime))
    _safe_print("=" * 60 + "\n")
    sys.stdout.flush()

    settings.init_settings()
    logging.info("=== NDRC-FZGGWL crawler started ===")

    output_dir = args.output_dir or os.path.join(
        _PROJECT_ROOT, "rag", args.task_name.strip()
    )
    os.makedirs(output_dir, exist_ok=True)
    _safe_print("[NDRC-FZGGWL] Output: {}\n".format(output_dir))
    sys.stdout.flush()

    # ── State ──────────────────────────────────────────────────────────
    state = _load_state(output_dir) if not args.full else {"processed_ids": []}
    processed_ids = set(state.get("processed_ids", []))

    _safe_print("[NDRC-FZGGWL] Already processed: {} article(s)".format(len(processed_ids)))
    sys.stdout.flush()

    crawl_start = time.time()

    # ── Step 1: Crawl all five listing sections ─────────────────────────
    _safe_print("\n[NDRC-FZGGWL] Step 1/3: Crawling listing pages...")
    sys.stdout.flush()

    seen_ids = set()
    all_articles = []

    for section_path, section_label in _LISTING_SECTIONS:
        elapsed = time.time() - crawl_start
        remaining = args.max_runtime - elapsed
        if remaining < 120:
            _safe_print("[NDRC-FZGGWL] Runtime limit approaching, skipping remaining sections")
            sys.stdout.flush()
            break

        _safe_print("[NDRC-FZGGWL]   Section: {} ({})".format(section_label, section_path))
        sys.stdout.flush()

        articles = _crawl_section_paginated(
            section_path, section_label, _SITE_ROOT + "/", seen_ids
        )
        _safe_print("[NDRC-FZGGWL]     -> {} new article(s)".format(len(articles)))
        sys.stdout.flush()
        all_articles.extend(articles)

    _safe_print("[NDRC-FZGGWL]   Total: {} unique article(s) across {} section(s)".format(
        len(all_articles), len(_LISTING_SECTIONS)))
    sys.stdout.flush()

    if not all_articles:
        _safe_print("[NDRC-FZGGWL] No articles found, done.")
        sys.stdout.flush()
        return

    # Section breakdown
    section_counts = {}
    for a in all_articles:
        sec = a.get("section", "unknown")
        section_counts[sec] = section_counts.get(sec, 0) + 1
    for sec, cnt in sorted(section_counts.items()):
        _safe_print("[NDRC-FZGGWL]     {}: {} article(s)".format(sec, cnt))
    sys.stdout.flush()

    # Filter already-processed
    new_articles = [
        a for a in all_articles
        if a.get("article_id") and a["article_id"] not in processed_ids
    ]
    skipped = len(all_articles) - len(new_articles)
    if skipped:
        _safe_print("[NDRC-FZGGWL] {} already processed, {} new".format(skipped, len(new_articles)))
        sys.stdout.flush()

    if not new_articles:
        _safe_print("[NDRC-FZGGWL] All available articles already processed.")
        sys.stdout.flush()
        return

    # ── Step 2: Process each article ───────────────────────────────────
    _safe_print("\n[NDRC-FZGGWL] Step 2/3: Processing {} article(s)...\n".format(len(new_articles)))
    sys.stdout.flush()

    processed_count = 0
    stopped_early = False
    downloads_dir = os.path.join(output_dir, "downloads")

    for idx, article in enumerate(new_articles, 1):
        # ── Time-bounded check ─────────────────────────────────────────
        elapsed = time.time() - crawl_start
        remaining = args.max_runtime - elapsed
        if remaining < 120:
            _safe_print(
                "\n[NDRC-FZGGWL] Runtime {:.0f}s, {:.0f}s remaining (limit {}s), "
                "stopping gracefully. {} processed. "
                "Next run will resume.".format(elapsed, remaining, args.max_runtime, processed_count))
            sys.stdout.flush()
            stopped_early = True
            break

        article_id = article["article_id"]
        title = article["title"]
        detail_url = article["url"]
        date_str = article.get("date_str", "")
        section_label = article.get("section", "")

        _safe_print("[NDRC-FZGGWL] [{}/{}] [{}] {}...".format(
            idx, len(new_articles), section_label, title[:50]))
        sys.stdout.flush()

        # Handle PDF files directly
        if detail_url.lower().endswith(".pdf"):
            dl_name = "{}_{}".format(article_id[:12], _sanitize_filename(title[:30], 40))
            article_dl_dir = os.path.join(downloads_dir, dl_name)
            os.makedirs(article_dl_dir, exist_ok=True)

            fname = os.path.basename(urllib.parse.urlparse(detail_url).path)
            pdf_path = os.path.join(article_dl_dir, fname)
            pdf_data = _download_binary(detail_url, referer=_SITE_ROOT + "/")
            if pdf_data:
                with open(pdf_path, "wb") as f:
                    f.write(pdf_data)
                pdf_text = _extract_text_from_file(pdf_path)
                md_content = "\n".join([
                    f"# {title}",
                    "",
                    f"**数据来源:** 国家发展和改革委员会",
                    f"**页面地址:** {detail_url}",
                    f"**抓取时间:** {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
                    f"**发布时间:** {date_str}",
                    "",
                    "---",
                    "",
                    "## 正文（PDF提取）",
                    "",
                    pdf_text or "（无法提取PDF文本内容）",
                ])
                date_for_name = date_str or datetime.now().strftime("%Y-%m-%d")
                folder_name = _sanitize_filename(
                    "{}_{}_{}".format(date_for_name, article_id[:12], title[:40]), max_len=120)
                md_path = os.path.join(output_dir, f"{folder_name}.md")
                with open(md_path, "w", encoding="utf-8") as f:
                    f.write(md_content)
                _safe_print("[NDRC-FZGGWL]   Saved PDF ({} chars)".format(len(md_content)))
                sys.stdout.flush()

                if args.kb_id:
                    try:
                        _upload_to_kb(md_content, [pdf_path], args.kb_id,
                                     args.tenant_id, folder_name)
                        _safe_print("[NDRC-FZGGWL]   Uploaded to KB {}".format(args.kb_id))
                        sys.stdout.flush()
                    except Exception as e:
                        logging.error("KB upload failed: %s", e)

                processed_ids.add(article_id)
                processed_count += 1
                time.sleep(random.uniform(_REQUEST_DELAY_MIN, _REQUEST_DELAY_MAX))
                continue

        # Fetch HTML detail page
        time.sleep(random.uniform(_REQUEST_DELAY_MIN, _REQUEST_DELAY_MAX))
        html_bytes = _http_get(detail_url, referer=_SITE_ROOT + "/xxgk/zcfb/")
        if not html_bytes:
            logging.warning("Failed to fetch detail: %s", detail_url)
            processed_ids.add(article_id)
            continue

        detail = _parse_detail_html(html_bytes, detail_url)
        if not detail:
            logging.warning("Failed to parse detail: %s", detail_url)
            processed_ids.add(article_id)
            continue

        # Prefer detail page title if longer
        if detail.get("title") and len(detail["title"]) > len(title):
            title = detail["title"]

        # Download attachments
        attachments = detail.get("attachments", [])
        local_files = []
        article_dl_dir = ""
        if attachments:
            dl_name = "{}_{}".format(article_id[:12], _sanitize_filename(title[:30], 40))
            article_dl_dir = os.path.join(downloads_dir, dl_name)
            local_files = _download_attachments(attachments, article_dl_dir)
            # Extract ZIPs
            for fp in list(local_files):
                is_zip = fp.lower().endswith(".zip")
                if not is_zip and os.path.exists(fp) and os.path.getsize(fp) >= 4:
                    with open(fp, "rb") as f:
                        is_zip = f.read(4) == b"PK\x03\x04"
                if is_zip:
                    extracted = _extract_zip(fp)
                    local_files.remove(fp)
                    local_files.extend(extracted)

        # Build markdown
        md_content = _build_markdown(detail, article_dl_dir, detail_url)

        # Save markdown locally
        date_for_name = date_str or datetime.now().strftime("%Y-%m-%d")
        folder_name = _sanitize_filename(
            "{}_{}_{}".format(date_for_name, article_id[:12], title[:40]), max_len=120)
        md_path = os.path.join(output_dir, f"{folder_name}.md")
        with open(md_path, "w", encoding="utf-8") as f:
            f.write(md_content)
        _safe_print("[NDRC-FZGGWL]   Saved ({} chars, {} attachments)".format(
            len(md_content), len(local_files)))
        sys.stdout.flush()

        # Upload to KB
        if args.kb_id:
            try:
                _upload_to_kb(md_content, local_files, args.kb_id,
                             args.tenant_id, folder_name)
                _safe_print("[NDRC-FZGGWL]   Uploaded to KB {}".format(args.kb_id))
                sys.stdout.flush()
            except Exception as e:
                logging.error("KB upload failed: %s", e)
                _safe_print("[NDRC-FZGGWL]   Upload error: {}".format(e))
                sys.stdout.flush()

        processed_ids.add(article_id)
        processed_count += 1

        # Checkpoint every batch
        if processed_count % _BATCH_SIZE == 0:
            _save_state(output_dir, {"processed_ids": list(processed_ids)})
            _safe_print("[NDRC-FZGGWL]   Checkpoint ({} processed)".format(processed_count))
            sys.stdout.flush()

    # ── Final state ────────────────────────────────────────────────────
    state["processed_ids"] = list(processed_ids)
    _save_state(output_dir, state)

    _safe_print("\n" + "=" * 60)
    _safe_print("[NDRC-FZGGWL] Crawl complete — {} new article(s)".format(processed_count))
    if stopped_early:
        _safe_print("[NDRC-FZGGWL] Stopped early, will resume next run")
    _safe_print("[NDRC-FZGGWL] Target: ALL data across {} sections".format(len(_LISTING_SECTIONS)))
    _safe_print("=" * 60 + "\n")
    sys.stdout.flush()
    logging.info("=== NDRC-FZGGWL crawler finished ===")


if __name__ == "__main__":
    CONSUMER_NAME = "ndrc_fzggwl_crawler"
    init_root_logger(CONSUMER_NAME)
    main()
