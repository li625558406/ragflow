# rag/svr/document_rewrite/sections.py
"""heading 切节与目录提取（纯函数，python-docx）。

切节协议（设计 §4）：
- 按 heading 1（样式名 Heading 1/标题 1，或段落直接大纲级别 0）切顶层节；
- section_no 即文档顺序编号，与 Word 自动编号无关；
- 顶层节标题段落本身永不参与正文替换（docx_edit 侧负责跳过）；
- 无任何 heading 的文档抛 NoSectionError，由工具层引导 LLM 建议全文重新生成。
"""
from __future__ import annotations

from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


class NoSectionError(ValueError):
    """文档没有任何一级标题结构。"""


_HEADING1_STYLE_NAMES = {"heading 1", "标题 1"}


def _direct_outline_level(p: Paragraph) -> int | None:
    """段落 pPr 上的直接大纲级别（w:outlineLvl），无则 None。"""
    pPr = p._p.pPr
    if pPr is None:
        return None
    el = pPr.find(qn("w:outlineLvl"))
    if el is None:
        return None
    try:
        return int(el.get(qn("w:val")))
    except (TypeError, ValueError):
        return None


def _is_heading1(p: Paragraph) -> bool:
    name = ""
    try:
        name = (p.style.name or "").strip().lower()
    except Exception:
        pass
    if name in _HEADING1_STYLE_NAMES:
        return True
    return _direct_outline_level(p) == 0


def split_sections(doc) -> list[dict]:
    """切节。返回按文档顺序的节列表：
    {section_no, title, para_start, para_end, preview, word_count}
    - para_start：标题段落在 doc.paragraphs 中的索引（含）；
    - para_end：本节最后一个正文段落索引（含）；邻接标题时 == para_start；
    - preview：正文首行前 40 字；word_count：正文总字数（表格内容不计）。
    注意：索引空间是 doc.paragraphs（body 顶层段落），表格元素在该空间不可见，
    docx_edit 按同一索引空间操作即天然保留表格。
    """
    paras = list(doc.paragraphs)
    headings = [
        (i, p) for i, p in enumerate(paras)
        if _is_heading1(p) and (p.text or "").strip()
    ]
    if not headings:
        raise NoSectionError(
            "该文档没有任何一级标题（Heading 1）结构，无法按节定位。"
            "建议整体重新生成文档，或先为文档补充标题样式。"
        )
    sections = []
    for idx, (i, p) in enumerate(headings):
        end = headings[idx + 1][0] - 1 if idx + 1 < len(headings) else len(paras) - 1
        body_text = "\n".join(
            (pp.text or "").strip() for pp in paras[i + 1:end + 1] if (pp.text or "").strip()
        )
        sections.append({
            "section_no": idx + 1,
            "title": p.text.strip(),
            "para_start": i,
            "para_end": end,
            "preview": body_text[:40],
            "word_count": len(body_text),
        })
    return sections


def get_section(sections: list[dict], section_no: int) -> dict | None:
    for s in sections:
        if s["section_no"] == section_no:
            return s
    return None


def build_outline(sections: list[dict]) -> str:
    return "\n".join(
        f"第{s['section_no']}节 {s['title']}（约{s['word_count']}字）" for s in sections
    )
