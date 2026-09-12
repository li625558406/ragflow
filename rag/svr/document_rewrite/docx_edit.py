# rag/svr/document_rewrite/docx_edit.py
# -*- coding: utf-8 -*-
"""docx 段落区间替换与格式保真（设计 §6）。

- 只操作 doc.paragraphs 索引空间里的 w:p 元素，区间内 w:tbl（表格）原样保留；
- 顶层节标题段落（para_start）永不删除、永不改写；
- 新段落 pPr 拷贝自首个被删正文段（样式/缩进/行距/对齐）；
- 新 run rPr 拷贝自「正文特征 run」（被删段落中非加粗且文本最长的 run，
  跳过加粗标题 run；全加粗时退化为最长 run）；
- 支持 1段→N段、N段→1段、N段→0段、空节→N段。
原 Document 对象就地修改；版本化由调用方负责（先读 blob → 改 → 存新 blob）。
"""
from __future__ import annotations

import copy

from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


def _pick_body_run_rpr(old_paras: list[Paragraph]):
    """从被删段落中挑「正文特征 run」的 rPr 深拷贝。"""
    best = None
    best_len = -1
    fallback = None
    fallback_len = -1
    for p in old_paras:
        for r in p.runs:
            n = len(r.text or "")
            if n > fallback_len:
                fallback, fallback_len = r, n
            if not r.bold and n > best_len:
                best, best_len = r, n
    chosen = best if best is not None else fallback
    if chosen is None or chosen._r.rPr is None:
        return None
    return copy.deepcopy(chosen._r.rPr)


def _insert_new_paragraph(anchor: Paragraph | None, doc, pPr_tmpl, rPr_tmpl, text: str):
    """在 anchor 段之前插入一个新段落（anchor 为 None 时追加到文档末尾）。"""
    if anchor is not None:
        p = anchor.insert_paragraph_before("")
    else:
        p = doc.add_paragraph("")
    if pPr_tmpl is not None:
        # 空段落无 pPr，直接把拷贝插为第一个子元素（w:pPr 必须是 w:p 首子元素）
        p._p.insert(0, copy.deepcopy(pPr_tmpl))
    run = p.add_run(text)
    if rPr_tmpl is not None:
        # w:rPr 必须是 w:r 首子元素
        run._r.insert(0, copy.deepcopy(rPr_tmpl))
    return p


def replace_section_paragraphs(doc, section: dict, new_paragraphs: list[str]) -> None:
    """把 section（sections.py 切出的节）的正文段替换为 new_paragraphs。
    就地修改 doc；标题段与表格不动。

    契约：section 的 para_start/para_end 必须来自**同一个 doc 当前状态**的
    split_sections——跨节连续替换时前一次替换会使后续节索引漂移，调用方须
    每次替换后重新切节（或倒序替换）。stale 索引不报错，会静默错切。
    """
    paras = list(doc.paragraphs)
    start = section["para_start"] + 1
    end = section["para_end"]

    if end < start:
        # 邻接标题的空节：无旧段可拷贝格式，插到下一节标题之前（即本节末尾）
        anchor = paras[start] if start < len(paras) else None
        for text in new_paragraphs:
            _insert_new_paragraph(anchor, doc, None, None, text)
        return

    old = paras[start:end + 1]
    first_pPr = old[0]._p.pPr
    if first_pPr is not None:
        pPr_tmpl = copy.deepcopy(first_pPr)
        # 段落级分节符绝不拷贝：N 段各带一份 sectPr 会让 Word 版面错乱；
        # 旧段（连同其分节符）随后被删除，剔除是唯一安全选择
        sect = pPr_tmpl.find(qn("w:sectPr"))
        if sect is not None:
            pPr_tmpl.remove(sect)
    else:
        pPr_tmpl = None
    rPr_tmpl = _pick_body_run_rpr(old)

    # 逐条插到首个旧段之前（保持顺序），再删旧段；表格元素全程不被触碰
    for text in new_paragraphs:
        _insert_new_paragraph(old[0], doc, pPr_tmpl, rPr_tmpl, text)
    for p in old:
        p._p.getparent().remove(p._p)
