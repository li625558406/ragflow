# test/test_doc_rewrite_docx_edit.py
# -*- coding: utf-8 -*-
"""段落区间替换对抗测试：格式保真/1↔N段/表格保留/首末节边界/空节插入/标题不可改写。"""
import io

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

from rag.svr.document_rewrite.docx_edit import replace_section_paragraphs
from rag.svr.document_rewrite.sections import split_sections


def _build_doc():
    doc = Document()
    doc.add_heading("第一节", level=1)
    body = doc.add_paragraph("原始正文段落，这是第一节的正文内容。")
    body.add_run("补充run")
    doc.add_heading("第二节", level=1)
    doc.add_paragraph("第二节正文。")
    return doc


def _roundtrip(doc) -> Document:
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return Document(buf)


def test_replace_keeps_heading_and_copies_style():
    doc = _build_doc()
    sections = split_sections(doc)
    replace_section_paragraphs(doc, sections[0], ["新正文一段。", "新正文二段。"])
    out = _roundtrip(doc)
    paras = out.paragraphs
    # 标题段永不被改写
    assert paras[0].text == "第一节"
    assert paras[1].text == "新正文一段。"
    assert paras[2].text == "新正文二段。"
    assert paras[3].text == "第二节"
    # pPr 样式拷贝：新段样式名 == 原正文段样式名
    assert paras[1].style.name == "Normal"
    # 1段 → 2段（N段插入）成功
    assert len(paras) == 5


def test_n_to_one():
    doc = _build_doc()
    doc.add_paragraph("原始正文第二段。")
    sections = split_sections(doc)
    replace_section_paragraphs(doc, sections[0], ["合并后的单段。"])
    out = _roundtrip(doc)
    assert out.paragraphs[1].text == "合并后的单段。"
    assert out.paragraphs[2].text == "第二节"


def test_run_rpr_copied_from_body_run_not_bold_title_run():
    doc = Document()
    doc.add_heading("节", level=1)
    p = doc.add_paragraph("")
    r_bold = p.add_run("加粗")
    r_bold.bold = True
    r_body = p.add_run("正文特征文本，较长一些。")
    r_body.font.size = Pt(12)
    sections = split_sections(doc)
    replace_section_paragraphs(doc, sections[0], ["替换后的正文。"])
    out = _roundtrip(doc)
    new_p = out.paragraphs[1]
    assert new_p.text == "替换后的正文。"
    # rPr 拷贝自「正文特征 run」（非加粗、最长文本），新段不应继承加粗
    assert all(not (r.bold or False) for r in new_p.runs)


def test_table_inside_section_preserved():
    doc = Document()
    doc.add_heading("节", level=1)
    doc.add_paragraph("表前段落。")
    tbl = doc.add_table(rows=2, cols=2)
    tbl.cell(0, 0).text = "表格内容"
    doc.add_paragraph("表后段落。")
    sections = split_sections(doc)
    replace_section_paragraphs(doc, sections[0], ["新表前段。", "新表后段。"])
    out = _roundtrip(doc)
    texts = [p.text for p in out.paragraphs]
    assert "新表前段。" in texts and "新表后段。" in texts
    assert len(out.tables) == 1
    assert out.tables[0].cell(0, 0).text == "表格内容"


def test_last_section_boundary():
    doc = _build_doc()
    sections = split_sections(doc)
    replace_section_paragraphs(doc, sections[-1], ["第二节新正文。"])
    out = _roundtrip(doc)
    assert out.paragraphs[-1].text == "第二节新正文。"
    assert out.paragraphs[-2].text == "第二节"


def test_empty_section_inserts_after_heading():
    doc = Document()
    doc.add_heading("空节", level=1)
    doc.add_heading("末节", level=1)
    doc.add_paragraph("末节正文。")
    sections = split_sections(doc)
    assert sections[0]["para_end"] == sections[0]["para_start"]  # 邻接标题空节
    replace_section_paragraphs(doc, sections[0], ["补进空节的正文。"])
    out = _roundtrip(doc)
    texts = [p.text for p in out.paragraphs]
    assert texts.index("补进空节的正文。") < texts.index("末节")


def test_empty_paragraphs_list_clears_section():
    doc = _build_doc()
    sections = split_sections(doc)
    replace_section_paragraphs(doc, sections[0], [])
    out = _roundtrip(doc)
    texts = [p.text for p in out.paragraphs]
    assert texts == ["第一节", "第二节", "第二节正文。"]


def test_sectpr_in_template_paragraph_not_duplicated():
    """分节符（w:pPr/w:sectPr）不得被 pPr 模板复制 N 份——否则 Word 版面错乱。"""
    doc = Document()
    doc.add_heading("第一节", level=1)
    p = doc.add_paragraph("带分节符的正文段。")
    pPr = p._p.get_or_add_pPr()
    pPr.append(pPr.makeelement(qn("w:sectPr"), {}))
    doc.add_heading("第二节", level=1)
    doc.add_paragraph("第二节正文。")
    sections = split_sections(doc)
    replace_section_paragraphs(doc, sections[0], ["新段一。", "新段二。"])
    out = _roundtrip(doc)
    sect_count = sum(
        1 for pp in out.paragraphs
        if pp._p.pPr is not None and pp._p.pPr.find(qn("w:sectPr")) is not None
    )
    assert sect_count == 0
    texts = [pp.text for pp in out.paragraphs]
    assert "新段一。" in texts and "新段二。" in texts


# ---------- find_and_replace 精准替换（2026-09-25 事故：整章重写替换单行修改）----------

def _fr_import():
    from rag.svr.document_rewrite.docx_edit import find_and_replace
    return find_and_replace


def test_find_replace_basic_count_and_text():
    fr = _fr_import()
    doc = _build_doc()
    count, hits = fr(doc, "第二节正文", "LG11111")
    assert count == 1
    out = _roundtrip(doc)
    texts = [p.text for p in out.paragraphs]
    assert "LG11111。" in texts
    assert hits or count == 0


def test_find_replace_cross_run_match():
    """跨 run 场景：目标文本被拆在多个 run 里，逐字子串仍命中并整段重写。"""
    fr = _fr_import()
    doc = Document()
    doc.add_heading("节", level=1)
    p = doc.add_paragraph("前缀")
    p.add_run("符合《政府采购法》第二十")
    p.add_run("二条规定的条件。后缀")
    count, _ = fr(doc, "符合《政府采购法》第二十二条规定的条件。", "LG11111")
    assert count == 1
    out = _roundtrip(doc)
    assert out.paragraphs[1].text == "前缀LG11111后缀"


def test_find_replace_no_match_returns_zero_and_doc_untouched():
    """对抗：找不到时零改动（事务语义），调用方据此给引导文案。"""
    fr = _fr_import()
    doc = _build_doc()
    before = _roundtrip(doc)
    count, hits = fr(doc, "根本不存在的原文内容XYZ", "新文本")
    assert count == 0 and hits == []
    after_texts = [p.text for p in _roundtrip(doc).paragraphs]
    before_texts = [p.text for p in before.paragraphs]
    assert after_texts == before_texts


def test_find_replace_multiple_occurrences_all_replaced():
    fr = _fr_import()
    doc = Document()
    doc.add_heading("节", level=1)
    doc.add_paragraph("甲方：A公司。乙方见证甲方履约。")
    doc.add_paragraph("甲方义务另列。")
    count, hits = fr(doc, "甲方", "采购人")
    assert count == 3
    assert len(hits) == 2
    out = _roundtrip(doc)
    assert out.paragraphs[1].text == "采购人：A公司。乙方见证采购人履约。"
    assert out.paragraphs[2].text == "采购人义务另列。"


def test_find_replace_keeps_first_run_format_and_drops_hyperlink():
    """整段重写保留首 run 字符格式；超链接/域先移除防新旧文本拼接
    （同 api/utils/docx_edit._replace_para_text 口径）。"""
    from docx.oxml import OxmlElement

    fr = _fr_import()
    doc = Document()
    doc.add_heading("节", level=1)
    p = doc.add_paragraph("")
    r1 = p.add_run("链接前")
    r1.font.size = Pt(14)
    hl = OxmlElement("w:hyperlink")
    r_in = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "链接内文本"
    r_in.append(t)
    hl.append(r_in)
    p._p.append(hl)
    p.add_run("链接后")
    count, _ = fr(doc, "链接前链接内文本链接后", "替换完成")
    assert count == 1
    out = _roundtrip(doc)
    np = out.paragraphs[1]
    assert np.text == "替换完成"
    assert np.runs[0].font.size.pt == 14


def test_find_replace_empty_doc_no_crash():
    """对抗：空文档/无段落不崩溃，返回零。"""
    fr = _fr_import()
    doc = Document()
    count, hits = fr(doc, "任意", "替换")
    assert count == 0 and hits == []
