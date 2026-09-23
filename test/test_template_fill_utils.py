import io
import logging

import pytest
from docx import Document


def _make_docx(paras, table=None):
    """paras: list[str]; table: list[list[str]] (1 表追加在末尾)"""
    doc = Document()
    for t in paras:
        doc.add_paragraph(t)
    if table:
        tbl = doc.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, text in enumerate(row):
                tbl.rows[r].cells[c].paragraphs[0].text = text
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture
def sample_docx():
    return _make_docx(
        ["项目名称：____________", "无填写点的普通段落"],
        table=[["投标单位", ""], ["日期", "____年____月____日"]],
    )


def test_iter_docx_paragraphs(sample_docx):
    from rag.svr.template_fill.docx_utils import iter_docx_paragraphs
    items = iter_docx_paragraphs(sample_docx)
    texts = [it["text"] for it in items]
    assert texts[0] == "项目名称：____________"
    # 表格内段落扁平编号继续递增
    assert "日期" in texts
    cell_items = [it for it in items if it["addr"].startswith("cell:")]
    assert any(it["addr"] == "cell:0:1:1:0" for it in cell_items)


def test_extract_docx_candidates(sample_docx):
    from rag.svr.template_fill.docx_utils import extract_docx_candidates
    cands = extract_docx_candidates(sample_docx)
    # 普通段落（无填写特征）被过滤
    assert all("无填写点" not in c["text"] for c in cands)
    assert any("项目名称" in c["text"] for c in cands)
    assert any(c["addr"] == "cell:0:1:1:0" for c in cands)


def test_apply_docx_placeholders(sample_docx):
    from rag.svr.template_fill.docx_utils import apply_docx_placeholders
    out = apply_docx_placeholders(sample_docx, [
        {"addr": "para:0", "anchor": "____________", "key": "project_name"},
        {"addr": "cell:0:1:1:0", "anchor": "____年____月____日", "key": "sign_date"},
    ])
    from rag.svr.template_fill.docx_utils import iter_docx_paragraphs
    texts = [it["text"] for it in iter_docx_paragraphs(out)]
    assert "{{project_name}}" in texts[0]
    # cell:1:1:0 扁平序号 = 2 正文段 + 4 单元格段中排第 6（idx 5），按 addr 定位更稳健
    by_addr = {it["addr"]: it["text"] for it in iter_docx_paragraphs(out)}
    assert "{{sign_date}}" in by_addr["cell:0:1:1:0"]
    # 原 anchor 消失
    assert "____________" not in texts[0]


def test_apply_docx_anchor_not_found_is_noop(sample_docx):
    from rag.svr.template_fill.docx_utils import apply_docx_placeholders
    out = apply_docx_placeholders(sample_docx, [
        {"addr": "para:0", "anchor": "不存在的锚文本", "key": "x"},
    ])
    from rag.svr.template_fill.docx_utils import iter_docx_paragraphs
    assert "{{x}}" not in "".join(it["text"] for it in iter_docx_paragraphs(out))


def test_apply_docx_empty_replacements_returns_original(sample_docx):
    """空替换列表：文档内容等价（python-docx 重新写 zip 时时间戳精度为 2 秒，
    字节相等会在跨秒边界时随机失败，故按段落 (addr, text) 比对）。"""
    from rag.svr.template_fill.docx_utils import apply_docx_placeholders, iter_docx_paragraphs
    out = apply_docx_placeholders(sample_docx, [])
    src_items = [(it["addr"], it["text"]) for it in iter_docx_paragraphs(sample_docx)]
    out_items = [(it["addr"], it["text"]) for it in iter_docx_paragraphs(out)]
    assert out_items == src_items


# ---------- 对抗性边界用例 ----------

def test_docx_multi_table_same_position_addr_unique():
    """回归：addr 曾缺表序号（cell:<r>:<c>:<pi>），两个表格同 (r,c,p) 段落撞号——
    parse 按 index 选了表 A 的段落，validate/渲染按 addr 查到表 B 的段落，
    锚文本错位导致校验失败/替换落错表。addr 必须带表序号保证全局唯一。"""
    doc = Document()
    doc.add_paragraph("标题：")
    for i, tbl_texts in enumerate([["5.工程特征：", ""], ["2）道路工程：", ""]]):
        tbl = doc.add_table(rows=1, cols=2)
        for c, text in enumerate(tbl_texts):
            tbl.rows[0].cells[c].paragraphs[0].text = text
        if i == 0:
            doc.add_paragraph("中间分隔段")
    buf = io.BytesIO()
    doc.save(buf)
    data = buf.getvalue()

    from rag.svr.template_fill.docx_utils import iter_docx_paragraphs
    items = iter_docx_paragraphs(data)
    addrs = [it["addr"] for it in items if it["addr"].startswith("cell:")]
    # 两表同位置段落 addr 不同
    assert len(addrs) == len(set(addrs))
    by_text = {it["text"]: it["addr"] for it in items if it["text"]}
    assert by_text["5.工程特征："] == "cell:0:0:0:0"
    assert by_text["2）道路工程："] == "cell:1:0:0:0"

    # 替换必须落到表 0 的段落，表 1 同位置不受影响
    from rag.svr.template_fill.docx_utils import apply_docx_placeholders
    out = apply_docx_placeholders(data, [
        {"addr": "cell:0:0:0:0", "anchor": "5.工程特征：", "key": "engineering_characteristics"},
    ])
    out_texts = {it["addr"]: it["text"] for it in iter_docx_paragraphs(out)}
    assert out_texts["cell:0:0:0:0"] == "{{engineering_characteristics}}"
    assert out_texts["cell:1:0:0:0"] == "2）道路工程："


def test_apply_docx_empty_document_no_crash():
    """空文档（无正文段落/表格）不崩溃。"""
    from rag.svr.template_fill.docx_utils import apply_docx_placeholders, iter_docx_paragraphs
    empty = _make_docx([])
    assert iter_docx_paragraphs(empty) == []
    out = apply_docx_placeholders(empty, [{"addr": "para:0", "anchor": "x", "key": "k"}])
    assert isinstance(out, bytes) and len(out) > 0


def test_apply_docx_anchor_twice_in_same_paragraph(sample_docx):
    """同一段落内锚文本出现两次：str.replace 语义应全部替换。"""
    from rag.svr.template_fill.docx_utils import apply_docx_placeholders, iter_docx_paragraphs
    src = _make_docx(["联系人：____ 电话：____"])
    out = apply_docx_placeholders(src, [
        {"addr": "para:0", "anchor": "____", "key": "v"},
    ])
    texts = [it["text"] for it in iter_docx_paragraphs(out)]
    assert "{{v}}" in texts[0]
    assert "____" not in texts[0]
    # 两次出现都被替换
    assert texts[0].count("{{v}}") == 2


def test_apply_docx_unknown_addr_is_ignored(sample_docx):
    """addr 不存在（越界/格式错）时静默跳过，不影响其他替换。"""
    from rag.svr.template_fill.docx_utils import apply_docx_placeholders, iter_docx_paragraphs
    out = apply_docx_placeholders(sample_docx, [
        {"addr": "para:999", "anchor": "____________", "key": "gone"},
        {"addr": "bogus", "anchor": "____________", "key": "gone2"},
        {"addr": "para:0", "anchor": "____________", "key": "project_name"},
    ])
    texts = [it["text"] for it in iter_docx_paragraphs(out)]
    assert "{{project_name}}" in texts[0]
    assert "{{gone}}" not in "".join(texts)
    assert "{{gone2}}" not in "".join(texts)


def _make_docx_with_hyperlink():
    """构造段落 = run("地址：") + run("____ 详见") + 尾部超链接("官网")，全程手工插
    XML，不依赖网络。锚文本 "：____" 跨前两个 run（不落在任何单一 run 内），
    强制触发跨 run 路径。

    复现原始 bug 的关键：python-docx 1.1+ 的 p.text 含超链接内文本，但 p.runs 不含，
    旧的整段重写路径会把 "官网" 复制进首 run → "地址{{k}}____ 详见官网官网"。
    """
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document()
    p = doc.add_paragraph("地址：")
    p.add_run("____ 详见")
    hl = OxmlElement("w:hyperlink")
    hl.set(qn("r:id"), "rIdLink")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "官网"
    r.append(t)
    hl.append(r)
    p._p.append(hl)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_apply_docx_hyperlink_paragraph_no_duplicate(sample_docx):
    """含 w:hyperlink 的段落（锚文本跨界）必须走 run 拼接替换：替换后超链接文本
    只出现一次，占位符正确落位，且原超链接文本不被复制进正文 run。"""
    from rag.svr.template_fill.docx_utils import apply_docx_placeholders, iter_docx_paragraphs
    src = _make_docx_with_hyperlink()
    # 前置确认构造成功：p.text 含超链接文本（python-docx 1.1+ 行为）
    assert next(it["text"] for it in iter_docx_paragraphs(src)) == "地址：____ 详见官网"

    out = apply_docx_placeholders(src, [{"addr": "para:0", "anchor": "：____", "key": "url"}])
    texts = [it["text"] for it in iter_docx_paragraphs(out)]
    assert texts[0] == "地址{{url}} 详见官网"
    # 修复前此处为 2（超链接文本被复制进首 run），修复后必须为 1
    assert texts[0].count("官网") == 1
    assert "____" not in texts[0]


def test_apply_docx_hyperlink_anchor_absent_is_noop():
    """含超链接的段落上锚文本不存在（且跨界拼不上）时返回 no-op，超链接文本不受影响。"""
    from rag.svr.template_fill.docx_utils import apply_docx_placeholders, iter_docx_paragraphs
    src = _make_docx_with_hyperlink()
    out = apply_docx_placeholders(src, [{"addr": "para:0", "anchor": "不存在的锚", "key": "k"}])
    texts = [it["text"] for it in iter_docx_paragraphs(out)]
    assert texts[0] == "地址：____ 详见官网"
    assert "{{k}}" not in texts[0]


def test_apply_docx_dirty_entry_missing_fields_skipped(sample_docx):
    """LLM 脏输入：addr/anchor/key 任一缺失或为空时跳过该条，不抛 KeyError，
    其余合法条目正常替换。"""
    from rag.svr.template_fill.docx_utils import apply_docx_placeholders, iter_docx_paragraphs
    out = apply_docx_placeholders(sample_docx, [
        {"addr": "para:0", "key": "no_anchor"},                 # 缺 anchor
        {"addr": "para:0", "anchor": "____________"},            # 缺 key
        {"anchor": "____年", "key": "no_addr"},                  # 缺 addr
        {"addr": "para:0", "anchor": "", "key": "empty_anchor"},  # anchor 为空
        {"addr": "cell:0:1:1:0", "anchor": "____年____月____日", "key": "sign_date"},
    ])
    items = iter_docx_paragraphs(out)
    texts = [it["text"] for it in items]
    by_addr = {it["addr"]: it["text"] for it in items}
    assert "{{sign_date}}" in by_addr["cell:0:1:1:0"]
    assert "{{no_anchor}}" not in "".join(texts)
    assert "{{no_addr}}" not in "".join(texts)
    assert "{{empty_anchor}}" not in "".join(texts)


def test_docx_nested_table_addr_and_apply():
    """嵌套表格必须进编址范围（此前只遍历顶层表格，范本表内表大量丢填写点）：
    addr 在父单元格后追加 :t<j> 段；候选提取可命中；apply 按 addr 精确替换。"""
    from rag.svr.template_fill.docx_utils import (
        apply_docx_placeholders,
        extract_docx_candidates,
        iter_docx_paragraphs,
    )
    doc = Document()
    doc.add_paragraph("封面段落")
    tbl = doc.add_table(rows=1, cols=2)
    tbl.rows[0].cells[0].paragraphs[0].text = "外层单元格"
    cell = tbl.rows[0].cells[1]
    nested = cell.add_table(rows=1, cols=2)
    nested.rows[0].cells[0].paragraphs[0].text = "内层：＿＿＿＿＿"
    nested.rows[0].cells[1].paragraphs[0].text = "（填写完整名称）"
    buf = io.BytesIO()
    doc.save(buf)
    data = buf.getvalue()

    items = iter_docx_paragraphs(data)
    addrs = [it["addr"] for it in items]
    assert len(addrs) == len(set(addrs))  # 全局唯一
    by_addr = {it["addr"]: it["text"] for it in items}
    assert by_addr["cell:0:0:0:0"] == "外层单元格"
    assert by_addr["cell:0:0:1:t0:0:0:0"] == "内层：＿＿＿＿＿"
    assert by_addr["cell:0:0:1:t0:0:1:0"] == "（填写完整名称）"

    # 嵌套单元格的填写点进候选（全角下划线特征 + 括号提示特征）
    cands = extract_docx_candidates(data)
    cand_addrs = {c["addr"] for c in cands}
    assert "cell:0:0:1:t0:0:0:0" in cand_addrs
    assert "cell:0:0:1:t0:0:1:0" in cand_addrs

    # 按 addr 替换落到嵌套段落，外层不受影响
    out = apply_docx_placeholders(data, [
        {"addr": "cell:0:0:1:t0:0:0:0", "anchor": "＿＿＿＿＿", "key": "inner_field"},
    ])
    out_by_addr = {it["addr"]: it["text"] for it in iter_docx_paragraphs(out)}
    assert out_by_addr["cell:0:0:1:t0:0:0:0"] == "内层：{{inner_field}}"
    assert out_by_addr["cell:0:0:0:0"] == "外层单元格"


def test_docx_merged_cell_dedup():
    """横向合并单元格（gridSpan）：row.cells 会重复返回同一 tc，去重后同一文本
    只编址一次——否则重复候选导致 anchor 反查「匹配到多处」歧义拒绝。"""
    from rag.svr.template_fill.docx_utils import (
        apply_docx_placeholders,
        extract_docx_candidates,
        iter_docx_paragraphs,
    )
    doc = Document()
    tbl = doc.add_table(rows=2, cols=2)
    merged = tbl.rows[0].cells[0].merge(tbl.rows[0].cells[1])
    merged.paragraphs[0].text = "合并单元格：＿＿＿＿"
    tbl.rows[1].cells[0].paragraphs[0].text = "甲方：____"
    tbl.rows[1].cells[1].paragraphs[0].text = "乙方：____"
    buf = io.BytesIO()
    doc.save(buf)
    data = buf.getvalue()

    items = iter_docx_paragraphs(data)
    merged_texts = [it for it in items if "合并单元格" in it["text"]]
    assert len(merged_texts) == 1  # 旧逻辑此处为 2（c=0 与 c=1 各编址一次）
    addrs = [it["addr"] for it in items]
    assert len(addrs) == len(set(addrs))

    cands = extract_docx_candidates(data)
    assert sum(1 for c in cands if "合并单元格" in c["text"]) == 1

    # 合并单元格自身可正常替换
    merged_addr = merged_texts[0]["addr"]
    out = apply_docx_placeholders(data, [
        {"addr": merged_addr, "anchor": "＿＿＿＿", "key": "merged_field"},
    ])
    out_items = iter_docx_paragraphs(out)
    assert any("{{merged_field}}" in it["text"] for it in out_items)


def test_docx_plain_table_addr_unchanged_after_nested_support():
    """回归守护：无嵌套表格的普通文档编址必须与旧逻辑完全一致（存量模板 addr
    兼容性依赖这一点）：顶层表 tbl_no 只数顶层，扁平序号连续递增。"""
    from rag.svr.template_fill.docx_utils import iter_docx_paragraphs
    data = _make_docx(
        ["段落A", "段落B"],
        table=[["投标单位", ""], ["日期", "____年____月____日"]],
    )
    items = iter_docx_paragraphs(data)
    assert [it["index"] for it in items] == list(range(len(items)))  # 扁平序号连续
    by_addr = {it["addr"]: it["text"] for it in items}
    assert by_addr["para:0"] == "段落A"
    assert by_addr["para:1"] == "段落B"
    assert by_addr["cell:0:0:0:0"] == "投标单位"
    assert by_addr["cell:0:1:1:0"] == "____年____月____日"


# ---------- xlsx 工具 ----------

def _make_xlsx(sheets: dict):
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(title=name)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_iter_xlsx_cells():
    from rag.svr.template_fill.xlsx_utils import iter_xlsx_cells
    blob = _make_xlsx({"封面": [["项目名称", "________"], [None, "空行跳过"]], "签章页": [["签字", "（）"]]})
    items = iter_xlsx_cells(blob)
    addrs = [it["addr"] for it in items]
    assert "封面!A1" in addrs and "封面!B1" in addrs
    assert "封面!A2" not in addrs  # None 值跳过
    assert "签章页!B1" in addrs


def test_apply_xlsx_placeholders():
    from rag.svr.template_fill.xlsx_utils import apply_xlsx_placeholders, iter_xlsx_cells
    blob = _make_xlsx({"封面": [["项目名称", "________"]]})
    out = apply_xlsx_placeholders(blob, [
        {"sheet": "封面", "coord": "B1", "addr": "封面!B1", "anchor": "________", "key": "project_name"},
    ])
    texts = {it["addr"]: it["text"] for it in iter_xlsx_cells(out)}
    assert texts["封面!B1"] == "{{project_name}}"


def test_apply_xlsx_dirty_entry_skipped():
    """LLM 脏输入：sheet/coord/anchor/key 任一缺失或为空时跳过该条不抛异常，
    其余合法条目正常替换。"""
    from rag.svr.template_fill.xlsx_utils import apply_xlsx_placeholders, iter_xlsx_cells
    blob = _make_xlsx({"封面": [["项目名称", "________"], ["日期", "____年____月____日"]]})
    out = apply_xlsx_placeholders(blob, [
        {"sheet": "封面", "coord": "B1", "anchor": "________"},                       # 缺 key
        {"sheet": "封面", "coord": "B1", "key": "no_anchor"},                          # 缺 anchor
        {"sheet": "封面", "anchor": "________", "key": "no_coord"},                    # 缺 coord
        {"coord": "B1", "anchor": "________", "key": "no_sheet"},                      # 缺 sheet
        {"sheet": "封面", "coord": "B1", "anchor": "", "key": "empty_anchor"},         # anchor 为空
        {"sheet": "封面", "coord": "B1", "anchor": "________", "key": ""},             # key 为空
        {"sheet": "封面", "coord": "B1", "anchor": "________", "key": "project_name"},  # 合法
        {"sheet": "封面", "coord": "B2", "anchor": "____年", "key": "sign_date"},       # 合法
    ])
    texts = {it["addr"]: it["text"] for it in iter_xlsx_cells(out)}
    assert texts["封面!B1"] == "{{project_name}}"
    assert texts["封面!B2"] == "{{sign_date}}____月____日"  # anchor "____年" 整体被替换


def test_apply_xlsx_bad_sheet_or_coord_skipped():
    """不存在的 sheet 名 / 非法 coord（openpyxl 会抛异常）跳过该条，不中断整批。"""
    from rag.svr.template_fill.xlsx_utils import apply_xlsx_placeholders, iter_xlsx_cells
    blob = _make_xlsx({"封面": [["项目名称", "________"]]})
    out = apply_xlsx_placeholders(blob, [
        {"sheet": "不存在的页", "coord": "B1", "anchor": "________", "key": "gone1"},   # KeyError 路径
        {"sheet": "封面", "coord": "不是坐标", "anchor": "________", "key": "gone2"},   # ValueError 路径
        {"sheet": "封面", "coord": "B1", "anchor": "________", "key": "project_name"},
    ])
    texts = {it["addr"]: it["text"] for it in iter_xlsx_cells(out)}
    assert texts["封面!B1"] == "{{project_name}}"
    assert "{{gone1}}" not in "".join(texts.values())
    assert "{{gone2}}" not in "".join(texts.values())


def test_apply_xlsx_range_coord_skipped():
    """对抗：coord 传 range 语法（如 "A1:B2"）时 openpyxl 合法接受并返回 cell 元组，
    后续 .value 会抛 AttributeError 中断整批。修复后该条被静默跳过，其余条目正常替换。"""
    from rag.svr.template_fill.xlsx_utils import apply_xlsx_placeholders, iter_xlsx_cells
    blob = _make_xlsx({"封面": [["项目名称", "________"], ["日期", "____年____月____日"]]})
    out = apply_xlsx_placeholders(blob, [
        {"sheet": "封面", "coord": "A1:B2", "anchor": "________", "key": "gone_range"},
        {"sheet": "封面", "coord": "B1", "anchor": "________", "key": "project_name"},
        {"sheet": "封面", "coord": "B2", "anchor": "____年____月____日", "key": "sign_date"},
    ])
    texts = {it["addr"]: it["text"] for it in iter_xlsx_cells(out)}
    assert texts["封面!B1"] == "{{project_name}}"
    assert texts["封面!B2"] == "{{sign_date}}"
    assert "{{gone_range}}" not in "".join(texts.values())


def test_apply_xlsx_addr_fallback():
    """addr 兜底：条目缺 sheet/coord 但 addr 为 "<sheet>!<coord>" 时按最后一个 "!" 拆分定位。"""
    from rag.svr.template_fill.xlsx_utils import apply_xlsx_placeholders, iter_xlsx_cells
    blob = _make_xlsx({"封面": [["项目名称", "________"]]})
    out = apply_xlsx_placeholders(blob, [
        {"addr": "封面!B1", "anchor": "________", "key": "project_name"},
    ])
    texts = {it["addr"]: it["text"] for it in iter_xlsx_cells(out)}
    assert texts["封面!B1"] == "{{project_name}}"


def test_apply_xlsx_explicit_sheet_coord_over_addr():
    """显式 sheet/coord 优先于 addr 兜底：addr 写错不影响按显式字段定位。"""
    from rag.svr.template_fill.xlsx_utils import apply_xlsx_placeholders, iter_xlsx_cells
    blob = _make_xlsx({"封面": [["项目名称", "________"]]})
    out = apply_xlsx_placeholders(blob, [
        {"sheet": "封面", "coord": "B1", "addr": "不存在页!Z99", "anchor": "________", "key": "project_name"},
    ])
    texts = {it["addr"]: it["text"] for it in iter_xlsx_cells(out)}
    assert texts["封面!B1"] == "{{project_name}}"


def test_extract_xlsx_skips_formula_cells():
    """以 "=" 开头的公式单元格不进候选（防止 LLM 把公式文本当填写点、渲染时覆盖公式），
    普通文本格正常提取。"""
    from openpyxl import Workbook

    from rag.svr.template_fill.xlsx_utils import extract_xlsx_candidates
    wb = Workbook()
    ws = wb.active
    ws.title = "封面"
    ws["A1"] = "=SUM(A2:A3)"          # 公式文本，恰好含 "____" 也不该被提取
    ws["A2"] = 10
    ws["A3"] = 20
    ws["B1"] = "合计：____元"          # 普通文本填写点
    ws["B2"] = "=IF(A2>5, \"____\", \"\")"  # 公式结果含填写特征同样排除
    buf = io.BytesIO()
    wb.save(buf)
    cands = extract_xlsx_candidates(buf.getvalue())
    texts = [c["text"] for c in cands]
    assert all(not t.startswith("=") for t in texts)
    assert any("合计：____元" in t for t in texts)


def test_apply_xlsx_preserves_cell_style():
    """对抗：替换只改 value，单元格样式（字体加粗/填充色）必须保留。"""
    from openpyxl import load_workbook
    from openpyxl.styles import PatternFill

    from rag.svr.template_fill.xlsx_utils import apply_xlsx_placeholders, iter_xlsx_cells
    blob = _make_xlsx({"封面": [["项目名称", "________"]]})
    # 给 B1 设样式后重新序列化
    wb = load_workbook(io.BytesIO(blob))
    ws = wb["封面"]
    ws["B1"].font = ws["B1"].font.copy(bold=True)
    ws["B1"].fill = PatternFill(fill_type="solid", start_color="FFFF00")
    buf = io.BytesIO()
    wb.save(buf)
    styled = buf.getvalue()

    out = apply_xlsx_placeholders(styled, [
        {"sheet": "封面", "coord": "B1", "anchor": "________", "key": "project_name"},
    ])
    texts = {it["addr"]: it["text"] for it in iter_xlsx_cells(out)}
    assert texts["封面!B1"] == "{{project_name}}"
    wb2 = load_workbook(io.BytesIO(out))
    cell = wb2["封面"]["B1"]
    assert cell.font.bold is True
    assert cell.fill.start_color.rgb == "00FFFF00" or cell.fill.start_color.rgb == "FFFF00"


# ---------- detector：LLM 填写点识别解析/校验（纯函数部分） ----------

CANDS = [
    {"index": 0, "text": "项目名称：____________", "addr": "para:0"},
    {"index": 2, "text": "投标单位（　　　）", "addr": "para:2"},
    {"index": 4, "text": "签字日期：____年____月____日", "addr": "para:4"},
]


def test_parse_detection_response_valid():
    from rag.svr.template_fill.detector import parse_detection_response
    raw = '```json\n[{"line": 0, "anchor": "____________", "key": "Project Name", "name": "项目名称", "description": "投标项目全称", "retrieval_query": "项目名称 概况", "fill_mode": "llm", "required": true}, {"line": 4, "anchor": "____年____月____日", "key": "sign_date", "name": "签字日期", "description": "", "retrieval_query": "", "fill_mode": "manual", "required": false}]\n```'
    out = parse_detection_response(raw, CANDS)
    assert len(out) == 2
    assert out[0]["key"] == "project_name"  # 归一化：小写+下划线
    assert out[0]["addr"] == "para:0"
    assert out[0]["top_k"] == 6
    # fill_mode 代码层强制 llm：LLM 输出 manual 也被规整（识别产物统一 AI 填写）
    assert all(it["fill_mode"] == "llm" for it in out)


def test_parse_detection_response_drops_invalid_anchor():
    from rag.svr.template_fill.detector import parse_detection_response
    raw = '[{"line": 0, "anchor": "不存在的锚", "key": "x", "name": "X", "fill_mode": "llm", "required": true}]'
    assert parse_detection_response(raw, CANDS) == []


def test_parse_detection_response_dedupes_keys():
    from rag.svr.template_fill.detector import parse_detection_response
    raw = '[{"line": 0, "anchor": "____________", "key": "date", "name": "A", "fill_mode": "llm", "required": true},{"line": 4, "anchor": "____年", "key": "date", "name": "B", "fill_mode": "manual", "required": false}]'
    out = parse_detection_response(raw, CANDS)
    assert [it["key"] for it in out] == ["date", "date_2"]


def test_parse_detection_response_garbage_returns_empty():
    from rag.svr.template_fill.detector import parse_detection_response
    assert parse_detection_response("我不明白你的意思", CANDS) == []


def test_validate_placeholders():
    from rag.svr.template_fill.detector import validate_placeholders
    ok, err = validate_placeholders([
        {"key": "project_name", "name": "项目名称", "addr": "para:0", "anchor": "____________",
         "fill_mode": "llm", "required": True, "retrieval_query": "q", "description": "", "top_k": 6},
    ], CANDS)
    assert ok and err == ""
    ok, err = validate_placeholders([
        {"key": "Bad Key!", "name": "x", "addr": "para:0", "anchor": "____________", "fill_mode": "llm",
         "required": True, "retrieval_query": "", "description": "", "top_k": 6},
    ], CANDS)
    assert not ok and "key" in err
    # anchor 不在 addr 对应文本中 → 拒绝
    ok, err = validate_placeholders([
        {"key": "a", "name": "x", "addr": "para:0", "anchor": "瞎写的", "fill_mode": "llm",
         "required": True, "retrieval_query": "", "description": "", "top_k": 6},
    ], CANDS)
    assert not ok and "anchor" in err


# ---------- detector 对抗性用例 ----------

def test_parse_detection_response_non_int_line_skipped():
    """对抗：line 为字符串 "0" 或 None 的项必须被安全跳过（不抛异常），其余项正常保留。"""
    from rag.svr.template_fill.detector import parse_detection_response
    raw = '[{"line": "0", "anchor": "____________", "key": "a", "name": "A"},' \
          '{"line": null, "anchor": "____年", "key": "b", "name": "B"},' \
          '{"line": 4, "anchor": "____年____月____日", "key": "c", "name": "C"}]'
    out = parse_detection_response(raw, CANDS)
    assert [it["key"] for it in out] == ["c"]


def test_parse_detection_response_noisy_surroundings():
    """对抗：JSON 数组前后有大量噪音文本（含代码块/说明）仍能解析。"""
    from rag.svr.template_fill.detector import parse_detection_response
    raw = "好的，我分析了这份模板。以下是识别结果：\n```json\n" \
          '[{"line": 0, "anchor": "____________", "key": "project_name", "name": "项目名称"}]' \
          "\n```\n以上共 1 个填写点，请确认。"
    out = parse_detection_response(raw, CANDS)
    assert len(out) == 1 and out[0]["key"] == "project_name"


def test_parse_detection_response_oversize_anchor_dropped():
    """对抗：超长 anchor（>500 字符）在 parse 阶段即被丢弃——约束选在 parse：
    识别阶段就该拦住异常项，避免脏数据流入人工确认/apply 链路。"""
    from rag.svr.template_fill.detector import parse_detection_response
    long_anchor = "长" * 501
    raw = f'[{{"line": 0, "anchor": "{long_anchor}", "key": "big", "name": "B"}},' \
          '{"line": 0, "anchor": "____________", "key": "ok", "name": "O"}]'
    out = parse_detection_response(raw, CANDS)
    assert [it["key"] for it in out] == ["ok"]


def test_parse_detection_response_overlong_key_truncated():
    """对抗：LLM 照中文长字段名直译出 >64 字符 key（生产实测 71 字符导致整次识别
    判死），parse 阶段必须截断到 ≤64 且仍为合法 snake_case，其余建议不受牵连。"""
    from rag.svr.template_fill.detector import parse_detection_response
    long_key = "liability_for_refusing_to_replace_key_construction_management_personnel"
    assert len(long_key) == 71
    raw = f'[{{"line": 0, "anchor": "____________", "key": "{long_key}", "name": "拒换_KEY责任"}},' \
          '{"line": 4, "anchor": "____年____月____日", "key": "sign_date", "name": "签字日期"}]'
    out = parse_detection_response(raw, CANDS)
    assert len(out) == 2
    k = out[0]["key"]
    assert len(k) <= 64 and k.startswith("liability_for_refusing")
    # 截断后仍过 validate_placeholders（生产故障链路的终审）
    from rag.svr.template_fill.detector import validate_placeholders
    ok, err = validate_placeholders(out, CANDS)
    assert ok, err


def test_parse_detection_response_truncated_key_dedup_keeps_length():
    """对抗：两条超长 key 截断后前 64 字符相同 → 去重后缀拼接时必须同步收缩
    基串，保证含后缀总长仍 ≤64（旧逻辑 key+"_2" 会溢出到 66）。"""
    from rag.svr.template_fill.detector import parse_detection_response
    long_key = "liability_for_refusing_to_replace_key_construction_management_personnel"
    raw = f'[{{"line": 0, "anchor": "____________", "key": "{long_key}", "name": "A"}},' \
          f'{{"line": 0, "anchor": "____________", "key": "{long_key}", "name": "B"}}]'
    out = parse_detection_response(raw, CANDS)
    assert len(out) == 2
    k1, k2 = out[0]["key"], out[1]["key"]
    assert k1 != k2
    for k in (k1, k2):
        assert len(k) <= 64
        import re as _re
        assert _re.fullmatch(r"[a-z][a-z0-9_]{0,63}", k)


def test_parse_detection_response_64char_key_dedup_no_overflow():
    """对抗：恰好 64 字符的 key 撞车去重，加后缀后不得超限。"""
    from rag.svr.template_fill.detector import parse_detection_response
    k64 = "a" + "b" * 63
    assert len(k64) == 64
    raw = f'[{{"line": 0, "anchor": "____________", "key": "{k64}", "name": "A"}},' \
          f'{{"line": 0, "anchor": "____________", "key": "{k64}", "name": "B"}}]'
    out = parse_detection_response(raw, CANDS)
    assert len(out) == 2
    for it in out:
        assert len(it["key"]) <= 64
    assert out[1]["key"] == "a" + "b" * 61 + "_2"


def test_validate_placeholders_rejects_oversize_key():
    """对抗：超过 64 字符的合法 snake_case key 在 validate 阶段被拒
    （parse 的 normalize 不截断长度，长度约束统一收口在 validate）。"""
    from rag.svr.template_fill.detector import validate_placeholders
    ok, err = validate_placeholders([
        {"key": "k" * 65, "name": "x", "addr": "para:0", "anchor": "____________", "fill_mode": "llm",
         "required": True, "retrieval_query": "", "description": "", "top_k": 6},
    ], CANDS)
    assert not ok and "key" in err


def test_validate_placeholders_anchor_length_boundary():
    """对抗：anchor 长度约束在 validate 与 parse（MAX_ANCHOR_LEN=500）对齐——
    恰好 500 通过（含 membership 校验需真实落在候选文本内），501 拒绝且报错可读。"""
    from rag.svr.template_fill.detector import MAX_ANCHOR_LEN, validate_placeholders
    anchor_500 = "长" * MAX_ANCHOR_LEN
    cands = [{"index": 0, "addr": "para:0", "text": "前缀" + anchor_500}]
    base = {"key": "k", "name": "n", "addr": "para:0", "fill_mode": "llm"}
    ok, err = validate_placeholders([{**base, "anchor": anchor_500}], cands)
    assert ok and err == ""
    ok, err = validate_placeholders([{**base, "anchor": "长" * (MAX_ANCHOR_LEN + 1)}], cands)
    assert not ok and "anchor" in err and str(MAX_ANCHOR_LEN) in err


def test_validate_placeholders_non_dict_item_no_crash():
    """对抗：items 元素非 dict（字符串/None）不抛异常，返回校验失败。"""
    from rag.svr.template_fill.detector import validate_placeholders
    ok, err = validate_placeholders(["不是字典"], CANDS)
    assert not ok and err
    ok, err = validate_placeholders([None], CANDS)
    assert not ok and err


# ---------- validate_placeholders：空 addr 按 anchor 反查推导（手动添加行） ----------

def _manual_item(**kw):
    base = {"key": "m", "name": "手动", "addr": "", "anchor": "项目名称",
            "fill_mode": "manual", "required": True}
    base.update(kw)
    return base


def test_validate_empty_addr_unique_anchor_hit_backfills():
    """空 addr + anchor 恰好命中 1 个候选：通过，且 addr 被回填进 item。"""
    from rag.svr.template_fill.detector import validate_placeholders
    item = _manual_item(anchor="项目名称")
    ok, err = validate_placeholders([item], CANDS)
    assert ok and err == ""
    assert item["addr"] == "para:0"  # 回填生效，落库前地址已补齐


def test_validate_empty_addr_anchor_not_found_rejected():
    """空 addr + anchor 未命中任何候选：拒绝且报错面向用户。"""
    from rag.svr.template_fill.detector import validate_placeholders
    ok, err = validate_placeholders([_manual_item(anchor="不存在的锚")], CANDS)
    assert not ok and "未找到" in err


def test_validate_empty_addr_anchor_ambiguous_rejected():
    """空 addr + anchor 多处命中（"____" 在 para:0 与 para:4 都出现）：歧义拒绝，防回填错位。"""
    from rag.svr.template_fill.detector import validate_placeholders
    ok, err = validate_placeholders([_manual_item(anchor="____")], CANDS)
    assert not ok and "匹配到" in err and "2" in err


def test_validate_empty_addr_empty_anchor_rejected():
    """空 addr + 空 anchor：`"" in text` 恒真会命中全部候选，须先行拦截为缺少锚文本。"""
    from rag.svr.template_fill.detector import validate_placeholders
    ok, err = validate_placeholders([_manual_item(anchor="")], CANDS)
    assert not ok and "锚文本" in err


def test_validate_empty_addr_case_sensitive_match():
    """anchor 反查为精确子串匹配、区分大小写，与 addr 非空分支的 membership 校验一致。"""
    from rag.svr.template_fill.detector import validate_placeholders
    cands = [{"index": 0, "addr": "para:0", "text": "Project Name: ______"}]
    item = _manual_item(anchor="Project Name")
    ok, err = validate_placeholders([item], cands)
    assert ok and item["addr"] == "para:0"
    ok, err = validate_placeholders([_manual_item(anchor="project name")], cands)
    assert not ok and "未找到" in err


def test_validate_non_empty_addr_keeps_original_logic():
    """addr 非空走原逻辑：定位不存在 / anchor 不在该定位文本中均拒绝，不做反查回填。"""
    from rag.svr.template_fill.detector import validate_placeholders
    item = {"key": "a", "name": "x", "addr": "para:9", "anchor": "项目名称",
            "fill_mode": "llm", "required": True}
    ok, err = validate_placeholders([item], CANDS)
    assert not ok and "不存在" in err
    assert item["addr"] == "para:9"  # 未被反查改写
    item2 = {"key": "b", "name": "x", "addr": "para:0", "anchor": "____年",  # "____年" 只在 para:4
             "fill_mode": "llm", "required": True}
    ok, err = validate_placeholders([item2], CANDS)
    assert not ok and "不在" in err


def test_normalize_key():
    """归一化：大写/空格/特殊字符 → snake_case；全非法字符兜底为 field。"""
    from rag.svr.template_fill.detector import normalize_key
    assert normalize_key("Project Name") == "project_name"
    assert normalize_key("  Sign-Date! ") == "sign_date"
    assert normalize_key("___") == "field"
    assert normalize_key("") == "field"


# ---------- service：_sanitize_filename 纯函数（不依赖 DB/MinIO） ----------
# 说明：api.db.db_models 的 DB 对象为懒连接，import 链不会立即触库，
# 故直接 `from api.db.services.template_fill_service import _sanitize_filename`
# 模块级直测即可，无需 stub-load。

def _sanitize(name, max_len=128):
    from api.db.services.template_fill_service import _sanitize_filename
    return _sanitize_filename(name, max_len)


def test_sanitize_forward_slash_takes_basename():
    assert _sanitize("a/b/c.docx") == "c.docx"


def test_sanitize_backslash_takes_basename():
    """Windows 路径分隔符：统一按 '/' 处理后取最后一段，防止 '/' 进 MinIO 当目录前缀。"""
    assert _sanitize("a\\b\\c.docx") == "c.docx"


def test_sanitize_empty_returns_fallback():
    """空串/纯空白/None 兜底 'template'，避免空 object 名片段。"""
    assert _sanitize("") == "template"
    assert _sanitize("   ") == "template"
    assert _sanitize(None) == "template"


def test_sanitize_trailing_dot_kept_as_is():
    """'file.' 清洗后 stem 非空但 ext 为空 → 不走扩展名保留分支，原样保留。"""
    assert _sanitize("file.") == "file."


def test_sanitize_hidden_file_dotfile_kept():
    """.hidden：rpartition 后 stem 为空 → 不截断分支命中，dotfile 原样保留。"""
    assert _sanitize(".hidden") == ".hidden"


def test_sanitize_oversize_chinese_name_truncated():
    """恰好超长（129 个汉字）→ 截断到 128，无扩展名按硬截断处理。"""
    name = "标" * 129
    out = _sanitize(name)
    assert out == "标" * 128
    assert len(out) == 128


def test_sanitize_path_traversal_reduced_to_basename():
    """对抗：'../../etc/passwd' 型穿越只留 base 名，无法借 object name 逃出 bucket 前缀。"""
    assert _sanitize("../../etc/passwd") == "passwd"
    assert _sanitize("..\\..\\windows\\system32\\config") == "config"
    # 整串都是点/斜杠组合时清洗后为空 → 兜底
    assert _sanitize("../../") == "template"


def test_sanitize_ext_exactly_10_chars_preserved_on_truncate():
    """扩展名恰好 10 字符（边界内）：超长截断时保留 '.xxxxxxxxxx' 扩展名。"""
    ext = "x" * 10
    name = "a" * 200 + "." + ext
    out = _sanitize(name)
    assert out == "a" * (128 - 10 - 1) + "." + ext
    assert len(out) == 128
    assert out.endswith("." + ext)


def test_sanitize_ext_11_chars_not_preserved_hard_truncate():
    """扩展名恰好 11 字符（越界）：不保留扩展名，硬截断到 128。"""
    ext = "x" * 11
    name = "a" * 200 + "." + ext
    out = _sanitize(name)
    assert out == "a" * 128
    assert "." not in out


def test_sanitize_custom_max_len_boundary():
    """自定义 max_len 边界：恰好等于 max_len 不截断；超 1 字符才截。"""
    assert _sanitize("a" * 20, max_len=20) == "a" * 20
    assert len(_sanitize("a" * 21, max_len=20)) == 20


# ---------- P2 renderer ----------

def _mk_docx_with_placeholder():
    """用 python-docx 造一个含 {{name}} 的 docx blob。"""
    from docx import Document
    doc = Document()
    doc.add_paragraph("项目名称：{{name}}")
    buf = io.BytesIO(); doc.save(buf)
    return buf.getvalue()


def test_render_docx_replaces_placeholder():
    from docx import Document as Docx

    from rag.svr.template_fill.renderer import render_docx
    out = render_docx(_mk_docx_with_placeholder(), {"name": "测试项目"})
    text = "\n".join(p.text for p in Docx(io.BytesIO(out)).paragraphs)
    assert "测试项目" in text and "{{" not in text


def test_render_docx_missing_value_renders_empty():
    """缺值占位符渲染为空串（人工二次加工留空），不再落【待人工】标记。"""
    from docx import Document as Docx

    from rag.svr.template_fill.renderer import render_docx
    out = render_docx(_mk_docx_with_placeholder(), {"name": ""})
    text = "\n".join(p.text for p in Docx(io.BytesIO(out)).paragraphs)
    assert "【待人工" not in text and "{{" not in text


def test_render_xlsx_by_addr():
    from openpyxl import Workbook, load_workbook

    from rag.svr.template_fill.renderer import render_xlsx
    wb = Workbook(); ws = wb.active; ws.title = "封面"; ws["B1"] = "{{name}}"
    buf = io.BytesIO(); wb.save(buf)
    out = render_xlsx(buf.getvalue(), {"name": "测试项目"}, {"name": "封面!B1"})
    ws2 = load_workbook(io.BytesIO(out))["封面"]
    assert ws2["B1"].value == "测试项目"


def test_render_xlsx_bad_addr_skipped():
    """addr 非法 → 跳过该格不抛异常，且好格不被误伤。"""
    from openpyxl import Workbook, load_workbook

    from rag.svr.template_fill.renderer import render_xlsx
    wb = Workbook(); ws = wb.active; ws.title = "S1"; ws["B2"] = "keep"
    buf = io.BytesIO(); wb.save(buf)
    dirty = {"a": "不存在的表!ZZ99",       # sheet 不存在 → KeyError
             "b": "S1!not-a-coord",        # coord 非法 → ValueError/IndexError
             "c": "S1!",                   # coord 空串 → IndexError
             "d": "S1!A1:B2",              # range 语法 → AttributeError
             "e": "S1!A1048577",           # 行越界 → ValueError
             "ok": "S1!B2"}                # 好格正常写
    out = render_xlsx(buf.getvalue(), dict.fromkeys(dirty, "x"), dirty)
    ws2 = load_workbook(io.BytesIO(out))["S1"]
    assert ws2["B2"].value == "x"          # 好格被写入
    assert ws2.max_row <= 2 and ws2.max_column <= 2  # 脏格全部未落值


def test_render_dispatch():
    from rag.svr.template_fill.renderer import render
    blob = _mk_docx_with_placeholder()
    out = render("docx", blob, {"name": "X"})
    assert out and out != blob


# ---------- detector：anchor 派生默认值（纯函数） ----------

def test_derive_default_from_anchor_blank_markers():
    from rag.svr.template_fill.detector import derive_default_from_anchor
    assert derive_default_from_anchor("") == ""
    assert derive_default_from_anchor(None) == ""
    assert derive_default_from_anchor("______") == ""
    assert derive_default_from_anchor("＿＿＿＿") == ""   # 全角下划线
    assert derive_default_from_anchor("　　") == ""       # 全角空格
    assert derive_default_from_anchor("---") == ""
    assert derive_default_from_anchor("………") == ""
    assert derive_default_from_anchor("N/A") == ""
    assert derive_default_from_anchor("▁▁▁▁") == ""   # 下八分之一块
    assert derive_default_from_anchor("＊＊＊") == ""  # 全角星号
    assert derive_default_from_anchor("⋯⋯") == ""     # 中点省略号
    assert derive_default_from_anchor("none") == ""   # 小写 none 同样视为留空标记


def test_derive_default_from_anchor_filled_values():
    from rag.svr.template_fill.detector import derive_default_from_anchor
    assert derive_default_from_anchor("XX建设工程有限公司") == "XX建设工程有限公司"
    assert derive_default_from_anchor("2026-09-10") == "2026-09-10"
    assert derive_default_from_anchor("  100万元  ") == "100万元"
    assert derive_default_from_anchor("____2026") == "____2026"  # 混合内容：含数字不误判为留空


def test_derive_default_from_anchor_truncates_and_strips_ctrl():
    """超长截断对齐 MAX_ANCHOR_LEN；控制字符剥离（防污染 docx XML）。"""
    from rag.svr.template_fill.detector import MAX_ANCHOR_LEN, derive_default_from_anchor
    assert len(derive_default_from_anchor("A" * 501)) == MAX_ANCHOR_LEN
    assert derive_default_from_anchor("值\x00\x01名") == "值名"


# ---------- service：save_placeholders 默认值合并（纯函数） ----------

def test_merge_defaults_explicit_wins():
    from api.db.services.template_fill_service import TplTemplateVersionService as S
    items = [{"key": "a", "anchor": "旧值", "default_value": "人工改的", "default_source": "manual"},
             {"key": "b", "anchor": "甲公司"}]
    out = S._merge_defaults(items, [{"key": "a", "default_value": "上版默认", "default_source": "auto"}])
    assert out[0]["default_value"] == "人工改的" and out[0]["default_source"] == "manual"
    assert out[1]["default_value"] == "甲公司" and out[1]["default_source"] == "detected"


def test_merge_defaults_inherit_by_key():
    from api.db.services.template_fill_service import TplTemplateVersionService as S
    items = [{"key": "a", "anchor": "＿＿＿"}]
    out = S._merge_defaults(items, [{"key": "a", "default_value": "历史沉淀", "default_source": "auto"}])
    assert out[0]["default_value"] == "历史沉淀" and out[0]["default_source"] == "auto"


def test_merge_defaults_echo_preserves_source():
    from api.db.services.template_fill_service import TplTemplateVersionService as S
    items = [{"key": "a", "anchor": "现值", "default_value": "", "default_source": ""}]
    out = S._merge_defaults(items, [])
    assert out[0]["default_value"] == "" and out[0]["default_source"] == ""


def test_merge_defaults_new_key_from_blank_anchor():
    from api.db.services.template_fill_service import TplTemplateVersionService as S
    out = S._merge_defaults([{"key": "k", "anchor": "______"}], [])
    assert out[0]["default_value"] == "" and out[0]["default_source"] == ""


def test_merge_defaults_cleared_sticky_across_saves():
    from api.db.services.template_fill_service import TplTemplateVersionService as S
    # 用户清空后（manual+空值），下次保存前端不回显 default_value 字段 → 不得复活
    out = S._merge_defaults([{"key": "a", "anchor": "现值"}],
                            [{"key": "a", "default_value": "", "default_source": "manual"}])
    assert out[0]["default_value"] == "" and out[0]["default_source"] == "manual"


def test_merge_defaults_explicit_truncated_and_stripped():
    from api.db.services.template_fill_service import TplTemplateVersionService as S
    out = S._merge_defaults([{"key": "a", "anchor": "", "default_value": "  " + "长" * 600, "default_source": ""}], [])
    assert out[0]["default_value"] == "长" * 500


def test_merge_defaults_explicit_null_treated_as_clear():
    from api.db.services.template_fill_service import TplTemplateVersionService as S
    # 前端序列化 null → 显式清空，不走派生
    out = S._merge_defaults([{"key": "a", "anchor": "现值", "default_value": None}], [])
    assert out[0]["default_value"] == "" and out[0]["default_source"] == ""


def test_merge_defaults_prev_dirty_items_ignored():
    from api.db.services.template_fill_service import TplTemplateVersionService as S
    out = S._merge_defaults([{"key": "a", "anchor": "甲"}],
                            ["dirty", {"no_key": 1}, {"key": "a", "default_value": "好值", "default_source": "auto"}])
    assert out[0]["default_value"] == "好值"


def test_merge_defaults_v2_explicit_empty_hard_gate():
    """V2 硬闸（设计 §4.5）：V2 条目显式带 default_value=""（键存在）→ 无论 anchor
    形态都不派生默认值——即使 anchor 是本会触发混合形态泄漏的骨架文本。"""
    from api.db.services.template_fill_service import TplTemplateVersionService as S
    out = S._merge_defaults([{"key": "proj", "anchor": "＿＿＿（项目名称）＿＿", "default_value": ""}], [])
    assert out[0]["default_value"] == ""
    assert out[0]["default_source"] == ""
    # 对照：V1 条目（无 default_value 键）同 anchor 现在也不派生（层1 混合形态骨架分支）
    out2 = S._merge_defaults([{"key": "proj", "anchor": "＿＿＿（项目名称）＿＿"}], [])
    assert out2[0]["default_value"] == ""


# ---------- service：产值沉淀默认值（纯函数） ----------

def test_sediment_into_placeholders():
    from api.db.services.template_fill_service import TplTemplateVersionService as S
    ph = [{"key": "a", "default_value": "旧", "default_source": "auto"},
          {"key": "b", "default_value": "人工", "default_source": "manual"},
          {"key": "c", "default_value": "", "default_source": ""},
          # 有历史默认值的占位符，本次产值留空 → 不得被空值抹掉
          {"key": "d", "default_value": "历史值", "default_source": "auto"},
          # 无 key 占位符 → 直接跳过，不报错也不被修改
          {"name": "无key"},
          {"key": "", "default_value": "旧空key", "default_source": "auto"}]
    values = {"a": "新A", "b": "新B", "c": "新C", "empty": "", "d": "", "e": None}
    changed = S._sediment_into_placeholders(ph, values, {"a", "b", "c", "d", "e"})
    assert changed is True
    assert ph[0]["default_value"] == "新A" and ph[0]["default_source"] == "auto"
    assert ph[1]["default_value"] == "人工"    # manual 不被覆盖
    assert ph[2]["default_value"] == "新C"
    # 空值（""/None）不沉淀：历史默认值原样保留
    assert ph[3]["default_value"] == "历史值" and ph[3]["default_source"] == "auto"
    # 无 key / key 为空串的占位符被跳过：内容不变、不报错
    assert ph[4] == {"name": "无key"}
    assert ph[5]["default_value"] == "旧空key" and ph[5]["default_source"] == "auto"


def test_sediment_override_manual():
    from api.db.services.template_fill_service import TplTemplateVersionService as S
    ph = [{"key": "b", "default_value": "人工", "default_source": "manual"}]
    S._sediment_into_placeholders(ph, {"b": "用户直填"}, {"b"}, override_keys={"b"})
    assert ph[0]["default_value"] == "用户直填" and ph[0]["default_source"] == "auto"


def test_sediment_no_change_returns_false():
    from api.db.services.template_fill_service import TplTemplateVersionService as S
    ph = [{"key": "a", "default_value": "同值", "default_source": "auto"}]
    assert S._sediment_into_placeholders(ph, {"a": "同值"}, {"a"}) is False
    assert S._sediment_into_placeholders(ph, {}, {"a"}) is False


def test_sediment_long_value_truncated_to_max_anchor_len():
    from api.db.services.template_fill_service import TplTemplateVersionService as S
    from rag.svr.template_fill.detector import MAX_ANCHOR_LEN
    ph = [{"key": "a", "default_value": "", "default_source": ""}]
    # 超长 LLM 输出（600 字）：截到 500，且算变更
    long_val = "长" * (MAX_ANCHOR_LEN + 100)
    assert S._sediment_into_placeholders(ph, {"a": long_val}, {"a"}) is True
    assert len(ph[0]["default_value"]) == MAX_ANCHOR_LEN
    assert ph[0]["default_source"] == "auto"


# ---------- service：沉淀 only_keys 白名单（按钮触发，对抗性） ----------


def test_sediment_only_keys_whitelist_excludes_others():
    """白名单是硬闸：不在 only_keys 的字段即使产值非空也一个字节都不许动。
    这是本需求的核心不变量——上一版正是「无差别全量写」把范本基线污染了。"""
    from api.db.services.template_fill_service import TplTemplateVersionService as S
    ph = [{"key": "in", "default_value": "旧", "default_source": "auto"},
          {"key": "out", "default_value": "旧", "default_source": "auto"},
          {"key": "out_manual", "default_value": "人工", "default_source": "manual"}]
    snapshot = [dict(it) for it in ph]
    assert S._sediment_into_placeholders(
        ph, {"in": "新", "out": "不该写", "out_manual": "也不该写"}, {"in"}) is True
    assert ph[0]["default_value"] == "新"
    assert ph[1] == snapshot[1], "白名单外字段被写入——硬闸失效"
    assert ph[2] == snapshot[2], "白名单外 manual 字段被写入——硬闸失效"


def test_sediment_only_keys_empty_set_writes_nothing():
    """only_keys=set()（noop 轮）→ 全量非空产值也不写，且返回 False。
    契约上「空集合」必须是「一个都不写」，绝不能退化成「不限制」。"""
    from api.db.services.template_fill_service import TplTemplateVersionService as S
    ph = [{"key": "a", "default_value": "旧", "default_source": "auto"},
          {"key": "b", "default_value": "", "default_source": ""}]
    snapshot = [dict(it) for it in ph]
    assert S._sediment_into_placeholders(ph, {"a": "新A", "b": "新B"}, set()) is False
    assert ph == snapshot


def test_sediment_only_keys_unknown_key_ignored():
    """only_keys 含 placeholders 里不存在的 key（画布旧 key / 脏数据）→
    不报错、不新建项、其余正常写。"""
    from api.db.services.template_fill_service import TplTemplateVersionService as S
    ph = [{"key": "a", "default_value": "", "default_source": ""}]
    assert S._sediment_into_placeholders(
        ph, {"a": "新A", "ghost": "幽灵值"}, {"a", "ghost"}) is True
    assert len(ph) == 1 and ph[0]["default_value"] == "新A"


def test_sediment_override_requires_both_sets():
    """覆盖 manual 是 only_keys 与 override_keys 的**与**关系：
    只在白名单内但不在 override 里（LLM 产的值）不许覆盖 manual；
    只在 override 里但不在白名单内（白名单外）也不许碰。"""
    from api.db.services.template_fill_service import TplTemplateVersionService as S
    ph = [{"key": "llm_key", "default_value": "人工", "default_source": "manual"},
          {"key": "direct_key", "default_value": "人工", "default_source": "manual"},
          {"key": "outside", "default_value": "人工", "default_source": "manual"}]
    S._sediment_into_placeholders(
        ph, {"llm_key": "LLM值", "direct_key": "直填值", "outside": "白名单外"},
        {"llm_key", "direct_key"}, override_keys={"direct_key", "outside"})
    assert ph[0]["default_value"] == "人工", "非直填键不得覆盖 manual"
    assert ph[1]["default_value"] == "直填值", "直填键应覆盖 manual"
    assert ph[2]["default_value"] == "人工", "白名单外不得覆盖 manual"


def test_sediment_only_keys_non_str_items():
    """only_keys 含非字符串项（前端/画布数据不可信）→ 不抛异常、不误命中。
    注意非 hashable 项（list/dict）会在成员判断处抛 TypeError——与
    values.get(非hashable key) 同源的既有失败模式，此处固化现状不额外兜底。"""
    import pytest
    from api.db.services.template_fill_service import TplTemplateVersionService as S
    ph = [{"key": "a", "default_value": "", "default_source": ""}]
    # 非字符串但 hashable：天然不匹配字符串 key
    assert S._sediment_into_placeholders(ph, {"a": "新A"}, {"a", 1, None}) is True
    assert ph[0]["default_value"] == "新A"
    # 非 hashable：抛 TypeError（固化现状，防后续静默改成「吞掉」而掩盖真 bug）
    with pytest.raises(TypeError):
        S._sediment_into_placeholders([{"key": "a"}], {"a": "新A"}, {["bad"]})


def test_sediment_placeholders_dirty_items_not_touched():
    """placeholders 混入非 dict / None / 无 key 项 → 跳过不报错，正常项照写。"""
    from api.db.services.template_fill_service import TplTemplateVersionService as S
    ph = ["junk", None, {"name": "无key"}, {"key": "a", "default_value": "", "default_source": ""}]
    assert S._sediment_into_placeholders(ph, {"a": "新A", "junk": "x"}, {"a", "junk"}) is True
    assert ph[0] == "junk" and ph[1] is None and ph[2] == {"name": "无key"}
    assert ph[3]["default_value"] == "新A"


# ---------- service：B端默认值编辑（纯函数） ----------

def test_update_defaults_validation():
    from api.db.services.template_fill_service import TplTemplateVersionService as S
    ph = [{"key": "a", "default_value": "", "default_source": ""},
          {"key": "b", "default_value": "x", "default_source": "auto"}]
    # 未知 key（如手改请求的错别字）整体拒绝，防编辑静默丢失
    ok, err = S._apply_defaults_edits(ph, {"a": "新值", "b": "", "ghost": "y"})
    assert ok is False and "ghost" in err
    ok, err = S._apply_defaults_edits(ph, {"a": "新值", "b": ""})
    assert ok is True
    assert ph[0]["default_value"] == "新值" and ph[0]["default_source"] == "manual"
    assert ph[1]["default_value"] == "" and ph[1]["default_source"] == ""  # 空串=清空
    # 超长值（600 字）截断到 MAX_ANCHOR_LEN=500，维持 default_value ≤ 500 不变量
    from rag.svr.template_fill.detector import MAX_ANCHOR_LEN
    ok, err = S._apply_defaults_edits(ph, {"a": "长" * (MAX_ANCHOR_LEN + 100)})
    assert ok is True
    assert ph[0]["default_value"] == "长" * MAX_ANCHOR_LEN
    assert ph[0]["default_source"] == "manual"
    # 纯空白值：剥空白后为空 → 视为清空（default_value="" 且 source=""）
    ok, err = S._apply_defaults_edits(ph, {"a": "   "})
    assert ok is True
    assert ph[0]["default_value"] == "" and ph[0]["default_source"] == ""


# ---------- FILL_HINT_RE 补漏报模式（标准范本实测） ----------

def test_fill_hint_colon_blank_then_content():
    """冒号+留白+后续文字：日期行/盖章行等冒号不在行尾的填写点必须命中。"""
    from rag.svr.template_fill.docx_utils import FILL_HINT_RE
    assert FILL_HINT_RE.search("招标文件编制日期：\u3000 \u3000  年   \u3000月   \u3000日")
    assert FILL_HINT_RE.search("招标人：\u3000\u3000\u3000\u3000（盖单位电子公章）")
    assert FILL_HINT_RE.search("开标时间:      年  月  日  时（北京时间）")


def test_fill_hint_date_blank_placeholder():
    """「　年　月　日」空白日期占位：年前有留白才命中；真实日期不误报。"""
    from rag.svr.template_fill.docx_utils import FILL_HINT_RE
    assert FILL_HINT_RE.search("请于          年     月     日前递交")
    assert FILL_HINT_RE.search("编制日期：\u3000\u3000年\u3000\u3000月\u3000\u3000日")
    # 真实日期（无留白）不算填写特征
    assert not FILL_HINT_RE.search("竣工日期为2026年9月11日，工期180天")
    assert not FILL_HINT_RE.search("本工程计划2026年09月开工")


def test_fill_hint_inline_long_blank():
    """行内长空白占位（跨栏留白）：6 个以上连续空白且两侧有内容。"""
    from rag.svr.template_fill.docx_utils import FILL_HINT_RE
    assert FILL_HINT_RE.search("本招标项目        （项目名称）  已由")
    assert FILL_HINT_RE.search("以\u3000\u3000\u3000\u3000\u3000\u3000\u3000（审批机关名称）批准建设")
    # 普通正文短空格不命中
    assert not FILL_HINT_RE.search("本招标项目 已由有关机关批准建设")


def test_fill_hint_fullwidth_underscore():
    """全角下划线填空线（中文范本常用形态）必须命中——此前只认半角 _{2,}，
    含全角下划线的行根本进不了候选集，LLM 再强也识别不到。"""
    from rag.svr.template_fill.docx_utils import FILL_HINT_RE
    assert FILL_HINT_RE.search("投标人名称：＿＿＿＿＿＿＿＿")
    assert FILL_HINT_RE.search("＿＿＿＿年＿＿月＿＿日")
    # 单个全角下划线不算填写特征（防正文普通下划线误报）
    assert not FILL_HINT_RE.search("详见第＿章")


def test_fill_hint_checkbox():
    """勾选框 □（范本常见「□ 是 □ 否」评审项）必须命中。"""
    from rag.svr.template_fill.docx_utils import FILL_HINT_RE
    assert FILL_HINT_RE.search("□ 是 □ 否")
    assert FILL_HINT_RE.search("资格审查方式：□资格预审 □资格后审")
    # 普通正文不误报
    assert not FILL_HINT_RE.search("按照有关规定执行")


# ---------- 手动占位符 {{key}} 直通 ----------

def test_docx_candidates_include_manual_placeholder_paragraphs():
    """含 {{snake_key}} 的段落无条件进候选（用户显式标注，不经特征猜测）。"""
    from rag.svr.template_fill.docx_utils import extract_docx_candidates
    doc = _make_docx([
        "本工程项目名称为{{project_name}}，工期{{duration_days}}天",
        "没有任何填写特征的普通说明段落",
    ])
    cands = extract_docx_candidates(doc)
    texts = [c["text"] for c in cands]
    assert any("{{project_name}}" in t for t in texts)
    assert all("普通说明段落" not in t for t in texts)


def test_xlsx_candidates_include_manual_placeholder_cells():
    from rag.svr.template_fill.xlsx_utils import extract_xlsx_candidates
    blob = _make_xlsx({"封面": [["项目名称：{{project_name}}"], ["说明文字"]]})
    texts = [c["text"] for c in extract_xlsx_candidates(blob)]
    assert any("{{project_name}}" in t for t in texts)
    assert all("说明文字" not in t for t in texts)


def test_extract_explicit_placeholders_basic_and_multi():
    """直通提取：anchor 为 {{key}} 整串、addr/line 对齐候选行；同段多个占位符拆多条。"""
    from rag.svr.template_fill.detector import extract_explicit_placeholders
    cands = [{"index": 3, "text": "名称{{name}} 工期{{duration_days}}天", "addr": "para:3"}]
    out = extract_explicit_placeholders(cands)
    assert [(it["key"], it["anchor"]) for it in out] == [
        ("name", "{{name}}"), ("duration_days", "{{duration_days}}")]
    assert all(it["addr"] == "para:3" and it["line"] == 3 for it in out)
    assert all(it["fill_mode"] == "llm" and it["top_k"] == 6 for it in out)


def test_extract_explicit_placeholders_ignores_invalid_forms():
    """对抗：大写/数字开头/空 key 等非法占位符形态不命中（与渲染端 docxtpl 口径一致）。"""
    from rag.svr.template_fill.detector import extract_explicit_placeholders
    cands = [{"index": 0, "text": "{{Name}} {{9bad}} {{ok_key}} {{}}", "addr": "para:0"}]
    out = extract_explicit_placeholders(cands)
    assert [it["key"] for it in out] == ["ok_key"]


# ---------- LLM 分块识别 + 合并 ----------

@pytest.mark.asyncio
async def test_detect_chunked_splits_and_merges():
    """候选按 60 行/块切分多次调用；LLM 返回可解析时合并全部识别项。"""
    from rag.svr.template_fill.detector import DETECT_CHUNK_SIZE, _detect_chunked
    n = DETECT_CHUNK_SIZE + 10  # 2 块
    cands = [{"index": i, "text": f"字段{i}：____________", "addr": f"para:{i}"} for i in range(n)]
    calls = []

    async def fake_chat(system, messages):
        calls.append(len(messages[0]["content"]))
        lines = messages[0]["content"].split("编号行：\n")[1]
        idxs = [int(l.split("\t")[0]) for l in lines.strip().splitlines()]
        import json as _json
        return _json.dumps([{"line": i, "anchor": "____________", "key": f"k{i}",
                             "name": f"N{i}", "fill_mode": "llm"} for i in idxs])

    items, failed = await _detect_chunked(fake_chat, "docx", cands)
    assert len(calls) == 2 and failed == 0
    assert len(items) == n


@pytest.mark.asyncio
async def test_detect_chunked_single_chunk_failure_is_partial(monkeypatch):
    """对抗：单块 LLM 抛异常不拖垮整体，其余块结果保留（failed=1）。
    块大小压到 1，保证 2 个候选切成 2 块。"""
    from rag.svr.template_fill import detector
    monkeypatch.setattr(detector, "DETECT_CHUNK_SIZE", 1)
    _detect_chunked = detector._detect_chunked
    cands = [{"index": 0, "text": "A：____________", "addr": "para:0"},
             {"index": 1, "text": "B：____________", "addr": "para:1"}]

    async def flaky_chat(system, messages):
        if "B：" in messages[0]["content"]:
            raise RuntimeError("llm boom")
        return '[{"line": 0, "anchor": "____________", "key": "ka", "name": "A", "fill_mode": "llm"}]'

    items, failed = await _detect_chunked(flaky_chat, "docx", cands)
    assert failed == 1 and [it["key"] for it in items] == ["ka"]


def test_merge_detection_dedups_pos_and_keys():
    """合并：手动直通优先（同 (addr, anchor) 丢弃 LLM 重复项）；跨源 key 冲突加后缀。"""
    from rag.svr.template_fill.detector import _merge_detection
    explicit = [{"key": "name", "anchor": "{{name}}", "addr": "para:1", "line": 1}]
    llm = [{"key": "name", "anchor": "{{name}}", "addr": "para:1", "line": 1},   # 同位置重复 → 丢弃
           {"key": "name", "anchor": "____", "addr": "para:0", "line": 0},       # 同 key 不同位置 → name_2
           {"key": "date", "anchor": "{{date}}", "addr": "para:2", "line": 2}]
    out = _merge_detection(explicit, llm)
    assert [(it["key"], it["anchor"]) for it in out] == [
        ("name", "{{name}}"), ("name_2", "____"), ("date", "{{date}}")]


def test_merge_detection_empty_llm_keeps_explicit():
    """LLM 全军覆没时手动直通项仍保留（确定性标注不依赖 LLM 存活）。"""
    from rag.svr.template_fill.detector import _merge_detection
    explicit = [{"key": "k", "anchor": "{{k}}", "addr": "para:0", "line": 0}]
    assert _merge_detection(explicit, []) == explicit


# ---------- detector：标签/骨架型 anchor 拒绝派生（2026-09-12 福建通用本污染防御） ----------

def test_derive_default_from_anchor_rejects_labels():
    """对抗：冒号结尾标签是模板提示文字而非已填现值——派生成默认值会被 D−C
    条件执行原样回写成稿（用户视角「没填」）并经 sediment 固化污染基线。"""
    from rag.svr.template_fill.detector import derive_default_from_anchor
    assert derive_default_from_anchor("编号：") == ""
    assert derive_default_from_anchor("申请人:") == ""      # 半角冒号
    assert derive_default_from_anchor("招标人： ") == ""     # 尾随空白
    assert derive_default_from_anchor("一、投标保证金（大写）：") == ""  # 前缀+冒号结尾


def test_derive_default_from_anchor_rejects_date_skeletons():
    """对抗：年月日字+空白标点、无数字 = 日期骨架（空范本留空位），不派生默认值；
    含数字的真实日期（2026年9月28日）必须保留。"""
    from rag.svr.template_fill.detector import derive_default_from_anchor
    assert derive_default_from_anchor("年　　月　　日") == ""   # 全角空格骨架
    assert derive_default_from_anchor("年   月   日") == ""    # 半角空格骨架
    assert derive_default_from_anchor("月") == ""              # 单字骨架
    assert derive_default_from_anchor("＿＿年＿＿月＿＿日") == ""
    # 含数字 → 真实现值，不误杀
    assert derive_default_from_anchor("2026年9月28日") == "2026年9月28日"
    assert derive_default_from_anchor("2026-09-28") == "2026-09-28"


def test_derive_default_from_anchor_rejects_paren_hints():
    """对抗：括号提示（待填提示语）不是现值；（大写） 元 这类带金额后缀也要拦。"""
    from rag.svr.template_fill.detector import derive_default_from_anchor
    assert derive_default_from_anchor("（投标人名称）") == ""
    assert derive_default_from_anchor("（大写）") == ""
    assert derive_default_from_anchor("（元） 元") == ""
    assert derive_default_from_anchor("（　　）万元") == ""      # 货币后缀提示
    assert derive_default_from_anchor("(投标人名称)") == ""     # 半角括号
    # 非括号包裹的正常内容不误杀
    assert derive_default_from_anchor("（含）税金额100万元") == "（含）税金额100万元"


def test_derive_default_from_anchor_rejects_mixed_blank_hint():
    """对抗（Major-1 回归）：blank_slots 合并位「下划线+括号提示+下划线」混合形态——
    strip 后既非纯留白、也不 fullmatch _HINT_ANCHOR_RE，曾泄漏提示语为默认值
    （'＿＿＿（项目名称）＿＿' 整串回流 default_map → 字段被判已有值永不填写）。"""
    from rag.svr.template_fill.detector import derive_default_from_anchor
    assert derive_default_from_anchor("＿＿＿（项目名称）＿＿") == ""       # 全角下划线+提示
    assert derive_default_from_anchor("___（甲方）___（乙方）___") == ""   # 半角下划线+多提示组
    assert derive_default_from_anchor("＿＿＿（大写）＿＿  （盖章）") == ""  # 多提示组+空格混合


def test_derive_default_from_anchor_mixed_real_content_not_killed():
    """反向保护：剔除非括号骨架字符后剩真实内容（非括号 CJK/字母/数字）不误杀。"""
    from rag.svr.template_fill.detector import derive_default_from_anchor
    assert derive_default_from_anchor("＿＿＿实际内容＿＿") == "＿＿＿实际内容＿＿"
    assert derive_default_from_anchor("厦门××公司") == "厦门××公司"
    assert derive_default_from_anchor("（含）税金额100万元") == "（含）税金额100万元"
    assert derive_default_from_anchor("＿＿＿2026年9月＿＿") == "＿＿＿2026年9月＿＿"


# ---------- renderer：AI 填入值标蓝（渲染前占位 run 预染色） ----------

def _run_colors_of_first_para(blob: bytes):
    """取成稿第一个段落的 [(run文本, w:color val 或 None)] 列表。"""
    from docx import Document
    from docx.oxml.ns import qn
    p = Document(io.BytesIO(blob)).paragraphs[0]
    out = []
    for r in p.runs:
        rPr = r._r.find(qn("w:rPr"))
        color = None
        if rPr is not None:
            c = rPr.find(qn("w:color"))
            if c is not None:
                color = c.get(qn("w:val"))
        out.append((r.text, color))
    return out


def test_render_docx_value_turns_blue():
    """整 run 即占位符：渲染后 AI 填入值继承占位 run 的蓝色（0000FF）。"""
    from docx import Document
    from rag.svr.template_fill.renderer import render_docx
    doc = Document()
    doc.add_paragraph("{{name}}")
    buf = io.BytesIO(); doc.save(buf)
    out = render_docx(buf.getvalue(), {"name": "测试项目"})
    runs = _run_colors_of_first_para(out)
    assert any(t == "测试项目" and c == "0000FF" for t, c in runs)


def test_render_docx_mixed_run_only_value_blue():
    """对抗：混合 run「项目名称：{{name}}」拆分后只有值标蓝，标签保持无色——
    若拆分实现有误把标签一起染色，此用例即红。"""
    from rag.svr.template_fill.renderer import render_docx
    out = render_docx(_mk_docx_with_placeholder(), {"name": "测试项目"})
    runs = _run_colors_of_first_para(out)
    label_blue = [c for t, c in runs if "项目名称" in t and c == "0000FF"]
    value_blue = [c for t, c in runs if "测试项目" in t and c == "0000FF"]
    assert not label_blue and value_blue


def test_render_docx_no_placeholder_blob_unchanged():
    """零拷贝：无占位符的 blob 原样返回（省一次解析+序列化）。"""
    from docx import Document
    from rag.svr.template_fill.renderer import _colorize_placeholder_runs
    doc = Document()
    doc.add_paragraph("纯文本没有占位符")
    buf = io.BytesIO(); doc.save(buf)
    assert _colorize_placeholder_runs(buf.getvalue()) == buf.getvalue()


def test_render_docx_run_with_br_skipped_safely():
    """对抗：含 w:br 的 run 拆分会丢结构 → 整体跳过不标蓝，渲染仍正常完成
    （防御性降级：值正常回填，只是不标蓝，不抛异常不坏文档）。"""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from rag.svr.template_fill.renderer import render_docx
    doc = Document()
    p = doc.add_paragraph()
    r = p.add_run("{{name}}")
    r._r.append(OxmlElement("w:br"))  # 非 (rPr|t) 子节点 → 触发跳过守卫
    buf = io.BytesIO(); doc.save(buf)
    out = render_docx(buf.getvalue(), {"name": "测试项目"})
    text = "\n".join(pp.text for pp in Document(io.BytesIO(out)).paragraphs)
    assert "测试项目" in text and "{{" not in text


def test_render_docx_color_respects_rpr_schema_order():
    """对抗：占位 run 已带 rPr（w:sz/w:u）时，w:color 必须按 OOXML schema
    sequence 插入（位于 w:sz/w:u 之前）且保留既有属性——裸 append 会产出
    乱序 XML，严格校验器（部分 WPS/LibreOffice/PDF 转换链）会丢弃颜色。"""
    from docx import Document
    from docx.oxml.ns import qn
    from rag.svr.template_fill.renderer import render_docx
    doc = Document()
    p = doc.add_paragraph("编号：{{k}}")
    for run in p.runs:
        run.font.size = 1  # 写入 w:sz（half-point=1）
        run.font.underline = True
    buf = io.BytesIO(); doc.save(buf)
    out = render_docx(buf.getvalue(), {"k": "值"})
    p2 = Document(io.BytesIO(out)).paragraphs[0]
    order, sz_kept = [], False
    for r in p2.runs:
        rPr = r._r.find(qn("w:rPr"))
        if rPr is None:
            continue
        order = [c.tag.split("}")[1] for c in rPr]
        if "color" in order and order.index("color") > order.index("sz"):
            raise AssertionError(f"w:color 乱序: {order}")
        if "值" in (r.text or ""):
            sz_kept = rPr.find(qn("w:sz")) is not None and any(
                t == "值" and c == "0000FF" for t, c in _run_colors_of_first_para(out))
    assert sz_kept


# ---------- 识别后处理：标签/实心 anchor 收缩修正与低置信标记 ----------

def test_shrink_anchor_to_blank_finds_blank_after_label():
    from rag.svr.template_fill.detector import _shrink_anchor_to_blank
    line = "投标人名称：＿＿＿＿＿＿（盖章）"
    assert _shrink_anchor_to_blank("投标人名称：", line) == "＿＿＿＿＿＿"


def test_shrink_anchor_to_blank_no_blank_returns_empty():
    from rag.svr.template_fill.detector import _shrink_anchor_to_blank
    assert _shrink_anchor_to_blank("投标人名称：", "投标人名称：签字") == ""


def test_parse_label_anchor_shrunk_to_blank():
    from rag.svr.template_fill.detector import parse_detection_response
    cands = [{"index": 0, "text": "投标人名称：＿＿＿＿＿＿", "addr": "para:0"}]
    raw = ('[{"line":0,"anchor":"投标人名称：","key":"bidder_name",'
           '"name":"投标人名称","retrieval_query":"","fill_mode":"llm","required":true}]')
    out = parse_detection_response(raw, cands)
    assert len(out) == 1
    assert out[0]["anchor"] == "＿＿＿＿＿＿"
    assert out[0]["low_confidence"] is False


def test_parse_label_anchor_without_blank_dropped():
    from rag.svr.template_fill.detector import parse_detection_response
    cands = [{"index": 0, "text": "包括以下内容：", "addr": "para:0"}]
    raw = ('[{"line":0,"anchor":"包括以下内容：","key":"content",'
           '"name":"内容","retrieval_query":"","fill_mode":"llm","required":true}]')
    assert parse_detection_response(raw, cands) == []


def test_parse_solid_anchor_with_blank_in_line_shrunk():
    from rag.svr.template_fill.detector import parse_detection_response
    cands = [{"index": 0, "text": "工程名称　　（填写完整名称）　开工日期", "addr": "para:0"}]
    raw = ('[{"line":0,"anchor":"工程名称","key":"project_name",'
           '"name":"工程名称","retrieval_query":"","fill_mode":"llm","required":true}]')
    out = parse_detection_response(raw, cands)
    assert len(out) == 1
    assert out[0]["anchor"] == "（填写完整名称）"


def test_parse_solid_anchor_no_blank_in_line_low_confidence_kept():
    from rag.svr.template_fill.detector import parse_detection_response
    # 已填范本现值：整行无留白特征，anchor 合法保留但打低置信
    cands = [{"index": 0, "text": "合同金额为人民币壹佰万元整", "addr": "para:0"}]
    raw = ('[{"line":0,"anchor":"壹佰万元整","key":"amount",'
           '"name":"金额","retrieval_query":"","fill_mode":"llm","required":true}]')
    out = parse_detection_response(raw, cands)
    assert len(out) == 1
    assert out[0]["anchor"] == "壹佰万元整"
    assert out[0]["low_confidence"] is True


def test_parse_blank_anchor_no_low_confidence():
    from rag.svr.template_fill.detector import parse_detection_response
    cands = [{"index": 0, "text": "编号：＿＿＿＿＿＿", "addr": "para:0"}]
    raw = ('[{"line":0,"anchor":"＿＿＿＿＿＿","key":"code",'
           '"name":"编号","retrieval_query":"","fill_mode":"llm","required":true}]')
    out = parse_detection_response(raw, cands)
    assert len(out) == 1
    assert out[0]["anchor"] == "＿＿＿＿＿＿"
    assert out[0]["low_confidence"] is False


# ---------- 识别后处理质量修复：实心收缩打低置信/撞车回退原anchor/混合anchor收缩 ----------

def test_parse_solid_anchor_shrunk_in_mixed_line_low_confidence():
    """对抗（混合行）：实心 anchor 收缩成功也只是兜底猜测——
    "大写：壹佰万元整 小写：＿＿＿＿" 中 anchor=壹佰万元整 会被收缩到"小写"字段的空位，
    属静默错位，必须打低置信警示人工确认。"""
    from rag.svr.template_fill.detector import parse_detection_response
    cands = [{"index": 0, "text": "大写：壹佰万元整 小写：＿＿＿＿", "addr": "para:0"}]
    raw = ('[{"line":0,"anchor":"壹佰万元整","key":"amount",'
           '"name":"金额","retrieval_query":"","fill_mode":"llm","required":true}]')
    out = parse_detection_response(raw, cands)
    assert len(out) == 1
    assert out[0]["anchor"] == "＿＿＿＿"
    assert out[0]["low_confidence"] is True


def test_merge_collision_from_shrink_falls_back_to_orig_anchor():
    """对抗（同行双标签）：两个标签收缩后 anchor 撞车（均为同一留白串），
    (addr, anchor) 去重会把 party_b 静默丢弃——必须回退到原标签 anchor 保留该项
    并打低置信；内部字段 _orig_anchor 不得泄漏进产物。"""
    from rag.svr.template_fill.detector import _merge_detection, parse_detection_response
    cands = [{"index": 0, "text": "甲方：＿＿＿＿ 乙方：＿＿＿＿", "addr": "para:0"}]
    raw = ('[{"line":0,"anchor":"甲方：","key":"party_a","name":"甲方","fill_mode":"llm"},'
           '{"line":0,"anchor":"乙方：","key":"party_b","name":"乙方","fill_mode":"llm"}]')
    parsed = parse_detection_response(raw, cands)
    assert len(parsed) == 2
    out = _merge_detection([], parsed)
    assert len(out) == 2
    assert out[0]["key"] == "party_a" and out[0]["anchor"] == "＿＿＿＿"
    assert out[0]["low_confidence"] is False
    assert out[1]["key"] == "party_b" and out[1]["anchor"] == "乙方："
    assert out[1]["low_confidence"] is True
    assert all("_orig_anchor" not in it for it in out)


def test_parse_mixed_label_blank_anchor_shrinks_to_blank():
    """对抗（混合 anchor）："编号：＿＿＿" 结尾非冒号、含留白，旧逻辑原样放行 →
    渲染时整个"编号：＿＿＿"被替换丢标签。必须收缩为纯留白后缀（low_confidence=False：
    收缩后即纯留白，替换只动留白，无歧义）。"""
    from rag.svr.template_fill.detector import parse_detection_response
    cands = [{"index": 0, "text": "编号：＿＿＿＿＿", "addr": "para:0"}]
    raw = ('[{"line":0,"anchor":"编号：＿＿＿＿＿","key":"code",'
           '"name":"编号","retrieval_query":"","fill_mode":"llm","required":true}]')
    out = parse_detection_response(raw, cands)
    assert len(out) == 1
    assert out[0]["anchor"] == "＿＿＿＿＿"
    assert out[0]["low_confidence"] is False


# ---------- 跨 run 区间替换（格式保真） ----------

def _doc_with_runs(texts, bolds=None):
    """构造单段落多 run 文档：texts 为各 run 文本，bolds 为对应加粗标记。"""
    buf = io.BytesIO()
    doc = Document()
    p = doc.add_paragraph()
    for k, t in enumerate(texts):
        r = p.add_run(t)
        if bolds and bolds[k]:
            r.bold = True
    doc.save(buf)
    return buf.getvalue()


def test_cross_run_replace_preserves_outside_runs():
    from rag.svr.template_fill.docx_utils import apply_docx_placeholders
    blob = _doc_with_runs(["AB", "CD", "EF"], [False, False, True])
    out = apply_docx_placeholders(blob, [{"addr": "para:0", "anchor": "BC", "key": "k"}])
    p = Document(io.BytesIO(out)).paragraphs[0]
    assert [r.text for r in p.runs] == ["A{{k}}", "D", "EF"]
    assert p.runs[2].bold  # 区间外 run 的格式完整保留


def test_cross_run_replace_span_three_runs_clears_middle():
    from rag.svr.template_fill.docx_utils import apply_docx_placeholders
    blob = _doc_with_runs(["AB", "CD", "EF"])
    out = apply_docx_placeholders(blob, [{"addr": "para:0", "anchor": "BCDE", "key": "k"}])
    p = Document(io.BytesIO(out)).paragraphs[0]
    assert [r.text for r in p.runs] == ["A{{k}}", "", "F"]


def test_cross_run_replace_anchor_at_end_single_tail_run():
    from rag.svr.template_fill.docx_utils import apply_docx_placeholders
    blob = _doc_with_runs(["AB", "CD"])
    out = apply_docx_placeholders(blob, [{"addr": "para:0", "anchor": "CD", "key": "k"}])
    p = Document(io.BytesIO(out)).paragraphs[0]
    assert "".join(r.text for r in p.runs) == "AB{{k}}"


def test_cross_run_replace_anchor_whole_paragraph():
    from rag.svr.template_fill.docx_utils import apply_docx_placeholders
    blob = _doc_with_runs(["AB", "CD"])
    out = apply_docx_placeholders(blob, [{"addr": "para:0", "anchor": "ABCD", "key": "k"}])
    p = Document(io.BytesIO(out)).paragraphs[0]
    assert "".join(r.text for r in p.runs) == "{{k}}"


def test_cross_run_replace_multiple_occurrences_across_runs():
    from rag.svr.template_fill.docx_utils import apply_docx_placeholders
    # anchor "AA" 不落在任何单一 run 内（真正走跨 run 路径），且跨 run 边界出现 2 次；
    # str.replace 语义须全部替换。注意不能用 "_" 这类单 run 内完整出现的 anchor——
    # 那会被 _replace_in_paragraph 的单 run 快速路径拦截（首个命中 run 即 return）。
    blob = _doc_with_runs(["xA", "Ax", "xA", "Ax"])
    out = apply_docx_placeholders(blob, [{"addr": "para:0", "anchor": "AA", "key": "k"}])
    p = Document(io.BytesIO(out)).paragraphs[0]
    assert "".join(r.text for r in p.runs) == "x{{k}}xx{{k}}x"


def test_cross_run_replace_anchor_equals_replacement_no_loop():
    from rag.svr.template_fill.docx_utils import apply_docx_placeholders
    # 手动占位符直通场景：anchor == repl，必须幂等且不死循环
    blob = _doc_with_runs(["AB", "{{k}}", "CD"])
    out = apply_docx_placeholders(blob, [{"addr": "para:0", "anchor": "{{k}}", "key": "k"}])
    p = Document(io.BytesIO(out)).paragraphs[0]
    assert "".join(r.text for r in p.runs) == "AB{{k}}CD"


def test_cross_run_replace_anchor_substring_of_repl_no_rescan():
    """对抗（C1 回归）：anchor 是 repl 的子串（如 anchor="name"、repl="{{name}}"）时，
    每轮循环若从头重扫 find(anchor)，会命中刚写入的替换产物（"{{name}}" 内含 "name"），
    导致文本腐坏 + 第二处真实出现漏替。必须用偏移扫描跳过刚写入的 repl，
    严格复刻 str.replace 不重扫产物的语义。"""
    from rag.svr.template_fill.docx_utils import apply_docx_placeholders
    blob = _doc_with_runs(["na", "me", "na", "me"])
    out = apply_docx_placeholders(blob, [{"addr": "para:0", "anchor": "name", "key": "name"}])
    p = Document(io.BytesIO(out)).paragraphs[0]
    joined = "".join(r.text for r in p.runs)
    assert joined == "{{name}}{{name}}"
    assert joined.count("{{name}}") == 2


def test_cross_run_replace_preserves_structured_empty_middle_run():
    """对抗（I1 回归）：区间中段文本为空但含结构节点（w:fldChar/w:drawing 等）的 run，
    Run.text = "" setter 会清掉 rPr 外全部子节点，域字符/行内图被删。
    注意：python-docx 1.2.0 的 Run.text getter 会把 w:br 翻译成 "\\n"（非空），
    故此处用贡献空文本的 w:fldChar 构造「空文本+结构」run。
    仅当 run 有文本时才置空；空文本 run 跳过（置空唯一效果就是销毁结构）。"""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    from rag.svr.template_fill.docx_utils import apply_docx_placeholders
    buf = io.BytesIO()
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("AB")
    r_mid = p.add_run()      # 空文本 run
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "begin")
    r_mid._r.append(fc)      # 带 w:fldChar 结构节点（域字符/行内图代表）
    p.add_run("CD")
    doc.save(buf)
    out = apply_docx_placeholders(buf.getvalue(), [{"addr": "para:0", "anchor": "BC", "key": "k"}])
    p2 = Document(io.BytesIO(out)).paragraphs[0]
    assert "".join(r.text for r in p2.runs) == "A{{k}}D"
    # 空文本中段 run 的 w:fldChar 必须仍存在
    assert p2.runs[1]._r.findall(qn("w:fldChar")), "中段空文本 run 的 w:fldChar 被 text='' setter 销毁"


def test_cross_run_replace_empty_text_run_sandwich():
    """纯空文本中段 run（无结构节点）：替换正确落位、中间 run 保持空文本，
    两侧文本各自保留前缀/后缀。"""
    from rag.svr.template_fill.docx_utils import apply_docx_placeholders
    blob = _doc_with_runs(["AB", "", "CD"])
    out = apply_docx_placeholders(blob, [{"addr": "para:0", "anchor": "BC", "key": "k"}])
    p = Document(io.BytesIO(out)).paragraphs[0]
    assert [r.text for r in p.runs] == ["A{{k}}", "", "D"]


# ---------- 范本识别加固（2026-09-13）：文本框/页眉页脚/内容控件编址 ----------

_MC_NS = ('xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
          'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
          'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"')


def _make_docx_with_textbox(txbx_text):
    """正文段落内嵌浮动文本框，mc:AlternateContent 双份存储（Choice+Fallback 同文）。
    直接拼 XML 绕开 python-docx 无文本框 API 的限制（项目运行时同样按 XML 层遍历）。"""
    from docx.oxml import parse_xml
    doc = Document()
    doc.add_paragraph("无填写点的普通段落")
    p = doc.add_paragraph("含文本框的容器段落")
    xml = (
        f'<w:r {_MC_NS}>'
        '<mc:AlternateContent>'
        f'<mc:Choice Requires="wps"><w:txbxContent>'
        f'<w:p><w:r><w:t>{txbx_text}</w:t></w:r></w:p>'
        '</w:txbxContent></mc:Choice>'
        f'<mc:Fallback><w:txbxContent>'
        f'<w:p><w:r><w:t>{txbx_text}</w:t></w:r></w:p>'
        '</w:txbxContent></mc:Fallback>'
        '</mc:AlternateContent>'
        '</w:r>'
    )
    p._p.append(parse_xml(xml))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_textbox_paragraph_addressed_and_fallback_deduped():
    from rag.svr.template_fill.docx_utils import extract_docx_candidates, iter_docx_paragraphs
    blob = _make_docx_with_textbox("单位名称：＿＿＿＿＿＿")
    items = iter_docx_paragraphs(blob)
    hits = [it for it in items if "单位名称" in it["text"]]
    # Choice + Fallback 双份只编址一次
    assert len(hits) == 1
    assert hits[0]["addr"] == "para:1:tx0:0"
    # 容器段落自身无填写特征，不进候选；文本框段落进候选
    cands = extract_docx_candidates(blob)
    assert [c["addr"] for c in cands] == ["para:1:tx0:0"]


# ---------- I-1 返工（2026-09-13）：存量计数器隔离 + sdt 独立 sdt: 前缀 ----------

_W_NS = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


def _make_sdt_xml(inner_xml):
    """构造 body 直系 w:sdt（含 sdtContent），inner 为 sdtContent 内部块 XML。"""
    from docx.oxml import parse_xml
    return parse_xml(
        f'<w:sdt {_W_NS}><w:sdtPr><w:alias w:val="t"/></w:sdtPr>'
        f'<w:sdtContent>{inner_xml}</w:sdtContent></w:sdt>'
    )


def _sdt_para(text):
    return f'<w:p><w:r><w:t>{text}</w:t></w:r></w:p>'


def _add_body_sdt(doc, inner_xml):
    """构造 body 直系 w:sdt 并插入 sectPr 之前（body.append 会落到文档末尾，
    破坏文档顺序断言），返回插入的元素。"""
    from docx.oxml.ns import qn
    sdt = _make_sdt_xml(inner_xml)
    body = doc.element.body
    sect_pr = body.find(qn("w:sectPr"))
    if sect_pr is not None:
        sect_pr.addprevious(sdt)
    else:
        body.append(sdt)
    return sdt


def test_sdt_wrapped_paragraph_addressed():
    """body 直系 sdt 内段落 addr == sdt:0:0；前后普通段落仍连续 para:0/para:1
    （存量计数器不被挤占，I-1 红线）；extract_docx_candidates 命中 sdt 文本。"""
    from rag.svr.template_fill.docx_utils import extract_docx_candidates, iter_docx_paragraphs
    doc = Document()
    doc.add_paragraph("sdt 前面的段落")
    _add_body_sdt(doc, _sdt_para("内容控件里的日期：____"))
    doc.add_paragraph("sdt 之后的段落：____")
    buf = io.BytesIO()
    doc.save(buf)
    blob = buf.getvalue()

    items = iter_docx_paragraphs(blob)
    by_addr = {it["addr"]: it["text"] for it in items}
    assert by_addr["sdt:0:0"] == "内容控件里的日期：____"
    assert by_addr["para:0"] == "sdt 前面的段落"
    assert by_addr["para:1"] == "sdt 之后的段落：____"
    cands = extract_docx_candidates(blob)
    assert any(c["addr"] == "sdt:0:0" and "内容控件里的日期" in c["text"] for c in cands)


def test_legacy_para_addrs_stable_with_txbx_and_sdt():
    """「段落→含文本框段落→sdt→段落→顶层表格」文档：存量 addr 序列与升级前逐字节一致，
    新元素（文本框/sdt）只消耗扁平 index，零消耗 para:/cell: 计数器。"""
    from docx.oxml import parse_xml

    from rag.svr.template_fill.docx_utils import iter_docx_paragraphs
    doc = Document()
    doc.add_paragraph("第一段")
    p = doc.add_paragraph("含文本框的容器段落")
    p._p.append(parse_xml(
        f'<w:r {_MC_NS}>'
        '<mc:AlternateContent>'
        f'<mc:Choice Requires="wps"><w:txbxContent>'
        f'<w:p><w:r><w:t>文本框里：＿＿＿</w:t></w:r></w:p>'
        '</w:txbxContent></mc:Choice>'
        '</mc:AlternateContent>'
        '</w:r>'
    ))
    _add_body_sdt(doc, _sdt_para("内容控件：____"))
    doc.add_paragraph("sdt 后的段落")
    tbl = doc.add_table(rows=1, cols=1)
    tbl.rows[0].cells[0].paragraphs[0].text = "表格单元格"
    buf = io.BytesIO()
    doc.save(buf)

    addrs = [it["addr"] for it in iter_docx_paragraphs(buf.getvalue())]
    # 存量 addr 序列逐字节不变（新元素不挤占 para:/cell: 计数器）
    assert addrs[:2] == ["para:0", "para:1"]
    # 文本框/sdt 段落只在扁平 index 后追加（addr 用独立前缀）
    assert "para:1:tx0:0" in addrs
    assert "sdt:0:0" in addrs
    # sdt 之后的段落仍是 para:2（若计数器被挤占会错位成 para:3）
    assert addrs[4] == "para:2"
    # 顶层表格仍是 cell:0（存量 0 起编号不变）
    assert "cell:0:0:0:0" in addrs
    # 扁平 index 严格按文档顺序递增（含新区域）
    idxs = [it["index"] for it in iter_docx_paragraphs(buf.getvalue())]
    assert idxs == list(range(len(idxs)))


def test_sdt_in_cell_and_textbox_addressed():
    """cell 内 sdt → <cell_addr>:sdt<k>:<pi>；文本框内 sdt → <父addr>:tx<k>:sdt<j>:<pi>；
    sdt 嵌 sdt → :sdt<j>: 段追加递归。"""
    from docx.oxml import parse_xml

    from rag.svr.template_fill.docx_utils import iter_docx_paragraphs
    doc = Document()
    tbl = doc.add_table(rows=1, cols=1)
    cell = tbl.rows[0].cells[0]
    cell.paragraphs[0].text = "单元格普通段落"
    cell._tc.append(_make_sdt_xml(_sdt_para("单元格内控件：____")))
    p = doc.add_paragraph("含文本框的容器段落")
    p._p.append(parse_xml(
        f'<w:r {_MC_NS}>'
        '<mc:AlternateContent>'
        f'<mc:Choice Requires="wps"><w:txbxContent>'
        f'<w:sdt {_W_NS}><w:sdtContent>{_sdt_para("文本框内控件：____")}</w:sdtContent></w:sdt>'
        '</w:txbxContent></mc:Choice>'
        '</mc:AlternateContent>'
        '</w:r>'
    ))
    # sdt 嵌 sdt：body 直系 sdt 内再嵌一层（该文档第 0 个 body 直系 sdt → sdt:0:sdt0:0）
    _add_body_sdt(doc, f'<w:sdt {_W_NS}><w:sdtContent>{_sdt_para("二层控件：____")}</w:sdtContent></w:sdt>')
    # 第二个 body 直系 sdt：锁定 sdt 计数器独立递增
    _add_body_sdt(doc, _sdt_para("第二个控件：____"))
    buf = io.BytesIO()
    doc.save(buf)

    by_addr = {it["addr"]: it["text"] for it in iter_docx_paragraphs(buf.getvalue())}
    assert by_addr["cell:0:0:0:sdt0:0"] == "单元格内控件：____"
    # 文本框容器段是第 2 个 body 直系段落（表格 cell 直系段落消耗 para: 计数，
    # 属存量语义）→ para:1；cell 内 sdt 深度递增不改变 addr 形态
    assert by_addr["para:1:tx0:sdt0:0"] == "文本框内控件：____"
    assert by_addr["sdt:0:sdt0:0"] == "二层控件：____"
    assert by_addr["sdt:1:0"] == "第二个控件：____"


def test_legacy_para_addrs_stable_table_before_paragraphs():
    """回归（I-1 二次返工）：表格在前、段落在后是范本极常见布局——存量基线
    （a8531b1b~1 部署版）中 cell 段落本就消耗扁平序号，表格后正文段落编号必须
    把表格段落数算进去（首次返工误把 cell 段落排除在存量计数器外，此类模板
    升级后全部 para: 前移、DB 存量 addr 错位错填）。文本框/sdt 新区域仍零消耗。"""
    from docx.oxml import parse_xml

    from rag.svr.template_fill.docx_utils import iter_docx_paragraphs
    doc = Document()
    tbl = doc.add_table(rows=1, cols=1)
    cell = tbl.rows[0].cells[0]
    cell.paragraphs[0].text = "表格第一段：____"
    cell.add_paragraph("表格第二段：＿＿＿")
    doc.add_paragraph("段落A")
    doc.add_paragraph("段落B")
    p = doc.add_paragraph("含文本框的容器段落")
    p._p.append(parse_xml(
        f'<w:r {_MC_NS}>'
        '<mc:AlternateContent>'
        f'<mc:Choice Requires="wps"><w:txbxContent>'
        f'<w:p><w:r><w:t>文本框里：＿＿＿</w:t></w:r></w:p>'
        '</w:txbxContent></mc:Choice>'
        '</mc:AlternateContent>'
        '</w:r>'
    ))
    _add_body_sdt(doc, _sdt_para("内容控件：____"))
    doc.add_paragraph("段落C")
    buf = io.BytesIO()
    doc.save(buf)

    addrs = [it["addr"] for it in iter_docx_paragraphs(buf.getvalue())]
    # 与存量基线语义逐字节一致：表格两段消耗 para: 计数；文本框/sdt 零消耗
    assert addrs == [
        "cell:0:0:0:0", "cell:0:0:0:1",   # 表格段（消耗计数）
        "para:2", "para:3",               # 段落A/B：编号含表格段落数（缺陷版本此处为 para:0/para:1）
        "para:4", "para:4:tx0:0",         # 文本框容器段照常计数；框内段落零消耗
        "sdt:0:0",                        # sdt 新区域零消耗
        "para:5",                         # 段落C 不受文本框/sdt 挤占
    ]


def test_legacy_nested_table_paras_do_not_consume_para_seq():
    """回归（I-1 二次返工）：嵌套 ":t<j>" 表内段落属本次新增编址范围——存量基线
    （a8531b1b~1）根本不编址嵌套表、其段落不消耗任何计数器。若 legacy 透传进
    嵌套表，含嵌套表的存量范本升级后段落编号同样前移、DB 存量 addr 错位。"""
    from rag.svr.template_fill.docx_utils import iter_docx_paragraphs
    doc = Document()
    tbl = doc.add_table(rows=1, cols=1)
    cell = tbl.rows[0].cells[0]
    cell.paragraphs[0].text = "外层单元格"
    nested = cell.add_table(rows=1, cols=1)
    nested.rows[0].cells[0].paragraphs[0].text = "内层：＿＿＿＿＿"
    doc.add_paragraph("表格后的段落")
    buf = io.BytesIO()
    doc.save(buf)

    by_addr = {it["addr"]: it["text"] for it in iter_docx_paragraphs(buf.getvalue())}
    assert by_addr["cell:0:0:0:0"] == "外层单元格"
    assert by_addr["cell:0:0:0:t0:0:0:0"] == "内层：＿＿＿＿＿"   # 嵌套表新增编址照常可用
    # 嵌套表段落不消耗 para: 计数（存量基线不编址嵌套表）——段号 = cell 直系段落数
    # 2（"外层单元格" + python-docx 按规则在嵌套表后自动补的空段）；若 legacy 误透传
    # 进嵌套表，此处会前移成 para:3
    assert by_addr["cell:0:0:0:1"] == ""
    assert by_addr["para:2"] == "表格后的段落"


def test_sdt_without_content_skipped():
    """无 sdtContent 的畸形 sdt：不产生 addr、不崩溃，后续正常编址。"""
    from docx.oxml import parse_xml
    from docx.oxml.ns import qn

    from rag.svr.template_fill.docx_utils import iter_docx_paragraphs
    doc = Document()
    doc.add_paragraph("畸形控件前的段落")
    # 空 sdtContent 的合法 sdt（内部无任何块）
    _add_body_sdt(doc, "")
    # 完全无 sdtContent 子节点的畸形 sdt
    sdt_el = parse_xml(f'<w:sdt {_W_NS}><w:sdtPr><w:alias w:val="x"/></w:sdtPr></w:sdt>')
    doc.element.body.find(qn("w:sectPr")).addprevious(sdt_el)
    doc.add_paragraph("畸形控件后的段落：____")
    buf = io.BytesIO()
    doc.save(buf)

    items = iter_docx_paragraphs(buf.getvalue())  # 不抛异常
    addrs = [it["addr"] for it in items]
    assert addrs == ["para:0", "para:1"]


def test_sdt_containing_table_addressed():
    """body 直系 sdt 内表格：cell 段落 addr == sdt:<k>:cell:<tbl_k>:<r>:<c>:<pi>
    （2x2 表格锁定行/列序）；sdt 内表格属新区域零消耗——其后正文段落仍是 para:0、
    顶层存量表格仍是 cell:0（tbl_no/para_seq 均不被挤占）。"""
    from rag.svr.template_fill.docx_utils import iter_docx_paragraphs
    tbl_xml = (
        '<w:tbl><w:tblPr><w:tblW w:w="0" w:type="auto"/></w:tblPr>'
        '<w:tblGrid><w:gridCol w:w="100"/><w:gridCol w:w="100"/></w:tblGrid>'
        '<w:tr><w:tc><w:tcPr/><w:p><w:r><w:t>项目名称：____</w:t></w:r></w:p></w:tc>'
        '<w:tc><w:tcPr/><w:p><w:r><w:t>编号：＿＿</w:t></w:r></w:p></w:tc></w:tr>'
        '<w:tr><w:tc><w:tcPr/><w:p><w:r><w:t>sdt 表格第二行</w:t></w:r></w:p></w:tc>'
        '<w:tc><w:tcPr/><w:p/></w:tc></w:tr></w:tbl>'
    )
    doc = Document()
    _add_body_sdt(doc, tbl_xml)
    doc.add_paragraph("sdt 表格后的段落：____")
    legacy_tbl = doc.add_table(rows=1, cols=1)
    legacy_tbl.rows[0].cells[0].paragraphs[0].text = "顶层存量表格"
    buf = io.BytesIO()
    doc.save(buf)

    by_addr = {it["addr"]: it["text"] for it in iter_docx_paragraphs(buf.getvalue())}
    # sdt 内表格 cell 段落：sdt:0:cell:0:<r>:<c>:<pi>（行/列各就各位）
    assert by_addr["sdt:0:cell:0:0:0:0"] == "项目名称：____"
    assert by_addr["sdt:0:cell:0:0:1:0"] == "编号：＿＿"
    assert by_addr["sdt:0:cell:0:1:0:0"] == "sdt 表格第二行"
    assert by_addr["sdt:0:cell:0:1:1:0"] == ""      # 空段照常编址
    # sdt 内表格零消耗：正文段落 para:0、顶层表格 cell:0（存量计数器独立）
    assert by_addr["para:0"] == "sdt 表格后的段落：____"
    assert by_addr["cell:0:0:0:0"] == "顶层存量表格"


# ---------- 限深回归钉子（2026-09-14 质量审查收尾）----------

def _deep_sdt_table_layers(n, tag):
    """递归构造 n 层「sdt(含段落) → 表格 → cell → 嵌套 sdt」交替嵌套 XML 片段
    （作为最外层 sdt 的 sdtContent 内容使用，命名空间由 _make_sdt_xml 根节点
    声明继承）。第 n 层含段落 f"{tag}{n}"，其表格 cell 内嵌第 n-1 层 sdt；
    n==0 时只剩叶子段落。深度每往内一层 +1，便于按文本断言各层编址情况。"""
    para = f'<w:p><w:r><w:t>{tag}{n}</w:t></w:r></w:p>'
    if n == 0:
        return para
    tbl = (
        '<w:tbl><w:tblPr><w:tblW w:w="0" w:type="auto"/></w:tblPr>'
        '<w:tblGrid><w:gridCol w:w="100"/></w:tblGrid>'
        '<w:tr><w:tc><w:tcPr/>'
        '<w:p/><w:sdt><w:sdtPr><w:alias w:val="d"/></w:sdtPr>'
        f'<w:sdtContent>{_deep_sdt_table_layers(n - 1, tag)}</w:sdtContent>'
        '</w:sdt>'
        '</w:tc></w:tr></w:tbl>'
    )
    return para + tbl


def test_sdt_table_alternate_nesting_depth_limit():
    """回归钉子：_walk_sdt 的 TXBX_DEPTH_LIMIT 限深是畸形 XML 防爆栈的关键防御，
    此前零测试覆盖（删掉 depth+1 全部用例照绿）。构造超过限深（8）层的
    「sdt→表格→cell→sdt」交替嵌套文档，断言：不抛 RecursionError；
    限深内各层正常编址（depth 0..8 → 共 9 个 sdt 层）；超限子树整体不产生 addr。"""
    from rag.svr.template_fill.docx_utils import iter_docx_paragraphs
    doc = Document()
    doc.add_paragraph("正文段落")
    # 10 层交替嵌套：最外层 body sdt depth=0（文本 层10），每往内一层 depth+1，
    # depth 9/10 的 层1/层0 超限应被跳过
    _add_body_sdt(doc, _deep_sdt_table_layers(10, "层"))
    doc.add_paragraph("尾部段落：____")
    buf = io.BytesIO()
    doc.save(buf)
    blob = buf.getvalue()

    items = iter_docx_paragraphs(blob)  # 不抛 RecursionError 即通过第一关
    by_text = {it["text"]: it["addr"] for it in items}
    # 限深内层正常编址：depth 0 最外层 sdt 直系段落；depth 8 为最后一个被编址层
    assert by_text["层10"] == "sdt:0:0"
    depth8_prefix = "sdt:0" + ":cell:0:0:0:sdt0" * 8
    assert by_text["层2"] == f"{depth8_prefix}:0"
    # 超限子树（depth 9/10）不产生任何 addr——限深砍掉的子树必须静默整棵跳过
    assert "层1" not in by_text
    assert "层0" not in by_text
    # 扁平序号连续、addr 全局唯一（限深截断不破坏编址完整性）
    assert [it["index"] for it in items] == list(range(len(items)))
    addrs = [it["addr"] for it in items]
    assert len(addrs) == len(set(addrs))


def test_nested_table_depth_accumulates_for_cell_sdt_limit():
    """嵌套表 ":t<j>" 递归必须 depth + 1（与 sdt/txbx 路径防护语义一致）：
    否则深层嵌套表内的 cell sdt 深度恒为浅层值，TXBX_DEPTH_LIMIT 对
    「表套表→cell→sdt」路径失效。断言：深层表套表本身不崩、各层段落照常编址
    （表格自身不限深）；cell sdt 随嵌套深度累计，depth 8 为最后编址层，
    depth 9+ 的 sdt 超限跳过。"""
    from rag.svr.template_fill.docx_utils import iter_docx_paragraphs
    from docx.oxml.ns import qn
    doc = Document()
    tbl = doc.add_table(rows=1, cols=1)
    cell = tbl.rows[0].cells[0]
    for i in range(10):  # 10 层表套表：第 i 层 cell 含段落+sdt，depth=i+1
        cell.paragraphs[0].text = f"表层{i}"
        cell._tc.append(_make_sdt_xml(_sdt_para(f"控件层{i}")))
        sub = cell.add_table(rows=1, cols=1)
        cell = sub.rows[0].cells[0]
    cell.paragraphs[0].text = "最内层段落"
    buf = io.BytesIO()
    doc.save(buf)

    items = iter_docx_paragraphs(buf.getvalue())  # 深层嵌套表不崩
    by_text = {it["text"]: it["addr"] for it in items}
    # 各层段落（表格自身不限深）全部照常编址
    assert by_text["表层0"] == "cell:0:0:0:0"
    assert "表层9" in by_text and "最内层段落" in by_text
    # cell sdt 深度随嵌套表累计：level0 sdt depth=1，level7 sdt depth=8（限深边界内）
    assert by_text["控件层0"] == "cell:0:0:0:sdt0:0"
    lvl7_prefix = "cell:0" + ":0:0:t0" * 7
    assert by_text["控件层7"] == f"{lvl7_prefix}:0:0:sdt0:0"
    # level8/level9 sdt depth=9/10，超限跳过（若嵌套表未递增 depth，此处会误编址）
    assert "控件层8" not in by_text
    assert "控件层9" not in by_text


# ---------- Task 2（2026-09-13）：页眉/页脚编址（hdr:/ftr: 前缀） ----------


def _make_docx_with_hf():
    """两节文档：第 1 节有显式页眉（含表格）+页脚；第 2 节 header 默认 linked。"""
    from docx.enum.section import WD_SECTION_START
    from docx.shared import Emu
    doc = Document()
    doc.add_paragraph("正文：________")
    h0 = doc.sections[0].header
    h0.is_linked_to_previous = False
    h0.paragraphs[0].text = "投标人（盖章）：____________"
    tbl = h0.add_table(rows=1, cols=1, width=Emu(4000000))
    tbl.rows[0].cells[0].paragraphs[0].text = "页眉表格：________"
    f0 = doc.sections[0].footer
    f0.is_linked_to_previous = False
    f0.paragraphs[0].text = "日期：____年____月____日"
    doc.add_section(WD_SECTION_START.NEW_PAGE)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_corrupt_hf_part_skips_without_killing_doc(monkeypatch):
    """单个 header/footer part 编址抛异常：跳过该 part（告警日志），不拖垮整文档。"""
    import rag.svr.template_fill.docx_utils as du

    def boom(hf_el, prefix, doc, walk_paragraph, walk_table):
        raise ValueError("corrupt part")

    monkeypatch.setattr(du, "_walk_hf_blocks", boom)
    items = du.iter_docx_paragraphs(_make_docx_with_hf())
    addrs = [it["addr"] for it in items]
    assert "para:0" in addrs  # 正文编址不受影响
    assert not any(a.startswith(("hdr:", "ftr:")) for a in addrs)  # 异常 part 整体跳过


def test_header_footer_addressed_and_linked_deduped():
    from rag.svr.template_fill.docx_utils import extract_docx_candidates, iter_docx_paragraphs
    items = iter_docx_paragraphs(_make_docx_with_hf())
    by_addr = {it["addr"]: it["text"] for it in items}
    assert "投标人（盖章）" in by_addr.get("hdr:0:0", "")
    assert "页眉表格" in by_addr.get("hdr:0:cell:0:0:0:0", "")
    assert "____年____月____日" in by_addr.get("ftr:0:0", "")
    # 第 2 节 header linked 到第 1 节（同一 part）→ 不重复编址
    assert sum(1 for a in by_addr if a.startswith("hdr:")) == 2  # 段落+表格段各一
    # 页眉页脚段落进候选
    cands = extract_docx_candidates(_make_docx_with_hf())
    addrs = {c["addr"] for c in cands}
    assert "hdr:0:0" in addrs and "ftr:0:0" in addrs


def test_hf_zero_legacy_counter_consumption():
    """对抗（红线）：页眉页脚段落零消耗存量 para:/cell: 计数器——同一 body 内容
    加页眉页脚前后，存量前缀（para:/cell:/sdt:）addr 序列逐字节一致。"""
    from rag.svr.template_fill.docx_utils import iter_docx_paragraphs

    def _fill_body(doc):
        doc.add_paragraph("第一段：____")
        tbl = doc.add_table(rows=1, cols=1)
        tbl.rows[0].cells[0].paragraphs[0].text = "单元格：____"
        doc.add_paragraph("第二段")

    doc_a = Document()
    _fill_body(doc_a)
    doc_b = Document()
    _fill_body(doc_b)
    h = doc_b.sections[0].header
    h.is_linked_to_previous = False
    h.paragraphs[0].text = "页眉：____________"
    f = doc_b.sections[0].footer
    f.is_linked_to_previous = False
    f.paragraphs[0].text = "日期：____年____月____日"
    buf_a, buf_b = io.BytesIO(), io.BytesIO()
    doc_a.save(buf_a)
    doc_b.save(buf_b)

    def legacy_addrs(blob):
        return [it["addr"] for it in iter_docx_paragraphs(blob)
                if it["addr"].split(":", 1)[0] in ("para", "cell", "sdt")]

    addrs_a = legacy_addrs(buf_a.getvalue())
    addrs_b = legacy_addrs(buf_b.getvalue())
    # cell 段落本就消耗 para: 计数（存量语义），表格后段落为 para:2；
    # 红线断言：有无页眉页脚，存量序列逐字节一致
    assert addrs_a == addrs_b == ["para:0", "cell:0:0:0:0", "para:2"]
    # 页眉页脚段落有自己的前缀且扁平 index 严格递增
    items_b = iter_docx_paragraphs(buf_b.getvalue())
    assert [it["index"] for it in items_b] == list(range(len(items_b)))
    assert any(it["addr"] == "hdr:0:0" for it in items_b)
    assert any(it["addr"] == "ftr:0:0" for it in items_b)


def test_hf_first_even_variants_addressed_distinctly():
    """对抗：同节 default/first/even 三类页眉页脚同时 unlinked——各 part 独立编址
    且 addr 全局唯一（若三类共用 hdr:<sec>:<pi> 前缀会撞号，validate/render 按
    addr 反查段落将错位到错误 part）。"""
    from rag.svr.template_fill.docx_utils import iter_docx_paragraphs
    doc = Document()
    doc.add_paragraph("正文：________")
    sec = doc.sections[0]
    sec.different_first_page_header_footer = True
    h = sec.header
    h.is_linked_to_previous = False
    h.paragraphs[0].text = "默认页眉：____"
    hfp = sec.first_page_header
    hfp.is_linked_to_previous = False
    hfp.paragraphs[0].text = "首页页眉：____"
    he = sec.even_page_header
    he.is_linked_to_previous = False
    he.paragraphs[0].text = "偶数页页眉：____"
    f = sec.footer
    f.is_linked_to_previous = False
    f.paragraphs[0].text = "默认页脚：____"
    ffp = sec.first_page_footer
    ffp.is_linked_to_previous = False
    ffp.paragraphs[0].text = "首页页脚：____"
    fe = sec.even_page_footer
    fe.is_linked_to_previous = False
    fe.paragraphs[0].text = "偶数页页脚：____"
    buf = io.BytesIO()
    doc.save(buf)

    items = iter_docx_paragraphs(buf.getvalue())
    by_addr = {it["addr"]: it["text"] for it in items}
    assert by_addr["hdr:0:0"] == "默认页眉：____"
    assert by_addr["hdr:0:first:0"] == "首页页眉：____"
    assert by_addr["hdr:0:even:0"] == "偶数页页眉：____"
    assert by_addr["ftr:0:0"] == "默认页脚：____"
    assert by_addr["ftr:0:first:0"] == "首页页脚：____"
    assert by_addr["ftr:0:even:0"] == "偶数页页脚：____"
    # addr 全局唯一（撞号即失守）
    addrs = [it["addr"] for it in items]
    assert len(addrs) == len(set(addrs))


def test_hf_shared_part_addressed_once():
    """对抗：两节 headerReference 指向同一 part（手工改 rel 模拟畸形/拷贝文档）——
    partname 去重保证只编址一次，不产生 hdr:1 重复 addr。"""
    from docx.enum.section import WD_SECTION_START
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml.ns import qn

    from rag.svr.template_fill.docx_utils import iter_docx_paragraphs
    doc = Document()
    doc.add_paragraph("第一节正文")
    h0 = doc.sections[0].header
    h0.is_linked_to_previous = False
    h0.paragraphs[0].text = "共享页眉：____"
    doc.add_section(WD_SECTION_START.NEW_PAGE)
    doc.add_paragraph("第二节正文")
    h1 = doc.sections[1].header
    h1.is_linked_to_previous = False  # 先生成自己的 part
    # add_section 会把旧 sentinel sectPr 留给新节：旧 h0 代理绑定的元素已属第 2 节，
    # 必须重新取第 1 节 header 才能拿到含「共享页眉」内容的 part
    part0 = doc.sections[0].header.part
    r_id = doc.part.relate_to(part0, RT.HEADER)
    for ref in doc.sections[1]._sectPr.findall(qn("w:headerReference")):
        if ref.get(qn("w:type")) == "default":
            ref.set(qn("r:id"), r_id)
    buf = io.BytesIO()
    doc.save(buf)

    items = iter_docx_paragraphs(buf.getvalue())
    hdr_addrs = [it["addr"] for it in items if it["addr"].startswith("hdr:")]
    assert hdr_addrs == ["hdr:0:0"]  # 共享 part 只编址一次
    assert any(it["text"] == "共享页眉：____" and it["addr"] == "hdr:0:0" for it in items)


def test_textbox_in_header_paragraph_addressed():
    """对抗：页眉段落内嵌 mc:AlternateContent 文本框 → hdr:0:0:tx0:0；
    Choice+Fallback 双份存储只编址一次。"""
    from docx.oxml import parse_xml

    from rag.svr.template_fill.docx_utils import iter_docx_paragraphs
    doc = Document()
    doc.add_paragraph("正文：________")
    h = doc.sections[0].header
    h.is_linked_to_previous = False
    p = h.paragraphs[0]
    p.text = "页眉容器段落"
    p._p.append(parse_xml(
        f'<w:r {_MC_NS}>'
        '<mc:AlternateContent>'
        f'<mc:Choice Requires="wps"><w:txbxContent>'
        f'<w:p><w:r><w:t>页眉文本框：＿＿＿＿＿＿</w:t></w:r></w:p>'
        '</w:txbxContent></mc:Choice>'
        f'<mc:Fallback><w:txbxContent>'
        f'<w:p><w:r><w:t>页眉文本框：＿＿＿＿＿＿</w:t></w:r></w:p>'
        '</w:txbxContent></mc:Fallback>'
        '</mc:AlternateContent>'
        '</w:r>'
    ))
    buf = io.BytesIO()
    doc.save(buf)

    items = iter_docx_paragraphs(buf.getvalue())
    hits = [it for it in items if "页眉文本框" in it["text"]]
    assert len(hits) == 1
    assert hits[0]["addr"] == "hdr:0:0:tx0:0"


# ---------- occ 同段同形留白（第 N 次出现定位） ----------


def test_apply_same_shape_blanks_occ_positional():
    """同段三处同形留白 occ=1/2/3 各归其位：降序应用后互不挤位。"""
    blob = _make_docx(["甲方：＿＿＿ 乙方：＿＿＿ 丙方：＿＿＿"])
    from rag.svr.template_fill.docx_utils import apply_docx_placeholders, iter_docx_paragraphs
    out = apply_docx_placeholders(blob, [
        {"addr": "para:0", "anchor": "＿＿＿", "key": "party_a", "occ": 1},
        {"addr": "para:0", "anchor": "＿＿＿", "key": "party_b", "occ": 2},
        {"addr": "para:0", "anchor": "＿＿＿", "key": "party_c", "occ": 3},
    ])
    assert iter_docx_paragraphs(out)[0]["text"] == "甲方：{{party_a}} 乙方：{{party_b}} 丙方：{{party_c}}"


def test_apply_occ_beyond_count_is_noop():
    """occ 超过段内出现次数：no-op，原文保留。"""
    blob = _make_docx(["姓名：＿＿ 电话：＿＿"])
    from rag.svr.template_fill.docx_utils import apply_docx_placeholders, iter_docx_paragraphs
    out = apply_docx_placeholders(blob, [
        {"addr": "para:0", "anchor": "＿＿", "key": "x", "occ": 3},
    ])
    assert "＿＿" in iter_docx_paragraphs(out)[0]["text"]
    assert "{{x}}" not in iter_docx_paragraphs(out)[0]["text"]


def test_apply_legacy_no_occ_replaces_all_in_run():
    """occ 缺省 = 存量 replace-all 语义，行为与加固前完全一致。"""
    blob = _make_docx(["姓名：＿＿ 电话：＿＿"])
    from rag.svr.template_fill.docx_utils import apply_docx_placeholders, iter_docx_paragraphs
    out = apply_docx_placeholders(blob, [
        {"addr": "para:0", "anchor": "＿＿", "key": "x"},
    ])
    assert iter_docx_paragraphs(out)[0]["text"] == "姓名：{{x}} 电话：{{x}}"


def test_apply_occ_across_runs_picks_nth():
    """occ=2 目标出现在第 3 个 run 内（第 1 个 run 里的出现不计入）：只动该出现。"""
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("AA")
    p.add_run("BB")
    p.add_run("AA")
    buf = io.BytesIO()
    doc.save(buf)
    from rag.svr.template_fill.docx_utils import apply_docx_placeholders, iter_docx_paragraphs
    out = apply_docx_placeholders(buf.getvalue(), [
        {"addr": "para:0", "anchor": "AA", "key": "second", "occ": 2},
    ])
    assert iter_docx_paragraphs(out)[0]["text"] == "AABB{{second}}"


def test_apply_occ_spanning_runs_falls_to_cross_run():
    """occ 目标出现跨 run（单 run 计数找不到）→ 落到跨 run 路径，只重写覆盖区间。"""
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("A_")
    p.add_run("_B")
    buf = io.BytesIO()
    doc.save(buf)
    from rag.svr.template_fill.docx_utils import apply_docx_placeholders, iter_docx_paragraphs
    out = apply_docx_placeholders(buf.getvalue(), [
        {"addr": "para:0", "anchor": "__", "key": "span", "occ": 1},
    ])
    assert iter_docx_paragraphs(out)[0]["text"] == "A{{span}}B"


def test_apply_occ_non_overlapping_semantics():
    """对抗：非重叠语义。anchor="aaa"、文本 6 个 a，非重叠出现仅 2 次（下标 0、3），
    occ=2 必须命中下标 3——按重叠推进（find(start+1)）会错命中下标 1。"""
    blob = _make_docx(["aaaaaa"])
    from rag.svr.template_fill.docx_utils import apply_docx_placeholders, iter_docx_paragraphs
    out = apply_docx_placeholders(blob, [
        {"addr": "para:0", "anchor": "aaa", "key": "k", "occ": 2},
    ])
    assert iter_docx_paragraphs(out)[0]["text"] == "aaa{{k}}"


def test_apply_occ_skips_spanning_occurrence_when_counting():
    """对抗：更早的出现跨 run（不计入任何单 run 的 count）时，occ 仍须按段落级
    拼接文本的出现序号定位——p.text="aaa" 的第 1 次非重叠出现是跨界的下标 0-1，
    若按单 run 局部计数会错命中 run1 内部的 "aa"（段落级视角下它是被覆盖的重叠位）。"""
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("a")
    p.add_run("aa")
    buf = io.BytesIO()
    doc.save(buf)
    from rag.svr.template_fill.docx_utils import apply_docx_placeholders, iter_docx_paragraphs
    out = apply_docx_placeholders(buf.getvalue(), [
        {"addr": "para:0", "anchor": "aa", "key": "k", "occ": 1},
    ])
    assert iter_docx_paragraphs(out)[0]["text"] == "{{k}}a"


def test_apply_occ_second_in_hyperlink_paragraph_via_concat():
    """对抗：occ 与超链接段落组合。两次出现：第 1 次跨 run0/run1 边界、第 2 次在
    run1 内，段落尾部带 w:hyperlink。occ=2 必须精确命中第 2 次、第 1 次原样保留，
    且超链接文本不被复制进正文 run（拼接路径只作用于 p.runs 文本）。"""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document()
    p = doc.add_paragraph("编号：___")
    p.add_run("__日期：____附注")
    hl = OxmlElement("w:hyperlink")
    hl.set(qn("r:id"), "rIdLink")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "官网"
    r.append(t)
    hl.append(r)
    p._p.append(hl)
    buf = io.BytesIO()
    doc.save(buf)
    from rag.svr.template_fill.docx_utils import apply_docx_placeholders, iter_docx_paragraphs
    # 前置确认：p.text 含超链接文本，且 "____" 非重叠出现 2 次（第 1 次跨界）
    assert next(it["text"] for it in iter_docx_paragraphs(buf.getvalue())) \
        == "编号：_____日期：____附注官网"

    out = apply_docx_placeholders(buf.getvalue(), [
        {"addr": "para:0", "anchor": "____", "key": "k", "occ": 2},
    ])
    text = iter_docx_paragraphs(out)[0]["text"]
    assert text == "编号：_____日期：{{k}}附注官网"
    # 第 1 次出现原样保留；超链接文本不被复制
    assert text.count("____") == 1
    assert text.count("官网") == 1


def test_apply_occ_dirty_value_skips_entry():
    """对抗：occ 为非法值（0/负数/字符串/布尔）→ 跳过该条，绝不退化成 replace-all
    误伤其他出现位（occ=None 才是存量 replace-all 语义）。"""
    blob = _make_docx(["姓名：＿＿ 电话：＿＿"])
    from rag.svr.template_fill.docx_utils import apply_docx_placeholders, iter_docx_paragraphs
    out = apply_docx_placeholders(blob, [
        {"addr": "para:0", "anchor": "＿＿", "key": "zero", "occ": 0},
        {"addr": "para:0", "anchor": "＿＿", "key": "neg", "occ": -1},
        {"addr": "para:0", "anchor": "＿＿", "key": "str", "occ": "1"},
        {"addr": "para:0", "anchor": "＿＿", "key": "bool", "occ": True},
    ])
    text = iter_docx_paragraphs(out)[0]["text"]
    assert text == "姓名：＿＿ 电话：＿＿"
    assert "{{" not in text


def test_apply_occ_and_none_mixed_order_same_paragraph():
    """对抗：同段 occ 条目与 occ=None（replace-all）条目混排。降序稳定排序必须
    让 occ 先行、None 兜底吞剩余——None 先行会吞掉 occ 的目标出现，导致 occ=2
    找不到第 2 处而 no-op。且最终文本与提交顺序无关（排序键决定应用顺序）。"""
    from rag.svr.template_fill.docx_utils import apply_docx_placeholders, iter_docx_paragraphs
    expected = "姓名：{{name}} 电话：{{phone}}"
    # 提交顺序 A：None 条目在前、occ=2 在后（若按原序应用会错）
    blob = _make_docx(["姓名：＿＿ 电话：＿＿"])
    out = apply_docx_placeholders(blob, [
        {"addr": "para:0", "anchor": "＿＿", "key": "name"},
        {"addr": "para:0", "anchor": "＿＿", "key": "phone", "occ": 2},
    ])
    assert iter_docx_paragraphs(out)[0]["text"] == expected
    # 提交顺序 B：两条目原序对调，结果必须逐字一致
    blob = _make_docx(["姓名：＿＿ 电话：＿＿"])
    out = apply_docx_placeholders(blob, [
        {"addr": "para:0", "anchor": "＿＿", "key": "phone", "occ": 2},
        {"addr": "para:0", "anchor": "＿＿", "key": "name"},
    ])
    assert iter_docx_paragraphs(out)[0]["text"] == expected


def test_apply_occ_and_none_mixed_overlap_different_anchors():
    """对抗：不同 anchor 且文本交叠（None 条目 anchor="AB" 覆盖 occ 条目
    anchor="B_" 的起点）。occ 先行改写交叠区后 "AB" 不再存在 → None no-op；
    若排序被误改（None 先行）则会吞掉交叠区使 occ no-op，两者结果不同。"""
    from rag.svr.template_fill.docx_utils import apply_docx_placeholders, iter_docx_paragraphs
    expected = "A{{occ_key}}C"
    # 提交顺序 A：None 在前、occ 在后
    blob = _make_docx(["AB_C"])
    out = apply_docx_placeholders(blob, [
        {"addr": "para:0", "anchor": "AB", "key": "none_key"},
        {"addr": "para:0", "anchor": "B_", "key": "occ_key", "occ": 1},
    ])
    assert iter_docx_paragraphs(out)[0]["text"] == expected
    # 提交顺序 B：原序对调，结果必须逐字一致
    blob = _make_docx(["AB_C"])
    out = apply_docx_placeholders(blob, [
        {"addr": "para:0", "anchor": "B_", "key": "occ_key", "occ": 1},
        {"addr": "para:0", "anchor": "AB", "key": "none_key"},
    ])
    assert iter_docx_paragraphs(out)[0]["text"] == expected


# ---------- 替换层 no-op 可观测性（设计 §8：no-op + 告警，旧数据回放） ----------


def _collect_warnings(caplog):
    import logging
    caplog.set_level(logging.WARNING, logger="rag.svr.template_fill.docx_utils")
    return caplog


def test_apply_occ_beyond_count_warns(caplog):
    """occ 超界 no-op 必须留告警日志（含 addr/key/occ），不得静默——否则旧数据
    回放/锚错位时「填写点没生效」无从定位。"""
    _collect_warnings(caplog)
    blob = _make_docx(["姓名：＿＿ 电话：＿＿"])
    from rag.svr.template_fill.docx_utils import apply_docx_placeholders
    apply_docx_placeholders(blob, [
        {"addr": "para:0", "anchor": "＿＿", "key": "x", "occ": 3},
    ])
    msgs = [r.getMessage() for r in caplog.records
            if r.levelno == logging.WARNING]
    assert any("no-op" in m and "para:0" in m and "x" in m and "occ=3" in m
               for m in msgs), msgs


def test_apply_anchor_absent_in_paragraph_warns(caplog):
    """anchor 不在 addr 指向的段落文本（模板版本漂移/LLM 脏锚）同样值得告警，
    不得把「anchor 不在文本」漏成静默。"""
    _collect_warnings(caplog)
    blob = _make_docx(["项目名称：____________"])
    from rag.svr.template_fill.docx_utils import apply_docx_placeholders
    apply_docx_placeholders(blob, [
        {"addr": "para:0", "anchor": "不存在的锚", "key": "x"},
    ])
    msgs = [r.getMessage() for r in caplog.records
            if r.levelno == logging.WARNING]
    assert any("no-op" in m and "para:0" in m and "不存在的锚" in m
               for m in msgs), msgs


def test_apply_addr_missing_warns_once_not_duplicated(caplog):
    """回归：addr 悬空既有告警保留，且每条目只告警一次（替换层告警不得与
    addr 悬空告警叠加产生重复行）。"""
    _collect_warnings(caplog)
    blob = _make_docx(["正文"])
    from rag.svr.template_fill.docx_utils import apply_docx_placeholders
    apply_docx_placeholders(blob, [
        {"addr": "para:99", "anchor": "＿＿", "key": "x"},
    ])
    msgs = [r.getMessage() for r in caplog.records
            if r.levelno == logging.WARNING]
    assert any("not found" in m and "para:99" in m for m in msgs), msgs
    assert sum(1 for m in msgs if "para:99" in m) == 1, msgs


def test_apply_success_path_no_warning(caplog):
    """正常替换路径零告警——告警只属于异常/悬空情形，避免正常回放刷屏。"""
    _collect_warnings(caplog)
    blob = _make_docx(["姓名：＿＿ 电话：＿＿"])
    from rag.svr.template_fill.docx_utils import apply_docx_placeholders
    apply_docx_placeholders(blob, [
        {"addr": "para:0", "anchor": "＿＿", "key": "x", "occ": 2},
    ])
    assert not [r for r in caplog.records if r.levelno == logging.WARNING]


# ---------- occ 识别链：merge 分配 + validate 校验 ----------


def _mk_fill_item(key, anchor, addr="para:0"):
    return {"key": key, "name": key, "description": "", "retrieval_query": "",
            "fill_mode": "llm", "required": True, "addr": addr,
            "anchor": anchor, "line": 0, "top_k": 6}


def test_merge_llm_same_anchor_duplicates_get_occ_and_low_confidence():
    """LLM 组同段同形留白多项：occ 1..n 全部保留，occ≥2 打低置信（同形歧义需人工核对）。"""
    from rag.svr.template_fill.detector import _merge_detection
    merged = _merge_detection([], [_mk_fill_item("a", "＿＿"), _mk_fill_item("b", "＿＿")])
    assert [it.get("occ") for it in merged] == [1, 2]
    assert merged[0].get("low_confidence") in (None, False)
    assert merged[1]["low_confidence"] is True


def test_merge_explicit_duplicates_get_occ_without_low_confidence():
    """手动直通组同 (addr, anchor) 多份（{{key}} 出现多次）：occ 1..n 确定性保留，不打低置信。"""
    from rag.svr.template_fill.detector import _merge_detection, extract_explicit_placeholders
    cands = [{"index": 0, "text": "{{name}} {{name}}", "addr": "para:0"}]
    explicit = extract_explicit_placeholders(cands)
    merged = _merge_detection(explicit, [])
    assert [it.get("occ") for it in merged] == [1, 2]
    assert all(not it.get("low_confidence") for it in merged)


def test_merge_llm_echo_of_explicit_anchor_still_dropped():
    """LLM 回显手动占位符同 anchor：手动优先丢弃，且不分配 occ（防 occ=2 挤爆校验）。"""
    from rag.svr.template_fill.detector import _merge_detection, extract_explicit_placeholders
    cands = [{"index": 0, "text": "编号：{{code}}", "addr": "para:0"}]
    explicit = extract_explicit_placeholders(cands)
    llm = [_mk_fill_item("code2", "{{code}}")]
    merged = _merge_detection(explicit, llm)
    assert [it["key"] for it in merged] == ["code"]
    assert all(it.get("occ") is None for it in merged)


def test_merge_llm_echo_group_does_not_extend_explicit_occ_sequence():
    """对抗：LLM 组与 explicit 组各自独立计数——explicit 同形 1 项 + LLM 回显同位 2 项
    全部丢弃，occ 序列不得串组（否则 explicit 项会被虚标 occ）。"""
    from rag.svr.template_fill.detector import _merge_detection, extract_explicit_placeholders
    cands = [{"index": 0, "text": "编号：{{code}}", "addr": "para:0"}]
    explicit = extract_explicit_placeholders(cands)
    llm = [_mk_fill_item("echo1", "{{code}}"), _mk_fill_item("echo2", "{{code}}")]
    merged = _merge_detection(explicit, llm)
    assert [it["key"] for it in merged] == ["code"]
    assert "occ" not in merged[0]


def test_merge_shrink_rollback_strips_occ():
    """收缩撞车回退到 _orig_anchor 时 occ 作废（occ 属于收缩后 anchor 的组）。"""
    from rag.svr.template_fill.detector import _merge_detection
    a = _mk_fill_item("a", "＿＿")
    a["_orig_anchor"] = "甲方名称："
    b = _mk_fill_item("b", "＿＿")
    b["_orig_anchor"] = "乙方名称："
    merged = _merge_detection([], [a, b])
    by_key = {it["key"]: it for it in merged}
    assert by_key["b"]["anchor"] == "乙方名称："  # 回退保留
    assert "occ" not in by_key["b"]  # 回退项 occ 作废


def test_merge_three_same_blanks_llm_occ_123():
    """对抗：同段三个同形留白 LLM 组 → occ=1/2/3 全保留，第 2、3 项均打低置信。"""
    from rag.svr.template_fill.detector import _merge_detection
    llm = [_mk_fill_item(k, "＿＿") for k in ("a", "b", "c")]
    merged = _merge_detection([], llm)
    assert [it["occ"] for it in merged] == [1, 2, 3]
    assert [it["key"] for it in merged] == ["a", "b", "c"]
    assert not merged[0].get("low_confidence")  # occ=1 不打低置信（字段缺省/False 均视为未打）
    assert merged[1]["low_confidence"] is True
    assert merged[2]["low_confidence"] is True


def test_validate_occ_within_actual_count_passes():
    from rag.svr.template_fill.detector import validate_placeholders
    cands = [{"index": 0, "text": "甲：＿＿ 乙：＿＿", "addr": "para:0"}]
    ok, msg = validate_placeholders([_mk_fill_item("a", "＿＿") | {"occ": 2}], cands)
    assert ok, msg


def test_validate_occ_beyond_actual_count_rejected():
    from rag.svr.template_fill.detector import validate_placeholders
    cands = [{"index": 0, "text": "甲：＿＿", "addr": "para:0"}]
    ok, msg = validate_placeholders([_mk_fill_item("a", "＿＿") | {"occ": 2}], cands)
    assert not ok and "出现次数不足" in msg


def test_validate_manual_row_occ_gt_1_rejected():
    from rag.svr.template_fill.detector import validate_placeholders
    cands = [{"index": 0, "text": "甲：＿＿", "addr": "para:0"}]
    row = _mk_fill_item("a", "＿＿")
    row["addr"] = ""  # 手动行：靠 anchor 反查
    ok, msg = validate_placeholders([row | {"occ": 2}], cands)
    assert not ok


def test_validate_legacy_no_occ_still_passes():
    from rag.svr.template_fill.detector import validate_placeholders
    cands = [{"index": 0, "text": "甲：＿＿ 乙：＿＿", "addr": "para:0"}]
    ok, msg = validate_placeholders([_mk_fill_item("a", "＿＿")], cands)
    assert ok, msg


@pytest.mark.parametrize("bad_occ", [0, -1, True, "2", 1.5])
def test_validate_occ_dirty_values_rejected(bad_occ):
    """对抗（脏值）：occ=0/负数/bool（True==1 的陷阱）/字符串/浮点一律拒绝。"""
    from rag.svr.template_fill.detector import validate_placeholders
    cands = [{"index": 0, "text": "甲：＿＿ 乙：＿＿", "addr": "para:0"}]
    ok, msg = validate_placeholders([_mk_fill_item("a", "＿＿") | {"occ": bad_occ}], cands)
    assert not ok, f"occ={bad_occ!r} 应被拒绝"
    assert "填写位" in msg


def test_validate_occ_count_is_non_overlapping():
    """对抗：出现次数按非重叠语义——anchor="＿＿"（2字符）在"＿＿＿＿"中即 2 次出现：
    occ=2 放行、occ=3 拒绝（重叠计数会误判 3 次）。"""
    from rag.svr.template_fill.detector import validate_placeholders
    cands = [{"index": 0, "text": "空白：＿＿＿＿", "addr": "para:0"}]
    ok, _ = validate_placeholders([_mk_fill_item("a", "＿＿") | {"occ": 2}], cands)
    assert ok
    ok, msg = validate_placeholders([_mk_fill_item("a", "＿＿") | {"occ": 3}], cands)
    assert not ok and "出现次数不足" in msg


def test_validate_manual_row_occ_true_rejected():
    """对抗（手动行 bool 陷阱）：True == 1 在 == 比较下会漏过 occ=1 校验，须显式拒绝 bool。"""
    from rag.svr.template_fill.detector import validate_placeholders
    cands = [{"index": 0, "text": "甲：＿＿", "addr": "para:0"}]
    row = _mk_fill_item("a", "＿＿")
    row["addr"] = ""
    ok, msg = validate_placeholders([row | {"occ": True}], cands)
    assert not ok and "单填写位" in msg


# ---------- occ 识别链质量加固：xlsx 禁用预分配 + 组内文本偏移定序 + 钉子 ----------


def test_merge_xlsx_no_preassign_keeps_dedupe_drop_semantics():
    """xlsx（preassign_occ=False）：同 (addr, anchor) 重复项回到旧「去重丢弃、
    单 key replace-all」语义，不产生 occ 字段——xlsx 渲染层 replace-all 不识别
    occ，预分配会让同格重复占位符串值覆盖（第一个字段被吞）。"""
    from rag.svr.template_fill.detector import _merge_detection
    merged = _merge_detection(
        [], [_mk_fill_item("a", "＿＿"), _mk_fill_item("b", "＿＿")], preassign_occ=False)
    assert [it["key"] for it in merged] == ["a"]  # 旧去重语义：后者丢弃
    assert all("occ" not in it for it in merged)


def test_merge_occ_assigned_by_text_offset_not_llm_order():
    """LLM 组乱序输出（乙的项在前、甲的项在后，同段 anchor 同形）：occ 必须按
    parse 阶段记录的文本偏移 _anchor_pos 升序分配（甲=1、乙=2），而非 LLM 输出序；
    内部字段 _anchor_pos 不得泄漏进 merge 产物。"""
    from rag.svr.template_fill.detector import _merge_detection
    b = _mk_fill_item("party_b", "＿＿")
    b["_anchor_pos"] = 7  # parse 阶段记录：cand["text"].find(anchor)
    a = _mk_fill_item("party_a", "＿＿")
    a["_anchor_pos"] = 0
    merged = _merge_detection([], [b, a])
    by_key = {it["key"]: it for it in merged}
    assert by_key["party_a"]["occ"] == 1  # 文本在前者占 occ=1
    assert by_key["party_b"]["occ"] == 2
    assert all("_anchor_pos" not in it for it in merged)


def test_merge_middle_item_rollback_leaves_occ_gap_1_3():
    """钉子：三项同形组中间项收缩撞车回退后，occ 序列钉在 [1, 3] 跳号且三项全保留——
    渲染层按第 1、3 次出现落位，中间项以原 anchor 独立保留（防跳号被误压缩成 [1,2]）。"""
    from rag.svr.template_fill.detector import _merge_detection
    a = _mk_fill_item("a", "＿＿")
    b = _mk_fill_item("b", "＿＿")
    b["_orig_anchor"] = "乙方名称："
    c = _mk_fill_item("c", "＿＿")
    merged = _merge_detection([], [a, b, c])
    assert len(merged) == 3
    by_key = {it["key"]: it for it in merged}
    assert by_key["a"]["occ"] == 1
    assert by_key["b"]["anchor"] == "乙方名称：" and "occ" not in by_key["b"]  # 回退项 occ 作废
    assert by_key["c"]["occ"] == 3


def test_validate_manual_row_occ_1_passes():
    """钉子（正向）：手动行显式 occ=1 与缺省等价，放行且反查回填 addr。"""
    from rag.svr.template_fill.detector import validate_placeholders
    cands = [{"index": 0, "text": "甲：＿＿", "addr": "para:0"}]
    row = _mk_fill_item("a", "＿＿")
    row["addr"] = ""
    item = row | {"occ": 1}
    ok, msg = validate_placeholders([item], cands)
    assert ok, msg
    assert item["addr"] == "para:0"  # 反查回填


# ---------- occ 预分配溢出封顶：candidates 传入时超界项丢弃而非判死整次识别 ----------
# 真实事故（2026-09-14 电子招标投标示范文本）：LLM 对同段同形留白输出 2 项建议，
# 预分配 occ=2 但 anchor 实际仅出现 1 次 → validate 判死整次识别（全部建议丢弃）。


def test_merge_overflow_capped_by_actual_count():
    """对抗（真实事故复刻）：同组 3 项但 anchor 实际仅 1 次出现——只保留第 1 项
    occ=1，溢出 2 项丢弃；整组结果过 validate 终审（修复前整次识别 failed）。"""
    from rag.svr.template_fill.detector import _merge_detection, validate_placeholders
    cands = [{"index": 0, "text": "标段划分：＿＿", "addr": "para:5"}]
    llm = [_mk_fill_item(k, "＿＿", addr="para:5") for k in ("a", "station_range", "c")]
    merged = _merge_detection([], llm, candidates=cands)
    assert [it["key"] for it in merged] == ["a"]  # 溢出项丢弃
    assert merged[0]["occ"] == 1
    ok, msg = validate_placeholders(merged, cands)
    assert ok, msg


def test_merge_overflow_cap_is_non_overlapping_count():
    """对抗：封顶口径与 validate 终审一致（text.count 非重叠语义）——anchor="＿＿"
    在"＿＿＿＿"中算 2 次，组 3 项时保 2 丢 1，保留项过终审。"""
    from rag.svr.template_fill.detector import _merge_detection, validate_placeholders
    cands = [{"index": 0, "text": "空白：＿＿＿＿", "addr": "para:0"}]
    llm = [_mk_fill_item(k, "＿＿") for k in ("a", "b", "c")]
    merged = _merge_detection([], llm, candidates=cands)
    assert [it["occ"] for it in merged] == [1, 2]
    assert [it["key"] for it in merged] == ["a", "b"]
    ok, msg = validate_placeholders(merged, cands)
    assert ok, msg


def test_merge_overflow_no_false_positive_when_group_matches_actual():
    """回归：组大小 == 实际出现次数时不误丢——同段 2 个真实留白 2 项建议全部
    保留 occ=1/2，occ≥2 仍打低置信（原语义不变）。"""
    from rag.svr.template_fill.detector import _merge_detection
    cands = [{"index": 0, "text": "甲：＿＿ 乙：＿＿", "addr": "para:0"}]
    llm = [_mk_fill_item(k, "＿＿") for k in ("a", "b")]
    merged = _merge_detection([], llm, candidates=cands)
    assert [it["occ"] for it in merged] == [1, 2]
    assert not merged[0].get("low_confidence")
    assert merged[1]["low_confidence"] is True


def test_merge_overflow_per_addr_isolated():
    """对抗：封顶按 (addr, anchor) 组独立核对——para:0 超 1 项丢弃不得波及
    para:1 的同形组（该组项数未超实际次数，全保留）。"""
    from rag.svr.template_fill.detector import _merge_detection
    cands = [
        {"index": 0, "text": "甲：＿＿", "addr": "para:0"},
        {"index": 1, "text": "乙：＿＿ 丙：＿＿", "addr": "para:1"},
    ]
    llm = [
        _mk_fill_item("a", "＿＿", addr="para:0"),
        _mk_fill_item("bad", "＿＿", addr="para:0"),  # para:0 仅 1 处 → 溢出丢弃
        _mk_fill_item("b", "＿＿", addr="para:1"),
        _mk_fill_item("c", "＿＿", addr="para:1"),
    ]
    merged = _merge_detection([], llm, candidates=cands)
    assert [it["key"] for it in merged] == ["a", "b", "c"]
    assert [it.get("occ") for it in merged] == [1, 1, 2]


def test_merge_overflow_unknown_addr_dropped_conservatively():
    """对抗：直调方传入不存在的 addr（cap 查不到候选文本，count=0）——组内全部
    项保守丢弃（cap=0 时 n=1 也溢出），不产生任何超界 occ 流向 validate。
    生产链路 addr 由 parse 阶段从真实候选分配、不会幻觉，此为直调防御。"""
    from rag.svr.template_fill.detector import _merge_detection
    cands = [{"index": 0, "text": "甲：＿＿", "addr": "para:0"}]
    llm = [
        _mk_fill_item("a", "＿＿", addr="para:0"),
        _mk_fill_item("ghost", "＿＿", addr="para:999"),  # 幻觉 addr
        _mk_fill_item("ghost2", "＿＿", addr="para:999"),
    ]
    merged = _merge_detection([], llm, candidates=cands)
    assert [it["key"] for it in merged] == ["a"]  # ghost 组 cap=0 全溢出


def test_merge_without_candidates_keeps_legacy_occ_semantics():
    """回归（兼容直调）：不传 candidates 时行为同旧——组内全保留 occ=1..n，
    超界与否交给 validate 终审（存量单测与潜在直调方不受影响）。"""
    from rag.svr.template_fill.detector import _merge_detection
    llm = [_mk_fill_item(k, "＿＿") for k in ("a", "b")]
    merged = _merge_detection([], llm)
    assert [it["occ"] for it in merged] == [1, 2]
    assert [it["key"] for it in merged] == ["a", "b"]


def test_merge_overflow_with_explicit_group_capped_too():
    """对抗：explicit 直通组同样封顶——{{k}} 在文本中仅 1 次但直通提取异常产生
    2 项同位时，第 2 项丢弃（防 extract 层回归放大成整次识别失败）。"""
    from rag.svr.template_fill.detector import _merge_detection
    cands = [{"index": 0, "text": "编号：{{code}}", "addr": "para:0"}]
    explicit = [
        {"key": "code", "name": "code", "description": "", "retrieval_query": "",
         "fill_mode": "llm", "required": True, "addr": "para:0",
         "anchor": "{{code}}", "line": 0, "top_k": 6},
        dict(_mk_fill_item("code_dup", "{{code}}", addr="para:0")),  # 同位异常项
    ]
    explicit[0]["_anchor_pos"] = 0
    explicit[1]["_anchor_pos"] = 3
    merged = _merge_detection(explicit, [], candidates=cands)
    assert [it["key"] for it in merged] == ["code"]
    assert merged[0]["occ"] == 1


def test_merge_overflow_takes_precedence_over_shrink_rollback():
    """对抗（溢出封顶 × 收缩撞车回退交互）：带 _orig_anchor 的项因溢出被丢弃时
    不再进入回退保留分支——封顶优先于回退，宁缺勿错（docstring 已声明该优先级）。"""
    from rag.svr.template_fill.detector import _merge_detection
    cands = [{"index": 0, "text": "甲：＿＿", "addr": "para:0"}]
    a = _mk_fill_item("a", "＿＿", addr="para:0")
    b = _mk_fill_item("b", "＿＿", addr="para:0")
    b["_orig_anchor"] = "乙方名称："  # 旧版会回退保留+低置信；新版溢出直接丢弃
    merged = _merge_detection([], [a, b], candidates=cands)
    assert [it["key"] for it in merged] == ["a"]  # b 溢出丢弃，不走 _orig_anchor 回退


def test_merge_cross_source_echo_unaffected_by_cap():
    """回归：LLM 组撞 explicit_pos 的跨源项仍按旧「手动优先丢弃」逻辑处理，
    封顶机制不改变该路径（跨源项本就不参与预分配）。"""
    from rag.svr.template_fill.detector import _merge_detection, extract_explicit_placeholders
    cands = [{"index": 0, "text": "编号：{{code}}", "addr": "para:0"}]
    explicit = extract_explicit_placeholders(cands)
    llm = [_mk_fill_item("echo1", "{{code}}"), _mk_fill_item("echo2", "{{code}}")]
    merged = _merge_detection(explicit, llm, candidates=cands)
    assert [it["key"] for it in merged] == ["code"]
    assert "occ" not in merged[0]


# ---------- LLM 输出解析三级容错 ----------


def test_parse_bare_array_direct():
    """裸 JSON 数组（无围栏无前后缀）走最快直解路径。"""
    from rag.svr.template_fill.detector import parse_detection_response
    cands = [{"index": 0, "text": "姓名：＿＿", "addr": "para:0"}]
    items = parse_detection_response(
        '[{"line": 0, "anchor": "＿＿", "key": "name", "name": "姓名"}]', cands)
    assert len(items) == 1 and items[0]["key"] == "name"


def test_parse_fenced_array():
    """```json 围栏包裹 → 剥围栏直解。"""
    from rag.svr.template_fill.detector import parse_detection_response
    cands = [{"index": 0, "text": "姓名：＿＿", "addr": "para:0"}]
    raw = '```json\n[{"line": 0, "anchor": "＿＿", "key": "name", "name": "姓名"}]\n```'
    assert len(parse_detection_response(raw, cands)) == 1


def test_parse_array_with_surrounding_prose():
    """数组前后有说明文字 → 贪婪正则回退兜住。"""
    from rag.svr.template_fill.detector import parse_detection_response
    cands = [{"index": 0, "text": "姓名：＿＿", "addr": "para:0"}]
    raw = '识别结果如下：\n[{"line": 0, "anchor": "＿＿", "key": "name", "name": "姓名"}]\n以上。'
    assert len(parse_detection_response(raw, cands)) == 1


def test_parse_multiple_arrays_conservative_failure_via_greedy_fallback():
    """多数组输入贪婪正则会整体失败 → 返回 []（保守失败，不误采半截）。"""
    from rag.svr.template_fill.detector import parse_detection_response
    cands = [{"index": 0, "text": "姓名：＿＿", "addr": "para:0"}]
    raw = '[{"line": 0}] 中间文字 [{"line": 1}]'
    assert parse_detection_response(raw, cands) == []


def test_parse_garbage_returns_empty():
    """纯文字拒绝/None 输入 → []（None 不得抛 TypeError）。"""
    from rag.svr.template_fill.detector import parse_detection_response
    cands = [{"index": 0, "text": "姓名：＿＿", "addr": "para:0"}]
    assert parse_detection_response("我无法完成该任务", cands) == []
    assert parse_detection_response(None, cands) == []


def test_parse_fenced_variant_no_language_tag():
    """对抗：围栏无语言标注 ```\n[...]\n``` → 仍能剥围栏直解。"""
    from rag.svr.template_fill.detector import parse_detection_response
    cands = [{"index": 0, "text": "姓名：＿＿", "addr": "para:0"}]
    raw = '```\n[{"line": 0, "anchor": "＿＿", "key": "name", "name": "姓名"}]\n```'
    assert len(parse_detection_response(raw, cands)) == 1


def test_parse_fenced_variant_uppercase_language():
    """对抗：围栏语言大写 ```JSON → 仍能剥围栏直解。"""
    from rag.svr.template_fill.detector import parse_detection_response
    cands = [{"index": 0, "text": "姓名：＿＿", "addr": "para:0"}]
    raw = '```JSON\n[{"line": 0, "anchor": "＿＿", "key": "name", "name": "姓名"}]\n```'
    assert len(parse_detection_response(raw, cands)) == 1


def test_parse_fenced_nested_array_fields():
    """对抗：围栏内数组元素含嵌套数组字段（如 retrieval_query 列表）→
    剥围栏直解成功，不因贪婪正则跨界截断而失败；嵌套字段被逐项校验安全忽略。"""
    from rag.svr.template_fill.detector import parse_detection_response
    cands = [{"index": 0, "text": "姓名：＿＿", "addr": "para:0"}]
    raw = ('```json\n[{"line": 0, "anchor": "＿＿", "key": "name", "name": "姓名", '
           '"aliases": ["甲", "乙"], "extra": [1, 2, 3]}]\n```')
    assert len(parse_detection_response(raw, cands)) == 1


def test_parse_empty_and_whitespace_raw():
    """对抗：raw 为空串/纯空白 → []。"""
    from rag.svr.template_fill.detector import parse_detection_response
    cands = [{"index": 0, "text": "姓名：＿＿", "addr": "para:0"}]
    assert parse_detection_response("", cands) == []
    assert parse_detection_response("   \n\t  ", cands) == []


def test_parse_unbalanced_bracket_truncated():
    """对抗：`[` 与 `]` 不配平（截断输出）→ json.loads 失败且贪婪正则无 `]`
    可回退 → 保守返回 []，不误采半截数据。"""
    from rag.svr.template_fill.detector import parse_detection_response
    cands = [{"index": 0, "text": "姓名：＿＿", "addr": "para:0"}]
    raw = '[{"line": 0, "anchor": "＿＿", "key": "name"'
    assert parse_detection_response(raw, cands) == []


# ---------------------------------------------------------------------------
# 超链接段落统一走 cross 区间重写 + occ 嵌套剔除（2026-09-16 事故回归）
# 事故：_replace_via_run_concat 按原 run 长度静态切片回写，替换点之后 run 文本
# 整体后移，而 w:hyperlink 元素钉在 XML 固定位置 → URL 文本相对位置「瞬移」进
# 先前写入的 {{key}} 内部，产出 {{bhttp://...}} 之类脏 token，docxtpl Jinja 解析
# 报 "expected token 'end of print statement', got ':'"。
# ---------------------------------------------------------------------------


def _make_incident_docx():
    """复刻事故段落结构：run("递交截止时间：") + run("7空年7空月3空日3空时3空分3空秒前 ")
    + 尾部超链接(URL) + run("（详见公告）")——超链接前后都有正文 run。"""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document()
    p = doc.add_paragraph("递交截止时间：")
    p.add_run("       年       月   日   时   分   秒前 ")
    hl = OxmlElement("w:hyperlink")
    hl.set(qn("r:id"), "rIdLink")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "http://fjggzyjy.cn/xxx"
    r.append(t)
    hl.append(r)
    p._p.append(hl)
    p.add_run("（详见公告）")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_apply_incident_hyperlink_paragraph_no_dirty_token():
    """事故回归：超链接居中+后随 run 的年月日时分秒段落，6 条目替换后 URL 不得
    混入任何 {{}} token（旧 concat 路径会产出 {{bhttp://fjggzyjy.cn/x...}}）。"""
    from rag.svr.template_fill.docx_utils import apply_docx_placeholders, iter_docx_paragraphs
    src = _make_incident_docx()
    URL = "http://fjggzyjy.cn/xxx"
    out = apply_docx_placeholders(src, [
        {"addr": "para:0", "anchor": "       ", "key": "d_year", "occ": 1},
        {"addr": "para:0", "anchor": "       ", "key": "d_month", "occ": 2},
        {"addr": "para:0", "anchor": "   ", "key": "d_day", "occ": 1},
        {"addr": "para:0", "anchor": "   ", "key": "d_hour", "occ": 2},
        {"addr": "para:0", "anchor": "   ", "key": "d_minute", "occ": 3},
    ])
    text = iter_docx_paragraphs(out)[0]["text"]
    # URL 只在超链接原位出现一次，绝不在 {{}} 内
    assert text.count(URL) == 1
    assert "bhttp" not in text
    import re as _re
    toks = _re.findall(r"\{\{[^}]*\}\}", text)
    assert len(toks) == 5, text
    for tok in toks:
        assert _re.fullmatch(r"\{\{[a-z0-9_]+\}\}", tok), tok
    # 5 个占位符分别落在年/月/日/时/分空位，锚文本归零（仅剩「秒」前的 3 空格）
    assert text == ("递交截止时间：{{d_year}}年{{d_month}}月{{d_day}}日{{d_hour}}时"
                    "{{d_minute}}分   秒前 " + URL + "（详见公告）"), text


def test_apply_occ_nested_in_longer_anchor_skipped():
    """对抗（occ 嵌套口径）：3 空格锚的嵌套出现（藏在 7 空格空位内）必须被剔除，
    occ=1/2 落到独立空位上——识别端按独立留白位编号，渲染端不剔除会把序号数到
    长锚内部而错位。"""
    from rag.svr.template_fill.docx_utils import apply_docx_placeholders, iter_docx_paragraphs
    doc = Document()
    doc.add_paragraph("X       Y   Z   W")  # 7空 + 3空 + 3空
    buf = io.BytesIO()
    doc.save(buf)
    out = apply_docx_placeholders(buf.getvalue(), [
        {"addr": "para:0", "anchor": "   ", "key": "s1", "occ": 1},
        {"addr": "para:0", "anchor": "   ", "key": "s2", "occ": 2},
        {"addr": "para:0", "anchor": "       ", "key": "long", "occ": 1},
    ])
    text = iter_docx_paragraphs(out)[0]["text"]
    assert text == "X{{long}}Y{{s1}}Z{{s2}}W", text


def test_apply_occ_exceeds_free_slots_is_noop():
    """对抗：剔除嵌套后幸存空位数不足 occ → no-op（占位符不落位，绝不退化成
    replace-all 误伤长锚内部）。"""
    from rag.svr.template_fill.docx_utils import apply_docx_placeholders, iter_docx_paragraphs
    doc = Document()
    doc.add_paragraph("X       Y   Z")  # 独立 3 空格仅 1 处，另一处嵌在 7 空格内
    buf = io.BytesIO()
    doc.save(buf)
    out = apply_docx_placeholders(buf.getvalue(), [
        {"addr": "para:0", "anchor": "   ", "key": "s2", "occ": 2},
        {"addr": "para:0", "anchor": "       ", "key": "long", "occ": 1},
    ])
    text = iter_docx_paragraphs(out)[0]["text"]
    assert "{{s2}}" not in text
    assert text == "X{{long}}Y   Z", text


def test_apply_docx_fld_simple_paragraph_field_intact():
    """对抗：简单域（w:fldSimple）段落与超链接同构——域内 run 不在 p.runs 中，
    替换只作用正文 run，域文本不复制不破坏。"""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from rag.svr.template_fill.docx_utils import apply_docx_placeholders, iter_docx_paragraphs
    doc = Document()
    p = doc.add_paragraph("页码：")
    p.add_run("____ 共")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), " PAGE ")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "9"
    r.append(t)
    fld.append(r)
    p._p.append(fld)
    p.add_run("页")
    buf = io.BytesIO()
    doc.save(buf)
    out = apply_docx_placeholders(buf.getvalue(), [
        {"addr": "para:0", "anchor": "____", "key": "pages"},
    ])
    # p.text 含 hyperlink 文本但不含 fldSimple 域内文本（python-docx 1.2 行为），
    # 提取文本只含正文 run；域完整性（含 "9"）在 XML 层断言
    text = iter_docx_paragraphs(out)[0]["text"]
    assert text == "页码：{{pages}} 共页", text
    from docx import Document as _D
    p2 = _D(io.BytesIO(out)).paragraphs[0]
    flds = p2._p.findall(qn("w:fldSimple"))
    assert len(flds) == 1
    assert flds[0].find(qn("w:r")).find(qn("w:t")).text == "9"


def test_apply_docx_fld_simple_occ_path_token_intact():
    """对抗（事故同构）：fldSimple 段落 + occ=1 + 锚跨界 run 边界 + 替换值长于锚。
    旧 run 拼接路径按原 run 长度静态切片回写，会把 {{token}} 拦腰切进两个 run
    （产出 "{{pages}" 残缺）；区间重写路径 token 原子落在首 run，域内文本不受
    影响。"""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx import Document as _D
    from rag.svr.template_fill.docx_utils import apply_docx_placeholders, iter_docx_paragraphs
    doc = Document()
    p = doc.add_paragraph("页码：_")
    p.add_run("___共")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), " PAGE ")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "9"
    r.append(t)
    fld.append(r)
    p._p.append(fld)
    p.add_run("终")
    buf = io.BytesIO()
    doc.save(buf)
    # 前置：runs 拼接 = "页码：____共终"，锚 "：____" 跨 run0/run1
    out = apply_docx_placeholders(buf.getvalue(), [
        {"addr": "para:0", "anchor": "：____", "key": "pages", "occ": 1},
    ])
    text = iter_docx_paragraphs(out)[0]["text"]
    assert text == "页码{{pages}}共终", text
    # 域完整性 + token 原子性须在 document 序全量 w:t 上断言（p.text 不含域内
    # 文本，旧 concat 切碎的 token 在 p.text 视图恰好拼回完整串，看不到腐坏；
    # docxtpl 按 document 序解析，域内 "9" 会插进被切碎的 token 中间）
    p2 = _D(io.BytesIO(out)).paragraphs[0]
    flow = "".join(t.text or "" for t in p2._p.iter() if t.tag == qn("w:t"))
    assert flow == "页码{{pages}}共9终", flow
    flds = p2._p.findall(qn("w:fldSimple"))
    assert len(flds) == 1
    assert flds[0].find(qn("w:r")).find(qn("w:t")).text == "9"


def _make_docx_with_underline_runs(para_specs):
    """para_specs: [ [(text, underline_or_None), ...] ]，每项一段。"""
    doc = Document()
    for runs_spec in para_specs:
        p = doc.add_paragraph()
        for text, u in runs_spec:
            r = p.add_run(text)
            if u is not None:
                r.font.underline = u
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_extract_docx_candidates_slots_attached():
    from rag.svr.template_fill.docx_utils import extract_docx_candidates
    blob = _make_docx_with_underline_runs([
        [("本招标项目", None), ("    （招标项目名称） ", True), ("已由", None),
         (" （审批机关） ", True)],
    ])
    cands = extract_docx_candidates(blob)
    assert len(cands) == 1
    slots = cands[0]["slots"]
    assert [s["kind"] for s in slots] == ["hint", "hint"]
    assert slots[0]["hint"] == "招标项目名称"
    assert slots[1]["hint"] == "审批机关"


def test_extract_docx_candidates_slot_line_included_without_hint_re():
    """有位的行即使不命中 FILL_HINT_RE 也入候选（结构性修漏识别）。
    注意间隙用 3 空格：6+ 空格会命中 FILL_HINT_RE 的 \\S[ \\u3000]{6,}\\S，测不出 slots 兜底。"""
    from rag.svr.template_fill.docx_utils import extract_docx_candidates
    blob = _make_docx_with_underline_runs([
        [("合同金额", None), ("   ", True), ("万元，工期", None), ("   ", True), ("天。", None)],
    ])
    cands = extract_docx_candidates(blob)
    assert len(cands) == 1
    assert len(cands[0]["slots"]) == 2


def test_extract_docx_candidates_blank_underline_only_paragraph_included():
    """整行纯空白下划线段落（text.strip() 为空但有位）入候选——真实范本 57 行漏识别根修。"""
    from rag.svr.template_fill.docx_utils import extract_docx_candidates
    blob = _make_docx_with_underline_runs([
        [("             ", True)],           # 纯空白下划线独立填写行
        [("普通空段落", None)],              # 无位无特征 → 不入
    ])
    cands = extract_docx_candidates(blob)
    assert len(cands) == 1
    assert cands[0]["slots"] and cands[0]["slots"][0]["kind"] == "blank"


def test_extract_docx_candidates_char_underscore_slots():
    """普通特征行（字符下划线）同样产出切位结果（kind=blank）。"""
    from rag.svr.template_fill.docx_utils import extract_docx_candidates
    cands = extract_docx_candidates(_make_docx(["项目名称：____________", "无填写点的普通段落"]))
    by_text = {c["text"]: c for c in cands}
    assert by_text["项目名称：____________"]["slots"][0]["kind"] == "blank"


def test_extract_docx_candidates_no_slots_hint_re_still_matched():
    """规格③：run 层无填写位（slots==[]）但文本命中 FILL_HINT_RE 的行仍入候选。
    "（此处填写）" 命中 [（(][^（）()]*填写[^（）()]*[)）] 分支，且无 underline run、
    无连续字符下划线（blank_slots 仅从这两处切位）→ slots 恒为 []；
    对照：无任何特征的段落不入候选。"""
    from rag.svr.template_fill.docx_utils import extract_docx_candidates
    cands = extract_docx_candidates(_make_docx(["（此处填写）", "普通正文段落无任何特征"]))
    by_text = {c["text"]: c for c in cands}
    assert "（此处填写）" in by_text
    assert by_text["（此处填写）"]["slots"] == []
    assert "普通正文段落无任何特征" not in by_text


# ---- detector V2「位编号→语义」契约 ----

def _slot_cand():
    return {
        "index": 7,
        "addr": "para:7",
        "text": "本招标项目    （招标项目名称） 已由 （审批机关） 以批准建设",
        "slots": [
            {"start": 5, "end": 18, "text": "    （招标项目名称） ", "kind": "hint", "hint": "招标项目名称"},
            {"start": 20, "end": 28, "text": " （审批机关） ", "kind": "hint", "hint": "审批机关"},
        ],
    }


def test_parse_slot_response_valid():
    from rag.svr.template_fill.detector import parse_slot_response
    cand = _slot_cand()
    raw = '[{"line":7,"slot":1,"key":"project_name","name":"招标项目名称","required":true},' \
          '{"line":7,"slot":2,"key":"approve_org","name":"审批机关","required":true}]'
    items, covered = parse_slot_response(raw, [cand])
    assert len(items) == 2
    assert covered == {(7, 1), (7, 2)}
    # anchor 由切位区间精确切片，LLM 无权决定
    assert items[0]["anchor"] == "    （招标项目名称） "
    assert items[0]["addr"] == "para:7"
    assert items[0]["_anchor_pos"] == 5
    assert items[1]["anchor"] == " （审批机关） "


def test_parse_slot_response_invalid_slot_dropped():
    from rag.svr.template_fill.detector import parse_slot_response
    raw = '[{"line":7,"slot":9,"key":"x","name":"x"},' \
          '{"line":7,"slot":"1","key":"y","name":"y"},' \
          '{"line":"7","slot":1,"key":"z","name":"z"}]'
    items, covered = parse_slot_response(raw, [_slot_cand()])
    assert items == []
    assert covered == set()


def test_parse_slot_response_duplicate_slot_first_wins():
    from rag.svr.template_fill.detector import parse_slot_response
    raw = '[{"line":7,"slot":1,"key":"first_key","name":"甲"},' \
          '{"line":7,"slot":1,"key":"second_key","name":"乙"}]'
    items, covered = parse_slot_response(raw, [_slot_cand()])
    assert len(items) == 1
    assert items[0]["key"] == "first_key"
    assert covered == {(7, 1)}


def test_parse_slot_response_garbage_returns_empty():
    from rag.svr.template_fill.detector import parse_slot_response
    items, covered = parse_slot_response("不是JSON", [_slot_cand()])
    assert items == [] and covered == set()


def test_slot_fallback_items_hint_and_blank():
    from rag.svr.template_fill.detector import slot_fallback_items
    cand = {
        "index": 3, "addr": "para:3", "text": "业主为        ，金额      万元",
        "slots": [
            {"start": 3, "end": 11, "text": "        ", "kind": "blank", "hint": ""},
            {"start": 14, "end": 20, "text": "      ", "kind": "blank", "hint": ""},
        ],
    }
    items = slot_fallback_items([cand], covered=set())
    assert len(items) == 2
    assert items[0]["key"] == "blank_1" and items[1]["key"] == "blank_2"
    assert items[0]["name"] == "未命名填写位"
    assert items[0]["low_confidence"] is True
    assert items[0]["anchor"] == "        "
    assert items[0]["_anchor_pos"] == 3


def test_slot_fallback_items_skips_covered_and_uses_hint():
    from rag.svr.template_fill.detector import slot_fallback_items
    cand = _slot_cand()
    items = slot_fallback_items([cand], covered=set())  # 未覆盖 → 兜底生成
    assert len(items) == 2
    assert items[0]["name"] == "招标项目名称"
    assert items[0]["low_confidence"] is False  # hint 位名字高可信
    # 覆盖后不再生成
    assert slot_fallback_items([cand], covered={(cand["index"], 1), (cand["index"], 2)}) == []


def test_v2_items_carry_explicit_empty_default():
    """V2 硬闸（设计 §4.5）：LLM 标注项与兜底项都必须显式带 default_value=""
    ——_merge_defaults 分支1按「键存在」触发，V2 条目无论 anchor 形态不派生默认值。"""
    from rag.svr.template_fill.detector import parse_slot_response, slot_fallback_items
    cand = _slot_cand()
    items, _ = parse_slot_response(
        '[{"line":7,"slot":1,"key":"project_name","name":"招标项目名称"}]', [cand])
    assert items[0]["default_value"] == ""
    fb = slot_fallback_items([cand], covered=set())
    assert len(fb) == 2
    assert all(it["default_value"] == "" for it in fb)


def test_merge_v2_fallback_same_shape_slots_occ_no_overflow():
    """同段两个同形空白位 + 各一条条目 → occ 预分配 1/2，无溢出丢弃。"""
    from rag.svr.template_fill.detector import _merge_detection
    cand = {
        "index": 3, "addr": "para:3", "text": "业主为        ，金额      万元",
        "slots": [
            {"start": 3, "end": 11, "text": "        ", "kind": "blank", "hint": ""},
            {"start": 14, "end": 20, "text": "      ", "kind": "blank", "hint": ""},
        ],
    }
    a = {"key": "owner", "addr": "para:3", "anchor": "        ", "_anchor_pos": 3,
         "low_confidence": False, "name": "业主", "line": 3}
    b = {"key": "amount", "addr": "para:3", "anchor": "      ", "_anchor_pos": 14,
         "low_confidence": False, "name": "金额", "line": 3}
    # 两个不同 anchor 文本：不同组，无 occ；再验证同形场景
    merged = _merge_detection([], [a, b], candidates=[cand])
    assert len(merged) == 2
    # cand2 文本 "甲     乙   "：anchor "   " 实际出现于 pos 1 与 pos 7（非重叠计数 2 处）
    same_a = {"key": "x1", "addr": "para:3", "anchor": "   ", "_anchor_pos": 1,
              "low_confidence": False, "name": "甲", "line": 3}
    same_b = {"key": "x2", "addr": "para:3", "anchor": "   ", "_anchor_pos": 7,
              "low_confidence": False, "name": "乙", "line": 3}
    cand2 = {"index": 3, "addr": "para:3", "text": "甲     乙   ", "slots": []}
    merged2 = _merge_detection([], [same_a, same_b], candidates=[cand2])
    assert len(merged2) == 2
    occs = sorted(it["occ"] for it in merged2 if "occ" in it)
    assert occs == [1, 2]


# ---- Task 3 审查修复回归（F1 required 判型 / key 去重 / membership 防御 / 序号续接）----

def test_parse_slot_response_required_string_typed():
    """对抗：LLM 输出 "required":"false"（字符串），bool("false")==True 陷阱
    → 显式判型：仅布尔 True / 字符串 "true" 为 True，缺失默认 True。"""
    from rag.svr.template_fill.detector import parse_slot_response
    raw = '[{"line":7,"slot":1,"key":"a","name":"甲","required":"false"},' \
          '{"line":7,"slot":2,"key":"b","name":"乙","required":"true"}]'
    items, _ = parse_slot_response(raw, [_slot_cand()])
    assert items[0]["required"] is False
    assert items[1]["required"] is True
    # 缺失 → 默认 True
    items2, _ = parse_slot_response('[{"line":7,"slot":1,"key":"c","name":"丙"}]', [_slot_cand()])
    assert items2[0]["required"] is True


def test_parse_slot_response_duplicate_long_key_dedup_no_overflow():
    """对抗（F1）：V2 parse 级缺 key 去重 → 两个 71 字符相同 key 原样透传，
    merge 再加 _2 后缀得 66 字符，validate_placeholders 判死整次识别。
    修后与 V1 同口径：截断感知后缀，含后缀总长 ≤64，且过 validate 终审。"""
    from rag.svr.template_fill.detector import parse_slot_response, validate_placeholders
    long_key = "liability_for_refusing_to_replace_key_construction_management_personnel"
    assert len(long_key) == 71
    raw = f'[{{"line":7,"slot":1,"key":"{long_key}","name":"A"}},' \
          f'{{"line":7,"slot":2,"key":"{long_key}","name":"B"}}]'
    items, covered = parse_slot_response(raw, [_slot_cand()])
    assert len(items) == 2 and covered == {(7, 1), (7, 2)}
    k1, k2 = items[0]["key"], items[1]["key"]
    assert k1 != k2
    import re as _re
    for k in (k1, k2):
        assert len(k) <= 64
        assert _re.fullmatch(r"[a-z][a-z0-9_]{0,63}", k)
    ok, err = validate_placeholders(items, [_slot_cand()])
    assert ok, err


def test_parse_slot_response_zero_based_slots_all_dropped():
    """对抗：LLM 全部输出 0 起位序号（位编号契约从 1 开始）→ 查无此位全部丢弃，
    items==[] 且 covered==set()，下游 slot_fallback_items 对全量 slot 兜底
    （固化丢弃行为——降级路径不因脏序号部分产脏项）。"""
    from rag.svr.template_fill.detector import parse_slot_response, slot_fallback_items
    cand = _slot_cand()
    raw = '[{"line":7,"slot":0,"key":"a","name":"甲"},' \
          '{"line":7,"slot":0,"key":"b","name":"乙"}]'
    items, covered = parse_slot_response(raw, [cand])
    assert items == [] and covered == set()
    assert len(slot_fallback_items([cand], covered)) == 2


def test_parse_slot_response_membership_defense_drops_forged_slot():
    """对抗（F4）：候选 text 与 slot["text"] 坐标系不一致（数据腐坏/伪造候选）
    → 该项丢弃且不占 covered，不静默产脏 anchor。"""
    from rag.svr.template_fill.detector import parse_slot_response
    cand = dict(_slot_cand(), text="与 slots 区间完全无关的另一段文本")
    raw = '[{"line":7,"slot":1,"key":"a","name":"甲"}]'
    items, covered = parse_slot_response(raw, [cand])
    assert items == [] and covered == set()


def test_slot_fallback_items_seq_start_continues_numbering():
    """（F2）分块调用传续接 seq_start → blank_N 跨块全局唯一。"""
    from rag.svr.template_fill.detector import slot_fallback_items
    cand = _slot_cand()
    items = slot_fallback_items([cand], covered=set(), seq_start=5)
    assert [it["key"] for it in items] == ["blank_6", "blank_7"]
    # 不传仍从 blank_1 起（默认行为不变）
    items2 = slot_fallback_items([cand], covered=set())
    assert [it["key"] for it in items2] == ["blank_1", "blank_2"]


def test_detect_fill_points_routes_v1_v2(monkeypatch):
    """有位行走 V2、无位行走 V1，兜底补漏，合并进 _merge_detection。
    _get_chat_model 打桩绕开 DB/LLM；_detect_chunked/_detect_slot_chunked
    双双打桩断言各收到且仅收到自己的候选分组。"""
    import asyncio

    from rag.svr.template_fill import detector

    v1_cand = {"index": 1, "addr": "para:1", "text": "项目名称：____________", "slots": []}
    v2_cand = _slot_cand()  # index=7, addr=para:7, 两个 hint 位

    captured = {}

    async def fake_v1(chat, file_type, candidates):
        captured["v1"] = [c["index"] for c in candidates]
        return [], 0

    async def fake_v2(chat, candidates):
        captured["v2"] = [c["index"] for c in candidates]
        items = [{
            "key": "project_name", "name": "招标项目名称", "description": "",
            "retrieval_query": "", "fill_mode": "llm", "required": True,
            "addr": "para:7", "anchor": "    （招标项目名称） ", "line": 7,
            "top_k": 6, "low_confidence": False, "_anchor_pos": 5,
        }]
        return items, {(7, 1)}, 0

    class _FakeModel:
        async def async_chat(self, system, messages):
            return "[]"

    monkeypatch.setattr(detector, "_get_chat_model", lambda tid: _FakeModel(), raising=False)
    monkeypatch.setattr(detector, "_detect_chunked", fake_v1)
    monkeypatch.setattr(detector, "_detect_slot_chunked", fake_v2, raising=False)

    merged = asyncio.run(detector.detect_fill_points(
        "tenant", "docx", [v1_cand, v2_cand]))
    assert captured["v1"] == [1]
    assert captured["v2"] == [7]
    keys = sorted(it["key"] for it in merged)
    # 位1 由 LLM 标注；位2 未覆盖 → 兜底（blank_1 序号跨 cand 递增）
    assert "project_name" in keys
    assert "blank_1" in keys
    # 无 V1 撞车：兜底项 anchor 与 LLM 项不同位
    assert len(merged) == 2
    # 内部字段不外泄：V2 切位偏移经校验闸后必须被剥离
    assert all("_anchor_pos" not in it for it in merged)


def _blank_gap_cand():
    """同段三个「  」同形出现，中间 [4,6) 是**非位**同形（不在 slots 里）：
    甲[0] 空位1[1,3) 乙[3] 非位空格[4,6) 丙[6] 空位2[7,9) 丁[9]"""
    return {
        "index": 3,
        "addr": "para:3",
        "text": "甲  乙  丙  丁",
        "slots": [
            {"start": 1, "end": 3, "text": "  ", "kind": "blank", "hint": ""},
            {"start": 7, "end": 9, "text": "  ", "kind": "blank", "hint": ""},
        ],
    }


def _gap_item(key, pos):
    return {
        "key": key, "name": key, "description": "", "retrieval_query": "",
        "fill_mode": "llm", "required": True, "addr": "para:3",
        "anchor": "  ", "line": 3, "top_k": 6, "low_confidence": False,
        "_anchor_pos": pos,
    }


def test_verify_slot_occ_corrects_plain_text_insertion():
    """Major 3 主场景：同段非位同形空白插在两位之间，_merge_detection 按
    _anchor_pos 排序给的 occ（1、2）相对渲染层非重叠出现序错位（真位2是第 3 次
    出现）→ 校验闸按渲染层同款算法回写 occ。"""
    from rag.svr.template_fill.detector import _merge_detection, _verify_slot_occ
    cand = _blank_gap_cand()
    merged = _merge_detection([], [_gap_item("a", 1), _gap_item("b", 7)],
                              preassign_occ=True, candidates=[cand])
    occs = {it["key"]: it.get("occ") for it in merged}
    # 错位证据：merge 给 b 的 occ=2，而渲染层非重叠序里第 2 次出现是非位空格
    assert occs == {"a": 1, "b": 2}
    # 回挂切位偏移（detect_fill_points 内部行为），过闸后 occ 被纠正为真实出现序
    pos_by_id = {id(it): p for it, p in zip(merged, (1, 7))}
    for it in merged:
        it["_anchor_pos"] = pos_by_id[id(it)]
    out = _verify_slot_occ(merged, [cand])
    assert len(out) == 2
    occs = {it["key"]: it["occ"] for it in out}
    assert occs == {"a": 1, "b": 3}


def test_verify_slot_occ_drops_non_occurrence():
    """对抗：_anchor_pos 不落在 anchor 任何非重叠出现区间（坐标系分歧/数据腐坏）
    → 条目丢弃，宁可漏不写错位。"""
    from rag.svr.template_fill.detector import _verify_slot_occ
    cand = _blank_gap_cand()
    it = _gap_item("a", 0)  # pos=0 是「甲」，不属于任何「  」出现区间
    out = _verify_slot_occ([it], [cand])
    assert out == []


def test_verify_slot_occ_passes_v1_untouched():
    """无 slots 候选（V1/explicit）条目原样透传：V1 的 _anchor_pos 本就来自同一
    文本序 find，零回归；有 slots 候选但条目无 _anchor_pos（手动占位符直通）同样不动。"""
    from rag.svr.template_fill.detector import _verify_slot_occ
    v1_cand = {"index": 1, "addr": "para:1", "text": "甲  乙", "slots": []}
    it_v1 = {"key": "a", "addr": "para:1", "anchor": "  ", "occ": 5, "_anchor_pos": 999}
    explicit_cand = _blank_gap_cand()
    it_explicit = {"key": "b", "addr": "para:3", "anchor": "  ", "occ": 2}  # 无 _anchor_pos
    # 对抗判型：_anchor_pos=True（bool 是 int 子类）必须按非 int 偏移处理 → 原样
    # 透传（不回写 occ、不丢弃），走与 None 同一判型分支
    bool_cand = {"index": 9, "addr": "para:9", "text": "甲  乙",
                 "slots": [{"start": 1, "end": 3, "text": "  ", "kind": "blank", "hint": ""}]}
    it_bool = {"key": "c", "addr": "para:9", "anchor": "  ", "occ": 4, "_anchor_pos": True}
    out = _verify_slot_occ([it_v1, it_explicit, it_bool],
                           [v1_cand, explicit_cand, bool_cand])
    assert out[0] is it_v1 and out[0]["occ"] == 5
    assert out[1] is it_explicit and out[1]["occ"] == 2
    assert out[2] is it_bool and out[2]["occ"] == 4


def test_verify_slot_occ_drops_same_hit_collision():
    """对抗（切位数据腐坏）：同组两 V2 条目的 _anchor_pos 落进同一非重叠出现区间
    （真实切位互不重叠，不可能发生；发生即数据坏了）→ 保留靠前者，丢弃靠后者。"""
    from rag.svr.template_fill.detector import _verify_slot_occ
    cand = _blank_gap_cand()
    it_a = _gap_item("a", 1)
    it_b = _gap_item("b", 2)  # 与 a 同在出现区间 [1,3) 内
    out = _verify_slot_occ([it_a, it_b], [cand])
    assert len(out) == 1
    assert out[0]["key"] == "a" and out[0]["occ"] == 1


def _nested_cand():
    """嵌套场景（审查探针）：长空白位 5 空格 [1,6)，其内部藏两个「  」的非重叠
    出现 [1,3)/[3,5)；真短位 2 空格在 [7,9)/[10,12)/[13,15)。
    甲[0] 「     」[1,6) 乙[6] 「  」[7,9) 丙[9] 「  」[10,12) 丁[12] 「  」[13,15) 戊[15]"""
    return {
        "index": 5,
        "addr": "para:5",
        "text": "甲" + "     " + "乙" + "  " + "丙" + "  " + "丁" + "  " + "戊",
        "slots": [
            {"start": 1, "end": 6, "text": "     ", "kind": "blank", "hint": ""},
            {"start": 7, "end": 9, "text": "  ", "kind": "blank", "hint": ""},
            {"start": 10, "end": 12, "text": "  ", "kind": "blank", "hint": ""},
            {"start": 13, "end": 15, "text": "  ", "kind": "blank", "hint": ""},
        ],
    }


def _nested_item(key, anchor, pos):
    return {
        "key": key, "name": key, "description": "", "retrieval_query": "",
        "fill_mode": "llm", "required": True, "addr": "para:5",
        "anchor": anchor, "line": 5, "top_k": 6, "low_confidence": False,
        "_anchor_pos": pos,
    }


def test_verify_slot_occ_nested_longer_anchor_survivors():
    """M-1：同段 1 个长空白位（5 空格）+ 3 个短空白位（2 空格，anchor 同形）。
    渲染层 _replace_cross_run_in_place 取第 N 次出现前剔除嵌套在更长锚出现区间
    内的出现（survivors 口径）；校验闸此前只复刻非重叠出现序、没复刻嵌套剔除
    → 短位 occ 被写成 3/4/5（幸存者序实为 1/2/3，长位内部两个嵌套「  」占掉
    非重叠序 1/2）→ 值串位或 occ 超界 no-op 填写点丢失。修复后按幸存者序回写。"""
    from rag.svr.template_fill.detector import _verify_slot_occ
    cand = _nested_cand()
    merged = [
        _nested_item("long", "     ", 3),   # 长位 [1,6) 内部
        _nested_item("s1", "  ", 7),        # 短位1 [7,9)
        _nested_item("s2", "  ", 10),       # 短位2 [10,12)
        _nested_item("s3", "  ", 13),       # 短位3 [13,15)
    ]
    out = _verify_slot_occ(merged, [cand])
    assert len(out) == 4
    occs = {it["key"]: it["occ"] for it in out}
    # 非重叠出现序会把短位数成 3/4/5；幸存者序（嵌套剔除后）为 1/2/3
    assert occs == {"long": 1, "s1": 1, "s2": 2, "s3": 3}


def test_verify_slot_occ_drops_nested_blocked_pos():
    """对抗：_anchor_pos 落在被嵌套剔除的出现上（长位内部 [1,3) 处的短锚
    「出现」实际是长锚留白的一段，渲染层永不取它）→ 条目丢弃，宁可漏不写错。"""
    from rag.svr.template_fill.detector import _verify_slot_occ
    cand = _nested_cand()
    merged = [
        _nested_item("long", "     ", 3),
        _nested_item("ghost", "  ", 1),  # 落在长位内部嵌套出现 [1,3) 上
    ]
    out = _verify_slot_occ(merged, [cand])
    assert len(out) == 1
    assert out[0]["key"] == "long" and out[0]["occ"] == 1


# ── compute_anchor_positions 宽容回退（2026-09-18 事故回归：空白 anchor 非段内
#    完整 run 时跳过 pHash → 前端回退全文顺序匹配错位）──


def test_compute_anchor_positions_partial_ws_anchor_keeps_phash():
    """空白 anchor 落在更长空白 run 内（anchor 4 空格、段内 run 6 空格）→
    完整 run 校验计数为 0，此前 seq>occurrences 直接跳过不写 pHash → 前端
    回退全文顺序匹配分到别的段落（contract_no_94 错位事故）。修复后回退
    普通非重叠计数，仍写 p_idx/p_hash/a_occ/p_total。"""
    from rag.svr.template_fill.docx_utils import compute_anchor_positions
    blob = _make_docx(["合同编号为      。", "普通段落"])
    ph = {"key": "contract_no", "addr": "para:0", "anchor": "    "}  # 4 空格
    compute_anchor_positions(blob, [ph])
    assert ph.get("p_idx") == 0
    assert ph.get("p_hash")
    assert ph.get("a_occ") == 1
    assert ph.get("p_total") == 2


def test_compute_anchor_positions_absent_anchor_still_skipped():
    """anchor 在段内完全不存在 → 两通道都计 0 → 仍不写定位字段（不乱标）。"""
    from rag.svr.template_fill.docx_utils import compute_anchor_positions
    blob = _make_docx(["合同编号为      。", "普通段落"])
    ph = {"key": "nope", "addr": "para:0", "anchor": "       "}  # 7 空格，段内无
    compute_anchor_positions(blob, [ph])
    assert "p_hash" not in ph
    assert "a_occ" not in ph


def test_compute_anchor_positions_fallback_occ_ordering():
    """回退口径下同段多次出现按非重叠序分配：anchor 2 空格落在 4 空格 run 内
    （完整 run 校验 0 次）→ 非重叠 2 次 → 两条目 a_occ=1/2。"""
    from rag.svr.template_fill.docx_utils import compute_anchor_positions
    blob = _make_docx(["甲：      乙：      丙", "普通段落"])  # 两个 6 空格 run
    ph1 = {"key": "a", "addr": "para:0", "anchor": "  "}
    ph2 = {"key": "b", "addr": "para:0", "anchor": "  "}
    compute_anchor_positions(blob, [ph1, ph2])
    assert ph1.get("a_occ") == 1 and ph1.get("p_idx") == 0
    assert ph2.get("a_occ") == 2 and ph2.get("p_idx") == 0


def test_detect_fill_points_sorted_document_order(monkeypatch):
    """识别产物按文档序排序（2026-09-18 用户反馈：填写点列表顺序与预览不一致）。
    管线拼接序 = 手动直通 → V1（无位行）→ V2（有位行）→ 兜底，与文档序无关；
    B端列表按数组序渲染会与预览从上到下对不上。V1 项 line=9、V2 两项 line=2
    （段内偏移 5/1 乱序输出）→ 排序后 line2-pos1、line2-pos5、line9；
    内部 _anchor_pos 仍不外泄。"""
    import asyncio

    from rag.svr.template_fill import detector

    v1_cand = {"index": 9, "addr": "para:9", "text": "项目名称：____________", "slots": []}
    v2_cand = {
        "index": 2, "addr": "para:2", "text": "甲（提示甲）乙（提示乙）",
        "slots": [
            {"start": 1, "end": 6, "text": "（提示甲）", "kind": "hint", "hint": "提示甲"},
            {"start": 7, "end": 12, "text": "（提示乙）", "kind": "hint", "hint": "提示乙"},
        ],
    }

    def _v1_item():
        return {"key": "v1_field", "name": "项目名称", "description": "",
                "retrieval_query": "", "fill_mode": "llm", "required": True,
                "addr": "para:9", "anchor": "____________", "line": 9, "top_k": 6,
                "low_confidence": False}

    def _v2_item(key, name, pos):
        return {"key": key, "name": name, "description": "",
                "retrieval_query": "", "fill_mode": "llm", "required": True,
                "addr": "para:2", "line": 2, "top_k": 6, "low_confidence": False,
                "_anchor_pos": pos,
                "anchor": "（提示甲）" if key == "ka" else "（提示乙）"}

    async def fake_v1(chat, file_type, candidates):
        return [_v1_item()], 0

    async def fake_v2(chat, candidates):
        # 乱序输出：偏移 7 的项在前
        return [_v2_item("kb", "提示乙", 7), _v2_item("ka", "提示甲", 1)], {(2, 1), (2, 2)}, 0

    class _FakeModel:
        async def async_chat(self, system, messages):
            return "[]"

    monkeypatch.setattr(detector, "_get_chat_model", lambda tid: _FakeModel(), raising=False)
    monkeypatch.setattr(detector, "_detect_chunked", fake_v1)
    monkeypatch.setattr(detector, "_detect_slot_chunked", fake_v2, raising=False)

    merged = asyncio.run(detector.detect_fill_points("tenant", "docx", [v1_cand, v2_cand]))
    keys = [it["key"] for it in merged]
    assert keys == ["ka", "kb", "v1_field"]
    assert all("_anchor_pos" not in it for it in merged)


def test_parse_slot_response_blue_default_value():
    """blue 位豁免「不派生默认值」硬闸：default_value = 位文本（旧值即默认值）；
    其余位维持 "" 不回归。"""
    from rag.svr.template_fill.detector import parse_slot_response
    cand = _slot_cand()
    # membership 防御要求位文本是候选原文子串：把旧值真实嵌入候选文本 [30:36]
    cand["text"] = cand["text"][:30] + "福建省厦门市" + cand["text"][30:]
    cand["slots"].append({"start": 30, "end": 36, "text": "福建省厦门市",
                          "kind": "blue", "hint": ""})
    raw = '[{"line":7,"slot":1,"key":"project_name","name":"招标项目名称"},' \
          '{"line":7,"slot":3,"key":"city","name":"城市"}]'
    items, covered = parse_slot_response(raw, [cand])
    assert covered == {(7, 1), (7, 3)}
    assert items[0]["default_value"] == ""  # hint 位硬闸不回归
    assert items[1]["default_value"] == "福建省厦门市"
    assert items[1]["default_source"] == "detected"
    assert "default_source" not in items[0]  # hint 位不带来源键，_merge_defaults 照旧兜底
    assert items[1]["anchor"] == "福建省厦门市"
    assert items[1]["_anchor_pos"] == 30


def test_parse_slot_response_blue_slot_text_corrupt_dropped():
    """对抗：slot 文本不再是候选原文子串（坐标系腐坏）→ blue 位同样丢弃，
    不静默产脏 anchor/默认值。"""
    from rag.svr.template_fill.detector import parse_slot_response
    cand = _slot_cand()
    cand["slots"].append({"start": 30, "end": 36, "text": "不存在的旧值",
                          "kind": "blue", "hint": ""})
    raw = '[{"line":7,"slot":3,"key":"city","name":"城市"}]'
    items, covered = parse_slot_response(raw, [cand])
    assert items == [] and covered == set()
