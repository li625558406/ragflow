# test/test_file_edit_document.py
"""POST /files/<file_id>/edit（对话附件编辑）+ docx_edit 共享内核对抗性测试。

不依赖 Quart 运行时/DB/MinIO/Redis：直接从源文件加载 file_api 模块（stub 掉
api.apps 与 api.apps.services，避免 settings.init_settings），用假 manager 记录
路由后把目标端点挂到内存 Quart app 上，STORAGE_IMPL 换成进程内 dict 假存储。
内核测试照 test_flow_doc_table_edit.py 的零依赖模式加载 docx_edit.py。
"""

import io
import os
import sys
import types
import uuid
from importlib.util import module_from_spec, spec_from_file_location

import pytest
from docx import Document

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))


def _make_stub_module(name, **attrs):
    mod = types.ModuleType(name)
    for k, v in attrs.items():
        setattr(mod, k, v)
    sys.modules[name] = mod
    return mod


def _noop_decorator(f=None, *a, **kw):
    """login_required 透传桩：必须原样返回被装饰函数（否则端点名丢失）。"""
    return f


def _load_docx_edit():
    path = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "api", "utils", "docx_edit.py"))
    spec = spec_from_file_location("docx_edit_under_test_file_edit", path)
    mod = module_from_spec(spec)
    sys.modules["docx_edit_under_test_file_edit"] = mod
    spec.loader.exec_module(mod)
    return mod


class _FakeManager:
    def __init__(self):
        self.routes = {}

    def route(self, rule, methods=None):
        def deco(fn):
            self.routes[(rule, tuple(methods or []))] = fn
            return fn

        return deco


def _load_file_api():
    user_stub = types.SimpleNamespace(id="tenant-u1")
    _make_stub_module("api.apps", current_user=user_stub, login_required=_noop_decorator)
    _make_stub_module("api.apps.services", file_api_service=types.SimpleNamespace())
    path = os.path.abspath(
        os.path.join(os.path.dirname(__file__), "..", "api", "apps", "restful_apis", "file_api.py"))
    spec = spec_from_file_location("file_api_under_test", path)
    mod = module_from_spec(spec)
    # manager 在包 __init__ 里注入（ruff F821），独立加载须先塞进模块 dict
    mod.manager = _FakeManager()
    sys.modules["file_api_under_test"] = mod
    spec.loader.exec_module(mod)
    return mod


class FakeStorage:
    """同步 dict 存储：thread_pool_exec 走线程池，async 方法会拿到 coroutine 而非 bytes。"""

    def __init__(self):
        self.objects = {}
        self.get_calls = []

    def get(self, bucket, name):
        self.get_calls.append((bucket, name))
        return self.objects.get((bucket, name))

    def put(self, bucket, name, blob):
        self.objects[(bucket, name)] = blob


_docx_edit = _load_docx_edit()
parse_edit_payload = _docx_edit.parse_edit_payload
edit_docx_blob = _docx_edit.edit_docx_blob
safe_filename = _docx_edit.safe_filename

_file_api = _load_file_api()
EDIT_RULE = "/files/<file_id>/edit"

from quart import Quart

from common import settings as _settings


async def _json(resp):
    """resp.get_json() 在本环境偶发 None，直接读字节解析。"""
    import json as _json_mod

    return _json_mod.loads(await resp.get_data())


def _docx_blob(paras=("第一段", "第二段")) -> bytes:
    doc = Document()
    for t in paras:
        doc.add_paragraph(t)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _parse_blob_texts(blob) -> list:
    return [p.text for p in Document(io.BytesIO(blob)).paragraphs if (p.text or "").strip()]


# ---------------------------------------------------------------- 内核：parse_edit_payload 对抗矩阵


def test_parse_empty_rejected():
    with pytest.raises(ValueError, match="没有需要保存的改动"):
        parse_edit_payload({})
    with pytest.raises(ValueError, match="没有需要保存的改动"):
        parse_edit_payload({"edits": [], "deletes": [], "inserts": [], "table_edits": []})


def test_parse_non_list_and_limit():
    with pytest.raises(ValueError, match="必须是数组"):
        parse_edit_payload({"edits": "x"})
    ops = [{"para_index": i, "new_text": "t"} for i in range(201)]
    with pytest.raises(ValueError, match="单次最多修改 200 处"):
        parse_edit_payload({"edits": ops})


def test_parse_edit_item_adversarial():
    for bad, msg in [
        ("not-a-dict", "edits 项格式非法"),
        ({"para_index": None, "new_text": "t"}, "para_index 必须是整数"),
        ({"para_index": 0, "new_text": ""}, "段落内容不能为空"),
        ({"para_index": 0, "new_text": "  "}, "段落内容不能为空"),
        ({"para_index": 0, "new_text": "x" * 20001}, "单段内容不能超过 20000 字"),
        ({"para_index": 0, "new_text": "x", "runs": [{"text": "y"}]}, "runs 文本与 new_text 不一致"),
        ({"para_index": 0, "new_text": "x", "runs": []}, "runs 必须是非空数组或省略"),
        ({"para_index": 0, "new_text": "x", "runs": [{"text": "x", "color": "red"}]}, "格式非法"),
        ({"para_index": 0, "new_text": "x", "align": "diagonal"}, "align 非法"),
        ({"para_index": 0, "new_text": "x", "indent": 9}, "indent 超出范围"),
        ({"para_index": 0, "new_text": "x", "heading_level": 4}, "heading_level 超出范围"),
        # 控制字符被清洗后为空 → 拒绝
        ({"para_index": 0, "new_text": "\x00\x01"}, "段落内容不能为空"),
    ]:
        with pytest.raises(ValueError, match=msg):
            parse_edit_payload({"edits": [bad]})


def test_parse_delete_rules():
    with pytest.raises(ValueError, match="重复删除"):
        parse_edit_payload({"deletes": [3, 3]})
    with pytest.raises(ValueError, match="不能同时修改和删除"):
        parse_edit_payload({"edits": [{"para_index": 2, "new_text": "x"}], "deletes": [2]})
    with pytest.raises(ValueError, match="deletes 项必须是整数"):
        parse_edit_payload({"deletes": ["abc"]})


def test_parse_insert_rules():
    with pytest.raises(ValueError, match="新段落内容不能为空"):
        parse_edit_payload({"inserts": [{"after_para_index": 0, "new_text": ""}]})
    with pytest.raises(ValueError, match="after_para_index 必须是整数"):
        parse_edit_payload({"inserts": [{"after_para_index": "x", "new_text": "t"}]})
    with pytest.raises(ValueError, match="新段落 runs 文本与 new_text 不一致"):
        parse_edit_payload({"inserts": [{"after_para_index": -1, "new_text": "t", "runs": [{"text": "z"}]}]})


def test_parse_table_edits_adversarial():
    base = {"para_index": 1, "row": 0, "col": 0}
    cases = [
        ({"row": -1, "new_text": "x"}, "负数行列"),  # 负数
        ({"col": "a", "new_text": "x"}, "非整数"),  # 非整数
        ({"row": 0, "new_text": "x" * 20001}, "超长"),  # 超长
        ({"row": 0, "new_text": "x", "runs": [{"text": "y"}]}, "runs 不一致"),
        ({"row": 0, "new_text": "", "runs": [{"text": "y"}]}, "清空带 runs"),
    ]
    for extra, _label in cases:
        item = dict(base)
        item.update(extra)
        with pytest.raises(ValueError):
            parse_edit_payload({"table_edits": [item]})
    with pytest.raises(ValueError):
        parse_edit_payload({"table_edits": ["not-a-dict"]})


def test_parse_table_edits_empty_text_allowed():
    items = parse_edit_payload({"table_edits": [{"para_index": 1, "row": 0, "col": 0, "new_text": ""}]})
    assert items["table_edits"][0]["new_text"] == ""


def test_parse_ok_structure_and_ctrl_strip():
    parsed = parse_edit_payload(
        {
            "edits": [{"para_index": 0, "new_text": "a\x01b"}],
            "deletes": [1],
            "inserts": [{"after_para_index": -1, "new_text": "c"}],
            "table_edits": [],
        }
    )
    assert parsed["edits"] == [(0, "ab", None, {})]
    assert parsed["deletes"] == [1]
    assert parsed["inserts"] == [(-1, "c", None, {})]
    assert parsed["table_edits"] == []


# ---------------------------------------------------------------- 内核：edit_docx_blob


def test_edit_docx_blob_happy_path():
    blob = _docx_blob()
    parsed = parse_edit_payload({"edits": [{"para_index": 0, "new_text": "改后"}]})
    new_blob, root, converted = edit_docx_blob(blob, "测试文档.docx", parsed)
    assert converted is False
    assert root == "测试文档"
    texts = _parse_blob_texts(new_blob)
    assert texts == ["改后", "第二段"]


def test_edit_docx_blob_rejects_pdf():
    with pytest.raises(ValueError, match="仅支持编辑 doc/docx 文档"):
        edit_docx_blob(b"%PDF-1.4 fake", "x.pdf", parse_edit_payload({"deletes": [0]}))


def test_edit_docx_blob_locate_failure_input_untouched():
    blob = _docx_blob()
    parsed = parse_edit_payload({"edits": [{"para_index": 99, "new_text": "x"}]})
    with pytest.raises(ValueError, match="段落 99 定位失败"):
        edit_docx_blob(blob, "t.docx", parsed)
    # 定位失败不得污染输入（调用方还要用原 blob 回滚/报错）
    assert _parse_blob_texts(blob) == ["第一段", "第二段"]


def test_edit_docx_blob_atomic_block_rejected():
    doc = Document()
    doc.add_paragraph("前置段落")
    doc.add_table(rows=1, cols=1)
    out = io.BytesIO()
    doc.save(out)
    blob = out.getvalue()
    with pytest.raises(ValueError, match="是表格，不支持编辑"):
        edit_docx_blob(blob, "t.docx", parse_edit_payload({"edits": [{"para_index": 1, "new_text": "x"}]}))
    with pytest.raises(ValueError, match="是表格，不支持删除"):
        edit_docx_blob(blob, "t.docx", parse_edit_payload({"deletes": [1]}))


def test_edit_docx_blob_delete_and_insert():
    blob = _docx_blob(("甲", "乙", "丙"))
    parsed = parse_edit_payload({"deletes": [1], "inserts": [{"after_para_index": -1, "new_text": "文首"}]})
    new_blob, _, _ = edit_docx_blob(blob, "t.docx", parsed)
    assert _parse_blob_texts(new_blob) == ["文首", "甲", "丙"]


def test_safe_filename_traversal_and_long():
    assert safe_filename("../../etc/passwd") == "passwd"
    assert safe_filename("a\\b\\c.docx") == "c.docx"
    assert safe_filename("") == "unnamed"
    long = "长" * 300 + ".docx"
    cleaned = safe_filename(long)
    assert len(cleaned) <= 200 and cleaned.endswith(".docx")


# ---------------------------------------------------------------- 端点：/files/<file_id>/edit


@pytest.fixture()
def api_env(monkeypatch):
    fake = FakeStorage()
    monkeypatch.setattr(_settings, "STORAGE_IMPL", fake, raising=False)
    app = Quart(__name__)
    app.config["TESTING"] = True
    fn = _file_api.edit_file_document
    app.add_url_rule(EDIT_RULE, view_func=fn, methods=["POST"], endpoint="edit_file_document")
    return app, fake, app.test_client()


@pytest.mark.asyncio
async def test_endpoint_rejects_bad_file_id(api_env):
    _app, fake, client = api_env
    # 注：含 "/" 的穿越串与空串在路由层即 404，到不了端点守卫，不在本用例范围
    for bad in ["z" * 32, "a" * 31, "-" * 32, uuid.uuid4().hex.upper() + "!"]:
        resp = await client.post(f"/files/{bad}/edit", json={"edits": [{"para_index": 0, "new_text": "x"}]})
        data = await _json(resp)
        assert data["code"] != 0, bad
        assert "非法的文件 id" in data["message"], bad
    # 未触碰存储即拦截
    assert fake.get_calls == []


@pytest.mark.asyncio
async def test_endpoint_missing_object(api_env):
    _app, _fake, client = api_env
    fid = uuid.uuid4().hex
    resp = await client.post(f"/files/{fid}/edit", json={"deletes": [0]})
    data = await _json(resp)
    assert data["code"] != 0
    assert "文件不存在或已过期" in data["message"]


@pytest.mark.asyncio
async def test_endpoint_empty_ops(api_env):
    _app, fake, client = api_env
    fid = uuid.uuid4().hex
    fake.objects[("tenant-u1-downloads", fid)] = _docx_blob()
    resp = await client.post(f"/files/{fid}/edit", json={})
    data = await _json(resp)
    assert "没有需要保存的改动" in data["message"]


@pytest.mark.asyncio
async def test_endpoint_success_produces_new_file_original_untouched(api_env):
    _app, fake, client = api_env
    fid = uuid.uuid4().hex
    original = _docx_blob(("原始甲", "原始乙"))
    fake.objects[("tenant-u1-downloads", fid)] = original
    resp = await client.post(
        f"/files/{fid}/edit",
        json={"edits": [{"para_index": 0, "new_text": "编辑后"}], "file_name": "测试文档.docx"},
    )
    data = await _json(resp)
    assert data["code"] == 0, data
    assert data["data"]["file_id"] != fid
    assert data["data"]["file_name"] == "测试文档_编辑.docx"
    new_id = data["data"]["file_id"]
    # 新对象落进同一 bucket（租户隔离边界不漂移），内容生效
    assert ("tenant-u1-downloads", new_id) in fake.objects
    assert _parse_blob_texts(fake.objects[("tenant-u1-downloads", new_id)]) == ["编辑后", "原始乙"]
    # 原文件字节不变
    assert fake.objects[("tenant-u1-downloads", fid)] == original


@pytest.mark.asyncio
async def test_endpoint_pdf_rejected(api_env):
    _app, fake, client = api_env
    fid = uuid.uuid4().hex
    fake.objects[("tenant-u1-downloads", fid)] = b"%PDF-1.4 fake"
    resp = await client.post(
        f"/files/{fid}/edit", json={"deletes": [0], "file_name": "scan.pdf"}
    )
    data = await _json(resp)
    assert "仅支持编辑 doc/docx 文档" in data["message"]
    # 拒绝后不产生任何新对象
    assert len(fake.objects) == 1


@pytest.mark.asyncio
async def test_endpoint_locate_failure_no_side_effect(api_env):
    _app, fake, client = api_env
    fid = uuid.uuid4().hex
    original = _docx_blob()
    fake.objects[("tenant-u1-downloads", fid)] = original
    resp = await client.post(
        f"/files/{fid}/edit",
        json={"edits": [{"para_index": 5, "new_text": "x"}], "file_name": "t.docx"},
    )
    data = await _json(resp)
    assert "定位失败" in data["message"]
    assert len(fake.objects) == 1 and fake.objects[("tenant-u1-downloads", fid)] == original
