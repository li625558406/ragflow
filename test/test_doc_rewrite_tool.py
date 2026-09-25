# -*- coding: utf-8 -*-
"""DocumentRewrite 工具（agent/tools/document_rewrite.py）契约测试。

stub 画布 / 存储 / LLM / 版本链，验证 action 分发、download 契约、sys 容错。
工具实例用 object.__new__ 绕过 ToolBase.__init__（其要求真实 Canvas 实例，
同 test_template_fill_tool.py 惯例）。

patch 目标约定（对计划骨架的修正）：
- _load_chat_blob / _build_chat_mdl 是类方法 → patch.object(DocumentRewrite, ...)，
  patch 模块属性对实例方法查找无效；
- versions / rewriter 函数在工具模块顶层 import（versions.py/rewriter.py 自身零 DB
  依赖，注册期安全）→ 按模块属性 patch("agent.tools.document_rewrite.xxx") 即命中；
- FileService 为延迟 import → patch 源模块类属性
  patch("api.db.services.file_service.FileService.get_blob")。
"""
import io
import json
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

import pytest
from docx import Document

from agent.tools.document_rewrite import DocumentRewrite

# pytest-asyncio 插件与「timeout 装饰器线程内 asyncio.run」在 Windows Proactor
# 组合下，会遗留已停用 loop 的 GC ResourceWarning（纯测试环境伪影，禁用插件即消失，
# 生产路径不受影响）。本文件 warnings-as-errors 会把它升级成 unraisable 失败并
# 错记到下一个用例，故文件级抑制该伪影。
pytestmark = pytest.mark.filterwarnings("ignore::pytest.PytestUnraisableExceptionWarning")

DOC_ID = "tplfill-task1"
BASE_NAME = "方案"


def _docx_blob():
    doc = Document()
    doc.add_heading("第一节", level=1)
    doc.add_paragraph("第一节正文。")
    doc.add_heading("第二节", level=1)
    doc.add_paragraph("第二节正文。")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _no_heading_blob():
    doc = Document()
    doc.add_paragraph("没有标题的文档")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class FakeCanvas:
    def __init__(self, tenant_id="t1", sys_vars=None):
        self.globals = {"sys.pending_downloads": []}
        self.globals.update(sys_vars or {})
        self._tenant = tenant_id

    def get_tenant_id(self):
        return self._tenant

    def get_variable_value(self, name):
        if name in self.globals:
            return self.globals[name]
        raise KeyError(name)


def _make_tool(canvas):
    tool = object.__new__(DocumentRewrite)
    tool._param = SimpleNamespace(outputs={}, inputs={}, debug_inputs={})
    tool.set_output = lambda key, value=None: tool._param.outputs.update({key: {"value": value}})
    tool.check_if_canceled = lambda msg="": False
    tool._canvas = canvas
    return tool


def _fake_llm():
    class FakeLLM:
        async def async_chat(self, system, messages):
            return json.dumps({"paragraphs": ["重写后的新正文。"]}, ensure_ascii=False)

    return FakeLLM()


def _recent():
    return {"sys.recent_downloads": [{"doc_id": DOC_ID, "filename": BASE_NAME}]}


async def _fake_rewrite_section(*args, **kwargs):
    return ["重写后的新正文。"]


# ---------- meta 声明 ----------

def test_meta_declaration():
    assert DocumentRewrite.component_name == "DocumentRewrite"
    from agent.tools.document_rewrite import DocumentRewriteParam
    p = DocumentRewriteParam()
    assert p.meta["name"] == "DocumentRewrite"
    assert set(p.meta["parameters"]) >= {"action", "find_text", "replace_text",
                                         "section_no", "instruction", "version_no", "doc_id"}
    assert "replace" in p.meta["parameters"]["action"]["enum"]
    assert p.meta["parameters"]["action"]["required"] is True


# ---------- 文档来源解析 ----------

def test_missing_doc_context_returns_guidance():
    tool = _make_tool(FakeCanvas())  # 无 recent_downloads 无 doc_id 无 flow_version_id
    out = tool._invoke(action="outline")
    assert "无法确定" in out or "成稿" in out
    assert canvas_pending(tool) == []


def canvas_pending(tool):
    return tool._canvas.globals.get("sys.pending_downloads") or []


# ---------- outline ----------

def test_outline_action():
    tool = _make_tool(FakeCanvas(sys_vars=_recent()))
    with patch.object(DocumentRewrite, "_load_chat_blob",
                      MagicMock(return_value=(_docx_blob(), DOC_ID, BASE_NAME))):
        out = tool._invoke(action="outline")
    assert "第1节 第一节" in out
    assert "第2节 第二节" in out


# ---------- rewrite ----------

def test_rewrite_produces_download_contract():
    canvas = FakeCanvas(sys_vars=_recent())
    tool = _make_tool(canvas)
    blob = _docx_blob()
    reg_calls = {}

    def fake_reg(tenant_id, root_id, new_blob, file_type, base_file_name, **kw):
        reg_calls.update(tenant_id=tenant_id, root_id=root_id, blob=new_blob,
                         file_type=file_type, base_file_name=base_file_name, **kw)
        return {"version_no": 2, "obj": "rewrite-task1-v2", "file_name": "方案_v2.docx"}

    with patch.object(DocumentRewrite, "_load_chat_blob",
                      MagicMock(return_value=(blob, DOC_ID, BASE_NAME))), \
            patch.object(DocumentRewrite, "_build_chat_mdl", MagicMock(return_value=_fake_llm())), \
            patch("agent.tools.document_rewrite.rewrite_section", _fake_rewrite_section), \
            patch("agent.tools.document_rewrite.ensure_base_version") as ens, \
            patch("agent.tools.document_rewrite.register_chat_version", side_effect=fake_reg):
        out = tool._invoke(action="rewrite", section_no="2", instruction="补充进度安排")

    # download 契约进入 canvas 待发队列，恰 1 条
    dl = canvas.globals["sys.pending_downloads"]
    assert len(dl) == 1
    assert dl[0]["doc_id"] == "rewrite-task1-v2"
    assert dl[0]["filename"] == "方案_v2.docx"
    for key in ("doc_id", "filename", "name", "mime_type", "size", "url"):
        assert key in dl[0]
    assert dl[0]["url"].startswith("/api/v1/agents/download?id=rewrite-task1-v2")
    assert "created_by=t1" in dl[0]["url"]
    assert isinstance(dl[0]["size"], int) and dl[0]["size"] > 0
    # 文本说明可转述：含节标题与版本号
    assert "第二节" in out and "2" in out
    # 版本登记契约：root_id 剥 tplfill- 前缀、source_type/节标题正确
    assert reg_calls["tenant_id"] == "t1"
    assert reg_calls["root_id"] == "task1"
    assert reg_calls["source_type"] == "rewrite"
    assert reg_calls["section_title"] == "第二节"
    assert reg_calls["base_file_name"] == BASE_NAME
    # 首次重写补原始版本 v1（用原始 blob）
    ens.assert_called_once_with("t1", "task1", blob, "docx", BASE_NAME)


def test_rewrite_unknown_section_no():
    canvas = FakeCanvas(sys_vars=_recent())
    tool = _make_tool(canvas)
    with patch.object(DocumentRewrite, "_load_chat_blob",
                      MagicMock(return_value=(_docx_blob(), DOC_ID, BASE_NAME))):
        out = tool._invoke(action="rewrite", section_no="99", instruction="改")
    assert "99" in out and "不存在" in out
    assert canvas.globals["sys.pending_downloads"] == []


def test_rewrite_invalid_section_no_guides_outline():
    """对抗性：section_no 非数字（LLM 幻觉）不得击穿 int()。"""
    canvas = FakeCanvas(sys_vars=_recent())
    tool = _make_tool(canvas)
    with patch.object(DocumentRewrite, "_load_chat_blob",
                      MagicMock(return_value=(_docx_blob(), DOC_ID, BASE_NAME))):
        out = tool._invoke(action="rewrite", section_no="abc", instruction="改")
    assert "section_no" in out
    assert canvas.globals["sys.pending_downloads"] == []


def test_rewrite_missing_instruction():
    canvas = FakeCanvas(sys_vars=_recent())
    tool = _make_tool(canvas)
    with patch.object(DocumentRewrite, "_load_chat_blob",
                      MagicMock(return_value=(_docx_blob(), DOC_ID, BASE_NAME))):
        out = tool._invoke(action="rewrite", section_no="1", instruction="")
    assert "instruction" in out
    assert canvas.globals["sys.pending_downloads"] == []


def test_no_heading_document_guides_full_regen():
    canvas = FakeCanvas(sys_vars=_recent())
    tool = _make_tool(canvas)
    with patch.object(DocumentRewrite, "_load_chat_blob",
                      MagicMock(return_value=(_no_heading_blob(), DOC_ID, BASE_NAME))):
        out = tool._invoke(action="rewrite", section_no="1", instruction="改")
    assert "标题" in out or "重新生成" in out
    assert canvas.globals["sys.pending_downloads"] == []


# ---------- 链感知目标解析（C1：连续重写不丢内容、root_id 不断链）----------

_CHAIN = [
    {"version_no": 1, "source_type": "chat_fill", "obj": "rewrite-aaa",
     "file_name": "方案.docx", "section_title": "", "instruction": ""},
    {"version_no": 2, "source_type": "rewrite", "obj": "rewrite-bbb",
     "file_name": "方案.docx", "section_title": "第二节", "instruction": "补充进度安排"},
]


def _patch_real_blob_loader(blobs_by_obj=None):
    """不 patch _load_chat_blob（走真实链感知逻辑），只 patch 存储/版本链。

    FileService.get_blob 返回可解析 docx；list_versions 返回链；
    root_id_for_doc 模拟真实语义：rewrite- 产物经 DB 反查得 task1，
    tplfill- 前缀剥前缀（真实函数纯路径已由 versions 纯函数测试覆盖）。"""
    blob = _docx_blob()

    def fake_get_blob(tenant_id, obj):
        return (blobs_by_obj or {}).get(obj, blob)

    return (
        patch("api.db.services.file_service.FileService.get_blob", side_effect=fake_get_blob),
        patch("agent.tools.document_rewrite.root_id_for_doc",
              MagicMock(side_effect=lambda d: "task1")),
    )


def test_rewrite_from_tplfill_doc_switches_to_chain_latest():
    """C1 回归：doc_id 指向原始成稿但链上已有新版本 → 必须切到链内最新版内容，
    且版本登记 root_id 是 task1 而非 doc_id 原样（链不断）。"""
    canvas = FakeCanvas(sys_vars=_recent())
    tool = _make_tool(canvas)
    reg_calls = {}

    def fake_reg(tenant_id, root_id, new_blob, file_type, base_file_name, **kw):
        reg_calls.update(tenant_id=tenant_id, root_id=root_id, blob=new_blob,
                         file_type=file_type, base_file_name=base_file_name, **kw)
        return {"version_no": 3, "obj": "rewrite-ccc", "file_name": "方案_v3.docx"}

    blob_patch, root_patch = _patch_real_blob_loader()
    with blob_patch as gb, root_patch, \
            patch("agent.tools.document_rewrite.list_versions",
                  return_value=_CHAIN) as lv, \
            patch("agent.tools.document_rewrite.ensure_base_version") as ens, \
            patch("agent.tools.document_rewrite.register_chat_version", side_effect=fake_reg), \
            patch.object(DocumentRewrite, "_build_chat_mdl", MagicMock(return_value=_fake_llm())), \
            patch("agent.tools.document_rewrite.rewrite_section", _fake_rewrite_section):
        out = tool._invoke(action="rewrite", section_no="1", instruction="再改一次")

    # 先读显式/最近卡对象，再切到链内最新版 rewrite-bbb 的内容
    assert gb.call_args_list[0].args == ("t1", DOC_ID)
    assert ("t1", "rewrite-bbb") in [c.args for c in gb.call_args_list]
    lv.assert_called_once_with("task1")
    # 版本登记契约：root_id 是链锚 task1，而非切换后的 obj 或 doc_id 原样
    assert reg_calls["root_id"] == "task1"
    assert reg_calls["tenant_id"] == "t1"
    assert reg_calls["source_type"] == "rewrite"
    assert reg_calls["base_file_name"] == BASE_NAME
    assert isinstance(reg_calls["blob"], bytes) and reg_calls["blob"]
    # 链非空，ensure_base_version 仍被调用但真实实现本就空操作（此处 mock 不校验）
    ens.assert_called_once()
    # 成稿卡指向新对象
    dl = canvas.globals["sys.pending_downloads"]
    assert len(dl) == 1 and dl[0]["doc_id"] == "rewrite-ccc"
    assert "已完成" in out


def test_rewrite_from_chain_obj_resolves_root_via_db():
    """C1 反向：doc_id 本身是链上产物 rewrite-bbb → register 仍收到 root_id=task1（DB 反查）。"""
    canvas = FakeCanvas()  # 无 recent_downloads，显式指定链上产物
    tool = _make_tool(canvas)
    reg_calls = {}

    def fake_reg(tenant_id, root_id, new_blob, file_type, base_file_name, **kw):
        reg_calls.update(tenant_id=tenant_id, root_id=root_id, **kw)
        return {"version_no": 3, "obj": "rewrite-ccc", "file_name": "方案_v3.docx"}

    blob_patch, root_patch = _patch_real_blob_loader()
    with blob_patch as gb, root_patch, \
            patch("agent.tools.document_rewrite.list_versions", return_value=_CHAIN), \
            patch("agent.tools.document_rewrite.ensure_base_version"), \
            patch("agent.tools.document_rewrite.register_chat_version", side_effect=fake_reg), \
            patch.object(DocumentRewrite, "_build_chat_mdl", MagicMock(return_value=_fake_llm())), \
            patch("agent.tools.document_rewrite.rewrite_section", _fake_rewrite_section):
        out = tool._invoke(action="rewrite", doc_id="rewrite-bbb",
                           section_no="2", instruction="继续改")

    # 读取链上产物对象本身；链末位即该对象，无需切换
    gb.assert_called_once_with("t1", "rewrite-bbb")
    assert reg_calls["root_id"] == "task1"
    assert reg_calls["source_type"] == "rewrite"
    dl = canvas.globals["sys.pending_downloads"]
    assert len(dl) == 1 and dl[0]["doc_id"] == "rewrite-ccc"
    assert "已完成" in out


def test_rewrite_falls_back_when_latest_blob_missing():
    """对抗性：链末位对象 blob 丢失 → 保留原对象内容继续重写，不得整体失败。"""
    canvas = FakeCanvas(sys_vars=_recent())
    tool = _make_tool(canvas)
    reg_calls = {}

    def fake_reg(tenant_id, root_id, new_blob, file_type, base_file_name, **kw):
        reg_calls.update(tenant_id=tenant_id, root_id=root_id, **kw)
        return {"version_no": 3, "obj": "rewrite-ccc", "file_name": "方案_v3.docx"}

    blob = _docx_blob()

    def fake_get_blob(tenant_id, obj):
        return b"" if obj == "rewrite-bbb" else blob  # 链末位对象 blob 丢失

    with patch("api.db.services.file_service.FileService.get_blob",
               side_effect=fake_get_blob) as gb, \
            patch("agent.tools.document_rewrite.root_id_for_doc",
                  MagicMock(side_effect=lambda d: "task1")), \
            patch("agent.tools.document_rewrite.list_versions", return_value=_CHAIN), \
            patch("agent.tools.document_rewrite.ensure_base_version"), \
            patch("agent.tools.document_rewrite.register_chat_version", side_effect=fake_reg), \
            patch.object(DocumentRewrite, "_build_chat_mdl", MagicMock(return_value=_fake_llm())), \
            patch("agent.tools.document_rewrite.rewrite_section", _fake_rewrite_section):
        out = tool._invoke(action="rewrite", section_no="1", instruction="改")

    assert ("t1", "rewrite-bbb") in [c.args for c in gb.call_args_list]
    assert reg_calls["root_id"] == "task1"
    dl = canvas.globals["sys.pending_downloads"]
    assert len(dl) == 1 and dl[0]["doc_id"] == "rewrite-ccc"
    assert "已完成" in out


def test_emit_download_resets_corrupted_pending():
    """I2：sys.pending_downloads 被污染为非 list → 强制重置为恰含 1 条契约的 list。"""
    canvas = FakeCanvas(sys_vars=_recent())
    canvas.globals["sys.pending_downloads"] = "corrupt"
    tool = _make_tool(canvas)
    reg_calls = {}

    def fake_reg(tenant_id, root_id, new_blob, file_type, base_file_name, **kw):
        reg_calls.update(tenant_id=tenant_id, root_id=root_id, **kw)
        return {"version_no": 2, "obj": "rewrite-v2", "file_name": "方案_v2.docx"}

    blob_patch, root_patch = _patch_real_blob_loader()
    with blob_patch, root_patch, \
            patch("agent.tools.document_rewrite.list_versions", return_value=[]), \
            patch("agent.tools.document_rewrite.ensure_base_version"), \
            patch("agent.tools.document_rewrite.register_chat_version", side_effect=fake_reg), \
            patch.object(DocumentRewrite, "_build_chat_mdl", MagicMock(return_value=_fake_llm())), \
            patch("agent.tools.document_rewrite.rewrite_section", _fake_rewrite_section):
        out = tool._invoke(action="rewrite", section_no="1", instruction="改")

    assert "已完成" in out
    dl = canvas.globals["sys.pending_downloads"]
    assert isinstance(dl, list)
    assert len(dl) == 1
    assert dl[0]["doc_id"] == "rewrite-v2"
    for key in ("filename", "name", "mime_type", "size", "url"):
        assert key in dl[0]


# ---------- replace（精准替换：2026-09-25 事故「改成X」曾被设计成整章重写）----------

def test_replace_produces_contract_and_applies_change():
    """happy path：docx 实际被改 + 契约/版本登记 source_type=replace。"""
    canvas = FakeCanvas(sys_vars=_recent())
    tool = _make_tool(canvas)
    reg_calls = {}

    def fake_reg(tenant_id, root_id, new_blob, file_type, base_file_name, **kw):
        reg_calls.update(tenant_id=tenant_id, root_id=root_id, blob=new_blob, **kw)
        return {"version_no": 2, "obj": "rewrite-task1-v2", "file_name": "方案_v2.docx"}

    with patch.object(DocumentRewrite, "_load_chat_blob",
                      MagicMock(return_value=(_docx_blob(), DOC_ID, BASE_NAME))), \
            patch("agent.tools.document_rewrite.list_versions", return_value=[]), \
            patch("agent.tools.document_rewrite.ensure_base_version") as ens, \
            patch("agent.tools.document_rewrite.register_chat_version", side_effect=fake_reg):
        out = tool._invoke(action="replace",
                           find_text="第二节正文", replace_text="LG11111")

    # 文档真实改动：登记收到的新 blob 里第二节正文已被替换、其他内容原样
    changed = Document(io.BytesIO(reg_calls["blob"]))
    texts = [p.text for p in changed.paragraphs]
    assert "LG11111。" in texts and "第一节正文。" in texts
    assert reg_calls["source_type"] == "replace"
    assert reg_calls["root_id"] == "task1"
    assert "第二节正文" in reg_calls["instruction"]
    ens.assert_called_once()
    dl = canvas.globals["sys.pending_downloads"]
    assert len(dl) == 1 and dl[0]["doc_id"] == "rewrite-task1-v2"
    assert "精准替换" in out and "共 1 处" in out


def test_replace_no_match_zero_change_and_guidance():
    """对抗：找不到原文必须零改动（不落版本不出卡），给可执行的引导文案。"""
    canvas = FakeCanvas(sys_vars=_recent())
    tool = _make_tool(canvas)
    with patch.object(DocumentRewrite, "_load_chat_blob",
                      MagicMock(return_value=(_docx_blob(), DOC_ID, BASE_NAME))), \
            patch("agent.tools.document_rewrite.register_chat_version") as reg:
        out = tool._invoke(action="replace",
                           find_text="文档里根本不存在的句子XYZ", replace_text="新文本")
    reg.assert_not_called()
    assert canvas.globals["sys.pending_downloads"] == []
    assert "未找到" in out and "outline" in out


def test_replace_missing_params():
    canvas = FakeCanvas(sys_vars=_recent())
    tool = _make_tool(canvas)
    out1 = tool._invoke(action="replace", find_text="", replace_text="x")
    out2 = tool._invoke(action="replace", find_text="原文", replace_text="  ")
    assert "find_text" in out1
    assert "replace_text" in out2
    assert canvas.globals["sys.pending_downloads"] == []


def test_replace_short_find_text_rejected():
    """对抗：过短 find_text 全篇逐字替换极易误伤，直接拒绝。"""
    canvas = FakeCanvas(sys_vars=_recent())
    tool = _make_tool(canvas)
    out = tool._invoke(action="replace", find_text="的", replace_text="X")
    assert "过短" in out
    assert canvas.globals["sys.pending_downloads"] == []


# ---------- versions ----------

def test_versions_action_empty_chain():
    canvas = FakeCanvas(sys_vars=_recent())
    tool = _make_tool(canvas)
    with patch.object(DocumentRewrite, "_load_chat_blob",
                      MagicMock(return_value=(_docx_blob(), DOC_ID, BASE_NAME))), \
            patch("agent.tools.document_rewrite.list_versions", return_value=[]):
        out = tool._invoke(action="versions")
    assert "版本" in out


def test_versions_lists_chain():
    canvas = FakeCanvas(sys_vars=_recent())
    tool = _make_tool(canvas)
    rows = [
        {"version_no": 1, "source_type": "chat_fill", "file_name": "方案.docx",
         "section_title": "", "instruction": ""},
        {"version_no": 2, "source_type": "rewrite", "file_name": "方案.docx",
         "section_title": "第二节", "instruction": "补充进度安排"},
    ]
    with patch.object(DocumentRewrite, "_load_chat_blob",
                      MagicMock(return_value=(_docx_blob(), DOC_ID, BASE_NAME))), \
            patch("agent.tools.document_rewrite.list_versions", return_value=rows) as lv:
        out = tool._invoke(action="versions")
    lv.assert_called_once_with("task1")  # root_id 剥 tplfill- 前缀
    assert "v1" in out and "v2" in out and "第二节" in out


# ---------- rollback ----------

def test_rollback_copies_history_version():
    canvas = FakeCanvas(sys_vars=_recent())
    tool = _make_tool(canvas)
    hist = {"version_no": 1, "obj": DOC_ID, "bucket": "t1-downloads",
            "file_name": "方案.docx", "file_type": "docx",
            "section_title": "", "instruction": ""}
    reg_calls = {}

    def fake_reg(tenant_id, root_id, blob, file_type, base_file_name, **kw):
        reg_calls.update(tenant_id=tenant_id, root_id=root_id, blob=blob,
                         file_type=file_type, base_file_name=base_file_name, **kw)
        return {"version_no": 3, "obj": "rewrite-task1-v3", "file_name": "方案_v3.docx"}

    with patch.object(DocumentRewrite, "_load_chat_blob",
                      MagicMock(return_value=(_docx_blob(), DOC_ID, BASE_NAME))), \
            patch("agent.tools.document_rewrite.get_version", return_value=hist) as gv, \
            patch("api.db.services.file_service.FileService.get_blob",
                  return_value=_docx_blob()) as gb, \
            patch("agent.tools.document_rewrite.register_chat_version", side_effect=fake_reg), \
            patch("agent.tools.document_rewrite.list_versions", return_value=[hist]):
        out = tool._invoke(action="rollback", version_no="1")

    # get_version 以 (剥离 tplfill- 后的 root_id, int) 调用
    gv.assert_called_once_with("task1", 1)
    gb.assert_called_once_with("t1", DOC_ID)
    assert reg_calls["source_type"] == "rollback"
    assert reg_calls["root_id"] == "task1"
    dl = canvas.globals["sys.pending_downloads"]
    assert len(dl) == 1 and dl[0]["doc_id"] == "rewrite-task1-v3"
    assert dl[0]["filename"] == "方案_v3.docx"
    assert "v3" in out or "已回退" in out


def test_rollback_missing_version():
    canvas = FakeCanvas(sys_vars=_recent())
    tool = _make_tool(canvas)
    with patch.object(DocumentRewrite, "_load_chat_blob",
                      MagicMock(return_value=(_docx_blob(), DOC_ID, BASE_NAME))), \
            patch("agent.tools.document_rewrite.get_version", return_value=None), \
            patch("agent.tools.document_rewrite.list_versions", return_value=[]):
        out = tool._invoke(action="rollback", version_no="42")
    assert "42" in out and "不存在" in out
    assert canvas.globals["sys.pending_downloads"] == []


def test_rollback_missing_version_no():
    canvas = FakeCanvas(sys_vars=_recent())
    tool = _make_tool(canvas)
    out = tool._invoke(action="rollback")
    assert "version_no" in out
    assert canvas.globals["sys.pending_downloads"] == []


# ---------- 分发与兜底 ----------

def test_bad_action():
    tool = _make_tool(FakeCanvas())
    out = tool._invoke(action="destroy")
    assert "action" in out or "不支持" in out
    assert canvas_pending(tool) == []


def test_top_level_exception_returns_fallback_text():
    """对抗性：底层异常不得裸抛，须记日志 + set_output(_ERROR) + 兜底文本。"""
    canvas = FakeCanvas(sys_vars=_recent())
    tool = _make_tool(canvas)
    with patch.object(DocumentRewrite, "_load_chat_blob",
                      MagicMock(side_effect=RuntimeError("minio down"))):
        out = tool._invoke(action="outline")
    assert "失败" in out
    assert "_ERROR" in tool._param.outputs
