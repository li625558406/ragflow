"""blank_slots 切位器单测：run 层位判定与区间合并（对抗性全场景）。"""

from docx import Document
from docx.oxml.ns import qn


def _add_runs(p, runs_spec):
    """runs_spec: [(text, underline)]，underline=True 时设格式下划线。"""
    for text, u in runs_spec:
        r = p.add_run(text)
        if u is not None:
            r.font.underline = u
    return p


def _para(runs_spec):
    doc = Document()
    p = doc.add_paragraph()
    _add_runs(p, runs_spec)
    return p


def _add_hyperlink(p, text):
    """段落内追加一个 w:hyperlink（模拟超链接 run 场景）。"""
    hl = p.part.relate_to("https://example.com", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    el = p._p.makeelement(qn("w:hyperlink"), {qn("r:id"): hl})
    r = p._p.makeelement(qn("w:r"), {})
    t = p._p.makeelement(qn("w:t"), {})
    t.text = text
    r.append(t)
    el.append(r)
    p._p.append(el)


def test_underline_blank_run():
    """真实范本形态：下划线格式纯空格 run = blank 位。"""
    from rag.svr.template_fill.blank_slots import extract_paragraph_slots

    p = _para([("本招标项目", None), ("             ", True), ("，建设资金来自", None)])
    slots = extract_paragraph_slots(p)
    assert len(slots) == 1
    s = slots[0]
    assert s["kind"] == "blank"
    assert s["text"] == "             "
    # runs 拼接文本坐标系：区间切片与 text 一致
    runs_text = "".join(r.text for r in p.runs)
    assert runs_text[s["start"] : s["end"]] == s["text"]


def test_underline_hint_run():
    """下划线格式「空格+（提示）+空格」run = hint 位，hint 去空白。"""
    from rag.svr.template_fill.blank_slots import extract_paragraph_slots

    p = _para([("本招标项目", None), ("    （招标项目名称） ", True), ("已由", None)])
    slots = extract_paragraph_slots(p)
    assert len(slots) == 1
    assert slots[0]["kind"] == "hint"
    assert slots[0]["hint"] == "招标项目名称"


def test_char_underscores_no_format():
    """无格式的连续 _ 字符 = blank 位（手打下划线）。"""
    from rag.svr.template_fill.blank_slots import extract_paragraph_slots

    p = _para([("项目名称：", None), ("____________", None)])
    slots = extract_paragraph_slots(p)
    assert len(slots) == 1
    assert slots[0]["kind"] == "blank"
    assert slots[0]["text"] == "____________"


def test_solid_underlined_text_not_slot():
    """实心文字加下划线格式（标题词）不是位——防误报关键闸。"""
    from rag.svr.template_fill.blank_slots import extract_paragraph_slots

    p = _para([("第一章 ", None), ("招标公告", True), ("（加粗下划线标题）", None)])
    assert extract_paragraph_slots(p) == []


def test_underline_val_none_not_slot():
    """显式 val=none（font.underline=False）不是位。"""
    from rag.svr.template_fill.blank_slots import extract_paragraph_slots

    p = _para([("文字 ", None), ("   ", False), ("续", None)])
    assert extract_paragraph_slots(p) == []


def test_underline_inherited_not_slot():
    """underline=None（继承样式）不是位（首版只认 run 级显式设置）。"""
    from rag.svr.template_fill.blank_slots import extract_paragraph_slots

    p = _para([("   ", None)])
    assert extract_paragraph_slots(p) == []


def test_adjacent_underline_runs_merged():
    """相邻下划线 run（中间无非空白文字）合并为一个位。"""
    from rag.svr.template_fill.blank_slots import extract_paragraph_slots

    p = _para([("编号：", None), ("（项目编号）", True), ("  ", True), ("___", None)])
    slots = extract_paragraph_slots(p)
    assert len(slots) == 1
    assert slots[0]["kind"] == "hint"  # 合并组内有 hint 则整体 hint
    assert slots[0]["hint"] == "项目编号"
    runs_text = "".join(r.text for r in p.runs)
    assert runs_text[slots[0]["start"] : slots[0]["end"]] == "（项目编号）  ___"


def test_slots_separated_by_text_not_merged():
    """中间隔实心文字的位不合并，各自独立。"""
    from rag.svr.template_fill.blank_slots import extract_paragraph_slots

    p = _para([("    （甲） ", True), ("已由", None), (" （乙） ", True)])
    slots = extract_paragraph_slots(p)
    assert len(slots) == 2
    assert [s["hint"] for s in slots] == ["甲", "乙"]


def test_same_shape_blank_slots_both_kept():
    """同段两个同形纯空白位都保留（occ 由下游预分配）。"""
    from rag.svr.template_fill.blank_slots import extract_paragraph_slots

    p = _para([("为", None), ("     ", True), ("，金额", None), ("     ", True), ("元。", None)])
    slots = extract_paragraph_slots(p)
    assert len(slots) == 2
    assert slots[0]["start"] < slots[1]["start"]
    assert slots[0]["text"] == slots[1]["text"] == "     "


def test_hyperlink_paragraph_no_slots():
    """段落含超链接 → 不切位（坐标系不可靠，整行回退 V1）。"""
    from rag.svr.template_fill.blank_slots import extract_paragraph_slots

    doc = Document()
    p = doc.add_paragraph()
    _add_runs(p, [("前缀", None), ("     ", True)])
    _add_hyperlink(p, "链接文字")
    assert extract_paragraph_slots(p) == []


def test_empty_paragraph_no_slots():
    from rag.svr.template_fill.blank_slots import extract_paragraph_slots

    doc = Document()
    assert extract_paragraph_slots(doc.add_paragraph("")) == []


def test_oversize_slot_dropped():
    """超长留白 run（>500 字符）丢弃——下游 validate 会拒，源头不产脏位。"""
    from rag.svr.template_fill.blank_slots import extract_paragraph_slots

    p = _para([("x", None), (" " * 600, True)])
    assert extract_paragraph_slots(p) == []


def test_empty_paren_hint_not_slot():
    """空括号「（）」无语义：hint 文本 strip 后为空则该 run 不成位（宁可漏不误）。"""
    from rag.svr.template_fill.blank_slots import extract_paragraph_slots

    p = _para([("前缀", None), ("（）", True)])
    assert extract_paragraph_slots(p) == []


def test_hint_overlong_truncated_to_100():
    """hint 内文超 100 字符截断（沿用现有 name 100 上限）。"""
    from rag.svr.template_fill.blank_slots import extract_paragraph_slots

    p = _para([("x", None), ("（" + "项" * 120 + "）", True)])
    slots = extract_paragraph_slots(p)
    assert len(slots) == 1
    assert len(slots[0]["hint"]) == 100


# ── 簇级分类与空白扩展（2026-09-18 事故回归：逐 run 全匹配拆碎片/漏提示语）──


def test_split_run_hint_cluster_two_slots():
    """真实事故形态（机电监理范本 para85）：Word 把提示语拆成多 run——
    '   '+'('+'项目名称'+')   ' 连续下划线 run 必须按簇合并成一个 hint 位，
    而不是只把空白 run 切成碎片、丢掉提示语。"""
    from rag.svr.template_fill.blank_slots import extract_paragraph_slots

    p = _para([
        ("本招标项目", None),
        ("   ", True), ("(", True), ("项目名称", True), (")   ", True),
        ("已由", None),
        ("   ", True), ("(", True),
        ("项目审批、核准或备案机关名称", True), (")    ", True),
        ("以", None),
    ])
    slots = extract_paragraph_slots(p)
    assert len(slots) == 2
    assert [s["kind"] for s in slots] == ["hint", "hint"]
    assert [s["hint"] for s in slots] == ["项目名称", "项目审批、核准或备案机关名称"]
    runs_text = "".join(r.text for r in p.runs)
    for s in slots:
        assert runs_text[s["start"] : s["end"]] == s["text"]


def test_multi_hint_cluster_split_into_slots():
    """一簇连续下划线 run 含多个括号提示 → 拆成多个 hint 位（para919 形态）。"""
    from rag.svr.template_fill.blank_slots import extract_paragraph_slots

    p = _para([
        ("对该", None),
        ("（", True), ("招标", True), ("项目名称）", True), ("（合同段名称）", True),
        ("合同段的投标。", None),
    ])
    slots = extract_paragraph_slots(p)
    assert [s["hint"] for s in slots] == ["招标项目名称", "合同段名称"]
    runs_text = "".join(r.text for r in p.runs)
    assert runs_text[slots[0]["start"] : slots[0]["end"]] == "（招标项目名称）"
    assert runs_text[slots[1]["start"] : slots[1]["end"]] == "（合同段名称）"


def test_underscore_paren_mixed_single_run_one_hint_slot():
    """「＿＿＿（项目名称）＿＿」混合形态（单下划线 run）：一个 hint 位，
    不再拆成两个空白位把提示语夹丢。"""
    from rag.svr.template_fill.blank_slots import extract_paragraph_slots

    p = _para([("本招标项目", None), ("＿＿＿（项目名称）＿＿", True), ("已由", None)])
    slots = extract_paragraph_slots(p)
    assert len(slots) == 1
    assert slots[0]["kind"] == "hint"
    assert slots[0]["hint"] == "项目名称"
    assert slots[0]["text"] == "＿＿＿（项目名称）＿＿"


def test_cluster_paren_plus_solid_falls_back_per_run():
    """簇内括号提示与实心文字混排（簇分类失败）→ 逐 run 回退，提示 run 仍成位。"""
    from rag.svr.template_fill.blank_slots import extract_paragraph_slots

    p = _para([("（甲）", True), ("实心", True)])
    slots = extract_paragraph_slots(p)
    assert len(slots) == 1
    assert slots[0]["kind"] == "hint"
    assert slots[0]["hint"] == "甲"


def test_blank_slot_expands_to_maximal_ws_run():
    """真实事故形态（合同编号 para76）：下划线 run 4 空格 + 相邻非下划线空格
    → anchor 扩展为段内最大空白 run（完整 run 才能被 compute_anchor_positions /
    前端 raw 通道按完整空白出现定位）。"""
    from rag.svr.template_fill.blank_slots import extract_paragraph_slots

    p = _para([("合同编号为 ", None), ("    ", True), ("。", None)])
    slots = extract_paragraph_slots(p)
    assert len(slots) == 1
    assert slots[0]["kind"] == "blank"
    assert slots[0]["text"] == "     "  # 5 空格 = 下划线 4 + 相邻 1
    runs_text = "".join(r.text for r in p.runs)
    assert runs_text[slots[0]["start"] : slots[0]["end"]] == slots[0]["text"]


def test_expansion_stops_at_non_whitespace_both_sides():
    """扩展只在空白内进行：两侧非空白（或已有位边界）即停，多位互不侵占。"""
    from rag.svr.template_fill.blank_slots import extract_paragraph_slots

    p = _para([
        ("为", None), (" ", None), ("   ", True), ("，金额", None),
        ("     ", True), (" 元", None),
    ])
    slots = extract_paragraph_slots(p)
    assert len(slots) == 2
    assert slots[0]["text"] == "    "   # 3 下划线空格 + 前邻 1 空格
    assert slots[1]["text"] == "      "  # 5 下划线空格 + 后邻 1 空格


def test_ws_gap_between_blank_and_hint_merges_into_hint():
    """空白位与提示位只隔空白文本 → 合并为一个 hint 位（既有 gap 合并语义回归闸）。"""
    from rag.svr.template_fill.blank_slots import extract_paragraph_slots

    p = _para([("编号", None), ("   ", True), (" ", None), ("（名称）  ", True)])
    slots = extract_paragraph_slots(p)
    assert len(slots) == 1
    assert slots[0]["kind"] == "hint"
    assert slots[0]["text"] == "    （名称）  "
